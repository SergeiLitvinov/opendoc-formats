"""DOCX → Text.

Единый ридер DOCX; раньше было две копии (extract/text.py и
organize/extractors/docx.py).
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any, Union

from opendoc_model.document_model import (
    Block as RichBlock,
)
from opendoc_model.document_model import (
    ConversionMode,
    DocumentModel,
    PackageGraph,
    Paragraph,
    Provenance,
    ProvenanceEvent,
)
from opendoc_model.document_model import Table as RichTable
from opendoc_model.units import canonical_coordinate_contract

from opendoc_formats.ooxml.package import (
    RELATIONSHIP_TYPE,
    load_package_graph,
)
from opendoc_formats.readers.docx_features import inspect_docx_features
from opendoc_formats.readers.docx_section import read_section
from opendoc_formats.readers.docx_style import (
    read_document_defaults,
    read_document_styles,
)
from opendoc_formats.readers.docx_table import read_table
from opendoc_formats.readers.docx_text import read_block_ooxml, read_paragraph
from opendoc_formats.support.io import check_archive_safety
from opendoc_formats.types import Block, BlockType, DocFormat, Table, Text


def read_docx(path: Union[str, Path], *, include_tables: bool = True) -> Text:
    from docx import Document  # type: ignore[import-not-found]

    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(p)
    check_archive_safety(p)
    doc = Document(str(p))
    blocks: list[Block] = []
    tables: list[Table] = []
    plain: list[str] = []

    for para in doc.paragraphs:
        t = para.text
        blocks.append(Block(type=BlockType.PARAGRAPH, text=t))
        if t:
            plain.append(t)

    if include_tables:
        for tbl in doc.tables:
            rows = [[cell.text for cell in row.cells] for row in tbl.rows]
            tables.append(Table(rows=rows))
            for row in rows:
                plain.append(" | ".join(row))

    return Text(
        blocks=blocks,
        tables=tables,
        plain="\n".join(plain),
        source_format=DocFormat.DOCX,
        engine="python-docx",
    )


def read_docx_model(
    path: Union[str, Path],
    *,
    mode: ConversionMode = ConversionMode.BALANCED,
) -> DocumentModel:
    """Импортировать DOCX в богатую модель с оформлением и ресурсами."""

    from docx import Document as OpenDocument  # type: ignore[import-not-found]
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    check_archive_safety(source)
    document = OpenDocument(str(source))
    model = DocumentModel(
        mode=mode,
        source_format=DocFormat.DOCX.value,
        metadata=_document_metadata(document, source),
    )
    model.package = _load_docx_package_graph(document)
    model.metadata["docx_features"] = inspect_docx_features(document)
    model.styles = read_document_styles(document, model)
    document_defaults = read_document_defaults(document, model)
    if document_defaults is not None:
        model.styles["__doc_defaults__"] = document_defaults

    section_blocks: list[list[RichBlock]] = [[]]
    for item in _iter_container_content(document):
        if isinstance(item, DocxParagraph):
            section_blocks[-1].append(read_paragraph(item, model))
            if item._p.pPr is not None and item._p.pPr.sectPr is not None:
                section_blocks.append([])
        elif isinstance(item, DocxTable):
            section_blocks[-1].append(read_table(item, model, _container_blocks))
        else:
            section_blocks[-1].append(read_block_ooxml(item))

    if len(section_blocks) > 1 and not section_blocks[-1]:
        section_blocks.pop()
    for index, blocks in enumerate(section_blocks):
        source_section = document.sections[min(index, len(document.sections) - 1)]
        model.sections.append(
            read_section(
                source_section,
                blocks,
                model,
                index,
                _container_blocks,
                odd_and_even_pages=bool(document.settings.odd_and_even_pages_header_footer),
            )
        )
    _attach_docx_provenance(model, source)
    return model


def _attach_docx_provenance(model: DocumentModel, source: Path) -> None:
    """Annotate semantic nodes after parsing without leaking OOXML fields into their contracts."""

    def origin(object_id: str, package_part: str, detail: str) -> Provenance:
        return Provenance(
            source_format=DocFormat.DOCX.value,
            source_path=str(source),
            object_id=object_id,
            package_part=package_part,
            events=[ProvenanceEvent("import.docx", detail)],
        )

    def annotate(block: RichBlock, object_id: str, package_part: str) -> None:
        if getattr(block, "provenance", None) is None:
            block.provenance = origin(object_id, package_part, f"parsed {type(block).__name__.lower()}")
        surrogate = getattr(block, "visual_surrogate", None)
        if surrogate is not None and not any(event.fallback_reason == surrogate.reason for event in block.provenance.events):
            block.provenance = block.provenance.transformed(
                "fallback.visual-surrogate",
                detail=f"linked visual surrogate {surrogate.resource_id}",
                fallback_reason=surrogate.reason,
            )
        if isinstance(block, Paragraph):
            for inline_index, item in enumerate(block.content):
                if getattr(item, "provenance", None) is None:
                    item.provenance = origin(f"{object_id}-inline-{inline_index}", package_part, "parsed inline content")
                item_surrogate = getattr(item, "visual_surrogate", None)
                if item_surrogate is not None:
                    item.provenance = item.provenance.transformed(
                        "fallback.visual-surrogate",
                        detail=f"linked visual surrogate {item_surrogate.resource_id}",
                        fallback_reason=item_surrogate.reason,
                    )
        elif isinstance(block, RichTable):
            for row_index, row in enumerate(block.rows):
                for cell_index, cell in enumerate(row.cells):
                    for child_index, child in enumerate(cell.blocks):
                        annotate(child, f"{object_id}-r{row_index}-c{cell_index}-b{child_index}", package_part)

    collections = (
        ("blocks", "/word/document.xml"),
        ("headers", "/word/header.xml"),
        ("footers", "/word/footer.xml"),
        ("first_page_headers", "/word/header.xml"),
        ("first_page_footers", "/word/footer.xml"),
        ("even_page_headers", "/word/header.xml"),
        ("even_page_footers", "/word/footer.xml"),
    )
    for section_index, section in enumerate(model.sections):
        if section.provenance is None:
            section.provenance = origin(f"section-{section_index}", "/word/document.xml", "parsed section properties")
        for collection_name, package_part in collections:
            for block_index, block in enumerate(getattr(section, collection_name)):
                annotate(block, f"section-{section_index}-{collection_name}-{block_index}", package_part)
    for resource_id, resource in model.resources.items():
        if resource.provenance is None:
            package_part = f"/word/media/{resource.filename}" if resource.filename else "/word/media"
            resource.provenance = origin(resource_id, package_part, "copied embedded resource")


def _load_docx_package_graph(document: Any) -> PackageGraph | None:
    supported = set(RELATIONSHIP_TYPE.values())
    recursive = {
        RELATIONSHIP_TYPE["footnotes"],
        RELATIONSHIP_TYPE["endnotes"],
        RELATIONSHIP_TYPE["comments"],
        RELATIONSHIP_TYPE["diagram_data"],
        RELATIONSHIP_TYPE["ole_object"],
        RELATIONSHIP_TYPE["package"],
    }
    return load_package_graph(
        document.part,
        format_name="ooxml",
        supported_relationships=supported,
        recursive_relationships=recursive,
        include=lambda relationship: relationship.reltype != RELATIONSHIP_TYPE["numbering"] or _document_uses_numbering(document),
    )


def _document_uses_numbering(document: Any) -> bool:
    from docx.oxml.ns import qn

    if document.element.xpath(".//w:pPr/w:numPr"):
        return True
    used_style_ids = {node.get(qn("w:val")) for node in document.element.xpath(".//w:pPr/w:pStyle") if node.get(qn("w:val"))}
    return any(style.style_id in used_style_ids and bool(style.element.xpath("./w:pPr/w:numPr")) for style in document.styles)


def _document_metadata(document: Any, source: Path) -> dict[str, Any]:
    props = document.core_properties
    return {
        "title": props.title or "",
        "subject": props.subject or "",
        "author": props.author or "",
        "keywords": props.keywords or "",
        "comments": props.comments or "",
        "created": props.created.isoformat() if props.created else None,
        "modified": props.modified.isoformat() if props.modified else None,
        "source_name": source.name,
        "coordinate_system": canonical_coordinate_contract(),
    }


def _container_blocks(container: Any, model: DocumentModel) -> list[RichBlock]:
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    result: list[RichBlock] = []
    for item in _iter_container_content(container):
        if isinstance(item, DocxParagraph):
            result.append(read_paragraph(item, model))
        elif isinstance(item, DocxTable):
            result.append(read_table(item, model, _container_blocks))
        else:
            result.append(read_block_ooxml(item))
    return result


def _iter_container_content(container: Any) -> Iterator[Any]:
    """Yield paragraphs, tables, and opaque block-level content controls in order."""

    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph
    from lxml import etree

    parent = container._body if hasattr(container, "_body") else container
    root = parent._element
    for element in root.iterchildren():
        if isinstance(element, CT_P):
            yield DocxParagraph(element, parent)
        elif isinstance(element, CT_Tbl):
            yield DocxTable(element, parent)
        elif etree.QName(element).localname == "sdt":
            yield element


__all__ = ["read_docx", "read_docx_model"]
