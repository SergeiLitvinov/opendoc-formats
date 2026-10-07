"""DOCX drawing exporter for raster and vector images."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from opendoc_model.diagnostics import ConversionReport, IssueSeverity
from opendoc_model.document_model import VECTOR_IMAGE_MEDIA_TYPES, DocumentModel, Image, Resource
from opendoc_model.units import degrees_to_ooxml_angle

from opendoc_formats.ooxml.package import points_to_emu


def write_image(
    paragraph: Any,
    image: Image,
    document: DocumentModel,
    report: ConversionReport,
    location: str,
) -> None:
    from docx.shared import Pt

    resource = document.resources.get(image.resource_id)
    if resource is None:
        report.add(IssueSeverity.ERROR, "image", f"resource {image.resource_id!r} not found", location)
        paragraph.add_run(image.alt_text)
        return
    if resource.media_type in VECTOR_IMAGE_MEDIA_TYPES:
        try:
            _add_vector_picture(paragraph, image, resource, document)
            return
        except Exception as error:  # noqa: BLE001 - malformed vector resources must produce a diagnostic
            report.add(
                IssueSeverity.LOSS,
                "image",
                f"{resource.media_type} could not be embedded; alt text used: {error}",
                location,
            )
            paragraph.add_run(image.alt_text or f"[{resource.filename or resource.id}]")
            return
    stream = _resource_stream(resource)
    width = Pt(image.box.width) if image.box and image.box.width > 0 else None
    height = Pt(image.box.height) if image.box and image.box.height > 0 else None
    try:
        shape = paragraph.add_run().add_picture(stream, width=width, height=height)
        if image.alt_text:
            properties = shape._inline.docPr
            properties.set("descr", image.alt_text)
        _apply_image_transform(shape._inline, image)
        _apply_image_effects(shape._inline, image)
        _apply_image_placement(shape._inline, image)
    except Exception as error:  # noqa: BLE001 - image backends raise several library-specific errors
        report.add(
            IssueSeverity.LOSS,
            "image",
            f"{resource.media_type} could not be embedded; alt text used: {error}",
            location,
        )
        paragraph.add_run(image.alt_text or f"[{resource.filename or resource.id}]")


def _add_vector_picture(paragraph: Any, image: Image, resource: Resource, document: DocumentModel) -> None:
    """Embed a Word-supported vector without decoding it through Pillow."""

    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.opc.part import Part
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    extension = {
        "image/svg+xml": "svg",
        "image/x-emf": "emf",
        "image/x-wmf": "wmf",
        "image/emf": "emf",
        "image/wmf": "wmf",
    }[resource.media_type]
    package = paragraph.part.package
    partname = package.next_partname(f"/word/media/vector%d.{extension}")
    vector_part = Part(partname, resource.media_type, _resource_bytes(resource), package)
    relationship_id = paragraph.part.relate_to(vector_part, RELATIONSHIP_TYPE.IMAGE)
    primary_relationship_id = relationship_id
    svg_extension = ""
    fallback_id = image.properties.fallback_resource_id
    fallback = document.resources.get(fallback_id) if fallback_id else None
    fallback_extensions = {
        "image/png": "png",
        "image/jpeg": "jpeg",
        "image/gif": "gif",
        "image/bmp": "bmp",
        "image/tiff": "tiff",
    }
    if resource.media_type == "image/svg+xml" and fallback is not None:
        fallback_extension = fallback_extensions.get(fallback.media_type)
        if fallback_extension:
            fallback_partname = package.next_partname(f"/word/media/fallback%d.{fallback_extension}")
            fallback_part = Part(fallback_partname, fallback.media_type, _resource_bytes(fallback), package)
            primary_relationship_id = paragraph.part.relate_to(fallback_part, RELATIONSHIP_TYPE.IMAGE)
            svg_extension = (
                '<a:extLst><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
                '<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
                f'r:embed="{relationship_id}"/></a:ext></a:extLst>'
            )
    width_pt = image.box.width if image.box and image.box.width > 0 else 288.0
    height_pt = image.box.height if image.box and image.box.height > 0 else 144.0
    width_emu = points_to_emu(width_pt)
    height_emu = points_to_emu(height_pt)
    drawing_id = paragraph.part.next_id
    xml = (
        f"<w:r {nsdecls('w', 'wp', 'a', 'pic', 'r')}><w:drawing>"
        '<wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="{width_emu}" cy="{height_emu}"/>'
        f'<wp:docPr id="{drawing_id}" name="Vector {drawing_id}"/>'
        '<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        "<pic:pic><pic:nvPicPr>"
        f'<pic:cNvPr id="0" name="vector.{extension}"/><pic:cNvPicPr/>'
        "</pic:nvPicPr><pic:blipFill>"
        f'<a:blip r:embed="{primary_relationship_id}">{svg_extension}</a:blip>'
        "<a:stretch><a:fillRect/></a:stretch>"
        '</pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/>'
        f'<a:ext cx="{width_emu}" cy="{height_emu}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        "</pic:spPr></pic:pic></a:graphicData></a:graphic>"
        "</wp:inline></w:drawing></w:r>"
    )
    run = parse_xml(xml)
    if image.alt_text:
        run.xpath(".//wp:docPr")[0].set("descr", image.alt_text)
    inline = run.xpath(".//wp:inline")[0]
    _apply_image_transform(inline, image)
    _apply_image_effects(inline, image)
    _apply_image_placement(inline, image)
    paragraph._p.append(run)


def _resource_bytes(resource: Resource) -> bytes:
    if resource.data is not None:
        return resource.data
    if resource.source is None:
        raise ValueError(f"resource {resource.id!r} has no content")
    return Path(resource.source).read_bytes()


def _apply_image_transform(drawing: Any, image: Image) -> None:
    from docx.oxml import OxmlElement

    transforms = drawing.xpath(".//pic:spPr/a:xfrm")
    if transforms:
        if image.box is not None and image.box.rotation:
            transforms[0].set("rot", str(degrees_to_ooxml_angle(image.box.rotation)))
        if image.properties.flip_horizontal is not None:
            transforms[0].set("flipH", str(int(image.properties.flip_horizontal)))
        if image.properties.flip_vertical is not None:
            transforms[0].set("flipV", str(int(image.properties.flip_vertical)))
    if image.crop is None:
        return
    fills = drawing.xpath(".//pic:blipFill")
    if not fills:
        return
    source = OxmlElement("a:srcRect")
    for attribute, value in (
        ("l", image.crop.left),
        ("t", image.crop.top),
        ("r", image.crop.right),
        ("b", image.crop.bottom),
    ):
        if value:
            source.set(attribute, str(round(value * 100000)))
    blips = fills[0].xpath("./a:blip")
    fills[0].insert(fills[0].index(blips[0]) + 1 if blips else 0, source)


def _apply_image_effects(drawing: Any, image: Image) -> None:
    from docx.oxml import OxmlElement, parse_xml
    from lxml import etree

    blips = drawing.xpath(".//pic:blipFill/a:blip")
    if blips:
        blip = blips[0]
        raw_effects = image.properties.blip_effects_xml
        if raw_effects:
            for raw_xml in raw_effects:
                effect = parse_xml(raw_xml)
                extension_lists = blip.xpath("./a:extLst")
                blip.insert(blip.index(extension_lists[0]) if extension_lists else len(blip), effect)
        else:
            if image.properties.opacity is not None:
                alpha = OxmlElement("a:alphaModFix")
                alpha.set("amt", str(round(max(0.0, min(1.0, image.properties.opacity)) * 100000)))
                blip.append(alpha)
            if image.properties.grayscale:
                blip.append(OxmlElement("a:grayscl"))

    shape_properties = drawing.xpath(".//pic:spPr")
    if not shape_properties:
        return
    shape = shape_properties[0]
    for raw_xml, names in (
        (image.properties.line_xml, {"ln"}),
        (image.properties.shape_effects_xml, {"effectLst", "effectDag"}),
    ):
        if not raw_xml:
            continue
        for child in list(shape):
            if etree.QName(child).localname in names:
                shape.remove(child)
        shape.append(parse_xml(raw_xml))


def _apply_image_placement(drawing: Any, image: Image) -> None:
    properties = image.properties
    if properties.placement != "anchor":
        return
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    drawing.tag = qn("wp:anchor")
    for name, default in (
        ("distT", "0"),
        ("distB", "0"),
        ("distL", "0"),
        ("distR", "0"),
        ("simplePos", "0"),
        ("relativeHeight", str(properties.relative_height or 0)),
        ("behindDoc", str(int(properties.behind_doc or False))),
        ("locked", "0"),
        ("layoutInCell", str(int(properties.layout_in_cell if properties.layout_in_cell is not None else True))),
        ("allowOverlap", str(int(properties.allow_overlap if properties.allow_overlap is not None else True))),
    ):
        drawing.set(name, default)

    simple_position = OxmlElement("wp:simplePos")
    simple_position.set("x", "0")
    simple_position.set("y", "0")
    drawing.insert(0, simple_position)
    for index, (axis, coordinate) in enumerate((("H", "x"), ("V", "y")), 1):
        position = OxmlElement(f"wp:position{axis}")
        relative_from = properties.horizontal_relative_from if axis == "H" else properties.vertical_relative_from
        position.set("relativeFrom", relative_from or "page")
        alignment = properties.horizontal_align if axis == "H" else properties.vertical_align
        if alignment:
            node = OxmlElement("wp:align")
            node.text = str(alignment)
        else:
            node = OxmlElement("wp:posOffset")
            value = getattr(image.box, coordinate) if image.box is not None else 0
            node.text = str(points_to_emu(value))
        position.append(node)
        drawing.insert(index, position)

    wrap_name = str(properties.wrap or "none").lower()
    wrap_tags = {
        "none": "wp:wrapNone",
        "square": "wp:wrapSquare",
        "tight": "wp:wrapTight",
        "through": "wp:wrapThrough",
        "topandbottom": "wp:wrapTopAndBottom",
    }
    wrap = OxmlElement(wrap_tags.get(wrap_name, "wp:wrapNone"))
    if wrap_name in {"square", "tight", "through"}:
        wrap.set("wrapText", properties.wrap_text or "bothSides")
    polygon = properties.wrap_polygon
    if wrap_name in {"tight", "through"} and polygon is not None:
        points = polygon.get("points")
        if isinstance(points, list) and points:
            polygon_element = OxmlElement("wp:wrapPolygon")
            polygon_element.set("edited", str(int(bool(polygon.get("edited", False)))))
            for index, point in enumerate(points):
                if not isinstance(point, dict):
                    continue
                node = OxmlElement("wp:start" if index == 0 else "wp:lineTo")
                node.set("x", str(int(point.get("x", 0))))
                node.set("y", str(int(point.get("y", 0))))
                polygon_element.append(node)
            if len(polygon_element):
                wrap.append(polygon_element)
    doc_properties = drawing.find(qn("wp:docPr"))
    drawing.insert(drawing.index(doc_properties) if doc_properties is not None else 4, wrap)


def _resource_stream(resource: Resource) -> BytesIO | str:
    if resource.data is not None:
        return BytesIO(resource.data)
    if resource.source is None:
        raise ValueError(f"resource {resource.id!r} has no content")
    return resource.source


__all__ = ["write_image"]
