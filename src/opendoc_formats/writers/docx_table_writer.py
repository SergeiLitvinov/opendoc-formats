"""DOCX table exporter for spans, geometry, fills, margins, and styles."""

from __future__ import annotations

from collections.abc import Callable
from typing import Any

from opendoc_model import ArtifactLimitError, DocumentLimits
from opendoc_model.diagnostics import ConversionReport, IssueSeverity
from opendoc_model.document_model import Block, DocumentModel, Table
from opendoc_model.properties import TableCellProperties, TableProperties, TableRowProperties

from opendoc_formats.ooxml.borders import export_borders
from opendoc_formats.ooxml.widths import export_width
from opendoc_formats.support.cancellation import checkpoint

BlockWriter = Callable[
    [Any, list[Block], DocumentModel, ConversionReport, str, dict[str, int]],
    None,
]


def write_table(
    container: Any,
    source: Table,
    document: DocumentModel,
    report: ConversionReport,
    location: str,
    counters: dict[str, int],
    write_blocks: BlockWriter,
    *, limits: DocumentLimits | None = None,
) -> None:
    from opendoc_formats.ooxml.table_grid import native_grid

    try:
        column_count, placements = native_grid(source, limits=limits)
    except ArtifactLimitError:
        raise
    except ValueError as error:
        report.add(IssueSeverity.ERROR, 'docx.table-grid', str(error), location)
        return
    row_count = max(len(source.rows), 1)
    try:
        target = container.add_table(rows=row_count, cols=column_count)
    except TypeError:
        from docx.shared import Inches

        target = container.add_table(rows=row_count, cols=column_count, width=Inches(6))
    _apply_table_properties(target, source.properties)
    from opendoc_formats.ooxml.table_position import export_table_position

    if source.properties.get('docx_table_position_unsupported'):
        report.add(IssueSeverity.ERROR, 'docx.table-position', 'Unknown table positioning requires native resolution', location)
    export_table_position(source, target, report, location)
    from opendoc_formats.ooxml.layout import export_layout

    export_layout(source.properties, target._tbl.tblPr, report, location, table=True)
    export_width(source, target._tbl.tblPr, "w:tblW")
    export_borders(source, target._tbl.tblPr, report, location)
    if source.style_id:
        _set_table_style_id(target, source.style_id)
    elif source.properties.style_name:
        try:
            target.style = source.properties.style_name
        except KeyError:
            report.add(
                IssueSeverity.WARNING,
                "table-style",
                f"style {source.properties.style_name!r} is unavailable",
                location,
            )
    for row_index, column_index, cell in placements:
        checkpoint()
        if cell.row_span > 1 or cell.column_span > 1:
            target.cell(row_index, column_index).merge(target.cell(row_index + cell.row_span - 1,
                                                                  column_index + cell.column_span - 1))
    position_iter = iter(placements)
    for row_index, source_row in enumerate(source.rows):
        checkpoint()
        _apply_row_properties(target.rows[row_index], source_row.properties, report, f'{location}.rows[{row_index}]')
        for cell_index, source_cell in enumerate(source_row.cells):
            checkpoint()
            _, column_index, _ = next(position_iter)
            target_cell = target.cell(row_index, column_index)
            if document.source_format == 'pdf' and source_cell.properties.get('pdf_text_origin') is not None:
                from opendoc_formats.writers.pdf_docx_cells import prepare_cell_baselines

                try:
                    prepare_cell_baselines(source_cell)
                except ValueError as error:
                    report.add(IssueSeverity.ERROR, 'docx.table-baseline',
                               str(error), f'{location}.rows[{row_index}].cells[{cell_index}]')
            _apply_cell_properties(target_cell, source_cell.properties)
            export_width(source_cell, target_cell._tc.get_or_add_tcPr(), "w:tcW")
            export_borders(source_cell, target_cell._tc.get_or_add_tcPr(), report,
                           f"{location}.rows[{row_index}].cells[{cell_index}]")
            _clear_container(target_cell)
            write_blocks(
                target_cell,
                source_cell.blocks,
                document,
                report,
                f"{location}.rows[{row_index}].cells[{cell_index}]",
                counters,
            )
            if not source_cell.blocks:
                target_cell.add_paragraph()
    from opendoc_formats.ooxml.merged_cell_borders import export_merge_borders

    export_merge_borders(target, placements, report, location)
    from opendoc_formats.ooxml.source_grid import export_source_grid

    export_source_grid(source, target, report, location)


def _clear_container(container: Any) -> None:
    from lxml import etree

    element = container._element
    for child in list(element):
        checkpoint()
        if etree.QName(child).localname in {"p", "tbl"}:
            element.remove(child)


def _set_table_style_id(table: Any, style_id: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = table._tbl.tblPr
    style = properties.find(qn("w:tblStyle"))
    if style is None:
        style = OxmlElement("w:tblStyle")
        properties.insert(0, style)
    style.set(qn("w:val"), style_id)


def _apply_table_properties(table: Any, properties: TableProperties) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn

    if properties.autofit is not None:
        table.autofit = properties.autofit
    if properties.alignment:
        alignment = getattr(WD_TABLE_ALIGNMENT, properties.alignment.upper(), None)
        if alignment is not None:
            table.alignment = alignment
    columns = table._tbl.xpath("./w:tblGrid/w:gridCol")
    for column, width in zip(columns, properties.grid_widths_twips, strict=False):
        checkpoint()
        column.set(qn("w:w"), str(width))


def _apply_row_properties(row: Any, properties: TableRowProperties, report: ConversionReport, location: str) -> None:
    from docx.oxml import OxmlElement

    if properties.repeat_header:
        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    height = properties.get('docx_height_twips')
    if height is not None:
        if type(height) is not int or not 0 < height < 2**31 or properties.get('docx_height_rule') not in {'exact', 'atLeast'}:
            report.add(IssueSeverity.ERROR, 'docx.row-height', 'Invalid finite row height or rule', location)
            return
        from docx.oxml.ns import qn

        node = OxmlElement('w:trHeight')
        node.set(qn('w:val'), str(height))
        node.set(qn('w:hRule'), properties['docx_height_rule'])
        row._tr.get_or_add_trPr().append(node)


def _apply_cell_properties(cell: Any, properties: TableCellProperties) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    cell_properties = cell._tc.get_or_add_tcPr()
    if properties.fill:
        shading = cell_properties.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            cell_properties.append(shading)
        shading.set(qn("w:fill"), properties.fill)
    if properties.margins_twips:
        margin_node = cell_properties.first_child_found_in("w:tcMar")
        if margin_node is None:
            margin_node = OxmlElement("w:tcMar")
            cell_properties.append(margin_node)
        for edge, value in properties.margins_twips.items():
            checkpoint()
            node = margin_node.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                margin_node.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")
    if properties.vertical_alignment:
        alignment = getattr(WD_CELL_VERTICAL_ALIGNMENT, properties.vertical_alignment.upper(), None)
        if alignment is not None:
            cell.vertical_alignment = alignment


__all__ = ["write_table"]
