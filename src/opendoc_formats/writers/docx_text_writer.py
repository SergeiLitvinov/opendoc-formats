"""Запись абзацев и строчного содержимого в DOCX."""

from __future__ import annotations

from typing import Any

from opendoc_model import LINE_SPACING_PROPERTY
from opendoc_model.diagnostics import ConversionReport, IssueSeverity
from opendoc_model.document_model import DocumentModel, Formula, FormulaFormat, Image, Paragraph, TextRun
from opendoc_model.limits import DocumentLimits

from opendoc_formats.ooxml.headings import export_heading
from opendoc_formats.ooxml.package import RELATIONSHIP_TYPE
from opendoc_formats.writers.docx_drawing_writer import write_image
from opendoc_formats.writers.docx_notes_writer import write_note_reference
from opendoc_formats.writers.docx_style_writer import apply_numbering, apply_paragraph_format, apply_text_style


def add_paragraph(
    container: Any,
    source: Paragraph,
    document: DocumentModel,
    report: ConversionReport,
    location: str,
    *, limits: DocumentLimits | None = None,
) -> Any:
    """Создать и оформить абзац модели в python-docx container."""

    source_style = document.styles.get(source.style_id) if source.style_id else None
    style = (source_style.properties.get('style_name') if source_style is not None else None)
    style = style or source.properties.get("style_name") or source.style_id
    if source.style_id is None and 'docx_paragraph_style_imported' in source.properties:
        style = None
    try:
        paragraph = container.add_paragraph(style=style) if style else container.add_paragraph()
    except KeyError:
        paragraph = container.add_paragraph()
        report.add(IssueSeverity.WARNING, "paragraph-style", f"style {style!r} is unavailable", location)
    if source.alignment:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        alignment = getattr(WD_ALIGN_PARAGRAPH, source.alignment.upper(), None)
        if alignment is not None:
            paragraph.alignment = alignment
    apply_paragraph_format(paragraph.paragraph_format, source.properties, report, location, limits=limits)
    if (document.source_format == 'docx' and 'docx_line_spacing_rule' not in source.properties
            and any(source.properties.get(key) is not None for key in ('line_spacing', 'line_spacing_pt'))):
        report.add(IssueSeverity.LOSS, 'docx.line-spacing-rule',
                   'Legacy DOCX model has no source line-spacing rule; units supply a default', location)
    from opendoc_model import get_tab_stops

    from opendoc_formats.ooxml.layout import export_layout, inherited_tab_positions

    tabs = get_tab_stops(source)
    try:
        inherited = inherited_tab_positions(source, document) if tabs is not None and not tabs.inherit else ()
    except (ValueError, TypeError, KeyError) as error:
        report.add(IssueSeverity.ERROR, 'docx.tab-stops', str(error), location)
        inherited = ()

    export_layout(source.properties, paragraph._p.get_or_add_pPr(), report, location, inherited=inherited)
    original_style_id = paragraph.style.style_id
    export_heading(source, paragraph)
    if (source.style_id in document.styles and document.package is not None
            and document.package.related_part(document.package.root, RELATIONSHIP_TYPE['styles']) is not None):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        if paragraph.style.style_id == original_style_id:
            owner = paragraph._p.get_or_add_pPr()
            nodes = owner.findall(qn('w:pStyle'))
            node = nodes[0] if nodes else OxmlElement('w:pStyle')
            node.set(qn('w:val'), source.style_id)
            if not nodes:
                owner.insert(0, node)
        else:
            # A newly derived role style must reference the retained source ID.
            paragraph.style.element.basedOn_val = source.style_id
    preserves_style_numbering = (
        source.properties.get("numbering_source") == "style"
        and source.properties.get("numbering_source_style_id") == source.style_id
        and document.package is not None
        and document.package.related_part(document.package.root, RELATIONSHIP_TYPE["styles"]) is not None
    )
    if not preserves_style_numbering:
        apply_numbering(paragraph._p.get_or_add_pPr(), source.properties)
    if "html_list_id" in source.properties:
        from opendoc_formats.writers.docx_html_lists import apply_html_list

        apply_html_list(paragraph, source.properties)
    if document.source_format == "html" and source.properties.get("anchor_id"):
        from opendoc_formats.writers.docx_html_links import add_bookmark

        add_bookmark(paragraph, source.properties["anchor_id"])
    return paragraph


def write_paragraph_content(
    paragraph: Any,
    source: Paragraph,
    document: DocumentModel,
    report: ConversionReport,
    location: str,
    counters: dict[str, int],
) -> None:
    """Записать упорядоченное содержимое абзаца."""

    for index, item in enumerate(source.content):
        item_location = f"{location}.content[{index}]"
        if isinstance(item, TextRun):
            from opendoc_formats.ooxml.paragraph_flow import FIELDS

            if any(item.style.properties.get(key) is not None for key in FIELDS):
                report.add(IssueSeverity.ERROR, 'docx.paragraph-flow',
                           'An inline run style does not declare paragraph flow in Word', item_location + '.style')
                continue
            from opendoc_formats.ooxml.paragraph_indent import ATTRIBUTES

            if any(item.style.properties.get(key) is not None for key in ATTRIBUTES):
                report.add(IssueSeverity.ERROR, 'docx.paragraph-indent',
                           'An inline run style does not declare paragraph indentation in Word', item_location + '.style')
                continue
            if LINE_SPACING_PROPERTY in item.style.properties:
                report.add(IssueSeverity.ERROR, 'docx.line-spacing',
                           'An inline run style does not declare paragraph line spacing in Word', item_location + '.style')
                continue
            if document.source_format == "html" and item.properties.get("anchor_id"):
                from opendoc_formats.writers.docx_html_links import add_bookmark

                add_bookmark(paragraph, item.properties["anchor_id"])
            if item.properties.get("docx_raw_inline_xml"):
                _write_raw_inline(paragraph, str(item.properties["docx_raw_inline_xml"]), report, item_location)
            elif item.properties.get("bookmark_start") is not None:
                _write_bookmark_start(paragraph, item.properties["bookmark_start"])
            elif item.properties.get("bookmark_end_id") is not None:
                _write_bookmark_end(paragraph, item.properties["bookmark_end_id"])
            elif write_note_reference(paragraph, item):
                continue
            elif item.properties.get("field_xml"):
                _write_raw_field(paragraph, item.properties["field_xml"])
            elif item.properties.get("field_instruction") is not None:
                _write_simple_field(paragraph, item)
            else:
                run = paragraph.add_run(item.text)
                character_style = (item.style.properties.base_style_id if 'base_style_id' in item.style.properties
                                   else item.style.properties.get('docx_character_style_id'))
                if character_style:
                    from docx.oxml import OxmlElement
                    from docx.oxml.ns import qn

                    if character_style not in document.styles:
                        report.add(IssueSeverity.ERROR, 'docx.character-style',
                                   'Imported character style is missing', item_location)
                    elif document.package is not None:
                        node = OxmlElement('w:rStyle')
                        node.set(qn('w:val'), character_style)
                        run._r.get_or_add_rPr().insert(0, node)
                    else:
                        run.style = document.styles[character_style].properties.get('style_name') or character_style
                apply_text_style(run, item.style)
                if item.link and item.link.startswith("#") and document.source_format == "html":
                    from urllib.parse import unquote

                    from opendoc_formats.writers.docx_html_links import bookmark_name

                    _wrap_internal_hyperlink(paragraph, run, bookmark_name(unquote(item.link[1:])))
                elif item.link:
                    _wrap_hyperlink(paragraph, run, item.link)
                elif item.properties.get("hyperlink_anchor"):
                    _wrap_internal_hyperlink(paragraph, run, str(item.properties["hyperlink_anchor"]))
        elif isinstance(item, Formula):
            write_formula(paragraph, item, report, item_location)
            counters["formulas"] += 1
        elif isinstance(item, Image):
            write_image(paragraph, item, document, report, item_location)
            counters["images"] += 1


def write_formula(paragraph: Any, formula: Formula, report: ConversionReport, location: str) -> None:
    """Записать OMML или диагностированный текстовый fallback."""

    if formula.format is FormulaFormat.MATHML:
        from opendoc_formats.writers.mathml_to_omml import mathml_to_omml

        try:
            from docx.oxml import parse_xml

            paragraph._p.append(parse_xml(mathml_to_omml(formula.value)))
            return
        except ValueError as error:
            report.add(IssueSeverity.LOSS, "formula", f"MathML fallback: {error}", location)
            paragraph.add_run(formula.fallback_text or formula.value)
            return
    if formula.format is FormulaFormat.OMML:
        from docx.oxml import parse_xml

        try:
            paragraph._p.append(parse_xml(formula.value))
            return
        except Exception as error:  # noqa: BLE001 - malformed foreign XML is reported as a loss
            report.add(IssueSeverity.LOSS, "formula", f"invalid OMML replaced by fallback: {error}", location)
    else:
        report.add(
            IssueSeverity.LOSS,
            "formula",
            f"{formula.format.value} is not natively convertible to OMML; fallback text used",
            location,
        )
    paragraph.add_run(formula.fallback_text or formula.value)


def _wrap_hyperlink(paragraph: Any, run: Any, url: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as DOCX_RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relationship_id = paragraph.part.relate_to(url, DOCX_RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run._r.getparent().remove(run._r)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def _wrap_internal_hyperlink(paragraph: Any, run: Any, anchor: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run._r.getparent().remove(run._r)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def _write_bookmark_start(paragraph: Any, bookmark: dict[str, Any]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark.get("id") or "0"))
    start.set(qn("w:name"), str(bookmark.get("name") or "_OpenDocFormats"))
    paragraph._p.append(start)


def _write_bookmark_end(paragraph: Any, bookmark_id: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.append(end)


def _write_simple_field(paragraph: Any, item: TextRun) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), str(item.properties["field_instruction"]))
    if item.text:
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = item.text
        run.append(text)
        field.append(run)
    paragraph._p.append(field)


def _write_raw_field(paragraph: Any, raw_elements: list[str]) -> None:
    from docx.oxml import parse_xml

    for raw_xml in raw_elements:
        paragraph._p.append(parse_xml(raw_xml))


def _write_raw_inline(paragraph: Any, raw_xml: str, report: ConversionReport, location: str) -> None:
    from docx.oxml import parse_xml

    try:
        element = parse_xml(raw_xml)
        paragraph._p.append(element)
        roots = getattr(paragraph.part, "_opendoc_opaque_roots", [])
        roots.append(element)
        paragraph.part._opendoc_opaque_roots = roots
    except Exception as error:  # noqa: BLE001 - foreign OOXML is a diagnostic boundary
        report.add(IssueSeverity.LOSS, "docx-inline-ooxml", f"raw inline OOXML could not be restored: {error}", location)


__all__ = ["add_paragraph", "write_formula", "write_paragraph_content"]
