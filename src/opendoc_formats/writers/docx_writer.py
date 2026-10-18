"""Экспорт богатой промежуточной модели в редактируемый DOCX."""

from __future__ import annotations

from functools import partial
from pathlib import Path
from typing import Any

from opendoc_model import ArtifactLimitError
from opendoc_model.diagnostics import ConversionReport, IssueSeverity
from opendoc_model.document_model import (
    Block,
    DocumentModel,
    Formula,
    Image,
    Paragraph,
    Section,
    Table,
)
from opendoc_model.limits import DocumentLimits

from opendoc_formats.fonts import FontResolver
from opendoc_formats.fonts.docx_embedding import embed_docx_fonts, verify_docx_font_embedding
from opendoc_formats.ooxml.headings import retain_derived_styles
from opendoc_formats.ooxml.package import restore_package_graph
from opendoc_formats.writers.color_preflight import preflight_colors
from opendoc_formats.writers.docx_drawing_writer import write_image
from opendoc_formats.writers.docx_section_writer import configure_section, section_start_type
from opendoc_formats.writers.docx_style_writer import write_styles
from opendoc_formats.writers.docx_table_writer import write_table
from opendoc_formats.writers.docx_text_writer import add_paragraph, write_formula, write_paragraph_content
from opendoc_formats.writers.font_preflight import prepare_fonts


def write_docx_model(
    document: DocumentModel, output_path: str | Path, *, limits: DocumentLimits | None = None,
) -> ConversionReport:
    """Записать ``DocumentModel`` в DOCX и вернуть отчёт о потерях."""

    output = Path(output_path)
    report = ConversionReport(output)
    errors = document.validate(limits=limits)
    if errors:
        from opendoc_formats.writers.html_resources import document_images

        for image, location in document_images(document):
            if image.resource_id not in document.resources:
                report.add(IssueSeverity.ERROR, 'image', f'Resource {image.resource_id!r} is missing', location)
    for error in errors:
        report.add(IssueSeverity.ERROR, 'model', error)
    if not report.success:
        return report
    native_pdf = (document.source_format == 'pdf' and any(
        section.properties.get('pdf', {}).get('coordinate_space') == 'pymupdf-unrotated' for section in document.sections))
    if native_pdf:
        report.metrics.update({'pdf_docx_layout_profile': 'opendoc-v1', 'assessment_complete': False})
        _check_pdf_projection_source(document, report, limits=limits)
        if not report.success:
            return report
        if not document.metadata.get('pdf_docx_layout'):
            from opendoc_formats.writers.pdf_docx_model import prepare_pdf_docx

            try:
                document = prepare_pdf_docx(document, limits=limits)
            except ArtifactLimitError:
                raise
            except ValueError as error:
                report.add(IssueSeverity.ERROR, 'pdf-docx.layout', str(error), 'document')
                return report
        from opendoc_model import TextRun, walk_model

        projections = document.metadata.get('pdf_docx_group_projection', [])
        if projections:
            report.metrics['pdf_docx_group_projection'] = projections
            report.add(IssueSeverity.LOSS, 'pdf-docx.group-projection',
                       'RGB group is represented by a vector SVG; native child/group editing is not reconstructed. '
                       'Typed source scopes remain in the input Model.', 'document')
        seen_fonts = set()
        for reference in walk_model(document, limits=limits):
            run = reference.node
            name = run.style.properties.get('pdf_source_font') if isinstance(run, TextRun) else None
            if name is not None and name not in seen_fonts:
                seen_fonts.add(name)
                report.add(IssueSeverity.WARNING, 'font-substitution',
                           f'Unembedded PDF font {name!r} maps to Word {run.style.font_family!r}; visual comparison is required',
                           reference.path)
    from opendoc_formats.fonts.resource_fonts import document_font_resolver

    with document_font_resolver(document, report, limits=limits) as resolver:
        return _write_validated_docx(document, output, report, resolver, limits=limits)


def _write_validated_docx(
    document: DocumentModel, output: Path, report: ConversionReport, resolver: FontResolver,
    *, limits: DocumentLimits | None = None,
) -> ConversionReport:
    from docx import Document
    from docx.enum.section import WD_SECTION

    document = prepare_fonts(document, report, limits=limits, resolver=resolver)
    _reject_pdf_losses(report)
    if not report.success:
        return report
    from opendoc_formats.writers.pdf_docx_cells import fit_virtual_tables

    fit_virtual_tables(document, limits=limits)
    preflight_colors(document, report, target="docx")
    target = Document()
    _write_metadata(target, document.metadata)
    _restore_document_protection(target, document.metadata, report)
    write_styles(target, document.styles, report, limits=limits)
    target.settings.odd_and_even_pages_header_footer = any(
        bool(section.properties.odd_and_even_pages_header_footer)
        or bool(section.even_page_headers)
        or bool(section.even_page_footers)
        for section in document.sections
    )

    sections = document.sections or [Section()]
    counters = {"paragraphs": 0, "tables": 0, "images": 0, "formulas": 0}
    for section_index, source_section in enumerate(sections):
        if section_index == 0:
            target_section = target.sections[0]
        else:
            target_section = target.add_section(section_start_type(source_section, WD_SECTION))
        configure_section(target_section, source_section, document, report, section_index, counters,
                          partial(_write_blocks, limits=limits))
        _write_blocks(target, source_section.blocks, document, report, f"sections[{section_index}]", counters, limits=limits)

    embed_docx_fonts(target, document, report)
    _restore_package_graph(target, document, report, limits=limits)
    _reject_pdf_losses(report)
    if not report.success:
        return report
    try:
        from opendoc_formats.support.artifacts import ArtifactWorkspace
        from opendoc_formats.support.io import atomic_copy

        with ArtifactWorkspace(prefix="opendoc_formats_docx_writer_") as workspace:
            staged = workspace.artifact_path("output.docx")
            target.save(staged)
            workspace.validate_artifact(staged)
            verify_docx_font_embedding(staged, report)
            _reject_pdf_losses(report)
            if not report.success:
                return report
            atomic_copy(staged, output)
    except Exception as error:  # noqa: BLE001 - diagnostic boundary must return a report
        report.add(IssueSeverity.ERROR, "docx-write", str(error))
        return report
    report.metrics.update(counters)
    report.metrics["sections"] = len(sections)
    report.metrics["resources"] = len(document.resources)
    return report


def _reject_pdf_losses(report: ConversionReport) -> None:
    if report.metrics.get('pdf_docx_layout_profile') != 'opendoc-v1':
        return
    losses = [issue for issue in report.issues if issue.severity is IssueSeverity.LOSS]
    if losses and report.success:
        report.add(IssueSeverity.ERROR, 'pdf-docx.loss-rejected',
                   'The measured PDF → Word profile rejects known preservation losses', losses[0].location)


def _check_pdf_projection_source(
    document: DocumentModel, report: ConversionReport, *, limits: DocumentLimits | None,
) -> None:
    from opendoc_model import PreservationState, get_integration

    integration = get_integration(document, limits=limits)
    if integration is None:
        return
    for name in ('annotations', 'forms'):
        if getattr(integration, name):
            report.add(IssueSeverity.ERROR, 'pdf-docx.interactions',
                       'PDF interactions are outside the measured native Word profile', f'integration.{name}')
    if integration.extra.get('pdf_outline'):
        report.add(IssueSeverity.ERROR, 'pdf-docx.outline',
                   'PDF outline projection is outside the measured native Word profile', 'integration.extra.pdf_outline')
    for record in integration.preservation:
        if record.state in {PreservationState.OPAQUE, PreservationState.LOST, PreservationState.REJECTED}:
            report.add(IssueSeverity.ERROR, 'pdf-docx.source-unknown', record.issue.message, record.issue.location)


def _write_metadata(target: Any, metadata: dict[str, Any]) -> None:
    properties = target.core_properties
    for name in ("title", "subject", "author", "keywords", "comments"):
        value = metadata.get(name)
        if value is not None:
            setattr(properties, name, str(value))


def _restore_package_graph(
    target: Any, document: DocumentModel, report: ConversionReport, *, limits: DocumentLimits | None = None,
) -> None:
    graph = document.package
    source_parts = {
        part._opendoc_source_part: part for part in target.part.package.iter_parts()
        if getattr(part, "_opendoc_source_part", None)
    }
    relationship_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for part in target.part.package.iter_parts():
        has_references = any(
            key.startswith(relationship_namespace)
            for root in getattr(part, "_opendoc_opaque_roots", ()) for node in root.iter() for key in node.attrib
        )
        if part is not target.part and has_references and not getattr(part, "_opendoc_source_part", None):
            report.add(IssueSeverity.ERROR, "package-graph", "Opaque running content has no source part binding")
            return
    if graph is None or graph.format != "ooxml":
        if any(
            key.startswith(relationship_namespace)
            for part in target.part.package.iter_parts()
            for root in getattr(part, "_opendoc_opaque_roots", ())
            for node in root.iter() for key in node.attrib
        ):
            report.add(IssueSeverity.ERROR, "package-graph", "Opaque OOXML relationships require their source package graph")
        return
    try:
        from opendoc_formats.ooxml.layout import retain_layout_styles

        restore_package_graph(
            target.part, retain_layout_styles(retain_derived_styles(target.part, graph), document.styles, report, limits=limits),
            protected_roots=getattr(target.part, "_opendoc_opaque_roots", ()),
            source_parts=source_parts,
        )
    except ArtifactLimitError:
        raise
    except Exception as error:  # noqa: BLE001 - foreign package parts are a diagnostic boundary
        report.add(IssueSeverity.ERROR, "package-graph", f"OOXML package graph could not be restored: {error}")


def _restore_document_protection(target: Any, metadata: dict[str, Any], report: ConversionReport) -> None:
    raw_xml = metadata.get("docx_features", {}).get("document_protection_xml")
    if not raw_xml:
        return
    from docx.oxml import parse_xml

    try:
        protection = parse_xml(str(raw_xml))
        settings = target.settings._element
        for existing in settings.xpath("./w:documentProtection"):
            settings.remove(existing)
        settings.append(protection)
    except Exception as error:  # noqa: BLE001 - foreign OOXML is a diagnostic boundary
        report.add(
            IssueSeverity.LOSS,
            "document-protection",
            f"document protection could not be restored: {error}",
        )


def _write_blocks(
    container: Any,
    blocks: list[Block],
    document: DocumentModel,
    report: ConversionReport,
    location: str,
    counters: dict[str, int],
    *, limits: DocumentLimits | None = None,
) -> None:
    for index, block in enumerate(blocks):
        block_location = f"{location}.blocks[{index}]"
        if isinstance(block, Paragraph):
            if block.properties.get("docx_raw_block_xml"):
                _write_raw_block(container, str(block.properties["docx_raw_block_xml"]), report, block_location)
                counters["paragraphs"] += 1
                continue
            paragraph = add_paragraph(container, block, document, report, block_location, limits=limits)
            write_paragraph_content(paragraph, block, document, report, block_location, counters)
            counters["paragraphs"] += 1
        elif isinstance(block, Table):
            write_table(container, block, document, report, block_location, counters,
                        partial(_write_blocks, limits=limits), limits=limits)
            counters["tables"] += 1
        elif isinstance(block, Formula):
            paragraph = container.add_paragraph()
            write_formula(paragraph, block, report, block_location)
            counters["formulas"] += 1
        elif isinstance(block, Image):
            paragraph = container.add_paragraph()
            write_image(paragraph, block, document, report, block_location)
            counters["images"] += 1


def _write_raw_block(container: Any, raw_xml: str, report: ConversionReport, location: str) -> None:
    from docx.oxml import parse_xml

    try:
        element = parse_xml(raw_xml)
        parent = container._body if hasattr(container, "_body") else container
        root = parent._element
        section_properties = root.sectPr if hasattr(root, "sectPr") else None
        if section_properties is None:
            root.append(element)
        else:
            section_properties.addprevious(element)
        roots = getattr(container.part, "_opendoc_opaque_roots", [])
        roots.append(element)
        container.part._opendoc_opaque_roots = roots
    except Exception as error:  # noqa: BLE001 - foreign OOXML is a diagnostic boundary
        report.add(IssueSeverity.LOSS, "docx-block-ooxml", f"raw block OOXML could not be restored: {error}", location)


__all__ = ["write_docx_model"]
