from __future__ import annotations

import os
import subprocess
from pathlib import Path

from opendoc_formats.errors import ConvertError
from opendoc_formats.writers.base import BaseConverter, ConversionResult


class Pdf2DocxConverter(BaseConverter):
    @property
    def name(self) -> str:
        return "pdf2docx"

    def convert(self, input_path: str | Path, output_path: str | Path) -> ConversionResult:
        prepared = self._prepare(input_path, output_path)
        if prepared is None:
            return ConversionResult(Path(input_path), Path(output_path), False, "Input file not found")
        input_path, output_path = prepared
        try:
            import pdf2docx

            from opendoc_formats.support.artifacts import ArtifactWorkspace
            from opendoc_formats.support.io import atomic_copy

            with ArtifactWorkspace(prefix="opendoc_formats_pdf2docx_") as workspace:
                staged = workspace.artifact_path("output.docx")
                pdf2docx.parse(str(input_path), str(staged), multi_processing=True)
                workspace.validate_artifact(staged)
                atomic_copy(staged, output_path)
        except Exception as e:
            return ConversionResult(input_path, output_path, False, str(e))
        return ConversionResult(input_path, output_path, output_path.exists())


class PyMuPdfConverter(BaseConverter):
    """PDF → DOCX через PyMuPDF.

    Стратегия (best-effort для «сложных» PDF):
    * Извлечь текст каждой страницы и положить в DOCX.
    * Если на странице мало текста (< 50 символов) — отрендерить её в PNG
      и вставить как изображение (для сканов и формул).
    """

    @property
    def name(self) -> str:
        return "pymupdf"

    def __init__(self, render_dpi: int = 150, text_threshold: int = 50) -> None:
        self.render_dpi = render_dpi
        self.text_threshold = text_threshold

    def convert(self, input_path: str | Path, output_path: str | Path) -> ConversionResult:
        prepared = self._prepare(input_path, output_path)
        if prepared is None:
            return ConversionResult(Path(input_path), Path(output_path), False, "Input file not found")
        input_path, output_path = prepared
        from opendoc_formats.support.artifacts import ArtifactWorkspace
        from opendoc_formats.support.io import atomic_copy

        try:
            import fitz
            from docx import Document as DocxDocument
            from docx.shared import Inches, Pt

            with ArtifactWorkspace(prefix="opendoc_formats_pymupdf_") as workspace:
                d = DocxDocument()
                with fitz.open(str(input_path)) as doc:
                    for page_num in range(len(doc)):
                        page = doc[page_num]
                        text_dict = page.get_text("dict") or {}
                        blocks = text_dict.get("blocks", []) if isinstance(text_dict, dict) else []
                        visible_text = "".join(
                            span.get("text", "")
                            for block in blocks
                            if block.get("type") == 0
                            for line in block.get("lines", [])
                            for span in line.get("spans", [])
                        )
                        has_text = len(visible_text.strip()) >= self.text_threshold
                        if has_text:
                            # Текстовая страница — текстом, с сохранением размера шрифта.
                            for block in blocks:
                                if block.get("type") != 0:
                                    continue
                                for line in block.get("lines", []):
                                    spans = line.get("spans", [])
                                    if not "".join(span.get("text", "") for span in spans).strip():
                                        continue
                                    paragraph = d.add_paragraph()
                                    for span in spans:
                                        run = paragraph.add_run(span.get("text", ""))
                                        size = span.get("size")
                                        if size:
                                            try:
                                                run.font.size = Pt(float(size))
                                            except Exception:  # noqa: BLE001
                                                pass
                        else:
                            # Реально пустая страница (скан, формулы) — картинкой.
                            pix = page.get_pixmap(dpi=self.render_dpi)
                            image_path = workspace.artifact_path(f"page_{page_num}.png")
                            pix.save(str(image_path))
                            workspace.validate_artifact(image_path)
                            width_in = 6.0
                            ratio = pix.height / max(pix.width, 1)
                            d.add_picture(
                                str(image_path),
                                width=Inches(width_in),
                                height=Inches(width_in * ratio),
                            )
                            image_path.unlink(missing_ok=True)
                            workspace.validate_artifact(workspace.path)
                        d.add_page_break()
                staged = workspace.artifact_path("output.docx")
                d.save(str(staged))
                workspace.validate_artifact(staged)
                atomic_copy(staged, output_path)
        except Exception as e:
            return ConversionResult(input_path, output_path, False, str(e))
        return ConversionResult(input_path, output_path, output_path.exists())


class LibreOfficeConverter(BaseConverter):
    """Legacy facade retained for CLI compatibility.

    LibreOffice imports PDF as a Draw document and does not expose a reliable
    Draw/PDF → Writer DOCX export filter. Advertising it as a working backend
    made runtime availability checks pass while real documents failed later.
    """

    @property
    def name(self) -> str:
        return "libreoffice"

    def __init__(self, libreoffice_path: str | None = None) -> None:
        self.libreoffice_path = libreoffice_path or self._find_libreoffice_path()

    def _find_libreoffice_path(self) -> str | None:
        possible = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
        for path in possible:
            if os.path.exists(path):
                return path
        try:
            result = subprocess.run(["soffice", "--version"], capture_output=True, text=True)
            if result.returncode == 0:
                return "soffice"
        except FileNotFoundError:
            pass
        return None

    def convert(self, input_path: str | Path, output_path: str | Path) -> ConversionResult:
        prepared = self._prepare(input_path, output_path)
        if prepared is None:
            return ConversionResult(Path(input_path), Path(output_path), False, "Input file not found")
        input_path, output_path = prepared
        if not self.libreoffice_path:
            return ConversionResult(input_path, output_path, False, "LibreOffice not found")
        return ConversionResult(
            input_path,
            output_path,
            False,
            "LibreOffice does not provide a reliable PDF to Writer DOCX export filter; use pdf2docx or pymupdf",
        )


class FanOutConverter(BaseConverter):
    """Пробует движки по очереди, пока один не выдаст валидный DOCX.

    Цепочка по умолчанию: ``pdf2docx → pymupdf``. Первый успех побеждает.
    LibreOffice не включён: PDF импортируется как Draw и не имеет надёжного
    фильтра экспорта в Writer DOCX.
    Используется, когда конкретный движок неизвестен или нужен best-effort.
    """

    @property
    def name(self) -> str:
        return "fanout"

    def __init__(self, tools: list[str] | None = None) -> None:
        # pdf2docx сохраняет текст и разметку, поэтому он приоритетный.
        # PyMuPDF остаётся надёжным fallback со снимками сканированных страниц.
        self.tools = tools or ["pdf2docx", "pymupdf"]

    def convert(self, input_path: str | Path, output_path: str | Path) -> ConversionResult:
        prepared = self._prepare(input_path, output_path)
        if prepared is None:
            return ConversionResult(Path(input_path), Path(output_path), False, "Input file not found")
        input_path, output_path = prepared
        last_error = ""
        tried: list[str] = []
        for tool in self.tools:
            try:
                converter = create_converter(tool)
            except ConvertError as e:
                tried.append(f"{tool}:unavailable")
                last_error = str(e)
                continue
            tried.append(tool)
            result = converter.convert(input_path, output_path)
            if result.success and result.output_path.exists() and result.output_path.stat().st_size > 0:
                self._postprocess(output_path)
                return ConversionResult(
                    input_path,
                    result.output_path,
                    True,
                    None,
                )
            last_error = f"{tool}: {result.error}"
        return ConversionResult(
            input_path,
            Path(output_path),
            False,
            f"all engines failed (tried: {tried}); last error: {last_error}",
        )

    @staticmethod
    def _postprocess(output_path: str | Path) -> None:
        """Восстановить структуру (заголовки) и нормализовать шрифт."""
        try:
            from opendoc_formats.writers.docx_postprocess import (
                apply_heading_styles,
                ensure_min_font,
            )

            ensure_min_font(output_path)
            apply_heading_styles(output_path)
        except Exception:  # noqa: BLE001
            # Пост-обработка — best-effort, не ломаем результат конвертации.
            pass


def create_converter(tool: str) -> BaseConverter:
    converters = {
        "pdf2docx": Pdf2DocxConverter,
        "pymupdf": PyMuPdfConverter,
        "libreoffice": LibreOfficeConverter,
        "fanout": FanOutConverter,
    }
    if tool not in converters:
        raise ConvertError(f"Unknown tool: {tool}. Available: {list(converters.keys())}")
    return converters[tool]()
