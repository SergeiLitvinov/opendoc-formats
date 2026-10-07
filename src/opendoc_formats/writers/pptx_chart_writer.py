"""Native chart data export, including stacking and independent XY series."""

from __future__ import annotations

import math
from typing import Any

import opendoc_model as od
from opendoc_model.diagnostics import IssueSeverity

from opendoc_formats.writers.pptx_chart_settings import configure_chart
from opendoc_formats.writers.pptx_combo_writer import compose_combo, install_combo, is_combo
from opendoc_formats.writers.pptx_label_writer import configure_labels
from opendoc_formats.writers.pptx_plot_writer import configure_plots
from opendoc_formats.writers.pptx_point_writer import configure_points
from opendoc_formats.writers.pptx_statistics_writer import configure_statistics
from opendoc_formats.writers.pptx_text_writer import set_color


def _number(value: str | int | float | None, *, missing: bool = False) -> float | None:
    if missing and value in (None, ""):
        return None
    number = float(value)
    if not math.isfinite(number):
        raise ValueError("non-finite chart value")
    return number


def _chart_kind(data: dict[str, Any]) -> Any:
    from pptx.enum.chart import XL_CHART_TYPE as K

    kind = data.get("chart_type")
    grouping = data.get("grouping", "standard")
    stacked = {"standard": 0, "clustered": 0, "stacked": 1, "percentStacked": 2}[grouping]
    category = {
        "barChart": (K.COLUMN_CLUSTERED, K.COLUMN_STACKED, K.COLUMN_STACKED_100),
        "lineChart": (K.LINE, K.LINE_STACKED, K.LINE_STACKED_100),
        "areaChart": (K.AREA, K.AREA_STACKED, K.AREA_STACKED_100),
        "pieChart": (K.PIE,) * 3,
        "doughnutChart": (K.DOUGHNUT,) * 3,
    }
    if kind == "barChart" and data.get("bar_direction") == "bar":
        return (K.BAR_CLUSTERED, K.BAR_STACKED, K.BAR_STACKED_100)[stacked]
    if kind == "scatterChart":
        return {
            "marker": K.XY_SCATTER,
            "line": K.XY_SCATTER_LINES_NO_MARKERS,
            "lineMarker": K.XY_SCATTER_LINES,
            "smooth": K.XY_SCATTER_SMOOTH_NO_MARKERS,
            "smoothMarker": K.XY_SCATTER_SMOOTH,
        }[data.get("scatter_style", "marker")]
    if kind == "bubbleChart":
        return K.BUBBLE
    return category[kind][stacked]


def _chart_data(data: dict[str, Any]) -> Any:
    from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData

    kind = data.get("chart_type")
    numeric = kind in ("scatterChart", "bubbleChart")
    result = (BubbleChartData() if kind == "bubbleChart" else XyChartData()) if numeric else CategoryChartData()
    if not data.get("series"):
        raise ValueError("empty chart")
    if not numeric:
        if not data.get("categories"):
            raise ValueError("missing categories")
        result.categories = data["categories"]
    for series in data["series"]:
        if "categories" in series and series["categories"] != data.get("categories"):
            raise ValueError("incompatible series categories")
        if not numeric:
            values = [_number(value, missing=True) for value in series["values"]]
            if len(values) != len(data["categories"]):
                raise ValueError("category/value mismatch")
            result.add_series(series.get("name", ""), values)
            continue
        dimensions = [series["x_values"], series["values"]]
        if kind == "bubbleChart":
            dimensions.append(series["bubble_sizes"])
        if not dimensions[0] or len({len(values) for values in dimensions}) != 1:
            raise ValueError("XY dimension mismatch")
        target = result.add_series(series.get("name", ""))
        for point in zip(*dimensions):
            values = tuple(_number(value) for value in point)
            if kind == "bubbleChart" and values[2] < 0:
                raise ValueError("negative bubble size")
            target.add_data_point(*values)
    return result


def write_chart(
    slide: Any, data: dict[str, Any], geometry: tuple[int, int, int, int], report: od.ConversionReport, location: str
) -> bool:
    from pptx.enum.chart import XL_MARKER_STYLE

    try:
        kind, chart_data = _chart_kind(data), _chart_data(data)
        scene = compose_combo(data, chart_data, _chart_kind) if is_combo(data) else None
    except (TypeError, ValueError, KeyError):
        report.add(IssueSeverity.LOSS, "charts", "Тип или данные диаграммы не поддержаны; сохранена текстовая сводка.", location)
        return False
    chart = slide.shapes.add_chart(kind, *geometry, chart_data).chart
    if scene is not None:
        install_combo(chart, scene)
    configure_chart(chart, data, report, location)
    configure_plots(chart=chart, data=data, report=report, location=location)
    configure_labels(chart=chart, data=data, report=report, location=location)
    if data.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = data["title"]
    for source, target in zip(data["series"], chart.series):
        configure_statistics(source=source, target=target, options=data, report=report, location=location)
        configure_points(
            source=source, target=target, kind=source.get("chart_type", data.get("chart_type")), report=report, location=location
        )
        if source.get("marker_symbol") is not None and hasattr(target, "marker"):
            try:
                target.marker.style = XL_MARKER_STYLE.from_xml(source["marker_symbol"])
            except ValueError:
                report.add(IssueSeverity.LOSS, "styles", "Неизвестный маркер серии заменён стандартным.", location)
        if source.get("color"):
            target.format.fill.solid()
            set_color(target.format.fill.fore_color, source.get("color_value", source["color"]), report, location)
            if source.get("chart_type", data.get("chart_type")) in ("lineChart", "scatterChart"):
                set_color(target.format.line.color, source.get("color_value", source["color"]), report, location)
    if data.get("chart_3d"):
        report.add(IssueSeverity.LOSS, "charts", "Объёмная диаграмма преобразована в плоскую.", location)
    report.add(
        IssueSeverity.LOSS,
        "styles",
        "Диаграмма редактируема; оформление упрощено: цвета точек из темы, шрифты заголовков, "
        "сетка и размещение осей могут отличаться.",
        location,
    )
    return True
