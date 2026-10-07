"""Restore editable DrawingML shapes instead of substituting text boxes."""

from __future__ import annotations

from typing import Any

import opendoc_model as od
from opendoc_model.diagnostics import IssueSeverity

from opendoc_formats.writers.pptx_text_writer import set_color

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def write_shape(
    slide: Any, meta: dict[str, Any], geometry: tuple[int, int, int, int], report: od.ConversionReport, location: str
) -> Any:
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR_TYPE
    from pptx.util import Pt

    preset = meta.get("prst")
    shape_type = next((value for value in MSO_AUTO_SHAPE_TYPE if value.xml_value == preset), None) if preset else None
    if meta.get("kind") == "cxnSp":
        x, y, width, height = geometry
        shape = slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT, x, y, x + width, y + height)
    elif meta.get("text_box"):
        shape = slide.shapes.add_textbox(*geometry)
    elif shape_type is not None:
        shape = slide.shapes.add_shape(shape_type, *geometry)
    elif meta.get("geometry_xml"):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, *geometry)
    else:
        shape = slide.shapes.add_textbox(*geometry)
        if preset:
            report.add(IssueSeverity.LOSS, "vector_graphics", "Неизвестная фигура заменена текстовым блоком.", location)
    if meta.get("has_style") is False:
        # add_shape supplies theme effects even when the source had no style reference.
        style = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
        if style is not None:
            shape._element.remove(style)
    elif meta.get("effect_ref") == "0":
        # An explicit zero is no theme effect, not the preset shape's shadow.
        effect = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style/" + A + "effectRef")
        if effect is not None:
            effect.set("idx", "0")
    if meta.get("geometry_xml"):
        _restore_geometry(shape, meta["geometry_xml"], report, location)
    if meta.get("fill") == "none" and shape.has_text_frame:
        shape.fill.background()
    elif meta.get("fill") and shape.has_text_frame:
        shape.fill.solid()
        set_color(shape.fill.fore_color, meta.get("fill_color", meta["fill"]), report, location)
    line = meta.get("line") or {}
    if line.get("color"):
        set_color(shape.line.color, line.get("stroke_color", line["color"]), report, location)
    if line.get("width") is not None:
        shape.line.width = Pt(line["width"])
    _restore_line_geometry(shape, meta.get("line_geometry", {}))
    if meta.get("gradient_colors"):
        report.add(IssueSeverity.LOSS, "styles", "Градиент фигуры не перенесён.", location)
    return shape


def _restore_geometry(shape: Any, xml: str, report: od.ConversionReport, location: str) -> None:
    from lxml import etree

    try:
        root = etree.fromstring(xml.encode(), etree.XMLParser(resolve_entities=False, no_network=True))
        if root.getroottree().docinfo.doctype or root.tag not in (A + "prstGeom", A + "custGeom"):
            raise ValueError("invalid geometry root")
        # Geometry contains only local DrawingML data, never OPC relationships.
        for node in root.iter():
            if not isinstance(node.tag, str) or not node.tag.startswith(A) or any(key.startswith("{") for key in node.attrib):
                raise ValueError("non-local geometry")
        properties = shape._element.spPr
        for child in list(properties):
            if child.tag in (A + "prstGeom", A + "custGeom"):
                properties.remove(child)
        properties.insert(1 if properties.find(A + "xfrm") is not None else 0, root)
    except (ValueError, etree.XMLSyntaxError):
        report.add(IssueSeverity.LOSS, "vector_graphics", "Некорректная геометрия заменена стандартной.", location)


def _restore_line_geometry(shape: Any, values: dict[str, Any]) -> None:
    from lxml import etree

    if not values:
        return
    line = shape._element.spPr.get_or_add_ln()
    allowed = {"headEnd": {"type", "w", "len"}, "tailEnd": {"type", "w", "len"}, "prstDash": {"val"}, "noFill": set()}
    for tag, attributes in values.items():
        if tag not in allowed:
            continue
        for old in list(line):
            if old.tag == A + tag or (tag == "noFill" and old.tag in (A + "solidFill", A + "gradFill")):
                line.remove(old)
        child = etree.SubElement(line, A + tag)
        for key, value in attributes.items():
            if key in allowed[tag]:
                child.set(key, str(value))
    order = {
        name: index
        for index, name in enumerate(
            (
                "noFill",
                "solidFill",
                "gradFill",
                "pattFill",
                "prstDash",
                "custDash",
                "round",
                "bevel",
                "miter",
                "headEnd",
                "tailEnd",
                "extLst",
            )
        )
    }
    line[:] = sorted(line, key=lambda node: order.get(node.tag.removeprefix(A), 99))
