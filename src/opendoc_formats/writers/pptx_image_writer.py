"""Native raster and SVG package resources for PPTX."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import opendoc_model as od
from opendoc_model.diagnostics import IssueSeverity


def write_image(
    slide: Any,
    image: od.Image,
    document: od.DocumentModel,
    geometry: tuple[int, int, int, int],
    report: od.ConversionReport,
    location: str,
) -> Any:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls

    resource = document.resources[image.resource_id]
    data = resource.data if resource.data is not None else Path(resource.source).read_bytes()
    if resource.media_type == "image/svg+xml":
        package = slide.part.package
        vector = Part(package.next_partname("/ppt/media/vector%d.svg"), resource.media_type, package, data)
        rid = slide.part.relate_to(vector, RT.IMAGE)
        primary, extension = rid, ""
        fallback = document.resources.get(image.properties.fallback_resource_id)
        if fallback:
            preview = fallback.data if fallback.data is not None else Path(fallback.source).read_bytes()
            _, primary = slide.part.get_or_add_image_part(BytesIO(preview))
            extension = (
                '<a:extLst><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
                '<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
                f'r:embed="{rid}"/></a:ext></a:extLst>'
            )
        else:
            report.add(
                IssueSeverity.WARNING,
                "vector_graphics",
                "SVG без превью требует совместимой программы просмотра.",
                location,
            )
        x, y, width, height = geometry
        shape_id = slide.shapes._next_shape_id
        element = parse_xml(
            f'<p:pic {nsdecls("p", "a", "r")}><p:nvPicPr><p:cNvPr id="{shape_id}" name="SVG {shape_id}"/>'
            "<p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:blipFill>"
            f'<a:blip r:embed="{primary}">{extension}</a:blip><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
            f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{width}" cy="{height}"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
        )
        slide.shapes._spTree.insert_element_before(element, "p:extLst")
        picture = slide.shapes[-1]
    else:
        picture = slide.shapes.add_picture(BytesIO(data), *geometry)
    picture._element.nvPicPr.cNvPr.set("descr", image.alt_text)
    if image.crop:
        for side in ("left", "top", "right", "bottom"):
            setattr(picture, f"crop_{side}", getattr(image.crop, side))
