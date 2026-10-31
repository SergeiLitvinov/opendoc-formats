"""Native PPTX tables and category charts."""

from __future__ import annotations

from typing import Any

import opendoc_model as od
from opendoc_model.diagnostics import IssueSeverity
from opendoc_model.document_model import Paragraph

from opendoc_formats.ooxml.table_grid import native_grid
from opendoc_formats.support.cancellation import cancellable_items, checkpoint
from opendoc_formats.writers.pptx_cell_writer import configure_cell
from opendoc_formats.writers.pptx_text_writer import write_text


def write_table(
    slide: Any, block: od.Table, geometry: tuple[int, int, int, int], report: od.ConversionReport, location: str,
    *, limits: od.DocumentLimits | None = None,
) -> None:
    from pptx.util import Pt

    rows = len(block.rows)
    if not rows or not any(row.cells for row in block.rows):
        report.add(IssueSeverity.LOSS, "tables", "Пустая таблица пропущена.", location)
        return
    layout = block.properties.get('opendoc_formats.pptx_table_grid', {}).get('layout')
    legacy_physical = layout is None and all(
        cell.properties.get('opendoc_formats.pptx_cell') == {'row': r, 'column': c}
        for r, row in enumerate(block.rows) for c, cell in enumerate(row.cells)
    )
    if layout == 'physical-unverified' or legacy_physical:
        columns = max(len(row.cells) for row in block.rows)
        placements = [(r, c, cell) for r, row in enumerate(block.rows) for c, cell in enumerate(row.cells)]
        if layout == 'physical-unverified':
            report.add(IssueSeverity.LOSS, 'pptx.table-grid',
                       'Непроверенная физическая сетка: объединения и скрытое содержимое могут отличаться.', location)
    else:
        columns, placements = native_grid(block, limits=limits)
    indexed = {id(cell): ci for row in block.rows for ci, cell in enumerate(row.cells)}
    row_placements: dict[int, list[Any]] = {}
    owners = {}
    identities = {}
    for r, c, cell in cancellable_items(placements):
        row_placements.setdefault(r, []).append((c, cell))
        identities[r, c] = cell.properties.get('opendoc_formats.pptx_cell')
        for y in range(r, min(rows, r + cell.row_span)):
            for x in range(c, min(columns, c + cell.column_span)):
                checkpoint()
                owners.setdefault((y, x), (r, c))
    table = slide.shapes.add_table(rows, columns, *geometry).table
    widths = block.properties.get("column_widths_pt", [])
    if len(widths) == columns and all(width > 0 for width in cancellable_items(widths)):
        for column, width in zip(table.columns, widths):
            checkpoint()
            column.width = Pt(width)
    for r, row in enumerate(block.rows):
        checkpoint()
        if row.properties.get("height_pt", 0) > 0:
            table.rows[r].height = Pt(row.properties["height_pt"])
        for c, cell in row_placements.get(r, []):
            checkpoint()
            target = table.cell(r, c)
            cell_location = f"{location}.rows[{r}].cells[{indexed[id(cell)]}]"
            if target.is_spanned:
                configure_cell(target, cell.properties, report, cell_location)
                continue
            if cell.row_span > 1 or cell.column_span > 1:
                end_r, end_c = r + cell.row_span - 1, c + cell.column_span - 1
                if end_r < rows and end_c < columns:
                    target.merge(table.cell(end_r, end_c))
                else:
                    report.add(IssueSeverity.LOSS, "tables", "Некорректное объединение ячеек пропущено.", location)
            paragraph_count = 0
            for child in cell.blocks:
                checkpoint()
                if isinstance(child, Paragraph):
                    write_text(target.text_frame, child, report, cell_location, append=paragraph_count > 0)
                    paragraph_count += 1
                else:
                    report.add(IssueSeverity.LOSS, "tables", "Нетекстовый блок внутри ячейки не перенесён.", location)
            configure_cell(target, cell.properties, report, cell_location)
    for item in cancellable_items(block.properties.get('opendoc_formats.pptx_continuations', [])):
        keys = ('row', 'column', 'owner_row', 'owner_column')
        if (not isinstance(item, dict) or any(type(item.get(key)) is not int for key in keys)
                or not isinstance(item.get('properties'), dict)):
            report.add(IssueSeverity.LOSS, 'pptx.table-grid', 'Некорректные данные оформления продолжения пропущены.', location)
            continue
        r, c, owner_r, owner_c = (item[key] for key in keys)
        owner = (owner_r, owner_c)
        same_origin = identities.get(owner) == {'row': owner_r, 'column': owner_c}
        if (0 <= r < rows and 0 <= c < columns and same_origin
                and owners.get((r, c)) == owner and table.cell(r, c).is_spanned):
            configure_cell(table.cell(r, c), item['properties'], report, location)
        else:
            report.add(IssueSeverity.LOSS, 'pptx.table-grid',
                       'Оформление продолжения не применено: структура объединения изменилась.', location)
    if block.properties:
        report.add(
            IssueSeverity.LOSS,
            "styles",
            "Тема таблицы не восстановлена: унаследованные цвет и начертание текста, заливки и границы могут отличаться; "
            "явное оформление ячеек сохранено.",
            location,
        )


def write_chart(
    slide: Any, data: dict[str, Any], geometry: tuple[int, int, int, int], report: od.ConversionReport, location: str
) -> bool:
    from opendoc_formats.writers.pptx_chart_writer import write_chart as export_chart

    return export_chart(slide, data, geometry, report, location)
