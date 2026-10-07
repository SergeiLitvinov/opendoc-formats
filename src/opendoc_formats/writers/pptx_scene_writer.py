"""Slide-local transforms, connector identities and native group hierarchy."""

from __future__ import annotations

import math
from collections import defaultdict
from itertools import groupby
from typing import Any

import opendoc_model as od
from opendoc_model.diagnostics import IssueSeverity

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def apply_transform(shape: Any, block: od.Block, report: od.ConversionReport, location: str) -> None:
    """Decompose orthogonal affine maps; Box edits override the imported frame."""
    from pptx.util import Pt

    box = block.box
    if box is None:
        return
    shape.rotation = box.rotation
    props = block.properties.get("pptx", {})
    transform, original = props.get("transform"), props.get("imported_box")
    if not transform:
        return
    a, b, c, d, e, f = transform["matrix"]
    width, height = transform["width_pt"], transform["height_pt"]
    if original:
        sx = box.width / original["width"] if original["width"] else 1
        sy = box.height / original["height"] if original["height"] else 1
        a, b, c, d = a * sx, b * sy, c * sx, d * sy
        e = box.x + (e - original["x"]) * sx
        f = box.y + (f - original["y"]) * sy
    horizontal, vertical = math.hypot(a, b), math.hypot(c, d)
    if not horizontal or not vertical or not math.isclose(a * c + b * d, 0, abs_tol=1e-7 * horizontal * vertical):
        report.add(IssueSeverity.LOSS, "page_geometry", "Сдвиговая трансформация заменена рамкой объекта.", location)
        return
    cx, cy = e + (a * width + c * height) / 2, f + (b * width + d * height) / 2
    target_width, target_height = horizontal * width, vertical * height
    shape.left, shape.top = Pt(cx - target_width / 2), Pt(cy - target_height / 2)
    shape.width, shape.height = Pt(target_width), Pt(target_height)
    rotation_delta = box.rotation - original["rotation"] if original else 0
    shape.rotation = (math.degrees(math.atan2(b, a)) + rotation_delta) % 360
    xfrm = shape._element.find(".//" + A + "xfrm")
    if xfrm is not None:
        xfrm.set("flipH", "0")
        xfrm.set("flipV", "1" if a * d - b * c < 0 else "0")


def restore_connections(
    entries: list[tuple[od.Block, Any, str]], report: od.ConversionReport, groups: dict[str, list[Any]] | None = None
) -> None:
    """Rebind source IDs to new slide IDs; ambiguous references remain detached."""
    from lxml import etree

    identities = defaultdict(list)
    for identity, shapes in (groups or {}).items():
        identities[identity].extend(shapes)
    for block, shape, _ in entries:
        identity = block.properties.get("pptx", {}).get("object_id")
        if identity is not None:
            identities[str(identity)].append(shape)
    for block, shape, location in entries:
        links = block.properties.get("pptx", {}).get("shape", {}).get("connections", {})
        parent = shape._element.find(P + "nvCxnSpPr/" + P + "cNvCxnSpPr")
        if not links or parent is None:
            continue
        for tag in ("stCxn", "endCxn"):
            link = links.get(tag)
            if link is None:
                continue
            targets = identities.get(str(link.get("id")), [])
            site = str(link.get("idx", ""))
            if len(targets) != 1 or not site.isdecimal():
                report.add(IssueSeverity.LOSS, "relationships", "Привязка коннектора отсутствует или неоднозначна.", location)
                continue
            etree.SubElement(parent, A + tag, id=str(targets[0].shape_id), idx=site)


def restore_groups(slide: Any, entries: list[tuple[od.Block, Any, str]], report: od.ConversionReport) -> dict[str, list[Any]]:
    """Restore nesting with identity group transforms and absolute child geometry."""
    items = [(shape, block.properties.get("pptx", {}).get("groups", []), location) for block, shape, location in entries]
    seen = set()
    identities = defaultdict(list)

    def assemble(rows: list[tuple[Any, list[dict[str, Any]], str]], depth: int = 0, prefix: tuple[str, ...] = ()) -> list[Any]:
        result = []
        for identity, run in groupby(rows, key=lambda row: str(row[1][depth].get("id")) if len(row[1]) > depth else None):
            run = list(run)
            if identity is None:
                result.extend(row[0] for row in run)
                continue
            key = (*prefix, identity)
            if key in seen:
                report.add(IssueSeverity.LOSS, "groups", "Разорванная порядком объектов группа разделена.", run[0][2])
            seen.add(key)
            members = assemble(run, depth + 1, key)
            position = list(slide.shapes._spTree).index(members[0]._element)
            group = slide.shapes.add_group_shape(members)
            group.name = run[0][1][depth].get("name") or "Группа"
            identities[identity].append(group)
            slide.shapes._spTree.remove(group._element)
            slide.shapes._spTree.insert(position, group._element)
            result.append(group)
        return result

    assemble(items)
    return identities
