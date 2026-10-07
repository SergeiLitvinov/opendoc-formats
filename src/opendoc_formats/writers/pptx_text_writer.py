"""Native DrawingML text and Office Math, without importing PPTX at package import."""

from __future__ import annotations

from typing import Any

import opendoc_model as od
from opendoc_model.color import ColorValue
from opendoc_model.diagnostics import IssueSeverity
from opendoc_model.document_model import Formula, FormulaFormat, TextRun

from opendoc_formats.writers.pptx_paragraph_writer import configure_frame, configure_paragraph


def set_color(
    color_format: Any, value: od.ColorValue | str | dict[str, Any] | None, report: od.ConversionReport, location: str
) -> None:
    from pptx.dml.color import RGBColor

    if value is None:
        return
    try:
        if isinstance(value, dict):
            value = ColorValue.from_dict(value)
        color = value if isinstance(value, ColorValue) else ColorValue.from_hex(value)
        color_format.rgb = RGBColor.from_string(color.to_hex().lstrip("#"))
        if color.alpha != 1 or color.icc_profile or color.blend_mode != "normal":
            report.add(IssueSeverity.LOSS, "color", "Цвет приведён к непрозрачному RGB.", location)
    except ValueError:
        report.add(IssueSeverity.LOSS, "color", "Цвет не удалось перенести.", location)


def write_text(frame: Any, block: od.Block, report: od.ConversionReport, location: str, *, append: bool = False) -> None:
    from pptx.util import Pt

    if append:
        paragraph = frame.add_paragraph()
    else:
        frame.clear()
        configure_frame(frame, block, report, location)
        paragraph = frame.paragraphs[0]
    paragraph_index = 0
    configure_paragraph(paragraph, block, paragraph_index)
    for item in block.content:
        if isinstance(item, Formula):
            write_formula(paragraph, item, report, location)
        elif isinstance(item, TextRun):
            for index, text in enumerate(item.text.split("\n")):
                if index:
                    if item.properties.get("pptx_break") == "line":
                        paragraph.add_line_break()
                    else:
                        paragraph = frame.add_paragraph()
                        paragraph_index += 1
                        configure_paragraph(paragraph, block, paragraph_index)
                run = paragraph.add_run()
                run.text = text
                style = item.style
                run.font.name = style.font_family
                if style.font_family and style.font_family.startswith(("+mj-", "+mn-")):
                    report.add(
                        IssueSeverity.LOSS,
                        "fonts",
                        "Ссылка на шрифт темы не разрешена; гарнитура зависит от темы результата.",
                        location,
                    )
                run.font.size = Pt(style.font_size.pt) if style.font_size else None
                run.font.bold, run.font.italic = style.bold, style.italic
                run.font.underline = style.underline
                set_color(run.font.color, style.color, report, location)
                if style.superscript or style.subscript:
                    run._r.get_or_add_rPr().set("baseline", "30000" if style.superscript else "-25000")
                if item.link:
                    run.hyperlink.address = item.link
                if style.background:
                    report.add(IssueSeverity.LOSS, "styles", "Фон текста не перенесён.", location)
        else:
            report.add(IssueSeverity.LOSS, "raster_images", "Встроенное в строку изображение не перенесено.", location)


def write_formula(paragraph: Any, formula: od.Formula, report: od.ConversionReport, location: str) -> None:
    from lxml import etree

    value = formula.value
    if formula.format is FormulaFormat.MATHML:
        from opendoc_formats.writers.mathml_to_omml import mathml_to_omml

        try:
            value = mathml_to_omml(value)
        except ValueError as error:
            paragraph.add_run().text = formula.fallback_text or formula.value
            report.add(IssueSeverity.LOSS, "formulas", f"MathML сохранён как текст: {error}.", location)
            return
    if formula.format in {FormulaFormat.OMML, FormulaFormat.MATHML}:
        try:
            root = etree.fromstring(value.encode(), etree.XMLParser(resolve_entities=False, no_network=True))
            ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
            math = root if root.tag == f"{{{ns}}}oMath" else root.find(f".//{{{ns}}}oMath")
            if math is not None:
                wrapper = etree.Element(
                    "{http://schemas.microsoft.com/office/drawing/2010/main}m",
                    nsmap={
                        "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
                    },
                )
                wrapper.append(math)
                paragraph._p.append(wrapper)
                return
        except etree.XMLSyntaxError:
            pass
    paragraph.add_run().text = formula.fallback_text or formula.value
    report.add(IssueSeverity.LOSS, "formulas", "Формула сохранена как текст; необходим корректный OMML.", location)
