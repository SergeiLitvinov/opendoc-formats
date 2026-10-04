"""DocumentModel → editable PPTX. One model section becomes one slide."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import opendoc as od
from opendoc.diagnostics import ConversionReport, IssueSeverity
from opendoc.document_model import Box, DocumentModel, Formula, Image, Paragraph, Section, Table

from opendoc_formats.writers.pptx_chart_writer import write_chart
from opendoc_formats.writers.pptx_geometry_writer import write_shape
from opendoc_formats.writers.pptx_image_writer import write_image
from opendoc_formats.writers.pptx_objects_writer import write_table
from opendoc_formats.writers.pptx_scene_writer import apply_transform, restore_connections, restore_groups
from opendoc_formats.writers.pptx_text_writer import set_color, write_text


def write_pptx_model(document: DocumentModel, output_path: str | Path) -> ConversionReport:
    from pptx import Presentation
    from pptx.util import Pt

    report = ConversionReport(Path(output_path))
    errors = document.validate()
    if errors:
        for error in errors:
            report.add(IssueSeverity.ERROR, "model", error)
        return report
    target = Presentation()
    sections = document.sections or [Section()]
    page = sections[0].page
    target.slide_width, target.slide_height = Pt(page.width.pt), Pt(page.height.pt)
    for index, section in enumerate(sections):
        location = f"sections[{index}]"
        slide = target.slides.add_slide(target.slide_layouts[6])
        if (section.page.width.pt, section.page.height.pt) != (page.width.pt, page.height.pt):
            report.add(IssueSeverity.LOSS, "page_geometry", "Все слайды используют размер первой секции.", location)
        background = section.properties.get("background_fill")
        if background and background != "none":
            slide.background.fill.solid()
            set_color(slide.background.fill.fore_color, section.properties.get("background_color", background), report, location)
        if section.properties.get("background_gradient_colors"):
            report.add(IssueSeverity.LOSS, "styles", "Градиент фона заменён сплошной заливкой.", location)
        if section.properties.get("notes"):
            slide.notes_slide.notes_text_frame.text = section.properties["notes"]
        y = section.page.margin_top.pt
        entries = []
        for number, block in enumerate(section.blocks):
            box = block.box or Box(
                section.page.margin_left.pt,
                y,
                max(1, page.width.pt - section.page.margin_left.pt - section.page.margin_right.pt),
                48,
            )
            connector = block.properties.get("pptx", {}).get("shape", {}).get("kind") == "cxnSp"
            minimum = 0 if connector else 1
            geometry = tuple(Pt(value) for value in (box.x, box.y, max(minimum, box.width), max(minimum, box.height)))
            block_location = f"{location}.blocks[{number}]"
            previous_count = len(slide.shapes)
            _write_block(slide, block, document, geometry, report, block_location)
            if len(slide.shapes) > previous_count:
                shape = slide.shapes[-1]
                apply_transform(shape, block, report, block_location)
                entries.append((block, shape, block_location))
            y = max(y, box.y + box.height + 12)
            if box.y + box.height > page.height.pt:
                report.add(IssueSeverity.LOSS, "page_geometry", "Объект выходит за нижнюю границу слайда.", block_location)
        groups = restore_groups(slide, entries, report)
        restore_connections(entries, report, groups)
        if any(
            (
                section.headers,
                section.footers,
                section.first_page_headers,
                section.first_page_footers,
                section.even_page_headers,
                section.even_page_footers,
            )
        ):
            report.add(IssueSeverity.LOSS, "running_content", "Колонтитулы не перенесены в слайды.", location)
    if document.package:
        report.add(IssueSeverity.LOSS, "package", "Исходные темы, макеты, анимации и непрозрачные части пакета не перенесены.")
    report.metrics["slides"] = len(target.slides)
    report.output_path.parent.mkdir(parents=True, exist_ok=True)
    target.save(str(report.output_path))
    return report


def _write_block(
    slide: Any,
    block: od.Block,
    document: od.DocumentModel,
    geometry: tuple[int, int, int, int],
    report: od.ConversionReport,
    location: str,
) -> None:
    if isinstance(block, Table):
        write_table(slide, block, geometry, report, location)
    elif isinstance(block, Image):
        try:
            write_image(slide, block, document, geometry, report, location)
        except (ValueError, OSError):
            report.add(IssueSeverity.LOSS, "raster_images", "Изображение не удалось перенести.", location)
    elif isinstance(block, (Paragraph, Formula)):
        if isinstance(block, Formula):
            block = Paragraph([block])
        meta = block.properties.get("pptx", {})
        if meta.get("chart") and write_chart(slide, meta["chart"], geometry, report, location):
            return
        shape = write_shape(slide, meta.get("shape", {}), geometry, report, location)
        if shape.has_text_frame:
            write_text(shape.text_frame, block, report, location)
        elif block.content:
            report.add(IssueSeverity.LOSS, "text", "Текст внутри коннектора не поддержан.", location)
