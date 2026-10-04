"""Text extraction-result rendering, independent of application operations."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Union

from opendoc_formats.support.latex import PREAMBLE as _PREAMBLE
from opendoc_formats.support.latex import escape_latex as _escape_latex
from opendoc_formats.types import Text

logger = logging.getLogger(__name__)


def render_latex(*, text: Text, title: str = "Document", author: str = "Author") -> str:
    """Минимальный LaTeX-рендер: заголовок, абзацы, таблицы.

    Сложная вёрстка (стили, списки, изображения) намеренно упрощена —
    для научных текстов обычно достаточно. Если нужна полная конвертация,
    используйте ``render.latex.pandoc`` (требует установленный ``pandoc``).
    """
    parts: list[str] = [_PREAMBLE, r"\begin{document}", ""]
    parts.append(rf"\title{{{_escape_latex(title)}}}")
    parts.append(rf"\author{{{_escape_latex(author)}}}")
    parts.append(r"\date{\today}")
    parts.append(r"\maketitle")
    parts.append("")

    in_list = False
    for block in text.blocks:
        if block.type.value == "heading":
            if in_list:
                parts.append(r"\end{itemize}")
                parts.append("")
                in_list = False
            level = max(1, min(block.level or 1, 3))
            cmd = ("section", "subsection", "subsubsection")[level - 1]
            parts.append(rf"\{cmd}{{{_escape_latex(block.text)}}}")
            parts.append("")
        elif block.type.value == "list_item":
            if not in_list:
                parts.append(r"\begin{itemize}")
                parts.append("")
                in_list = True
            parts.append(rf"\item {_escape_latex(block.text)}")
        else:
            if in_list:
                parts.append(r"\end{itemize}")
                parts.append("")
                in_list = False
            if block.type.value == "code" or block.type.value == "equation":
                parts.append(block.text)
            else:
                parts.append(_escape_latex(block.text))
            parts.append("")
    if in_list:
        parts.append(r"\end{itemize}")
        parts.append("")

    for table in text.tables:
        if not table.rows:
            continue
        cols = max((len(r) for r in table.rows), default=0)
        spec = "|" + "|".join(["c"] * cols) + "|"
        parts.append(rf"\begin{{tabular}}{{{spec}}}")
        parts.append(r"\hline")
        for row in table.rows:
            cells = [_escape_latex(c) for c in row]
            parts.append(" & ".join(cells) + r" \\")
            parts.append(r"\hline")
        parts.append(r"\end{tabular}")
        parts.append("")

    parts.append(r"\end{document}")
    return "\n".join(parts)


def render_latex_pandoc(*, text: Text, input_path: Union[str, Path, None] = None) -> str:
    """Если есть ``input_path`` (DOCX), конвертирует pandoc-ом. Иначе fallback на ``render.latex``."""
    import shutil
    import subprocess

    if not shutil.which("pandoc"):
        logger.warning("pandoc не найден, fallback на render.latex")
        return render_latex(text=text)

    if input_path is None:
        return render_latex(text=text)

    from opendoc_formats.support.artifacts import ArtifactWorkspace

    try:
        with ArtifactWorkspace(prefix="opendoc_formats_pandoc_") as workspace:
            out = workspace.artifact_path("output.tex")
            subprocess.run(
                [
                    "pandoc",
                    str(input_path),
                    "-o",
                    str(out),
                    "--from",
                    "docx",
                    "--to",
                    "latex",
                    "--standalone",
                    "--top-level-division=chapter",
                ],
                check=True,
                capture_output=True,
                text=True,
                timeout=120,
            )
            workspace.validate_artifact(out)
            return out.read_text(encoding="utf-8")
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, FileNotFoundError) as error:
        logger.warning("pandoc failed: %s, fallback на render.latex", error)
        return render_latex(text=text)


def render_docx(*, text: Text, output_path: Union[str, Path]) -> Path:
    """Записать ``Text`` в DOCX. Возвращает путь к созданному файлу."""
    import io

    from docx import Document

    from opendoc_formats.support.io import atomic_write_bytes

    out = Path(output_path)
    d = Document()
    for block in text.blocks:
        if block.type.value == "heading":
            level = max(1, min(block.level or 1, 3))
            d.add_heading(block.text, level=level)
        else:
            d.add_paragraph(block.text)
    for table in text.tables:
        if not table.rows:
            continue
        cols = max((len(r) for r in table.rows), default=0)
        t = d.add_table(rows=len(table.rows), cols=cols)
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row):
                if j < cols:
                    t.cell(i, j).text = cell
    buffer = io.BytesIO()
    d.save(buffer)
    atomic_write_bytes(out, buffer.getvalue())
    return out
