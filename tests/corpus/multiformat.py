"""Воспроизводимый многоформатный corpus: PDF, PPTX и EMF."""

from __future__ import annotations

import hashlib
import io
import posixpath
import struct
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import fitz
from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

FIXED_ZIP_TIME = (2024, 1, 1, 12, 0, 0)
FIXED_DATETIME = datetime(2024, 1, 1, 12, 0, tzinfo=UTC)
EMF_SIGNATURE = 0x464D4520


def build_multiformat_corpus(output_dir: str | Path) -> dict[str, Path]:
    target = Path(output_dir)
    target.mkdir(parents=True, exist_ok=True)
    emf = build_emf(target / "scientific-diagram.emf")
    return {
        "emf": emf,
        "pdf": build_pdf(target / "scientific-layout.pdf"),
        "pptx": build_pptx(target / "scientific-slides.pptx", emf),
    }


def build_emf(output: str | Path) -> Path:
    """Создать минимальный валидный EMF с header и EOF records."""

    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    header = struct.pack(
        "<II4i4iIIIIHHIII2i2iIII2i",
        1,
        108,
        0,
        0,
        100,
        100,
        0,
        0,
        2540,
        2540,
        EMF_SIGNATURE,
        0x00010000,
        128,
        2,
        1,
        0,
        0,
        0,
        0,
        100,
        100,
        25,
        25,
        0,
        0,
        0,
        25000,
        25000,
    )
    eof = struct.pack("<IIIII", 14, 20, 0, 0, 20)
    path.write_bytes(header + eof)
    return path


def build_pdf(output: str | Path) -> Path:
    """Создать PDF с фиксированной геометрией, таблицей, raster и vector content."""

    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = fitz.open()
    page = document.new_page(width=480, height=640)
    page.insert_text((36, 48), "OpenDoc Formats PDF quality corpus", fontname="helv", fontsize=18, color=(0.18, 0.45, 0.71))
    page.insert_text((36, 78), "Left column: semantic text and reading order.", fontname="Times-Roman", fontsize=10)
    page.insert_text((258, 78), "Right column: geometry and fidelity metrics.", fontname="Times-Roman", fontsize=10)
    for index in range(4):
        y = 102 + index * 18
        page.insert_text((36, y), f"L{index + 1}: structured scientific paragraph", fontname="Times-Roman", fontsize=9)
        page.insert_text((258, y), f"R{index + 1}: measured conversion signal", fontname="Times-Roman", fontsize=9)

    table = fitz.Rect(36, 205, 444, 285)
    page.draw_rect(table, color=(0.18, 0.45, 0.71), width=1)
    for x in (180, 320):
        page.draw_line((x, table.y0), (x, table.y1), color=(0.35, 0.35, 0.35), width=0.7)
    for y in (232, 258):
        page.draw_line((table.x0, y), (table.x1, y), color=(0.35, 0.35, 0.35), width=0.7)
    for x, y, value in (
        (43, 222, "Feature"),
        (187, 222, "Expected"),
        (327, 222, "Observed"),
        (43, 249, "Text"),
        (187, 249, "Editable"),
        (327, 249, "Exact"),
        (43, 276, "Vector"),
        (187, 276, "Native"),
        (327, 276, "Preserved"),
    ):
        page.insert_text((x, y), value, fontname="helv", fontsize=8)

    shape = page.new_shape()
    shape.draw_bezier((55, 365), (150, 300), (270, 420), (420, 330))
    shape.draw_circle((150, 350), 22)
    shape.finish(color=(0.1, 0.35, 0.65), fill=(0.84, 0.92, 0.98), width=2)
    shape.commit()
    page.insert_text((36, 445), "Figure 1. Native PDF vector commands and raster control.", fontname="helv", fontsize=9)
    page.insert_image(fitz.Rect(330, 470, 430, 545), stream=_control_png())
    page.insert_text((36, 590), "Header/footer repetition candidate — page 1", fontname="helv", fontsize=8)

    second = document.new_page(width=640, height=480)
    second.insert_text((40, 50), "Landscape formulas and diagrams", fontname="helv", fontsize=18)
    second.insert_text((40, 105), "E = mc^2     integral_0^1 x^2 dx = 1/3", fontname="cour", fontsize=14)
    second.draw_rect(fitz.Rect(40, 150, 600, 370), color=(0.18, 0.45, 0.71), fill=(0.95, 0.98, 1.0), width=1.2)
    second.draw_line((90, 320), (550, 190), color=(0.85, 0.25, 0.2), width=3)
    second.insert_text((40, 438), "Header/footer repetition candidate — page 2", fontname="helv", fontsize=8)
    document.set_metadata(
        {
            "title": "OpenDoc Formats multi-format PDF corpus",
            "author": "OpenDoc Formats",
            "subject": "PDF geometry and semantic extraction",
            "keywords": "pdf, corpus, geometry, vector, table",
            "creator": "OpenDoc Formats corpus generator",
            "producer": "OpenDoc Formats / PyMuPDF",
            "creationDate": "D:20240101120000Z",
            "modDate": "D:20240101120000Z",
        }
    )
    document.save(path, garbage=4, deflate=True, no_new_id=True)
    document.close()
    return path


def build_pptx(output: str | Path, emf_path: str | Path) -> Path:
    """Создать OOXML presentation с notes, chart, table и нативным EMF part."""

    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    properties = presentation.core_properties
    properties.title = "OpenDoc Formats multi-format PPTX corpus"
    properties.subject = "Presentation semantics and vector fidelity"
    properties.author = "OpenDoc Formats"
    properties.keywords = "pptx, corpus, emf, chart, notes"
    properties.created = FIXED_DATETIME
    properties.modified = FIXED_DATETIME

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = title_slide.shapes.title
    title.text = "Scientific conversion quality"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    title.text_frame.paragraphs[0].runs[0].font.bold = True
    title_slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.4), Inches(1.0)).text = (
        "Editable semantics + faithful rendering"
    )
    table = title_slide.shapes.add_table(3, 3, Inches(0.8), Inches(2.8), Inches(5.8), Inches(2.1)).table
    table_rows = (
        ("Layer", "Metric", "Target"),
        ("Text", "Editability", "Exact"),
        ("Graphics", "Fidelity", "Native"),
    )
    for row, values in enumerate(table_rows):
        for column, value in enumerate(values):
            table.cell(row, column).text = value

    chart_data = ChartData()
    chart_data.categories = ["Text", "Tables", "Graphics"]
    chart_data.add_series("Retention", (100, 96, 92))
    title_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(7.0),
        Inches(1.4),
        Inches(5.4),
        Inches(4.2),
        chart_data,
    )
    title_slide.notes_slide.notes_text_frame.text = "Speaker note: verify chart, table, theme and editable text."

    vector_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    vector_slide.shapes.title.text = "Vector and formula controls"
    placeholder = path.parent / ".__emf_placeholder.png"
    placeholder.write_bytes(_control_png())
    picture = vector_slide.shapes.add_picture(str(placeholder), Inches(0.9), Inches(1.4), Inches(5.5), Inches(3.4))
    picture.name = "EMF_VECTOR_PLACEHOLDER"
    placeholder.unlink(missing_ok=True)
    formula = vector_slide.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.2), Inches(1.0))
    formula.text_frame.text = "E = mc²   |   ∫₀¹ x² dx = ⅓"
    formula.text_frame.paragraphs[0].runs[0].font.name = "Cambria Math"
    formula.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    vector_slide.notes_slide.notes_text_frame.text = "Speaker note: EMF must remain vector; formulas remain editable text."

    presentation.save(path)
    _patch_emf_picture(path, Path(emf_path).read_bytes())
    _normalise_archive(path)
    return path


def sha256(path: str | Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _control_png() -> bytes:
    image = Image.new("RGB", (160, 120), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((5, 5, 154, 114), outline="#2E74B5", width=4)
    draw.ellipse((28, 24, 132, 100), fill="#D9EAF7", outline="#2E74B5", width=3)
    draw.line((35, 92, 78, 48, 126, 82), fill="#C43D32", width=5)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=False, compress_level=9)
    return buffer.getvalue()


def _patch_emf_picture(path: Path, emf: bytes) -> None:
    entries = _read_archive(path)
    slide_name = "ppt/slides/slide2.xml"
    relationships_name = "ppt/slides/_rels/slide2.xml.rels"
    slide = etree.fromstring(entries[slide_name])
    namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    blips = slide.xpath(
        ".//p:pic[p:nvPicPr/p:cNvPr[@name='EMF_VECTOR_PLACEHOLDER']]//a:blip",
        namespaces=namespaces,
    )
    if len(blips) != 1:
        raise ValueError("EMF placeholder relationship was not found")
    relationship_id = blips[0].get(f"{{{namespaces['r']}}}embed")
    relationships = etree.fromstring(entries[relationships_name])
    package_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    relationship = relationships.find(f"{{{package_namespace}}}Relationship[@Id='{relationship_id}']")
    if relationship is None:
        raise ValueError("EMF placeholder package relationship was not found")
    old_target = relationship.get("Target", "")
    old_part = posixpath.normpath(posixpath.join("ppt/slides", old_target))
    relationship.set("Target", "../media/scientific-diagram.emf")
    entries[relationships_name] = etree.tostring(relationships, xml_declaration=True, encoding="UTF-8", standalone=True)
    entries.pop(old_part, None)
    entries["ppt/media/scientific-diagram.emf"] = emf

    content_types = etree.fromstring(entries["[Content_Types].xml"])
    content_namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    if not content_types.xpath("./ct:Default[@Extension='emf']", namespaces={"ct": content_namespace}):
        default = etree.SubElement(content_types, f"{{{content_namespace}}}Default")
        default.set("Extension", "emf")
        default.set("ContentType", "image/x-emf")
    entries["[Content_Types].xml"] = etree.tostring(
        content_types,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )
    _write_archive(path, entries)


def _normalise_archive(path: Path) -> None:
    _write_archive(path, _read_archive(path))


def _read_archive(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _write_archive(path: Path, entries: dict[str, bytes]) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(entries):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            payload = _normalise_embedded_archive(entries[name]) if name.endswith(".xlsx") else entries[name]
            archive.writestr(info, payload)
    path.write_bytes(buffer.getvalue())


def _normalise_embedded_archive(data: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    core_name = "docProps/core.xml"
    if core_name in entries:
        core = etree.fromstring(entries[core_name])
        date_namespace = "http://purl.org/dc/terms/"
        for name in ("created", "modified"):
            node = core.find(f"{{{date_namespace}}}{name}")
            if node is not None:
                node.text = "2024-01-01T12:00:00Z"
        entries[core_name] = etree.tostring(core, xml_declaration=True, encoding="UTF-8", standalone=True)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(entries):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, entries[name])
    return buffer.getvalue()


__all__ = ["EMF_SIGNATURE", "build_emf", "build_multiformat_corpus", "build_pdf", "build_pptx", "sha256"]


if __name__ == "__main__":
    for name, artifact in build_multiformat_corpus(Path("tmp/corpus/multiformat")).items():
        print(f"{name}: {artifact} ({sha256(artifact)})")
