"""Real DOCX transactions: original bytes, OOXML parts, styles and conservative editing."""

import copy
import io
import zipfile
from dataclasses import FrozenInstanceError, replace
from hashlib import sha256
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image

from opendoc_formats.docx import (
    DocxLimits,
    DocxPackage,
    InsertParagraph,
    InsertTableRow,
    ReplaceTextSpan,
    SetParagraphText,
    SetTableRowText,
)
from opendoc_formats.errors import (
    BackendUnavailableError,
    DocumentClosedError,
    InvalidDocumentError,
    OperationCancelledError,
    PatchConflictError,
    ResourceLimitError,
    UnsupportedDocumentError,
)


def fixture_bytes():
    doc = Document()
    doc.core_properties.title = "Native editing fixture"
    p = doc.add_paragraph(style="Title")
    p.add_run("Before ").italic = True
    p.add_run("{{ va").bold = True
    p.add_run("lue }}")
    p.add_run(" after").underline = True
    doc.add_paragraph("Second body paragraph")
    doc.sections[0].header.paragraphs[0].text = "Header {{ value }}"
    doc.sections[0].footer.paragraphs[0].text = "Footer value"
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cell {{ value }}"
    table.cell(0, 1).text = "Other cell"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Merged cell"
    nested = doc.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested table"
    image = io.BytesIO()
    Image.new("RGB", (12, 8), "blue").save(image, format="PNG")
    p = doc.add_paragraph("Picture ")
    p.add_run().add_picture(io.BytesIO(image.getvalue()))
    p = doc.add_paragraph("Math ")
    math = OxmlElement("m:oMath")
    math.append(OxmlElement("m:r"))
    p._p.append(math)
    doc.add_section()
    doc.add_paragraph("Last section")
    stream = io.BytesIO()
    doc.save(stream)
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source:
        infos = [(copy.copy(i), source.read(i.filename)) for i in source.infolist()]
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as target:
        target.comment = b"Preserved native package comment"
        for info, data in infos:
            info.comment = b"part comment"
            info.date_time = (2020, 1, 2, 3, 4, 6)
            info.external_attr = 0o100644 << 16
            target.writestr(info, data)
    return output.getvalue()


@pytest.fixture
def source():
    return fixture_bytes()


def transform(data, part, transform_bytes):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as archive, zipfile.ZipFile(out, "w") as target:
        target.comment = archive.comment
        for info in archive.infolist():
            value = archive.read(info.filename)
            target.writestr(info, transform_bytes(value) if info.filename == part else value)
    return out.getvalue()


def test_snapshot_ids_roles_body_tables_flags_and_immutability(source):
    with DocxPackage(source) as package:
        paragraphs = package.paragraphs
        assert {p.role for p in paragraphs} >= {"body", "header", "footer"}
        assert paragraphs[0].id == "word/document.xml:0"
        assert paragraphs[0].is_body and not paragraphs[0].in_table
        assert any(p.in_table and p.table_id and p.cell_id for p in paragraphs)
        assert any(p.has_image for p in paragraphs)
        assert any(p.has_embedded_objects for p in paragraphs)
        assert any(p.has_section_break for p in paragraphs)
        assert [v.kind for v in package.body][:3] == ["paragraph", "paragraph", "table"]
        table = next(t for t in package.tables if t.is_body)
        assert table.grid_columns == 2 and table.rows[0].simple
        assert table.rows[1].has_merge and not table.rows[1].simple
        assert table.rows[1].cells[0].grid_span == 2
        assert any(t.parent_cell_id and not t.is_body for t in package.tables)
        with pytest.raises(FrozenInstanceError):
            paragraphs[0].text = "mutated"
        assert package.to_bytes() == source
    assert package.closed
    assert paragraphs[0].text == "Before {{ value }} after"
    with pytest.raises(DocumentClosedError):
        _ = package.paragraphs
    with pytest.raises(DocumentClosedError):
        package.to_bytes()
    package.close()


def test_split_run_replacement_preserves_styles_parts_and_zip_info(source):
    with DocxPackage(source) as package:
        p = package.paragraphs[0]
        begin = p.text.index("{{")
        end = p.text.index("}}") + 2
        output = package.to_bytes([ReplaceTextSpan(p.id, begin, end, "VALUE\nNEXT\tEND")])
        assert package.paragraphs[0] == p
        assert package.to_bytes() == source
    with DocxPackage(output) as package:
        p = package.paragraphs[0]
        assert p.text == "Before VALUE\nNEXT\tEND after"
        assert p.runs[0].style.italic
        assert next(r for r in p.runs if "VALUE" in r.text).style.bold
        assert p.runs[-1].style.underline
    restored = Document(io.BytesIO(output))
    assert restored.paragraphs[0].text == "Before VALUE\nNEXT\tEND after"
    with zipfile.ZipFile(io.BytesIO(source)) as before, zipfile.ZipFile(io.BytesIO(output)) as after:
        assert before.comment == after.comment
        assert before.namelist() == after.namelist()
        for original in before.infolist():
            written = after.getinfo(original.filename)
            assert (written.date_time, written.external_attr, written.comment, written.compress_type) == (
                original.date_time,
                original.external_attr,
                original.comment,
                original.compress_type,
            )
            if original.filename != "word/document.xml":
                assert sha256(before.read(original)).digest() == sha256(after.read(original.filename)).digest()


def test_header_table_and_body_spans_are_independent(source):
    with DocxPackage(source) as package:
        selected = [p for p in package.paragraphs if "{{ value }}" in p.text]
        edits = [ReplaceTextSpan(p.id, p.text.index("{{"), p.text.index("}}") + 2, "ordinary {% text %}") for p in selected]
        out = package.to_bytes(edits)
    with DocxPackage(out) as package:
        assert len([p for p in package.paragraphs if "ordinary {% text %}" in p.text]) == len(selected)


def test_ranges_use_original_offsets_not_mutated_order(source):
    with DocxPackage(source) as package:
        p = package.paragraphs[0]
        out = package.to_bytes([ReplaceTextSpan(p.id, 0, 6, "X"), ReplaceTextSpan(p.id, 19, 24, "Y")])
    with DocxPackage(out) as package:
        assert package.paragraphs[0].text == "X {{ value }} Y"


@pytest.mark.parametrize("index", [0, 7, 24])
def test_zero_length_insertions_at_run_boundaries(source, index):
    with DocxPackage(source) as package:
        p = package.paragraphs[0]
        assert len(p.text) == 24
        out = package.to_bytes([ReplaceTextSpan(p.id, index, index, "INSERT")])
    with DocxPackage(out) as package:
        assert package.paragraphs[0].text == p.text[:index] + "INSERT" + p.text[index:]


def test_whole_and_append_preserve_paragraph_and_first_run_style(source):
    with DocxPackage(source) as package:
        p = package.paragraphs[0]
        replaced = package.to_bytes([SetParagraphText(p.id, "Replacement")])
        appended = package.to_bytes([SetParagraphText(p.id, " appended", append=True)])
    for out, text in [(replaced, "Replacement"), (appended, p.text + " appended")]:
        with DocxPackage(out) as package:
            paragraph = package.paragraphs[0]
            assert paragraph.text == text and paragraph.style_id == p.style_id
            assert paragraph.runs[-1].style.italic


def test_insert_and_update_body_and_table_markers(source):
    with DocxPackage(source) as package:
        p = package.paragraphs[0]
        row = package.tables[0].rows[0]
        out = package.to_bytes(
            [
                InsertParagraph(p.id, "OPEN"),
                InsertParagraph(p.id, "CLOSE", position="after"),
                InsertTableRow(row.id, "ROW OPEN"),
                InsertTableRow(row.id, "ROW CLOSE", position="after"),
            ]
        )
    with DocxPackage(out) as package:
        assert [p.text for p in package.paragraphs if p.is_body][:3] == ["OPEN", p.text, "CLOSE"]
        table = package.tables[0]
        assert [r.text for r in table.rows][:3] == ["ROW OPEN", row.text, "ROW CLOSE"]
        assert table.rows[0].cells[0].grid_span == 2
        updated = package.to_bytes([SetTableRowText(table.rows[0].id, "UPDATED")])
    with DocxPackage(updated) as package:
        assert package.tables[0].rows[0].text == "UPDATED"
        assert package.tables[0].rows[0].cells[0].grid_span == 2


@pytest.mark.parametrize(
    "case", ["overlap", "double-insertion", "rewrite-and-span", "unknown-id", "invalid-range", "row-overlap"]
)
def test_conflicting_transaction_does_not_modify_source_or_existing_output(source, tmp_path, case):
    output = tmp_path / "result.docx"
    output.write_bytes(b"previous")
    with DocxPackage(source) as package:
        p = package.paragraphs[0]
        cases = {
            "overlap": [ReplaceTextSpan(p.id, 0, 3, "a"), ReplaceTextSpan(p.id, 2, 5, "b")],
            "double-insertion": [InsertParagraph(p.id, "a"), InsertParagraph(p.id, "b")],
            "rewrite-and-span": [SetParagraphText(p.id, "a"), ReplaceTextSpan(p.id, 0, 1, "b")],
            "unknown-id": [SetParagraphText("stale", "a")],
            "invalid-range": [ReplaceTextSpan(p.id, 0, 1000, "a")],
            "row-overlap": [InsertTableRow(package.tables[0].rows[0].id, "a"), InsertTableRow(package.tables[0].rows[0].id, "b")],
        }
        with pytest.raises(PatchConflictError):
            package.write(output, cases[case])
        assert package.to_bytes() == source
    assert output.read_bytes() == b"previous"
    assert not list(tmp_path.glob("*.partial"))


def test_complex_rewrite_is_rejected_but_span_keeps_picture(source):
    with DocxPackage(source) as package:
        p = next(p for p in package.paragraphs if p.has_image)
        with pytest.raises(UnsupportedDocumentError):
            package.to_bytes([SetParagraphText(p.id, "destroy image")])
        out = package.to_bytes([ReplaceTextSpan(p.id, 0, 7, "Caption ")])
    with DocxPackage(out) as package:
        assert next(p for p in package.paragraphs if p.text.startswith("Caption")).has_image


@pytest.mark.parametrize(
    "markup,flag",
    [
        (b'<w:r><w:fldChar w:fldCharType="begin"/></w:r>', "has_fields"),
        (b"<w:ins><w:r><w:t>tracked</w:t></w:r></w:ins>", "has_revisions"),
        (b"<w:r><w:drawing><w:p><w:r><w:t>nested</w:t></w:r></w:p></w:drawing></w:r>", "has_nested_paragraphs"),
    ],
)
def test_complex_span_flags_and_rejection(source, markup, flag):
    data = transform(source, "word/document.xml", lambda b: b.replace(b"</w:p>", markup + b"</w:p>", 1))
    with DocxPackage(data) as package:
        p = package.paragraphs[0]
        assert getattr(p, flag)
        with pytest.raises(UnsupportedDocumentError):
            package.to_bytes([ReplaceTextSpan(p.id, 0, 1, "X")])


def test_macro_policy_is_explicit_and_never_executes(source):
    data = transform(
        source,
        "[Content_Types].xml",
        lambda b: b.replace(
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
            b"application/vnd.ms-word.document.macroEnabled.main+xml",
        ),
    )
    with pytest.raises(UnsupportedDocumentError):
        DocxPackage(data)
    with DocxPackage(data, allow_macros=True) as package:
        assert package.info.has_macros and package.to_bytes() == data


@pytest.mark.parametrize(
    "value",
    [
        b"not ZIP",
        b"PK\x03\x04 truncated",
    ],
)
def test_corrupt_input(value):
    with pytest.raises(InvalidDocumentError):
        DocxPackage(value)


def test_dtd_external_entities_and_duplicate_parts_are_rejected(source):
    xml = b'<!DOCTYPE x [<!ENTITY e SYSTEM "https://example.invalid/resource">]><x>&e;</x>'
    data = transform(source, "word/document.xml", lambda _: xml)
    with pytest.raises(UnsupportedDocumentError):
        DocxPackage(data)
    out = io.BytesIO(source)
    with zipfile.ZipFile(out, "a") as archive:
        with pytest.warns(UserWarning):
            archive.writestr("word/document.xml", b"duplicate")
    with pytest.raises(InvalidDocumentError):
        DocxPackage(out.getvalue())


@pytest.mark.parametrize(
    "limits",
    [
        replace(DocxLimits(), max_input_bytes=100),
        replace(DocxLimits(), max_entries=2),
        replace(DocxLimits(), max_uncompressed_bytes=100),
        replace(DocxLimits(), max_part_bytes=100),
        replace(DocxLimits(), max_compression_ratio=1),
        replace(DocxLimits(), max_paragraphs=1),
        replace(DocxLimits(), max_text_chars=1),
    ],
)
def test_input_budgets(source, limits):
    with pytest.raises(ResourceLimitError):
        DocxPackage(source, limits=limits)


def test_missing_optional_engine(source, monkeypatch):
    import opendoc_formats.native.common as common

    original = common.importlib.import_module

    def unavailable(name):
        if name == "lxml.etree":
            raise ImportError("blocked")
        return original(name)

    monkeypatch.setattr(common.importlib, "import_module", unavailable)
    with pytest.raises(BackendUnavailableError):
        DocxPackage(source)


def test_cancel_after_staging_preserves_existing_destination(source, tmp_path, monkeypatch):
    import opendoc_formats.native.common as common

    output = tmp_path / "result.docx"
    output.write_bytes(b"previous")
    state = {"cancel": False}
    original = common.Path.write_bytes

    def written(path, data):
        result = original(path, data)
        if path.suffix == ".partial":
            state["cancel"] = True
        return result

    monkeypatch.setattr(common.Path, "write_bytes", written)
    with DocxPackage(source, cancelled=lambda: state["cancel"]) as package:
        with pytest.raises(OperationCancelledError):
            package.write(output, [SetParagraphText(package.paragraphs[0].id, "Updated")])
    assert output.read_bytes() == b"previous"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["result.docx"]


def test_repository_fixture_can_be_read_and_edited():
    path = Path(__file__).parents[1] / "corpus/native/native-edit.docx"
    with DocxPackage(path) as package:
        p = package.paragraphs[0]
        edited = package.to_bytes([ReplaceTextSpan(p.id, 7, 18, "Unicode 🦊\tvalue\nnext")])
        assert package.to_bytes() == path.read_bytes()
    with DocxPackage(edited) as package:
        assert package.paragraphs[0].text == "Before Unicode 🦊\tvalue\nnext after"


def test_signed_package_is_readable_but_edits_are_rejected(source):
    out = io.BytesIO(source)
    with zipfile.ZipFile(out, "a") as archive:
        archive.writestr("_xmlsignatures/sig1.xml", b"<Signature/>")
    signed = out.getvalue()
    with DocxPackage(signed) as package:
        assert package.info.has_signatures and package.to_bytes() == signed
        with pytest.raises(UnsupportedDocumentError):
            package.to_bytes([SetParagraphText(package.paragraphs[0].id, "edited")])


@pytest.mark.parametrize("target", [b"missing.xml", b"../../../../escape.xml", b"C:/external.xml"])
def test_invalid_internal_relationship_never_opens_target(source, target):
    data = transform(source, "_rels/.rels", lambda value: value.replace(b"word/document.xml", target))
    with pytest.raises(InvalidDocumentError):
        DocxPackage(data)


def test_field_spanning_paragraphs_blocks_whole_and_span_edits(source):
    def fields(value):
        return value.replace(
            b"<w:t>Before </w:t>",
            b'<w:fldChar w:fldCharType="begin"/><w:t>Before </w:t>',
        ).replace(
            b"<w:t>Second body paragraph</w:t>",
            b'<w:t>Second body paragraph</w:t><w:fldChar w:fldCharType="end"/>',
        )

    with DocxPackage(transform(source, "word/document.xml", fields)) as package:
        p = package.paragraphs[1]
        assert p.has_fields
        for patch in (
            SetParagraphText(p.id, "replacement"),
            SetParagraphText(p.id, "append", append=True),
            ReplaceTextSpan(p.id, 0, 1, "X"),
        ):
            with pytest.raises(UnsupportedDocumentError):
                package.to_bytes([patch])


def test_output_and_patch_budgets_preserve_destination(source, tmp_path):
    output = tmp_path / "result.docx"
    output.write_bytes(b"old result")
    with DocxPackage(source, limits=replace(DocxLimits(), max_output_bytes=100)) as package:
        with pytest.raises(ResourceLimitError):
            package.write(output, [SetParagraphText(package.paragraphs[0].id, "New")])
    assert output.read_bytes() == b"old result"
    with DocxPackage(source, limits=replace(DocxLimits(), max_patches=1)) as package:
        with pytest.raises(ResourceLimitError):
            package.to_bytes([SetParagraphText(p.id, "New") for p in package.paragraphs[:2]])
