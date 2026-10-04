"""Managed office conversion: real executable acceptance plus failure boundaries."""

from pathlib import Path

import pytest
from docx import Document

from opendoc_formats.errors import BackendUnavailableError, InvalidDocumentError, OfficeTimeoutError, OperationCancelledError
from opendoc_formats.office import convert_office_to_pdf, find_libreoffice
from opendoc_formats.pdf import PdfDocument


def office_input(tmp_path):
    source = tmp_path / "Source with spaces.docx"
    doc = Document()
    doc.add_paragraph("Office preview acceptance")
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "Table content"
    doc.save(source)
    return source


@pytest.mark.skipif(find_libreoffice() is None, reason="LibreOffice executable not installed")
def test_real_office_conversion_isolated_profile_and_valid_pdf(tmp_path):
    source = office_input(tmp_path)
    original = source.read_bytes()
    output = tmp_path / "Output with spaces.pdf"
    result = convert_office_to_pdf(source, output, timeout=90)
    assert result.output == output and result.page_count >= 1 and result.bytes_written == output.stat().st_size
    assert source.read_bytes() == original
    with PdfDocument(output) as pdf:
        assert "Office preview acceptance" in pdf.page_info(0).text
        assert pdf.render_page(0, max_dimension=300).png.startswith(b"\x89PNG")
    assert not list(tmp_path.glob("*.partial"))


def test_missing_executable_and_corrupt_output_preserve_target(tmp_path, monkeypatch):
    import opendoc_formats.office as office

    source = office_input(tmp_path)
    output = tmp_path / "result.pdf"
    output.write_bytes(b"previous")
    with pytest.raises(BackendUnavailableError):
        convert_office_to_pdf(source, output, executable=tmp_path / "missing")

    class FailedProcess:
        returncode = 0

        def __init__(self, args, **kwargs):
            Path(args[args.index("--outdir") + 1], "input.pdf").write_bytes(b"not PDF")

        def poll(self):
            return 0

    monkeypatch.setattr(office.subprocess, "Popen", FailedProcess)
    with pytest.raises(InvalidDocumentError):
        convert_office_to_pdf(source, output, executable=Path(__file__))
    assert output.read_bytes() == b"previous"


@pytest.mark.parametrize("cancel", [False, True])
def test_timeout_and_running_cancellation_terminate_owned_process(tmp_path, monkeypatch, cancel):
    import opendoc_formats.office as office

    source = office_input(tmp_path)
    output = tmp_path / "result.pdf"
    output.write_bytes(b"previous")
    terminated = []
    state = {"running": False}

    class WaitingProcess:
        returncode = None

        def __init__(self, args, **kwargs):
            state["running"] = True
            profile_uri = next(arg.split("=", 1)[1] for arg in args if arg.startswith("-env:UserInstallation="))
            assert "profile" in profile_uri and " " not in profile_uri
            profile = Path(args[args.index("--outdir") + 1]) / "profile/user/registrymodifications.xcu"
            assert "DisableMacrosExecution" in profile.read_text() and "<value>true</value>" in profile.read_text()

        def poll(self):
            return None

    monkeypatch.setattr(office.subprocess, "Popen", WaitingProcess)
    monkeypatch.setattr(office, "_kill_process_tree", lambda process: terminated.append(process))
    error = OperationCancelledError if cancel else OfficeTimeoutError
    with pytest.raises(error):
        convert_office_to_pdf(
            source,
            output,
            executable=Path(__file__),
            timeout=0.01,
            cancelled=lambda: state["running"] if cancel else False,
        )
    assert len(terminated) == 1 and output.read_bytes() == b"previous"
