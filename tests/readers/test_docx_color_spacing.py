"""Own source regressions for automatic color and minimum/exact/relative spacing."""

import zipfile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree
from opendoc_model import IssueSeverity, document_from_json, document_to_json, effective_text_style

from opendoc_formats import ExportOptions, read_document, write_document


def _source(path):
    native = Document()
    auto = native.add_paragraph(style='Heading 3').add_run('Own automatic color')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), 'auto')
    auto._r.get_or_add_rPr().append(color)
    native.add_paragraph('Own inherited color', style='Heading 3')
    for name, value, rule in (
        ('minimum', Pt(14), WD_LINE_SPACING.AT_LEAST),
        ('exact', Pt(14), WD_LINE_SPACING.EXACTLY),
        ('relative', 1.15, WD_LINE_SPACING.MULTIPLE),
    ):
        paragraph = native.add_paragraph()
        paragraph.add_run(f'Own {name} first line\nOwn {name} second line').font.size = Pt(24)
        paragraph.paragraph_format.line_spacing = value
        paragraph.paragraph_format.line_spacing_rule = rule
    base = native.styles.add_style('Own spacing base', WD_STYLE_TYPE.PARAGRAPH)
    base.paragraph_format.line_spacing = Pt(20)
    base.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    derived = native.styles.add_style('Own spacing derived', WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = base
    native.add_paragraph('Own inherited spacing', style=derived)
    native.save(path)


def _cycle(model, path):
    restored = document_from_json(document_to_json(model))
    report = write_document(restored, path)
    assert report.success, report.issues
    result = read_document(path)
    assert result.success, result.issues
    return result.document


def _direct_spacing(path, index):
    paragraph = Document(path).paragraphs[index]
    nodes = paragraph._p.xpath('./w:pPr/w:spacing')
    return ({key: value for key, value in nodes[0].attrib.items() if key in {qn('w:line'), qn('w:lineRule')}}
            if nodes else {})


def test_source_auto_and_all_spacing_rules_survive_two_json_docx_cycles(tmp_path):
    source = tmp_path / 'source.docx'
    _source(source)
    before = source.read_bytes()
    model = read_document(source).document
    assert model.sections[0].blocks[0].content[0].style.color == 'auto'
    assert 'color' in model.sections[0].blocks[0].content[0].style.properties['direct_fields']
    assert 'color' not in model.sections[0].blocks[1].content[0].style.properties['direct_fields']
    inherited = model.sections[0].blocks[1]
    assert inherited.content[0].style.color is None
    assert effective_text_style(model, inherited, inherited.content[0]).color.to_hex() == '#4F81BD'
    expected = [_direct_spacing(source, index) for index in range(2, 6)]
    for cycle in range(2):
        target = tmp_path / f'cycle-{cycle}.docx'
        model = _cycle(model, target)
        native = Document(target)
        assert native.paragraphs[0].runs[0]._r.xpath('./w:rPr/w:color')[0].get(qn('w:val')) == 'auto'
        assert model.sections[0].blocks[0].content[0].style.color == 'auto'
        assert [_direct_spacing(target, index) for index in range(2, 6)] == expected
        assert model.sections[0].blocks[2].properties['docx_line_spacing_rule'] == 'atLeast'
        assert native.styles['Own spacing base'].paragraph_format.line_spacing_rule == WD_LINE_SPACING.AT_LEAST
    assert source.read_bytes() == before


def test_auto_edits_and_removal_do_not_restore_original_declaration(tmp_path):
    source = tmp_path / 'source.docx'
    _source(source)
    model = read_document(source).document
    for value in ('#123456', None, 'auto'):
        model.sections[0].blocks[0].content[0].style.color = value
        for cycle in range(2):
            target = tmp_path / f'color-{value}-{cycle}.docx'
            model = _cycle(model, target)
            nodes = Document(target).paragraphs[0].runs[0]._r.xpath('./w:rPr/w:color')
            if value is None:
                assert not nodes
            else:
                assert nodes
                if value is not None:
                    assert nodes[0].get(qn('w:val')) == value.removeprefix('#')
            if value is None:
                assert model.sections[0].blocks[0].content[0].style.color is None


def test_spacing_edits_deletion_and_inherited_style_edits_survive(tmp_path):
    source = tmp_path / 'source.docx'
    _source(source)
    model = read_document(source).document
    properties = model.sections[0].blocks[2].properties
    properties['line_spacing_pt'] = 18
    properties['docx_line_spacing_rule'] = 'exact'
    model.styles['Ownspacingbase'].properties['line_spacing_pt'] = 22
    for cycle in range(2):
        target = tmp_path / f'edited-{cycle}.docx'
        model = _cycle(model, target)
        assert _direct_spacing(target, 2) == {qn('w:line'): '360', qn('w:lineRule'): 'exact'}
        assert Document(target).styles['Own spacing base'].paragraph_format.line_spacing.pt == 22
    properties = model.sections[0].blocks[2].properties
    del properties['line_spacing_pt']
    del properties['docx_line_spacing_rule']
    for cycle in range(2):
        target = tmp_path / f'deleted-{cycle}.docx'
        model = _cycle(model, target)
        assert not _direct_spacing(target, 2)
    base = model.styles['Ownspacingbase'].properties
    base.pop('line_spacing_pt')
    base.pop('docx_line_spacing_rule')
    for cycle in range(2):
        target = tmp_path / f'style-deleted-{cycle}.docx'
        model = _cycle(model, target)
        assert Document(target).styles['Own spacing base'].paragraph_format.line_spacing is None


def test_inherited_character_auto_keeps_style_and_direct_color_absence(tmp_path):
    source = tmp_path / 'character.docx'
    native = Document()
    style = native.styles.add_style('Own automatic character', WD_STYLE_TYPE.CHARACTER)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), 'auto')
    style.element.get_or_add_rPr().append(color)
    run = native.add_paragraph(style='Heading 3').add_run('Own character automatic color')
    run.style = style
    native.save(source)
    model = read_document(source).document
    for cycle in range(2):
        target = tmp_path / f'character-{cycle}.docx'
        model = _cycle(model, target)
        run = Document(target).paragraphs[0].runs[0]
        assert run.style.name == 'Own automatic character'
        assert not run._r.xpath('./w:rPr/w:color')
        paragraph = model.sections[0].blocks[0]
        assert paragraph.content[0].style.color is None
        assert effective_text_style(model, paragraph, paragraph.content[0]).color == 'auto'


def test_edit_inherited_theme_color_removes_stale_theme_attributes(tmp_path):
    source = tmp_path / 'source.docx'
    _source(source)
    model = read_document(source).document
    model.sections[0].blocks[1].content[0].style.color = '#AABBCC'
    for cycle in range(2):
        target = tmp_path / f'theme-edited-{cycle}.docx'
        model = _cycle(model, target)
        color = Document(target).paragraphs[1].runs[0]._r.xpath('./w:rPr/w:color')[0]
        assert dict(color.attrib) == {qn('w:val'): 'AABBCC'}


def test_named_color_edits_and_removal_keep_live_inheritance(tmp_path):
    source = tmp_path / 'named.docx'
    _source(source)
    model = read_document(source).document
    for value in ('#123456', None, 'auto'):
        model.styles['Heading3'].color = value
        for cycle in range(2):
            target = tmp_path / f'named-{value}-{cycle}.docx'
            model = _cycle(model, target)
            paragraph = model.sections[0].blocks[1]
            assert paragraph.content[0].style.color is None
            effective = effective_text_style(model, paragraph, paragraph.content[0]).color
            assert (effective.to_hex() if hasattr(effective, 'to_hex') else effective) == value
            native = Document(target)
            assert not native.paragraphs[1].runs[0]._r.xpath('./w:rPr/w:color')
            declarations = native.styles['Heading 3'].element.xpath('./w:rPr/w:color')
            if value is None:
                assert not declarations
            else:
                assert dict(declarations[0].attrib) == {qn('w:val'): value.removeprefix('#')}


def test_document_default_color_edits_and_removal_survive(tmp_path):
    from opendoc_formats.ooxml.word_color import default_color_owner

    source = tmp_path / 'defaults.docx'
    native = Document()
    node = OxmlElement('w:color')
    node.set(qn('w:val'), 'auto')
    default_color_owner(native.styles.element).append(node)
    native.add_paragraph('Own defaults')
    native.save(source)
    model = read_document(source).document
    for value in ('auto', '#654321', None):
        model.styles['__doc_defaults__'].color = value
        for cycle in range(2):
            target = tmp_path / f'defaults-{value}-{cycle}.docx'
            model = _cycle(model, target)
            paragraph = model.sections[0].blocks[0]
            assert paragraph.content[0].style.color is None
            effective = effective_text_style(model, paragraph, paragraph.content[0]).color
            assert (effective.to_hex() if hasattr(effective, 'to_hex') else effective) == value
            declarations = Document(target).styles.element.xpath('./w:docDefaults/w:rPrDefault/w:rPr/w:color')
            assert ([dict(node.attrib) for node in declarations]
                    == ([{qn('w:val'): value.removeprefix('#')}] if value is not None else []))


def test_character_style_reference_deletion_does_not_restore_snapshot(tmp_path):
    source = tmp_path / 'character-reference.docx'
    native = Document()
    style = native.styles.add_style('Own auto', WD_STYLE_TYPE.CHARACTER)
    node = OxmlElement('w:color')
    node.set(qn('w:val'), 'auto')
    style.element.get_or_add_rPr().append(node)
    native.add_paragraph(style='Heading 3').add_run('Own character reference').style = style
    native.save(source)
    model = read_document(source).document
    model.sections[0].blocks[0].content[0].style.properties['base_style_id'] = None
    for cycle in range(2):
        target = tmp_path / f'character-reference-{cycle}.docx'
        model = _cycle(model, target)
        paragraph = model.sections[0].blocks[0]
        assert effective_text_style(model, paragraph, paragraph.content[0]).color.to_hex() == '#4F81BD'
        assert not Document(target).paragraphs[0].runs[0]._r.xpath('./w:rPr/w:rStyle')


def test_paragraph_style_reference_deletion_does_not_restore_snapshot(tmp_path):
    from opendoc_model import set_heading

    source = tmp_path / 'paragraph-reference.docx'
    native = Document()
    native.add_paragraph('Own paragraph reference', style='Heading 3')
    native.save(source)
    model = read_document(source).document
    model.sections[0].blocks[0].style_id = None
    set_heading(model.sections[0].blocks[0], None)
    for cycle in range(2):
        target = tmp_path / f'paragraph-reference-{cycle}.docx'
        model = _cycle(model, target)
        paragraph = model.sections[0].blocks[0]
        assert effective_text_style(model, paragraph, paragraph.content[0]).color is None
        assert Document(target).paragraphs[0].style.name == 'Normal'


@pytest.mark.parametrize('target_format', ['html', 'pdf', 'pptx'])
def test_cross_format_inherited_rgb_is_rendered_without_mutating_declarations(tmp_path, target_format):
    source = tmp_path / 'inherited.docx'
    native = Document()
    native.add_paragraph('Own inherited RGB', style='Heading 3')
    native.save(source)
    model = read_document(source).document
    before = document_to_json(model)
    target = tmp_path / ('inherited.' + target_format)
    report = write_document(model, target)
    assert report.success, report.issues
    assert document_to_json(model) == before
    assert model.sections[0].blocks[0].content[0].style.color is None
    if target_format == 'html':
        assert '#4f81bd' in target.read_text(encoding='utf-8').lower()
    elif target_format == 'pdf':
        import fitz

        with fitz.open(target) as pdf:
            spans = [span for block in pdf[0].get_text('dict')['blocks'] if block['type'] == 0
                     for line in block['lines'] for span in line['spans']]
            assert spans and all(span['color'] == 0x4F81BD for span in spans)
    else:
        from pptx import Presentation

        assert str(Presentation(target).slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.color.rgb) == '4F81BD'


def test_nonstandard_source_style_ids_and_derived_role_survive(tmp_path):
    from opendoc_model import Heading, get_heading, set_heading

    source = tmp_path / 'renamed.docx'
    _source(source)
    native = Document(source)
    native.styles['Heading 3'].element.set(qn('w:styleId'), 'own-heading-31')
    for node in native._element.findall('.//' + qn('w:pStyle')):
        if node.get(qn('w:val')) == 'Heading3':
            node.set(qn('w:val'), 'own-heading-31')
    native.save(source)
    model = read_document(source).document
    for cycle in range(2):
        target = tmp_path / f'renamed-{cycle}.docx'
        model = _cycle(model, target)
        assert Document(target).paragraphs[1].style.style_id == 'own-heading-31'
        assert not Document(target).paragraphs[1].runs[0]._r.xpath('./w:rPr/w:color')
    set_heading(model.sections[0].blocks[1], Heading(4))
    for cycle in range(2):
        target = tmp_path / f'derived-{cycle}.docx'
        model = _cycle(model, target)
        assert get_heading(model.sections[0].blocks[1]) == Heading(4)
        assert Document(target).paragraphs[1].style.base_style.style_id == 'own-heading-31'


def test_cross_format_auto_and_spacing_losses_are_located_and_gate_refuses(tmp_path):
    source = tmp_path / 'source.docx'
    _source(source)
    model = read_document(source).document
    target = tmp_path / 'output.txt'
    target.write_bytes(b'own sentinel')
    report = write_document(model, target, options=ExportOptions(
        quality_gate=lambda report, _: not any(issue.severity is IssueSeverity.LOSS for issue in report.issues),
    ))
    assert not report.success and target.read_bytes() == b'own sentinel'
    assert any(issue.feature == 'color-auto' and 'blocks[0]' in issue.location for issue in report.issues)
    assert any(issue.feature == 'docx.line-spacing-rule' and 'blocks[2]' in issue.location for issue in report.issues)


@pytest.mark.parametrize('changes', [
    {'docx_line_spacing_rule': 'unknown'},
    {'line_spacing_pt': .025},
    {'line_spacing_pt': True},
    {'line_spacing_pt': 2**31},
    {'line_spacing': 1.2},
])
def test_unrepresentable_spacing_refuses_atomically_with_location(tmp_path, changes):
    source = tmp_path / 'source.docx'
    _source(source)
    model = read_document(source).document
    model.sections[0].blocks[2].properties.update(changes)
    target = tmp_path / 'target.docx'
    target.write_bytes(b'own sentinel')
    report = write_document(model, target)
    assert not report.success and target.read_bytes() == b'own sentinel'
    assert any(issue.feature == 'docx.line-spacing' and issue.severity is IssueSeverity.ERROR
               and 'blocks[2]' in issue.location for issue in report.issues)


def test_unknown_native_rule_is_located_and_requires_explicit_resolution(tmp_path):
    source = tmp_path / 'source.docx'
    _source(source)
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(entries['word/document.xml'])
    spacing = root.find('.//' + qn('w:spacing'))
    spacing.set(qn('w:lineRule'), 'unknown')
    entries['word/document.xml'] = etree.tostring(root)
    with zipfile.ZipFile(source, 'w') as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    imported = read_document(source)
    assert imported.success
    assert any(issue.code == 'docx.line-spacing' and '#xpath=' in issue.location for issue in imported.issues)
    target = tmp_path / 'target.docx'
    target.write_bytes(b'own sentinel')
    assert not write_document(imported.document, target).success
    assert target.read_bytes() == b'own sentinel'
    imported.document.sections[0].blocks[2].properties['docx_line_spacing_rule'] = None
    assert write_document(imported.document, target).success


def test_legacy_missing_rule_has_loss_and_zero_budget_preserves_destination(tmp_path):
    source = tmp_path / 'source.docx'
    _source(source)
    model = read_document(source).document
    properties = model.sections[0].blocks[2].properties
    properties.pop('docx_line_spacing_rule')
    properties.pop('docx_line_spacing_imported')
    target = tmp_path / 'target.docx'
    target.write_bytes(b'own sentinel')
    report = write_document(model, target, options=ExportOptions(
        quality_gate=lambda report, _: not any(issue.severity is IssueSeverity.LOSS for issue in report.issues),
    ))
    assert not report.success and target.read_bytes() == b'own sentinel'
    assert any(issue.feature == 'docx.line-spacing-rule' and 'blocks[2]' in issue.location for issue in report.issues)
