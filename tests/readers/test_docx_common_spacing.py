"""Native own sources exercise the common Model spacing contract and live inheritance."""

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt
from opendoc_model import (
    DocumentLimits,
    IssueSeverity,
    LineSpacing,
    compare_documents,
    compare_inspections,
    document_from_json,
    document_to_json,
    get_line_spacing,
    inspect_document_model,
    resolve_line_spacing,
    set_line_spacing,
)

from opendoc_formats import ExportOptions, read_document, write_document


@pytest.fixture
def source(tmp_path):
    path = tmp_path / 'own-common.docx'
    native = Document()
    base = native.styles.add_style('Own base', WD_STYLE_TYPE.PARAGRAPH)
    base.paragraph_format.line_spacing = Pt(20)
    base.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    child = native.styles.add_style('Own child', WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    for value, rule in ((Pt(14), WD_LINE_SPACING.AT_LEAST), (Pt(18), WD_LINE_SPACING.EXACTLY),
                        (1.15, WD_LINE_SPACING.MULTIPLE)):
        p = native.add_paragraph('Own line one\nOwn line two', style=child)
        p.paragraph_format.line_spacing = value
        p.paragraph_format.line_spacing_rule = rule
    native.add_paragraph('Own inherited paragraph', style=child)
    native.save(path)
    return path


def cycle(model, path):
    restored = document_from_json(document_to_json(model))
    report = write_document(restored, path)
    assert report.success, report.issues
    return read_document(path).document


def test_import_direct_and_inherited_common_spacing_with_two_cycles(source, tmp_path):
    before = source.read_bytes()
    original = model = read_document(source).document
    for n in range(3):
        assert [get_line_spacing(p) for p in model.sections[0].blocks] == [
            LineSpacing('minimum', 14), LineSpacing('exact', 18), LineSpacing('relative', 1.15), None,
        ]
        assert get_line_spacing(model.styles['Ownchild']) is None
        assert get_line_spacing(model.styles['Ownbase']) == LineSpacing('minimum', 20)
        inherited = model.sections[0].blocks[3]
        assert resolve_line_spacing(inherited, document=model) == LineSpacing('minimum', 20)
        compared = compare_inspections(inspect_document_model(original), inspect_document_model(model))
        assert compared.object_diff['line_spacing_available'] is True
        assert not any(i.code.startswith('line-spacing-') for i in compare_documents(original, model).issues)
        model = cycle(model, tmp_path / f'own-cycle-{n}.docx') if n < 2 else model
    assert source.read_bytes() == before


@pytest.mark.parametrize('record', [LineSpacing('minimum', 22), LineSpacing('exact', 18), LineSpacing('relative', 1.25)])
def test_style_edits_remain_inherited_and_are_compared(source, tmp_path, record):
    original = read_document(source).document
    model = document_from_json(document_to_json(original))
    set_line_spacing(model.styles['Ownbase'], record)
    assert any(i.code.startswith('line-spacing-') for i in compare_documents(original, model).issues)
    for n in range(2):
        target = tmp_path / f'own-style-{n}.docx'
        model = cycle(model, target)
        paragraph = model.sections[0].blocks[3]
        assert get_line_spacing(paragraph) is None
        assert resolve_line_spacing(paragraph, document=model) == record
        assert not Document(target).paragraphs[3]._p.xpath('./w:pPr/w:spacing/@w:line')


def test_deleting_direct_spacing_returns_live_style_and_removes_native_override(source, tmp_path):
    model = read_document(source).document
    set_line_spacing(model.sections[0].blocks[0], None)
    set_line_spacing(model.styles['Ownbase'], LineSpacing('exact', 24))
    for n in range(2):
        target = tmp_path / f'own-remove-{n}.docx'
        model = cycle(model, target)
        paragraph = model.sections[0].blocks[0]
        assert get_line_spacing(paragraph) is None
        assert resolve_line_spacing(paragraph, document=model) == LineSpacing('exact', 24)
        assert not Document(target).paragraphs[0]._p.xpath('./w:pPr/w:spacing/@w:line')
    set_line_spacing(model.styles['Ownbase'], None)
    for n in range(2):
        target = tmp_path / f'own-style-remove-{n}.docx'
        model = cycle(model, target)
        assert get_line_spacing(model.styles['Ownbase']) is None
        assert not Document(target).styles['Own base'].element.xpath('./w:pPr/w:spacing/@w:line')


@pytest.mark.parametrize('record', [LineSpacing(), LineSpacing('exact', .025)])
def test_unknown_and_unrepresentable_spacing_refuse_atomically(source, tmp_path, record):
    model = read_document(source).document
    set_line_spacing(model.sections[0].blocks[0], record)
    snapshot = document_to_json(model)
    target = tmp_path / 'refused.docx'
    target.write_bytes(b'previous destination')
    report = write_document(model, target)
    assert not report.success and target.read_bytes() == b'previous destination'
    assert any(i.feature == 'docx.line-spacing' and 'blocks[0]' in i.location for i in report.issues)
    assert document_to_json(model) == snapshot


def test_spacing_extensions_report_loss_and_selected_budget_reaches_codec(source, tmp_path):
    model = read_document(source).document
    limits = DocumentLimits(max_nodes=300000)
    set_line_spacing(model.sections[0].blocks[0], LineSpacing('minimum', 14, extra={'own': [0] * 100001}), limits=limits)
    target = tmp_path / 'extension.docx'
    report = write_document(model, target, options=ExportOptions(document_limits=limits))
    assert report.success, report.issues
    assert any(i.feature == 'line-spacing.extra' for i in report.issues)
    target.write_bytes(b'previous destination')
    report = write_document(model, target, options=ExportOptions(document_limits=limits,
        quality_gate=lambda r, _: not any(i.severity is IssueSeverity.LOSS for i in r.issues)))
    assert not report.success and target.read_bytes() == b'previous destination'


def test_common_direct_rule_is_written_in_native_units(tmp_path):
    from opendoc_model import DocumentModel, Paragraph, Section, TextRun
    p = Paragraph([TextRun('Own common minimum')])
    set_line_spacing(p, LineSpacing('minimum', 14))
    path = tmp_path / 'own-new.docx'
    report = write_document(DocumentModel(sections=[Section(blocks=[p])]), path)
    assert report.success, report.issues
    spacing = Document(path).paragraphs[0]._p.xpath('./w:pPr/w:spacing')[0]
    assert spacing.get(qn('w:line')) == '280' and spacing.get(qn('w:lineRule')) == 'atLeast'


def test_defaults_edit_and_delete_use_common_inheritance(source, tmp_path):
    model = read_document(source).document
    set_line_spacing(model.styles['Ownbase'], None)
    set_line_spacing(model.styles['__doc_defaults__'], LineSpacing('exact', 19))
    for n in range(2):
        target = tmp_path / f'own-default-{n}.docx'
        model = cycle(model, target)
        assert resolve_line_spacing(model.sections[0].blocks[3], document=model) == LineSpacing('exact', 19)
        spacing = Document(target).styles.element.xpath('./w:docDefaults/w:pPrDefault/w:pPr/w:spacing')[0]
        assert spacing.get(qn('w:line')) == '380' and spacing.get(qn('w:lineRule')) == 'exact'
    set_line_spacing(model.styles['__doc_defaults__'], None)
    for n in range(2):
        target = tmp_path / f'own-default-remove-{n}.docx'
        model = cycle(model, target)
        assert resolve_line_spacing(model.sections[0].blocks[3], document=model) is None
        assert not Document(target).styles.element.xpath('./w:docDefaults/w:pPrDefault/w:pPr/w:spacing/@w:line')


def test_private_preview_preserves_legacy_edits_and_common_unknown_blocks_fallback(source):
    from opendoc_formats.writers.color_inheritance import materialize_docx_colors
    model = read_document(source).document
    model.styles['Ownbase'].properties['line_spacing_pt'] = 22
    model.sections[0].blocks[0].properties['line_spacing_pt'] = 16
    set_line_spacing(model.sections[0].blocks[1], LineSpacing())
    snapshot = document_to_json(model)
    view = materialize_docx_colors(model)
    assert view.sections[0].blocks[0].properties['line_spacing_pt'] == 16
    assert view.sections[0].blocks[3].properties['line_spacing_pt'] == 22
    assert 'line_spacing_pt' not in view.sections[0].blocks[1].properties
    assert document_to_json(model) == snapshot


def test_replacing_unknown_native_spacing_uses_common_setter(source, tmp_path):
    import zipfile

    from lxml import etree
    with zipfile.ZipFile(source) as z:
        entries = {name: z.read(name) for name in z.namelist()}
    root = etree.fromstring(entries['word/document.xml'])
    root.find('.//' + qn('w:spacing')).set(qn('w:lineRule'), 'own-unknown')
    entries['word/document.xml'] = etree.tostring(root)
    with zipfile.ZipFile(source, 'w') as z:
        for name, data in entries.items():
            z.writestr(name, data)
    model = read_document(source).document
    assert get_line_spacing(model.sections[0].blocks[0]) == LineSpacing()
    set_line_spacing(model.sections[0].blocks[0], LineSpacing('relative', 1.25))
    for n in range(2):
        model = cycle(model, tmp_path / f'own-unknown-replaced-{n}.docx')
        assert get_line_spacing(model.sections[0].blocks[0]) == LineSpacing('relative', 1.25)


@pytest.mark.parametrize('kind', ['character', 'table', 'inline'])
def test_unsupported_common_style_owner_refuses_atomically(tmp_path, kind):
    from opendoc_model import DocumentModel, Paragraph, Section, TextRun, TextStyle
    style = TextStyle(properties={'style_type': kind} if kind != 'inline' else {})
    set_line_spacing(style, LineSpacing('minimum', 14))
    model = DocumentModel(styles={'own': style} if kind != 'inline' else {}, sections=[Section(blocks=[Paragraph([
        TextRun('Own source', style=style) if kind == 'inline' else TextRun('Own source'),
    ])])])
    target = tmp_path / 'unsupported-style.docx'
    target.write_bytes(b'previous destination')
    report = write_document(model, target)
    assert not report.success and target.read_bytes() == b'previous destination'
    location = 'styles[own]' if kind != 'inline' else 'sections[0].blocks[0].content[0].style'
    assert any(i.feature == 'docx.line-spacing' and i.location == location for i in report.issues)
