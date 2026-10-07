"""Own DOCX source fixtures distinguish retained XML from lost objects."""

import zipfile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from lxml import etree
from opendoc_model import document_from_json, document_to_json, get_integration

from opendoc_formats import read_document, write_document


def _text(parent, text):
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    parent.append(run)


def _verify(source, result):
    assert result.success and not result.lossless and not result.assessment_complete
    ledger = get_integration(result.document)
    assert ledger and not ledger.assessment_complete
    with zipfile.ZipFile(source) as archive:
        for record in ledger.preservation:
            assert record.provenance.source_path == str(source)
            root = etree.fromstring(archive.read(record.provenance.package_part.lstrip("/")))
            if "#xpath=" in record.issue.location:
                xpath = record.issue.location.split("#xpath=", 1)[1]
                assert root.xpath(xpath, namespaces={key: value for key, value in root.nsmap.items() if key})
    restored = document_from_json(document_to_json(result.document))
    assert get_integration(restored) == ledger
    transport = source.with_suffix(".json")
    transport.write_text(document_to_json(restored), encoding="utf-8")
    assert read_document(transport).issues == result.issues
    return restored


def test_inline_opaque_and_block_revision_loss_are_distinct(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Before ")
    inline = OxmlElement("w:ins")
    inline.set(qn("w:id"), "7")
    _text(inline, "Retained inline")
    paragraph._p.append(inline)
    block = OxmlElement("w:ins")
    block.set(qn("w:id"), "8")
    block_paragraph = OxmlElement("w:p")
    _text(block_paragraph, "Lost block")
    block.append(block_paragraph)
    document._element.body.insert(1, block)
    source = tmp_path / "revisions.docx"
    document.save(source)
    result = read_document(source)
    records = [record for record in get_integration(result.document).preservation if record.issue.code == "docx.revisions"]
    assert {(record.provenance.object_id, record.state.value) for record in records} == {("7", "opaque"), ("8", "lost")}
    assert "Retained inline" in result.document.sections[0].blocks[0].plain_text
    assert "Lost block" not in str(result.document.sections)
    _verify(source, result)


def test_comments_fields_controls_and_header_objects_have_retention_evidence(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Reviewed")
    document.add_comment(paragraph.runs[0], text="Own comment", author="Fixture")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    _text(field, "1")
    paragraph._p.append(field)
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    _text(content, "Control text")
    control.append(content)
    paragraph._p.append(control)
    header = document.sections[0].header.paragraphs[0]
    insertion = OxmlElement("w:ins")
    insertion.set(qn("w:id"), "header-change")
    _text(insertion, "Header insertion")
    header._p.append(insertion)
    protection = OxmlElement("w:documentProtection")
    protection.set(qn("w:edit"), "readOnly")
    document.settings._element.append(protection)
    source = tmp_path / "features.docx"
    document.save(source)
    result = read_document(source)
    records = get_integration(result.document).preservation
    assert {record.issue.code for record in records} >= {
        "docx.comments", "docx.fields", "docx.content-controls", "docx.protection",
        "docx.relationship.comments", "docx.revisions",
    }
    assert all(record.state.value == "opaque" for record in records)
    assert any(record.provenance.package_part == "/word/header1.xml" for record in records)
    restored = _verify(source, result)
    output = tmp_path / "restored.docx"
    assert write_document(restored, output).success
    assert any(comment.text == "Own comment" for comment in Document(output).comments)


def test_office_ole_case_and_part_relationship_are_preserved_without_execution(tmp_path):
    document = Document()
    package = document.part.package
    payload = b"own inert OLE fixture"
    part = Part(PackURI("/word/embeddings/own.bin"), "application/vnd.openxmlformats-officedocument.oleObject", payload, package)
    relation_id = document.part.relate_to(part, RT.OLE_OBJECT)
    paragraph = document.add_paragraph()
    paragraph._p.append(parse_xml(f'''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
      xmlns:o="urn:schemas-microsoft-com:office:office"
      xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <w:object><o:OLEObject Type="Embed" r:id="{relation_id}"/></w:object></w:r>'''))
    source = tmp_path / "ole.docx"
    document.save(source)
    result = read_document(source)
    records = get_integration(result.document).preservation
    assert any(record.issue.code == "docx.embedded-ole" and record.state.value == "opaque" for record in records)
    assert any(record.issue.code == "docx.relationship.embedded_ole" and record.state.value == "opaque" for record in records)
    restored = _verify(source, result)
    assert restored.package.parts["/word/embeddings/own.bin"].data == payload
    output = tmp_path / "restored.docx"
    assert write_document(restored, output).success
    with zipfile.ZipFile(output) as archive:
        assert archive.read("word/embeddings/own.bin") == payload


def test_foreign_namespace_does_not_claim_word_revision_support(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Retained")
    paragraph._p.append(parse_xml('<other:ins xmlns:other="urn:own:foreign"/>'))
    paragraph._p.append(parse_xml('<other:documentProtection xmlns:other="urn:own:foreign"/>'))
    source = tmp_path / "foreign.docx"
    document.save(source)
    result = read_document(source)
    assert "tracked_insertions" not in result.document.metadata["docx_features"]["features"]
    assert "document_protection_xml" not in result.document.metadata["docx_features"]
    assert not get_integration(result.document).preservation


def test_external_ole_relationship_is_not_claimed_retained(tmp_path):
    document = Document()
    relation = document.part.relate_to("https://example.invalid/inert.bin", RT.OLE_OBJECT, is_external=True)
    source = tmp_path / "external.docx"
    document.save(source)
    result = read_document(source)
    records = get_integration(result.document).preservation
    assert len(records) == 1 and records[0].state.value == "lost"
    assert records[0].provenance.object_id == relation
    assert records[0].issue.measurement["external"] is True
    _verify(source, result)
