"""Independent native declarations, public cycles, edits, deletion and atomic refusal."""

import pytest
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from opendoc_model import (
    TableIndent,
    TabStop,
    TabStops,
    document_from_json,
    document_to_json,
    get_tab_stops,
    get_table_indent,
    set_tab_stops,
    set_table_indent,
)

from opendoc_formats import ExportOptions, read_document, write_document


def _source(path):
    doc = Document()
    paragraph = doc.add_paragraph('Own Unicode\tПример')
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(396), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Own cell'
    indent = OxmlElement('w:tblInd')
    indent.set(qn('w:type'), 'dxa')
    indent.set(qn('w:w'), '600')
    table._tbl.tblPr.append(indent)
    doc.save(path)


def _snapshot(path):
    doc = Document(path)
    indentation = doc.tables[0]._tbl.tblPr.find(qn('w:tblInd'))
    tabs = doc.paragraphs[0]._p.pPr.find(qn('w:tabs'))
    return (None if indentation is None else dict(indentation.attrib),
            [] if tabs is None else [dict(tab.attrib) for tab in tabs], doc.paragraphs[0].text)


def test_public_two_cycles_edit_and_delete_do_not_restore_native_values(tmp_path):
    source = tmp_path / 'own.docx'
    _source(source)
    original, expected = source.read_bytes(), _snapshot(source)
    current = source
    for cycle in range(2):
        imported = read_document(current)
        model = document_from_json(document_to_json(imported.document))
        paragraph, table = model.sections[0].blocks
        assert get_table_indent(table) == TableIndent(30)
        assert get_tab_stops(paragraph) == TabStops((TabStop(396, alignment='right', leader='line'),))
        current = tmp_path / f'cycle-{cycle}.docx'
        report = write_document(model, current)
        assert report.success, report.to_dict()
        assert _snapshot(current) == expected
    set_table_indent(table, TableIndent(-6.05))
    set_tab_stops(paragraph, TabStops((TabStop(72.05, alignment='decimal', leader='dot'),)))
    edited = tmp_path / 'edited.docx'
    assert write_document(model, edited).success
    model = read_document(edited).document
    paragraph, table = model.sections[0].blocks
    assert get_table_indent(table).value_pt == -6.05
    assert get_tab_stops(paragraph).stops[0].position_pt == 72.05
    set_table_indent(table, None)
    set_tab_stops(paragraph, None)
    for cycle in range(2):
        deleted = tmp_path / f'deleted-{cycle}.docx'
        assert write_document(document_from_json(document_to_json(model)), deleted).success
        assert _snapshot(deleted)[:2] == (None, [])
        model = read_document(deleted).document
    assert source.read_bytes() == original


def test_style_tabs_clear_and_direct_override_preserve_native_cascade(tmp_path):
    doc = Document()
    from docx.enum.style import WD_STYLE_TYPE

    base = doc.styles.add_style('Own base', WD_STYLE_TYPE.PARAGRAPH)
    base_id = base.style_id
    base.paragraph_format.tab_stops.add_tab_stop(Pt(72), WD_TAB_ALIGNMENT.LEFT)
    derived = doc.styles.add_style('Own derived', WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = base
    paragraph = doc.add_paragraph('Own\tПример', style=derived)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(72), WD_TAB_ALIGNMENT.CLEAR)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(396), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
    source = tmp_path / 'styles.docx'
    doc.save(source)
    current = source
    for cycle in range(2):
        model = document_from_json(document_to_json(read_document(current).document))
        assert len(get_tab_stops(model.sections[0].blocks[0]).stops) == 2
        current = tmp_path / f'styles-{cycle}.docx'
        assert write_document(model, current).success
        native = Document(current)
        direct = native.paragraphs[0]._p.pPr.find(qn('w:tabs'))
        assert [(tab.get(qn('w:pos')), tab.get(qn('w:val'))) for tab in direct] == [('1440', 'clear'), ('7920', 'right')]
        assert native.styles['Own base'].paragraph_format.tab_stops[0].position.pt == 72
    model.styles[base_id].properties['docx_tab_stops'] = [{'position_pt': 200, 'alignment': 'right', 'leader': 'dot'}]
    edited = tmp_path / 'style-edited.docx'
    assert write_document(model, edited).success
    assert Document(edited).styles['Own base'].paragraph_format.tab_stops[0].position.pt == 200
    model = read_document(edited).document
    del model.styles[base_id].properties['docx_tab_stops']
    for cycle in range(2):
        deleted = tmp_path / f'style-deleted-{cycle}.docx'
        assert write_document(document_from_json(document_to_json(model)), deleted).success
        assert len(Document(deleted).styles['Own base'].paragraph_format.tab_stops) == 0
        model = read_document(deleted).document


def test_cross_format_loss_is_located_and_zero_budget_keeps_destination(tmp_path):
    source = tmp_path / 'own.docx'
    _source(source)
    model = read_document(source).document
    output = tmp_path / 'output.txt'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(quality_gate=lambda report, staged: report.lossless))
    assert not report.success and output.read_bytes() == b'previous'
    assert {issue.feature for issue in report.issues if issue.location} >= {'docx.table-indent', 'docx.tab-stops'}


def test_replacing_set_materializes_inherited_positions_with_located_loss(tmp_path):
    source = tmp_path / 'source.docx'
    _source(source)
    doc = Document(source)
    doc.styles['Normal'].paragraph_format.tab_stops.add_tab_stop(Pt(72), WD_TAB_ALIGNMENT.LEFT)
    doc.save(source)
    model = read_document(source).document
    set_tab_stops(model.sections[0].blocks[0], TabStops(inherit=False))
    output = tmp_path / 'output.docx'
    report = write_document(model, output)
    assert report.success and any(issue.feature == 'docx.tab-stops-materialized' and issue.location for issue in report.issues)
    tabs = Document(output).paragraphs[0]._p.pPr.find(qn('w:tabs'))
    assert [(tab.get(qn('w:pos')), tab.get(qn('w:val'))) for tab in tabs] == [('1440', 'clear')]
    restored = read_document(output).document
    assert get_tab_stops(restored.sections[0].blocks[0]).stops[0].kind == 'clear'
    previous = output.read_bytes()
    refused = write_document(model, output, options=ExportOptions(quality_gate=lambda report, staged: report.lossless))
    assert not refused.success and output.read_bytes() == previous


@pytest.mark.parametrize('alignment', ['start', 'end', 'num'])
def test_native_unknown_common_alignment_is_retained_without_guessing(tmp_path, alignment):
    source = tmp_path / 'source.docx'
    _source(source)
    doc = Document(source)
    doc.paragraphs[0]._p.pPr.find(qn('w:tabs'))[0].set(qn('w:val'), alignment)
    doc.save(source)
    current = source
    for cycle in range(2):
        imported = read_document(current)
        assert any(issue.code == 'docx.tab-stops-semantics' for issue in imported.issues)
        model = document_from_json(document_to_json(imported.document))
        stop = get_tab_stops(model.sections[0].blocks[0]).stops[0]
        assert stop.alignment is None and stop.extra['docx_alignment'] == alignment
        current = tmp_path / f'cycle-{cycle}.docx'
        report = write_document(model, current)
        assert report.success and any(issue.feature == 'docx.tab-stops-semantics' for issue in report.issues)
        assert Document(current).paragraphs[0]._p.pPr.find(qn('w:tabs'))[0].get(qn('w:val')) == alignment


@pytest.mark.parametrize('kind,value', [
    ('table', TableIndent(0.025)), ('table', TableIndent()), ('table', TableIndent(2**31)),
    ('tabs', TabStops((TabStop(0.025, alignment='right', leader='none'),))),
    ('tabs', TabStops((TabStop(396, alignment=None, leader='none'),))),
    ('tabs', TabStops((TabStop(396, alignment='right', leader=None),))),
    ('tabs', TabStops(tuple(TabStop(index, alignment='right', leader='none') for index in range(257)))),
])
def test_unrepresentable_declarations_refuse_publication_with_location(tmp_path, kind, value):
    source = tmp_path / 'source.docx'
    _source(source)
    model = read_document(source).document
    owner = model.sections[0].blocks[1 if kind == 'table' else 0]
    (set_table_indent if kind == 'table' else set_tab_stops)(owner, value)
    output = tmp_path / 'output.docx'
    output.write_bytes(b'previous')
    report = write_document(model, output)
    assert not report.success and output.read_bytes() == b'previous'
    assert any(issue.feature in ('docx.table-indent', 'docx.tab-stops')
               and issue.location for issue in report.issues)


@pytest.mark.parametrize('kind', ['indent', 'tabs'])
def test_unsupported_source_has_native_location_and_refuses_until_explicit_removal(tmp_path, kind):
    source = tmp_path / 'source.docx'
    _source(source)
    doc = Document(source)
    if kind == 'indent':
        node = doc.tables[0]._tbl.tblPr.find(qn('w:tblInd'))
        node.set(qn('w:type'), 'pct')
    else:
        node = doc.paragraphs[0]._p.pPr.find(qn('w:tabs'))[0]
        node.set(qn('w:val'), 'unrecognized')
    doc.save(source)
    imported = read_document(source)
    assert any(issue.code == 'docx.' + ('table-indent' if kind == 'indent' else 'tab-stops')
               and 'xpath=' in issue.location for issue in imported.issues)
    model = imported.document
    output = tmp_path / 'output.docx'
    output.write_bytes(b'previous')
    report = write_document(model, output)
    assert not report.success and output.read_bytes() == b'previous'
    owner = model.sections[0].blocks[1 if kind == 'indent' else 0]
    (set_table_indent if kind == 'indent' else set_tab_stops)(owner, None)
    assert write_document(model, output).success
