"""Formal heading roles survive native Word XML and two JSON cycles."""

from copy import deepcopy

import pytest

pytest.importorskip("docx")

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from opendoc_model import DocumentModel, Heading, Paragraph, Section, TextRun, get_heading, set_heading
from opendoc_model.document_codec import document_from_json, document_to_json

from opendoc_formats import read_document, write_document


def outline(owner, level):
    node = OxmlElement("w:outlineLvl")
    node.set(qn("w:val"), str(level))
    owner.get_or_add_pPr().append(node)


@pytest.mark.parametrize("level", [1, 2, 3, 9])
def test_explicit_heading_without_named_style(tmp_path, level):
    heading = Paragraph([TextRun("Unique heading")])
    set_heading(heading, Heading(level))
    model = DocumentModel(sections=[Section(blocks=[heading, Paragraph([TextRun("Body")])])])
    original = deepcopy(model)
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        snapshot = document_to_json(model)
        target = tmp_path / f"cycle-{cycle}.docx"
        assert write_document(model, target).success
        assert document_to_json(model) == snapshot
        native = Document(target)
        assert native.paragraphs[0]._p.xpath("./w:pPr/w:outlineLvl")[0].get(qn("w:val")) == str(level - 1)
        result = read_document(target)
        assert result.success, result.issues
        model = result.document
        assert get_heading(model.sections[0].blocks[0]) == Heading(level)
        assert get_heading(model.sections[0].blocks[1]) is None
    assert original.sections[0].blocks[0] == heading


def test_native_inheritance_direct_override_and_body_reset(tmp_path):
    native = Document()
    native.add_paragraph("Native heading", "Heading 2")
    custom = native.styles.add_style("Own localized name", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = native.styles["Heading 3"]
    native.add_paragraph("Inherited heading", custom)
    direct = native.add_paragraph("Direct override", custom)
    outline(direct._p, 4)
    body = native.add_paragraph("Body reset", custom)
    outline(body._p, 9)
    visual = native.add_paragraph("Large bold body")
    visual.runs[0].bold = True
    source = tmp_path / "source.docx"
    native.save(source)
    original = source.read_bytes()
    for cycle in range(2):
        result = read_document(source)
        assert result.success, result.issues
        model = document_from_json(document_to_json(result.document))
        assert [get_heading(p).level if get_heading(p) else None for p in model.sections[0].blocks] == [2, 3, 5, None, None]
        target = tmp_path / f"native-{cycle}.docx"
        assert write_document(model, target).success
        source = target
    assert (tmp_path / "source.docx").read_bytes() == original


def test_heading_edit_and_removal_override_original_style(tmp_path):
    native = Document()
    native.add_paragraph("Changed", "Heading 2")
    native.add_paragraph("Removed", "Heading 3")
    source = tmp_path / "input.docx"
    native.save(source)
    model = read_document(source).document
    first, second = model.sections[0].blocks
    set_heading(first, Heading(7))
    set_heading(second, None)
    for cycle in range(2):
        target = tmp_path / ("edited.docx" if cycle == 0 else "edited-1.docx")
        assert write_document(document_from_json(document_to_json(model)), target).success
        native_result = Document(target)
        assert native_result.paragraphs[0].style.base_style.name == "Heading 2"
        assert native_result.paragraphs[1].style.base_style.name == "Heading 3"
        result = read_document(target)
        assert get_heading(result.document.sections[0].blocks[0]) == Heading(7)
        assert get_heading(result.document.sections[0].blocks[1]) is None
        model = result.document


def test_builtin_heading_has_fixed_native_word_level(tmp_path):
    native = Document()
    heading = native.add_paragraph("Fixed role", "Heading 2")
    outline(heading._p, 9)
    source = tmp_path / "fixed.docx"
    native.save(source)
    assert get_heading(read_document(source).document.sections[0].blocks[0]) == Heading(2)


def test_document_defaults_and_direct_override(tmp_path):
    native = Document()
    defaults = native.styles.element.xpath("./w:docDefaults/w:pPrDefault/w:pPr")[0]
    node = OxmlElement("w:outlineLvl")
    node.set(qn("w:val"), "2")
    defaults.append(node)
    first = native.styles.add_style("Own A", WD_STYLE_TYPE.PARAGRAPH)
    second = native.styles.add_style("Own B", WD_STYLE_TYPE.PARAGRAPH)
    first.base_style = second
    native.add_paragraph("Default role", first)
    reset = native.add_paragraph("Body override", first)
    outline(reset._p, 9)
    source = tmp_path / "defaults.docx"
    native.save(source)
    model = read_document(source).document
    assert get_heading(model.sections[0].blocks[0]) == Heading(3)
    assert get_heading(model.sections[0].blocks[1]) is None


@pytest.mark.parametrize("level", ["bad", "10", "-1"])
def test_invalid_native_outline_is_rejected(tmp_path, level):
    native = Document()
    paragraph = native.add_paragraph("Invalid outline")
    outline(paragraph._p, level)
    source = tmp_path / "invalid.docx"
    native.save(source)
    with pytest.raises(ValueError, match="outline level"):
        read_document(source)


def test_invalid_extension_preserves_existing_output(tmp_path):
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Text")], properties={"docx_outline_level": True})])])
    target = tmp_path / "old.docx"
    target.write_bytes(b"old output")
    with pytest.raises(ValueError, match="docx_outline_level"):
        write_document(model, target)
    assert target.read_bytes() == b"old output"
