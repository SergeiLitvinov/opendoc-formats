"""Own DOCX fixtures, checked against source XML across two JSON roundtrips."""

import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from opendoc_model import document_from_json, document_to_json, get_integration

from opendoc_formats.errors import ConvertError
from opendoc_formats.readers.docx import read_docx_model
from opendoc_formats.writers.docx_writer import write_docx_model

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def xml(path):
    with zipfile.ZipFile(path) as package:
        return ET.fromstring(package.read("word/document.xml"))


def test_smart_tag_text_and_wrapper_survive_two_json_cycles(tmp_path):
    source = Document()
    paragraph = source.add_paragraph("Before ")
    tag = OxmlElement("w:smartTag")
    tag.set(qn("w:uri"), "urn:opendoc-formats:fixture")
    tag.set(qn("w:element"), "OwnToken")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "OwnSmartToken"
    run.append(text)
    tag.append(run)
    paragraph._p.append(tag)
    paragraph.add_run(" After")
    path = tmp_path / "source.docx"
    source.save(path)
    original = path.read_bytes()
    for cycle in range(2):
        model = read_docx_model(path)
        assert "".join(item.text for item in model.sections[0].blocks[0].content) == "Before OwnSmartToken After"
        ledger = get_integration(model)
        assert any(record.issue.code == "docx.smart-tags" and record.state.value == "opaque"
                   for record in ledger.preservation)
        model = document_from_json(document_to_json(model))
        path = tmp_path / f"cycle-{cycle}.docx"
        write_docx_model(model, path)
        root = xml(path)
        assert "".join(node.text or "" for node in root.iter(W + "t")) == "Before OwnSmartToken After"
        assert root.find(".//" + W + "smartTag").get(W + "uri") == "urn:opendoc-formats:fixture"
    assert (tmp_path / "source.docx").read_bytes() == original


def test_percentage_table_widths_survive_two_json_cycles(tmp_path):
    source = Document()
    section = source.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.left_margin = section.right_margin = Cm(2)
    normal = source.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    table = source.add_table(rows=20, cols=3)
    table.autofit = False
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), "9600")
    for column, width in zip(table._tbl.tblGrid.gridCol_lst, (960, 7680, 960), strict=True):
        column.set(qn("w:w"), str(width))
    for index, row in enumerate(table.rows, 1):
        for cell, width, text in zip(
            row.cells, (500, 4000, 500),
            (str(index), "Own long middle column verifies table geometry through two JSON cycles. " * 2, "8"), strict=True,
        ):
            cell.text = text
            node = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            node.set(qn("w:type"), "pct")
            node.set(qn("w:w"), str(width))
    path = tmp_path / "source.docx"
    source.save(path)
    for cycle in range(2):
        model = document_from_json(document_to_json(read_docx_model(path)))
        path = tmp_path / f"cycle-{cycle}.docx"
        write_docx_model(model, path)
        root = xml(path)
        width = root.find(".//" + W + "tblPr/" + W + "tblW")
        assert width.attrib == {W + "type": "dxa", W + "w": "9600"}
        assert [node.get(W + "w") for node in root.findall(".//" + W + "gridCol")] == ["960", "7680", "960"]
        widths = root.findall(".//" + W + "tcW")
        assert [(node.get(W + "type"), node.get(W + "w")) for node in widths] == [
            ("pct", width) for _ in range(20) for width in ("500", "4000", "500")
        ]


@pytest.mark.parametrize("kind,value", [("pct", "5000"), ("pct", "100%"), ("auto", "0"), ("nil", "0")])
def test_preferred_width_units_and_explicit_absolute_edit(tmp_path, kind, value):
    source = Document()
    table = source.add_table(rows=1, cols=1)
    for node in (table._tbl.tblPr.find(qn("w:tblW")), table.cell(0, 0)._tc.tcPr.tcW):
        node.set(qn("w:type"), kind)
        node.set(qn("w:w"), value)
    path = tmp_path / "source.docx"
    source.save(path)
    model = document_from_json(document_to_json(read_docx_model(path)))
    target = tmp_path / "result.docx"
    write_docx_model(model, target)
    root = xml(target)
    for tag in ("tblW", "tcW"):
        node = root.find(".//" + W + tag)
        assert (node.get(W + "type"), node.get(W + "w")) == (kind, value)
    model.sections[0].blocks[0].rows[0].cells[0].properties.set_typed("width_twips", 2400)
    write_docx_model(model, target)
    node = xml(target).find(".//" + W + "tcW")
    assert (node.get(W + "type"), node.get(W + "w")) == ("dxa", "2400")


def test_invalid_preferred_width_preserves_previous_output(tmp_path):
    from opendoc_formats.api import read_document
    from opendoc_formats.export import write_document

    source = Document()
    source.add_table(rows=1, cols=1)
    path = tmp_path / "source.docx"
    source.save(path)
    model = read_document(path).document
    model.sections[0].blocks[0].properties["docx_preferred_width"] = {"type": "pct", "value": "bad"}
    target = tmp_path / "result.docx"
    target.write_bytes(b"previous")
    with pytest.raises(ConvertError, match="preferred width"):
        write_document(model, target)
    assert target.read_bytes() == b"previous"
