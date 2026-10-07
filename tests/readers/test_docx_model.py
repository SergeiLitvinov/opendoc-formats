"""Интеграционные тесты DOCX → DocumentModel."""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from opendoc_model.color import ColorValue
from opendoc_model.document_model import Formula, Image, Paragraph, ResourceKind, Table
from PIL import Image as PillowImage

from opendoc_formats.readers.docx import read_docx_model


def test_read_docx_model_preserves_order_styles_tables_and_metadata(tmp_path):
    path = tmp_path / "structured.docx"
    source = Document()
    source.core_properties.title = "Structured document"
    paragraph = source.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Formatted")
    run.bold = True
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(14)
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    source.add_paragraph("After table")
    source.save(path)

    model = read_docx_model(path)

    assert model.sections[0].provenance is not None
    assert model.sections[0].provenance.source_path == str(path)
    assert model.sections[0].provenance.package_part == "/word/document.xml"
    assert model.sections[0].blocks[0].provenance is not None

    assert model.metadata["title"] == "Structured document"
    assert model.source_format == "docx"
    assert [type(block) for block in model.sections[0].blocks] == [Paragraph, Table, Paragraph]
    imported = model.sections[0].blocks[0]
    assert isinstance(imported, Paragraph)
    assert imported.alignment == "center"
    assert imported.content[0].style.bold is True
    assert imported.content[0].style.italic is True
    assert imported.content[0].style.font_family == "Arial"
    assert imported.content[0].style.font_size.pt == 14


def test_read_docx_model_resolves_theme_color_tint_on_direct_run(tmp_path):
    from docx.oxml.ns import qn

    path = tmp_path / "theme-color.docx"
    source = Document()
    run = source.add_paragraph().add_run("Theme color")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "FF0000")
    color.set(qn("w:themeColor"), "accent1")
    color.set(qn("w:themeTint"), "80")
    run._r.get_or_add_rPr().append(color)
    source.save(path)

    model = read_docx_model(path)
    imported = model.sections[0].blocks[0].content[0]

    assert isinstance(imported.style.color, ColorValue)
    assert imported.style.color.to_hex() == "#A7C0DE"
    assert imported.style.properties["color_theme"] == "accent1"
    assert imported.style.properties["color_modifiers"]["tint"] == 128 / 255


def test_read_docx_model_extracts_image_formula_header_and_page_geometry(tmp_path):
    path = tmp_path / "media.docx"
    image_path = tmp_path / "pixel.png"
    PillowImage.new("RGB", (8, 6), "red").save(image_path)

    source = Document()
    source.sections[0].header.paragraphs[0].text = "Header text"
    source.sections[0].page_width = Inches(9)
    paragraph = source.add_paragraph("Before ")
    picture = paragraph.add_run()
    picture.add_picture(str(image_path), width=Inches(1))
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = "x+1"
    math_run.append(math_text)
    math.append(math_run)
    paragraph._p.append(math)
    source.save(path)

    model = read_docx_model(path)

    imported = model.sections[0].blocks[0]
    assert isinstance(imported, Paragraph)
    assert any(isinstance(item, Image) for item in imported.content)
    assert any(isinstance(item, Formula) and item.fallback_text == "x+1" for item in imported.content)
    image_resources = [resource for resource in model.resources.values() if resource.kind is ResourceKind.RASTER_IMAGE]
    assert len(image_resources) == 1
    resource = image_resources[0]
    assert resource.kind is ResourceKind.RASTER_IMAGE
    assert resource.media_type == "image/png"
    assert resource.data
    assert model.sections[0].headers[0].plain_text == "Header text"
    assert model.sections[0].page.width.pt == 648
    assert model.validate() == []


def test_read_docx_model_normalizes_drawingml_picture_outline_color(tmp_path):
    path = tmp_path / "picture-outline.docx"
    image_path = tmp_path / "outline.png"
    PillowImage.new("RGB", (8, 6), "red").save(image_path)
    source = Document()
    picture = source.add_paragraph().add_run().add_picture(str(image_path), width=Inches(1))
    shape_properties = picture._inline.xpath(".//pic:spPr")[0]
    line = OxmlElement("a:ln")
    solid = OxmlElement("a:solidFill")
    scheme = OxmlElement("a:schemeClr")
    scheme.set("val", "accent2")
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "50000")
    scheme.append(alpha)
    solid.append(scheme)
    line.append(solid)
    shape_properties.append(line)
    source.save(path)

    model = read_docx_model(path)
    imported = next(item for item in model.sections[0].blocks[0].content if isinstance(item, Image))
    color = ColorValue.from_dict(imported.properties["stroke_color"])

    assert color.alpha == 0.5
    assert imported.properties["stroke_color_source"]["color_theme"] == "accent2"


def test_read_docx_model_catalogs_vml_fill_and_stroke_colors(tmp_path):
    from lxml import etree

    path = tmp_path / "vml-colors.docx"
    source = Document()
    run = source.add_paragraph().add_run()
    pict = OxmlElement("w:pict")
    vml_namespace = "urn:schemas-microsoft-com:vml"
    shape = etree.Element(f"{{{vml_namespace}}}shape", nsmap={"v": vml_namespace})
    shape.set("fillcolor", "#336699")
    shape.set("strokecolor", "#FF0000")
    fill = etree.Element(f"{{{vml_namespace}}}fill")
    fill.set("color", "#00FF00")
    fill.set("opacity", "50%")
    shape.append(fill)
    pict.append(shape)
    run._r.append(pict)
    source.save(path)

    model = read_docx_model(path)
    imported = model.sections[0].blocks[0].content[0]
    colors = imported.properties["vml_colors"]

    assert ColorValue.from_dict(colors[0]["color"]).to_hex() == "#336699"
    assert ColorValue.from_dict(colors[1]["color"]).to_hex() == "#FF0000"
    assert ColorValue.from_dict(colors[2]["color"]).to_hex(include_alpha=True) == "#00FF0080"
    assert imported.properties["vml_xml"]


def test_read_docx_model_preserves_inline_revisions_and_content_control(tmp_path):
    from docx.oxml.ns import qn

    path = tmp_path / "revisions-sdt.docx"
    source = Document()
    paragraph = source.add_paragraph("Before ")
    insertion = OxmlElement("w:ins")
    insertion.set(qn("w:id"), "7")
    insertion.set(qn("w:author"), "Editor")
    inserted_run = OxmlElement("w:r")
    inserted_text = OxmlElement("w:t")
    inserted_text.text = "inserted"
    inserted_run.append(inserted_text)
    insertion.append(inserted_run)
    paragraph._p.append(insertion)
    deletion = OxmlElement("w:del")
    deletion.set(qn("w:id"), "8")
    deleted_run = OxmlElement("w:r")
    deleted_text = OxmlElement("w:delText")
    deleted_text.text = "deleted"
    deleted_run.append(deleted_text)
    deletion.append(deleted_run)
    paragraph._p.append(deletion)
    control = OxmlElement("w:sdt")
    control_properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "student-name")
    control_properties.append(tag)
    content = OxmlElement("w:sdtContent")
    control_run = OxmlElement("w:r")
    control_text = OxmlElement("w:t")
    control_text.text = "Alice"
    control_run.append(control_text)
    content.append(control_run)
    control.extend((control_properties, content))
    paragraph._p.append(control)
    source.save(path)

    model = read_docx_model(path)
    runs = model.sections[0].blocks[0].content

    assert [run.text for run in runs] == ["Before ", "inserted", "deleted", "Alice"]
    assert runs[1].properties["revision"] == {"type": "ins", "id": "7", "author": "Editor", "date": None}
    assert runs[2].properties["revision"]["type"] == "del"
    assert runs[3].properties["content_control"]["tag"] == "student-name"
    features = model.metadata["docx_features"]["features"]
    assert features["tracked_insertions"]["count"] == 1
    assert features["tracked_deletions"]["count"] == 1
    assert features["content_controls"]["count"] == 1


def test_read_docx_model_preserves_comment_markers_and_text_box_ooxml(tmp_path):
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    path = tmp_path / "comments-textbox.docx"
    source = Document()
    paragraph = source.add_paragraph()
    range_start = OxmlElement("w:commentRangeStart")
    range_start.set(qn("w:id"), "4")
    paragraph._p.append(range_start)
    paragraph.add_run("Reviewed")
    range_end = OxmlElement("w:commentRangeEnd")
    range_end.set(qn("w:id"), "4")
    paragraph._p.append(range_end)
    reference_run = OxmlElement("w:r")
    reference = OxmlElement("w:commentReference")
    reference.set(qn("w:id"), "4")
    reference_run.append(reference)
    paragraph._p.append(reference_run)
    paragraph._p.append(
        parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape><v:textbox>'
            "<w:txbxContent><w:p><w:r><w:t>Box text</w:t></w:r></w:p></w:txbxContent>"
            "</v:textbox></v:shape></w:pict></w:r>"
        )
    )
    source.save(path)

    model = read_docx_model(path)
    runs = model.sections[0].blocks[0].content

    assert [run.properties.get("docx_raw_inline_type") for run in runs] == [
        "commentRangeStart",
        None,
        "commentRangeEnd",
        "r",
        "r",
    ]
    assert runs[3].properties["docx_advanced_types"] == ["comments"]
    assert runs[4].text == "Box text"
    assert runs[4].properties["docx_advanced_types"] == ["text_boxes"]
    features = model.metadata["docx_features"]["features"]
    assert features["comments"]["count"] == 3
    assert features["text_boxes"]["preservation"] == "native-ooxml"


def test_read_docx_model_records_document_protection_xml(tmp_path):
    from docx.oxml.ns import qn

    path = tmp_path / "protected.docx"
    source = Document()
    protection = OxmlElement("w:documentProtection")
    protection.set(qn("w:edit"), "readOnly")
    protection.set(qn("w:enforcement"), "1")
    source.settings._element.append(protection)
    source.save(path)

    model = read_docx_model(path)

    features = model.metadata["docx_features"]
    assert features["features"]["protected_fields"]["count"] == 1
    assert "documentProtection" in features["document_protection_xml"]


def test_read_docx_model_preserves_block_content_control(tmp_path):
    path = tmp_path / "block-sdt.docx"
    source = Document()
    source.add_paragraph("Before")
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    from docx.oxml.ns import qn

    tag.set(qn("w:val"), "lesson")
    properties.append(tag)
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Controlled block"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    control.extend((properties, content))
    source._element.body.insert(-1, control)
    source.save(path)

    model = read_docx_model(path)

    block = model.sections[0].blocks[1]
    assert block.plain_text == "Controlled block"
    assert block.properties["docx_raw_block_type"] == "sdt"
    assert block.properties["content_control"]["tag"] == "lesson"


def test_read_docx_model_preserves_horizontal_cell_merge(tmp_path):
    path = tmp_path / "merged.docx"
    source = Document()
    table = source.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "merged"
    table.cell(0, 2).text = "single"
    source.save(path)

    model = read_docx_model(path)

    imported = model.sections[0].blocks[0]
    assert isinstance(imported, Table)
    assert len(imported.rows[0].cells) == 2
    assert imported.rows[0].cells[0].column_span == 2
    assert imported.rows[0].cells[0].blocks[0].plain_text == "merged"


def test_read_docx_model_resolves_style_inheritance_for_paragraphs_and_runs(tmp_path):
    path = tmp_path / "inherited-styles.docx"
    source = Document()
    base = source.styles.add_style("Scientific Base", WD_STYLE_TYPE.PARAGRAPH)
    base.font.name = "Arial"
    base.font.size = Pt(12)
    base.font.bold = True
    base.paragraph_format.space_after = Pt(9)
    child = source.styles.add_style("Scientific Result", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    child.font.color.rgb = source.styles["Heading 1"].font.color.rgb
    emphasis = source.styles.add_style("Scientific Emphasis", WD_STYLE_TYPE.CHARACTER)
    emphasis.font.italic = True
    paragraph = source.add_paragraph(style=child)
    run = paragraph.add_run("Inherited formatting")
    run.style = emphasis
    source.save(path)

    model = read_docx_model(path)

    result_style = next(style for style in model.styles.values() if style.properties["style_name"] == "Scientific Result")
    assert result_style.font_family == "Arial"
    assert result_style.font_size.pt == 12
    assert result_style.bold is True
    assert result_style.properties["space_after_pt"] == 9
    imported = model.sections[0].blocks[0]
    assert isinstance(imported, Paragraph)
    assert imported.properties["space_after_pt"] == 9
    assert imported.content[0].style.font_family == "Arial"
    assert imported.content[0].style.font_size.pt == 12
    assert imported.content[0].style.bold is True
    assert imported.content[0].style.italic is True
    assert imported.content[0].style.properties["direct_fields"] == []


def test_read_docx_model_preserves_continuous_section_and_header_linkage(tmp_path):
    path = tmp_path / "sections.docx"
    source = Document()
    source.sections[0].header.paragraphs[0].text = "Shared header"
    source.add_paragraph("First section")
    second = source.add_section(WD_SECTION.CONTINUOUS)
    second.header.is_linked_to_previous = True
    second.footer.is_linked_to_previous = True
    second.header_distance = Inches(0.35)
    second.footer_distance = Inches(0.4)
    source.add_paragraph("Second section")
    source.save(path)

    model = read_docx_model(path)

    assert len(model.sections) == 2
    assert model.sections[1].properties["start_type"] == "CONTINUOUS"
    assert model.sections[1].properties["header_linked_to_previous"] is True
    assert model.sections[1].properties["footer_linked_to_previous"] is True
    assert model.sections[1].properties["header_distance_pt"] == 25.2
    assert model.sections[1].properties["footer_distance_pt"] == 28.8


def test_read_docx_model_preserves_first_and_even_running_content(tmp_path):
    path = tmp_path / "alternate-running-content.docx"
    source = Document()
    section = source.sections[0]
    source.settings.odd_and_even_pages_header_footer = True
    section.different_first_page_header_footer = True
    section.header.paragraphs[0].text = "Odd header"
    section.footer.paragraphs[0].text = "Odd footer"
    section.first_page_header.paragraphs[0].text = "First header"
    section.first_page_footer.paragraphs[0].text = "First footer"
    section.even_page_header.paragraphs[0].text = "Even header"
    section.even_page_footer.paragraphs[0].text = "Even footer"
    source.add_paragraph("Body")
    source.save(path)

    model = read_docx_model(path)
    imported = model.sections[0]

    assert imported.properties["different_first_page_header_footer"] is True
    assert imported.properties["odd_and_even_pages_header_footer"] is True
    assert imported.headers[0].plain_text == "Odd header"
    assert imported.footers[0].plain_text == "Odd footer"
    assert imported.first_page_headers[0].plain_text == "First header"
    assert imported.first_page_footers[0].plain_text == "First footer"
    assert imported.even_page_headers[0].plain_text == "Even header"
    assert imported.even_page_footers[0].plain_text == "Even footer"
