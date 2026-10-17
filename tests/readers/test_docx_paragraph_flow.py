"""Native paragraph flow declarations remain editable through the style graph."""

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from opendoc_model import TextStyle, document_from_json, document_to_json, resolve_style

from opendoc_formats import ExportOptions, read_document, write_document

FIELDS = {
    'space_before_pt': ('space_before', 14, 22.05),
    'space_after_pt': ('space_after', 8, 32.05),
    'keep_together': ('keep_together', False, True),
    'keep_with_next': ('keep_with_next', False, True),
    'page_break_before': ('page_break_before', False, True),
    'widow_control': ('widow_control', True, False),
}


def own_source(path):
    native = Document()
    normal = native.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(11)
    base = native.styles.add_style('Own Flow Base', WD_STYLE_TYPE.PARAGRAPH)
    child = native.styles.add_style('Own Flow Child', WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    for key, (name, value, _) in FIELDS.items():
        setattr(base.paragraph_format, name, Pt(value) if key.endswith('_pt') else value)
    for index in range(3):
        paragraph = native.add_paragraph(f'Own paragraph {index + 1}. ' + 'Editable words with live style inheritance. ' * 8,
                                         style=child)
        if index:
            for key, (name, _, _) in FIELDS.items():
                value = (20 if index == 1 else 0) if key.endswith('_pt') else False
                setattr(paragraph.paragraph_format, name, Pt(value) if key.endswith('_pt') else value)
    native.save(path)


@pytest.fixture
def source(tmp_path):
    path = tmp_path / 'source.docx'
    own_source(path)
    return path


def cycle(model, path):
    before = document_to_json(model)
    report = write_document(document_from_json(before), path)
    assert report.success, report.issues
    assert document_to_json(model) == before
    return read_document(path).document


def test_direct_zero_false_and_inheritance_in_two_cycles(source, tmp_path):
    original = source.read_bytes()
    model = read_document(source).document
    for index in range(3):
        inherited, direct, zero = model.sections[0].blocks
        for key, (_, value, _) in FIELDS.items():
            assert key not in inherited.properties
            assert key not in model.styles['OwnFlowChild'].properties
            assert resolve_style(model, 'OwnFlowChild').properties[key] == value
            assert direct.properties[key] == (20 if key.endswith('_pt') else False)
            assert zero.properties[key] == (0 if key.endswith('_pt') else False)
        if index < 2:
            model = cycle(model, tmp_path / f'cycle-{index}.docx')
    assert source.read_bytes() == original


@pytest.mark.parametrize('key', FIELDS)
def test_style_edit_direct_deletion_and_false_override(source, tmp_path, key):
    model = read_document(source).document
    name, _, value = FIELDS[key]
    model.styles['OwnFlowBase'].properties[key] = value
    del model.sections[0].blocks[1].properties[key]
    for index in range(2):
        path = tmp_path / f'edit-{index}.docx'
        model = cycle(model, path)
        native = Document(path)
        actual = getattr(native.styles['Own Flow Base'].paragraph_format, name)
        assert (actual.pt if key.endswith('_pt') else actual) == pytest.approx(value)
        assert getattr(native.paragraphs[0].paragraph_format, name) is None
        assert getattr(native.paragraphs[1].paragraph_format, name) is None
        assert getattr(native.paragraphs[2].paragraph_format, name) == (Pt(0) if key.endswith('_pt') else False)
        assert resolve_style(model, 'OwnFlowChild').properties[key] == pytest.approx(value)


def test_style_deletion_defaults_edit_and_removal(source, tmp_path):
    model = read_document(source).document
    for key in FIELDS:
        del model.styles['OwnFlowBase'].properties[key]
        model.styles['__doc_defaults__'].properties[key] = 6 if key.endswith('_pt') else False
    for index in range(2):
        model = cycle(model, tmp_path / f'defaults-{index}.docx')
        for key in FIELDS:
            assert resolve_style(model, 'OwnFlowChild').properties[key] == (6 if key.endswith('_pt') else False)
            assert key not in model.styles['OwnFlowBase'].properties
    for key in FIELDS:
        del model.styles['__doc_defaults__'].properties[key]
    for index in range(2):
        model = cycle(model, tmp_path / f'deleted-{index}.docx')
        assert all(resolve_style(model, 'OwnFlowChild').properties.get(key) is None for key in FIELDS)


@pytest.mark.parametrize('key,value', [('space_before_pt', -1), ('space_after_pt', .025),
                                     ('space_before_pt', True), ('keep_together', 'false'),
                                     ('page_break_before', 1)])
def test_invalid_flow_refuses_before_publication(source, tmp_path, key, value):
    model = read_document(source).document
    model.sections[0].blocks[0].properties[key] = value
    target = tmp_path / 'previous.docx'
    target.write_bytes(b'previous output')
    report = write_document(model, target)
    assert not report.success and target.read_bytes() == b'previous output'
    assert any(issue.feature == 'docx.paragraph-flow' and key in issue.location for issue in report.issues)


def test_unknown_spacing_units_require_explicit_resolution(tmp_path):
    native = Document()
    paragraph = native.add_paragraph('Own unknown spacing')
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:beforeLines'), '100')
    paragraph._p.get_or_add_pPr().append(spacing)
    source = tmp_path / 'unknown.docx'
    native.save(source)
    result = read_document(source)
    assert any(issue.code == 'docx.paragraph-flow' for issue in result.issues)
    model = result.document
    target = tmp_path / 'resolved.docx'
    assert not write_document(model, target).success
    model.sections[0].blocks[0].properties['space_before_pt'] = 9.05
    assert write_document(model, target).success
    assert Document(target).paragraphs[0].paragraph_format.space_before.pt == pytest.approx(9.05)


def test_private_preview_resolves_spacing_without_changing_public_model(source):
    from opendoc_formats.writers.color_inheritance import materialize_docx_colors

    model = read_document(source).document
    model.styles['OwnFlowBase'].properties['space_before_pt'] = 22.05
    before = document_to_json(model)
    prepared = materialize_docx_colors(model)
    assert prepared.sections[0].blocks[0].properties['space_before_pt'] == pytest.approx(22.05)
    assert prepared.sections[0].blocks[2].properties['space_before_pt'] == 0
    assert document_to_json(model) == before


@pytest.mark.parametrize('owner', ['character', 'table', 'inline'])
def test_wrong_flow_owner_refuses(source, tmp_path, owner):
    model = read_document(source).document
    style = TextStyle(properties={'keep_together': True, 'style_type': owner})
    if owner == 'inline':
        model.sections[0].blocks[0].content[0].style = style
    else:
        model.styles['Own invalid owner'] = style
    report = write_document(model, tmp_path / 'invalid.docx')
    assert not report.success
    assert any(issue.feature == 'docx.paragraph-flow' for issue in report.issues)


@pytest.mark.parametrize('owner', ['paragraph', 'style'])
def test_unknown_flags_and_cross_format_loss_are_not_hidden(tmp_path, owner):
    native = Document()
    paragraph = native.add_paragraph('Own unknown pagination')
    properties = (paragraph._p.get_or_add_pPr() if owner == 'paragraph'
                  else native.styles['Normal'].element.get_or_add_pPr())
    node = OxmlElement('w:keepNext')
    node.set(qn('w:val'), 'unknown')
    properties.append(node)
    source = tmp_path / 'unknown.docx'
    native.save(source)
    model = read_document(source).document
    report = write_document(model, tmp_path / 'unknown.html', options=ExportOptions(
        quality_gate=lambda report, _: not any(issue.severity.value == 'loss' for issue in report.issues)))
    assert not report.success and not (tmp_path / 'unknown.html').exists()
    assert any(issue.feature == 'docx.paragraph-flow' for issue in report.issues)
    destination = tmp_path / 'resolved.docx'
    assert not write_document(model, destination).success
    target = model.sections[0].blocks[0].properties if owner == 'paragraph' else model.styles['Normal'].properties
    target['keep_with_next'] = None
    assert write_document(model, destination).success


def test_other_spacing_axes_and_line_rule_survive(source, tmp_path):
    model = read_document(source).document
    model.styles['OwnFlowBase'].properties['space_before_pt'] = 22.05
    model = cycle(model, tmp_path / 'one-axis.docx')
    assert model.styles['OwnFlowBase'].properties['space_after_pt'] == 8
    defaults = Document(tmp_path / 'one-axis.docx').styles.element.xpath('./w:docDefaults/w:pPrDefault/w:pPr/w:spacing')[0]
    assert defaults.get(qn('w:line')) == '276' and defaults.get(qn('w:lineRule')) == 'auto'


def test_run_character_spacing_is_not_paragraph_spacing(tmp_path):
    native = Document()
    paragraph = native.add_paragraph('Own character tracking')
    node = OxmlElement('w:spacing')
    node.set(qn('w:val'), '30')
    paragraph.runs[0]._r.get_or_add_rPr().append(node)
    paragraph.paragraph_format.space_before = Pt(8)
    source = tmp_path / 'tracking.docx'
    native.save(source)
    result = read_document(source)
    assert result.success
    assert result.document.sections[0].blocks[0].properties['space_before_pt'] == 8
    assert not any(issue.code == 'docx.paragraph-flow' for issue in result.issues)
