"""Own native indents retain declaration ownership across edits and inheritance."""

import math

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from opendoc_model import (
    DocumentModel,
    Paragraph,
    Section,
    TextRun,
    TextStyle,
    document_from_json,
    document_to_json,
    resolve_style,
)

from opendoc_formats import ExportOptions, read_document, write_document
from opendoc_formats.ooxml.paragraph_indent import ATTRIBUTES


@pytest.fixture
def source(tmp_path):
    native = Document()
    defaults = native.styles.element.xpath('./w:docDefaults/w:pPrDefault/w:pPr')[0]
    node = OxmlElement('w:ind')
    for name, value in [('left', '120'), ('right', '60'), ('firstLine', '0')]:
        node.set(qn('w:' + name), value)
    defaults.append(node)
    base = native.styles.add_style('Own base', WD_STYLE_TYPE.PARAGRAPH)
    for name, value in [('left_indent', 30), ('right_indent', 18), ('first_line_indent', 12)]:
        setattr(base.paragraph_format, name, Pt(value))
    child = native.styles.add_style('Own child', WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    for values in [None, (48, 24, -12), (0, 0, 0)]:
        paragraph = native.add_paragraph('Own words over multiple lines\nOwn second line', style=child)
        if values is not None:
            for name, value in zip(['left_indent', 'right_indent', 'first_line_indent'], values):
                setattr(paragraph.paragraph_format, name, Pt(value))
    path = tmp_path / 'own.docx'
    native.save(path)
    return path


def cycle(model, path):
    before = document_to_json(model)
    report = write_document(document_from_json(before), path)
    assert report.success, report.issues
    assert document_to_json(model) == before
    return read_document(path).document


def test_direct_zero_and_inherited_indents_in_two_native_json_cycles(source, tmp_path):
    before = source.read_bytes()
    result = read_document(source)
    model = result.document
    for index in range(3):
        first, direct, zero = model.sections[0].blocks
        assert all(key not in first.properties and key not in model.styles['Ownchild'].properties for key in ATTRIBUTES)
        assert [direct.properties[key] for key in ATTRIBUTES] == [48, 24, -12]
        assert [zero.properties[key] for key in ATTRIBUTES] == [0, 0, 0]
        assert [resolve_style(model, 'Ownchild').properties[key] for key in ATTRIBUTES] == [30, 18, 12]
        if index < 2:
            path = tmp_path / f'own-{index}.docx'
            model = cycle(model, path)
            assert not Document(path).paragraphs[0]._p.xpath('./w:pPr/w:ind/@w:left')
    assert source.read_bytes() == before


@pytest.mark.parametrize('key,value,native_name', [
    ('left_indent_pt', 50, 'left_indent'), ('right_indent_pt', 24.05, 'right_indent'),
    ('first_line_indent_pt', -6.05, 'first_line_indent'),
])
def test_style_edit_and_direct_deletion_stay_live(source, tmp_path, key, value, native_name):
    model = read_document(source).document
    model.styles['Ownbase'].properties[key] = value
    del model.sections[0].blocks[1].properties[key]
    for index in range(2):
        path = tmp_path / f'edit-{index}.docx'
        model = cycle(model, path)
        native = Document(path)
        assert getattr(native.styles['Own base'].paragraph_format, native_name).pt == pytest.approx(value)
        assert getattr(native.paragraphs[0].paragraph_format, native_name) is None
        assert getattr(native.paragraphs[1].paragraph_format, native_name) is None
        assert getattr(native.paragraphs[2].paragraph_format, native_name).pt == 0
        assert resolve_style(model, 'Ownchild').properties[key] == pytest.approx(value)


def test_style_deletion_and_defaults_edit_and_removal(source, tmp_path):
    model = read_document(source).document
    for key in ATTRIBUTES:
        del model.styles['Ownbase'].properties[key]
    for index in range(2):
        path = tmp_path / f'default-{index}.docx'
        model = cycle(model, path)
        assert [resolve_style(model, 'Ownchild').properties[key] for key in ATTRIBUTES] == [6, 3, 0]
        assert all(key not in model.styles['Ownbase'].properties for key in ATTRIBUTES)
    model.styles['__doc_defaults__'].properties['left_indent_pt'] = 9.05
    model = cycle(model, tmp_path / 'default-edit.docx')
    assert resolve_style(model, 'Ownchild').properties['left_indent_pt'] == pytest.approx(9.05)
    del model.styles['__doc_defaults__'].properties['left_indent_pt']
    for index in range(2):
        model = cycle(model, tmp_path / f'default-remove-{index}.docx')
        assert resolve_style(model, 'Ownchild').properties.get('left_indent_pt') is None


@pytest.mark.parametrize('value', [True, '12', float('nan'), float('inf'), 0.025, 2**40])
def test_unrepresentable_indent_refuses_atomically(tmp_path, value):
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun('Own')], properties={'left_indent_pt': value})])])
    path = tmp_path / 'previous.docx'
    path.write_bytes(b'previous destination')
    report = write_document(model, path)
    assert not report.success and path.read_bytes() == b'previous destination'
    if isinstance(value, float) and not math.isfinite(value):
        assert any('left_indent_pt' in issue.message for issue in report.issues)
        return
    issue = next(issue for issue in report.issues if issue.feature == 'docx.paragraph-indent')
    assert issue.location.endswith('.left_indent_pt')


def test_unknown_native_units_block_inheritance_until_resolved(source, tmp_path):
    native = Document(source)
    node = OxmlElement('w:ind')
    node.set(qn('w:leftChars'), '100')
    native.styles['Own child'].element.get_or_add_pPr().append(node)
    native.save(source)
    before = source.read_bytes()
    result = read_document(source)
    assert any(issue.code == 'docx.paragraph-indent' for issue in result.issues)
    model = result.document
    from opendoc_formats.writers.color_inheritance import materialize_docx_colors
    view = materialize_docx_colors(model)
    assert 'left_indent_pt' not in view.sections[0].blocks[0].properties
    path = tmp_path / 'unknown.docx'
    path.write_bytes(b'previous destination')
    assert not write_document(model, path).success and path.read_bytes() == b'previous destination'
    preview = tmp_path / 'unknown.html'
    preview.write_bytes(b'previous preview')
    report = write_document(model, preview, options=ExportOptions(
        quality_gate=lambda report, _: not any(issue.severity.value == 'loss' for issue in report.issues),
    ))
    assert not report.success and preview.read_bytes() == b'previous preview'
    assert any(i.feature == 'docx.paragraph-indent' and i.severity.value == 'loss' for i in report.issues)
    model.styles['Ownchild'].properties['left_indent_pt'] = 22.05
    for index in range(2):
        model = cycle(model, tmp_path / f'resolved-{index}.docx')
        assert resolve_style(model, 'Ownchild').properties['left_indent_pt'] == 22.05
    model.styles['Ownchild'].properties['left_indent_pt'] = None
    model = cycle(model, tmp_path / 'remove.docx')
    assert resolve_style(model, 'Ownchild').properties['left_indent_pt'] == 30
    assert source.read_bytes() == before


def test_private_preview_resolves_changes_without_mutation(source):
    model = read_document(source).document
    model.styles['Ownbase'].properties['left_indent_pt'] = 50
    del model.sections[0].blocks[1].properties['left_indent_pt']
    before = document_to_json(model)
    from opendoc_formats.writers.color_inheritance import materialize_docx_colors
    view = materialize_docx_colors(model)
    assert [p.properties['left_indent_pt'] for p in view.sections[0].blocks] == [50, 50, 0]
    assert document_to_json(model) == before


@pytest.mark.parametrize('kind', ['character', 'table'])
def test_unsupported_style_owners_refuse(tmp_path, kind):
    model = DocumentModel(styles={'Own': TextStyle(properties={'style_name': 'Own', 'style_type': kind, 'left_indent_pt': 30})},
                          sections=[Section(blocks=[Paragraph([TextRun('Own')])])])
    path = tmp_path / 'unsupported.docx'
    path.write_bytes(b'previous')
    report = write_document(model, path)
    assert not report.success and path.read_bytes() == b'previous'
    assert any(i.feature == 'docx.paragraph-indent' and i.location == 'styles[Own]' for i in report.issues)


def test_unsupported_inline_owner_refuses(tmp_path):
    model = DocumentModel(sections=[Section(blocks=[Paragraph([
        TextRun('Own', TextStyle(properties={'left_indent_pt': 30})),
    ])])])
    path = tmp_path / 'inline.docx'
    path.write_bytes(b'previous')
    report = write_document(model, path)
    assert not report.success and path.read_bytes() == b'previous'
    assert any(i.feature == 'docx.paragraph-indent' and i.location.endswith('.content[0].style') for i in report.issues)


@pytest.mark.parametrize('extension', ['txt', 'tex'])
def test_formats_without_indent_report_loss(tmp_path, extension):
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun('Own')], properties={'left_indent_pt': 30})])])
    path = tmp_path / ('own.' + extension)
    report = write_document(model, path)
    assert report.success
    assert any(i.feature == 'docx.paragraph-indent' and i.severity.value == 'loss' for i in report.issues)
