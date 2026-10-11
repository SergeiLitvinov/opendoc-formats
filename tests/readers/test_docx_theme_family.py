"""Own Latin theme precedence against literal fallbacks at each cascade level."""

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Pt

from opendoc_formats import read_document, write_document


def _source(path):
    native = Document()
    style = native.styles['Heading 1']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.bold = True
    native.add_paragraph('Own theme heading', style='Heading 1')
    native.save(path)


def test_theme_beats_literal_fallback_in_two_public_cycles(tmp_path, monkeypatch):
    from opendoc_formats.writers import docx_writer

    # Font discovery is a separate policy; this fixture exercises the native declarations.
    monkeypatch.setattr(docx_writer, 'prepare_fonts', lambda document, report: document)
    source = tmp_path / 'own.docx'
    _source(source)
    original = source.read_bytes()
    path = source
    for cycle in range(2):
        imported = read_document(path)
        run = imported.document.sections[0].blocks[0].content[0]
        assert run.style.font_family == 'Calibri'
        assert run.style.font_size.pt == 12 and run.style.bold
        if cycle == 0:
            assert 'font_family' not in run.style.properties['direct_fields']
        else:
            assert 'font_family' in run.style.properties['direct_fields']  # Effective family was materialized by the writer.
        transport = tmp_path / f'cycle-{cycle}.json'
        assert write_document(imported.document, transport).success
        path = transport.with_suffix('.docx')
        assert write_document(read_document(transport).document, path).success
        assert Document(path).paragraphs[0].runs[0].font.name == 'Calibri'
    assert source.read_bytes() == original


def test_direct_theme_and_character_style_cascade(tmp_path):
    source = tmp_path / 'own.docx'
    _source(source)
    native = Document(source)
    paragraph = native.paragraphs[0]
    character = native.styles.add_style('Own Character', WD_STYLE_TYPE.CHARACTER)
    character.font.name = 'Arial'
    paragraph.runs[0].style = character
    native.save(source)
    assert read_document(source).document.sections[0].blocks[0].content[0].style.font_family == 'Arial'
    fonts = paragraph.runs[0]._r.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:ascii'), 'Courier New')
    fonts.set(qn('w:asciiTheme'), 'majorHAnsi')
    native.save(source)
    run = read_document(source).document.sections[0].blocks[0].content[0]
    assert run.style.font_family == 'Calibri'
    assert 'font_family' in run.style.properties['direct_fields']


def test_hansi_theme_does_not_override_explicit_ascii_family(tmp_path):
    source = tmp_path / 'own-ascii.docx'
    native = Document()
    run = native.add_paragraph('Own ASCII').runs[0]
    run.font.name = 'Arial'
    fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:hAnsiTheme'), 'majorHAnsi')
    native.save(source)
    assert read_document(source).document.sections[0].blocks[0].content[0].style.font_family == 'Arial'
