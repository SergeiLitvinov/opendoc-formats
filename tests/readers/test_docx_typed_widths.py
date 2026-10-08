"""Typed preferences, legacy JSON migration and explicit edits against native XML."""

import pytest
from docx import Document
from docx.oxml.ns import qn
from opendoc_model import WidthMeasure, document_from_json, document_to_json, get_preferred_width, set_preferred_width

from opendoc_formats.readers.docx import read_docx_model
from opendoc_formats.writers.docx_writer import write_docx_model
from tests.readers.test_docx_layout_roundtrip import W, xml


@pytest.mark.parametrize("kind,value,expected", [
    ("dxa", "2400", WidthMeasure("absolute", 120, "pt")),
    ("dxa", "0", WidthMeasure("absolute", 0, "pt")),
    (None, "2400", WidthMeasure("absolute", 120, "pt")),
    ("dxa", "2.54cm", WidthMeasure("absolute", 72, "pt")),
    ("dxa", "0.125pt", WidthMeasure("absolute", 0.125, "pt")),
    ("pct", "1665", WidthMeasure("relative", 0.333, "ratio", "content")),
    ("pct", "33.3%", WidthMeasure("relative", 0.333, "ratio", "content")),
    ("pct", "6000", WidthMeasure("relative", 1.2, "ratio", "content")),
    ("pct", "0", WidthMeasure("relative", 0, "ratio", "content")),
    ("auto", "0", WidthMeasure("auto")),
    ("nil", "0", WidthMeasure("unspecified")),
])
def test_native_widths_fill_shared_contract_without_changing_source(tmp_path, kind, value, expected):
    source = Document()
    table = source.add_table(rows=1, cols=1)
    for node in (table._tbl.tblPr.find(qn("w:tblW")), table.cell(0, 0)._tc.tcPr.tcW):
        if kind is None:
            node.attrib.pop(qn("w:type"), None)
        else:
            node.set(qn("w:type"), kind)
        node.set(qn("w:w"), value)
    path = tmp_path / "source.docx"
    source.save(path)
    original = path.read_bytes()
    for cycle in range(2):
        model = document_from_json(document_to_json(read_docx_model(path)))
        parsed = model.sections[0].blocks[0]
        assert get_preferred_width(parsed) == expected
        cell_expected = WidthMeasure("relative", expected.value, "ratio", "table") if kind == "pct" else expected
        assert get_preferred_width(parsed.rows[0].cells[0]) == cell_expected
        before = document_to_json(model)
        path = tmp_path / f"cycle-{cycle}.docx"
        write_docx_model(model, path)
        assert document_to_json(model) == before
        for tag in ("tblW", "tcW"):
            node = xml(path).find(".//" + W + tag)
            assert (node.get(W + "type"), node.get(W + "w")) == (kind, value)
    assert (tmp_path / "source.docx").read_bytes() == original


@pytest.mark.parametrize("legacy", [False, True])
def test_typed_edit_then_remove_overrides_native_source_after_json(tmp_path, legacy):
    source = Document()
    source.add_table(rows=1, cols=1).cell(0, 0).text = "Own width edit"
    path = tmp_path / "source.docx"
    source.save(path)
    model = read_docx_model(path)
    table = model.sections[0].blocks[0]
    cell = table.rows[0].cells[0]
    if legacy:
        for node in (table, cell):
            node.properties.pop("preferred_width", None)
            node.properties.pop("docx_width_imported", None)
        old = tmp_path / "legacy.docx"
        write_docx_model(document_from_json(document_to_json(model)), old)
        assert xml(old).find(".//" + W + "tcW").get(W + "w") == cell.properties["docx_preferred_width"]["value"]
    set_preferred_width(table, WidthMeasure("relative", 0.75001, "ratio", "content"))
    set_preferred_width(cell, WidthMeasure("absolute", 12.125, "pt"))
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        path = tmp_path / f"edited-{cycle}.docx"
        write_docx_model(model, path)
        model = read_docx_model(path)
        table = model.sections[0].blocks[0]
        cell = table.rows[0].cells[0]
        assert get_preferred_width(table) == WidthMeasure("relative", 0.75001, "ratio", "content")
        assert get_preferred_width(cell) == WidthMeasure("absolute", 12.125, "pt")
    set_preferred_width(table, None)
    set_preferred_width(cell, None)
    path = tmp_path / "removed.docx"
    write_docx_model(document_from_json(document_to_json(model)), path)
    root = xml(path)
    assert root.find(".//" + W + "tblW") is None
    assert root.find(".//" + W + "tcW") is None
    parsed = read_docx_model(path).sections[0].blocks[0]
    assert get_preferred_width(parsed) is None
    assert get_preferred_width(parsed.rows[0].cells[0]) is None


def test_absent_widths_remain_distinct_from_explicit_unspecified(tmp_path):
    source = Document()
    table = source.add_table(rows=1, cols=1)
    table._tbl.tblPr.remove(table._tbl.tblPr.find(qn("w:tblW")))
    cell = table.cell(0, 0)
    cell._tc.tcPr.remove(cell._tc.tcPr.tcW)
    path = tmp_path / "source.docx"
    source.save(path)
    for cycle in range(2):
        model = document_from_json(document_to_json(read_docx_model(path)))
        parsed = model.sections[0].blocks[0]
        assert get_preferred_width(parsed) is None
        assert get_preferred_width(parsed.rows[0].cells[0]) is None
        path = tmp_path / f"absent-{cycle}.docx"
        write_docx_model(model, path)
        assert xml(path).find(".//" + W + "tcW") is None


def test_public_json_import_migrates_legacy_width_before_removal(tmp_path):
    from opendoc_formats.api import read_document

    source = Document()
    source.add_table(rows=1, cols=1)
    path = tmp_path / "source.docx"
    source.save(path)
    model = read_docx_model(path)
    table = model.sections[0].blocks[0]
    for node in (table, table.rows[0].cells[0]):
        node.properties.pop("preferred_width", None)
        node.properties.pop("docx_width_imported", None)
    legacy = tmp_path / "legacy.json"
    legacy.write_text(document_to_json(model), encoding="utf-8")
    before = legacy.read_bytes()
    migrated = read_document(legacy).document
    table = migrated.sections[0].blocks[0]
    assert get_preferred_width(table).kind == "auto"
    assert get_preferred_width(table.rows[0].cells[0]).kind == "absolute"
    for node in (table, table.rows[0].cells[0]):
        set_preferred_width(node, None)
    target = tmp_path / "removed.docx"
    write_docx_model(document_from_json(document_to_json(migrated)), target)
    assert xml(target).find(".//" + W + "tcW") is None
    assert legacy.read_bytes() == before


def test_negative_native_width_is_not_claimed_as_shared_nonnegative_preference(tmp_path):
    source = Document()
    table = source.add_table(rows=1, cols=1)
    for node in (table._tbl.tblPr.find(qn("w:tblW")), table.cell(0, 0)._tc.tcPr.tcW):
        node.set(qn("w:type"), "dxa")
        node.set(qn("w:w"), "-1")
    path = tmp_path / "source.docx"
    source.save(path)
    for cycle in range(2):
        model = document_from_json(document_to_json(read_docx_model(path)))
        parsed = model.sections[0].blocks[0]
        assert get_preferred_width(parsed) is None
        assert get_preferred_width(parsed.rows[0].cells[0]) is None
        path = tmp_path / f"negative-{cycle}.docx"
        write_docx_model(model, path)
        assert xml(path).find(".//" + W + "tcW").get(W + "w") == "-1"
