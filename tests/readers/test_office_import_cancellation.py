"""Own Office sources exercise mid-import cancellation and unchanged model output."""

import importlib
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Barrier, Event
from zipfile import ZipFile

import pytest
from opendoc_model import document_to_json

from opendoc_formats import ImportOptions, read_document
from opendoc_formats.errors import OperationCancelledError
from opendoc_formats.support.cancellation import cancellation_scope, checkpoint


def _image():
    from PIL import Image

    data = BytesIO()
    Image.new('RGB', (8, 8), (40, 110, 170)).save(data, format='PNG')
    data.seek(0)
    return data


@pytest.fixture
def docx_source(tmp_path):
    pytest.importorskip('docx')
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_heading('Own heading', level=1)
    paragraph = document.add_paragraph()
    for index in range(3):
        run = paragraph.add_run(f'Own run {index} ')
        run.bold = index == 1
        run.italic = index == 2
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged cell'
    for index, row in enumerate(table.rows):
        row.cells[2].text = f'Own row {index}'
    table.cell(1, 0).add_table(rows=1, cols=2).cell(0, 0).text = 'Nested table'
    for _ in range(2):
        document.add_picture(_image(), width=Inches(.2))
    document.sections[0].header.paragraphs[0].text = 'Own running header'
    document.sections[0].footer.paragraphs[0].text = 'Own running footer'
    path = tmp_path / 'own.docx'
    document.save(path)
    return path


@pytest.fixture
def pptx_source(tmp_path):
    pytest.importorskip('pptx')
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    for index in range(3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        for child in range(2):
            frame = group.shapes.add_textbox(Inches(child), Inches(.2), Inches(1), Inches(.5)).text_frame
            for run_index in range(3):
                run = frame.paragraphs[0].add_run()
                run.text = f'Slide {index} child {child} run {run_index} '
                run.font.bold = run_index == 1
        table = slide.shapes.add_table(2, 2, Inches(0), Inches(1), Inches(2), Inches(1)).table
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = 'Merged cell'
        table.cell(1, 0).text = 'Own cell'
        table.cell(1, 1).text = 'Other cell'
        for offset in range(2):
            slide.shapes.add_picture(_image(), Inches(offset), Inches(2.2), width=Inches(.2))
        chart = CategoryChartData()
        chart.categories = ['A', 'B', 'C']
        chart.add_series('Own series', (1, 2, 3))
        chart.add_series('Other series', (3, 2, 1))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(3), Inches(1), Inches(3), Inches(2), chart)
        slide.notes_slide.notes_text_frame.text = 'Own notes'
    path = tmp_path / 'own.pptx'
    presentation.save(path)
    return path


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_inactive_callback_preserves_complete_json(request, kind):
    path = request.getfixturevalue(kind + '_source')
    source = path.read_bytes()
    expected = read_document(path).document
    polls = []

    def inactive():
        polls.append(True)
        return False

    actual = read_document(path, options=ImportOptions(cancelled=inactive)).document
    assert actual is not None and expected is not None
    assert document_to_json(actual) == document_to_json(expected)
    assert len(polls) > 30
    assert path.read_bytes() == source


@pytest.mark.parametrize(('kind', 'module', 'owner', 'name'), [
    ('docx', 'docx_text', None, '_append_run_content'),
    ('docx', 'docx_table', None, '_cell_properties'),
    ('docx', 'docx_drawing', None, '_add_image_resource'),
    ('docx', 'docx_style', None, '_resolved_font_style'),
    ('docx', 'ooxml.package', None, '_add_part'),
    ('docx', 'docx', None, '_container_blocks'),
    ('docx', 'docx', None, 'attach_docx_diagnostics'),
    ('pptx', 'pptx', None, '_parse_run'),
    ('pptx', 'pptx', None, '_handle_text_shape'),
    ('pptx', 'pptx_table', None, 'paragraph_metadata'),
    ('pptx', 'pptx_chart_data', None, 'cached_values'),
    ('pptx', 'pptx', '_ImporterState', 'add_image_resource'),
    ('pptx', 'pptx_diagnostics', 'PptxDiagnostics', 'slide'),
    ('pptx', 'pptx_diagnostics', 'PptxDiagnostics', 'attach'),
])
def test_stop_inside_native_work_prevents_next_object_and_keeps_source(request, monkeypatch, kind, module, owner, name):
    path = request.getfixturevalue(kind + '_source')
    source = path.read_bytes()
    target = importlib.import_module('opendoc_formats.' + (module if '.' in module else 'readers.' + module))
    if owner:
        target = getattr(target, owner)
    native = getattr(target, name)
    stop, completed = Event(), []

    def first(*args, **kwargs):
        assert not completed, 'Native processing continued after cancellation'
        result = native(*args, **kwargs)
        completed.append(True)
        stop.set()
        return result

    with monkeypatch.context() as context:
        context.setattr(target, name, first)
        result = read_document(path, options=ImportOptions(cancelled=stop.is_set))
    assert completed == [True]
    assert result.document is None
    assert any(issue.code == 'import.cancelled' for issue in result.issues)
    assert path.read_bytes() == source
    with ZipFile(path) as archive:
        assert archive.testzip() is None
    assert read_document(path).document is not None


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_cancel_after_backend_open_precedes_model_work(request, monkeypatch, kind):
    path = request.getfixturevalue(kind + '_source')
    backend = importlib.import_module(kind)
    name = 'Document' if kind == 'docx' else 'Presentation'
    native = getattr(backend, name)
    stop, opened = Event(), []

    def open_source(*args, **kwargs):
        result = native(*args, **kwargs)
        opened.append(True)
        stop.set()
        return result

    def forbidden(*args, **kwargs):
        pytest.fail('Model parsing started after cancelled backend open')

    with monkeypatch.context() as context:
        context.setattr(backend, name, open_source)
        reader = importlib.import_module('opendoc_formats.readers.' + kind)
        context.setattr(reader, '_document_metadata' if kind == 'docx' else '_presentation_metadata', forbidden)
        result = read_document(path, options=ImportOptions(cancelled=stop.is_set))
    assert opened == [True]
    assert result.document is None and any(issue.code == 'import.cancelled' for issue in result.issues)
    assert read_document(path).document is not None


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
@pytest.mark.parametrize('rich', [False, True])
def test_direct_readers_cancel_midway_and_restore_context(request, monkeypatch, kind, rich):
    path = request.getfixturevalue(kind + '_source')
    module = importlib.import_module('opendoc_formats.readers.' + kind)
    reader = getattr(module, 'read_' + kind + ('_model' if rich else ''))
    name = ('read_paragraph' if kind == 'docx' else '_parse_run') if rich else ('Block' if kind == 'docx' else 'Table')
    native = getattr(module, name)
    stop, completed = Event(), []

    def first(*args, **kwargs):
        assert not completed
        result = native(*args, **kwargs)
        completed.append(True)
        stop.set()
        return result

    with monkeypatch.context() as context:
        context.setattr(module, name, first)
        with pytest.raises(OperationCancelledError):
            reader(path, cancelled=stop.is_set)
    assert completed == [True]
    checkpoint()  # A cancelled import must not leave its callback active.
    assert reader(path) is not None
    with pytest.raises(ValueError, match='callable'):
        reader(path, cancelled=True)


def test_nested_scope_restores_parent_even_after_failure():
    stop = Event()
    with pytest.raises(OperationCancelledError):
        with cancellation_scope(stop.is_set, 'outer'):
            with pytest.raises(LookupError):
                with cancellation_scope(None, 'inner'):
                    stop.set()
                    checkpoint()
                    raise LookupError('own failure')
            checkpoint()
    checkpoint()


def test_theme_cancellation_is_not_swallowed_as_default_colors(pptx_source, monkeypatch):
    from pptx import Presentation

    from opendoc_formats.readers import pptx

    slide = Presentation(pptx_source).slides[0]
    stop = Event()
    parse = pptx.etree.fromstring

    def theme(data, *args, **kwargs):
        root = parse(data, *args, **kwargs)
        stop.set()
        return root

    with monkeypatch.context() as context:
        context.setattr(pptx.etree, 'fromstring', theme)
        with pytest.raises(OperationCancelledError):
            with cancellation_scope(stop.is_set, 'own theme import'):
                pptx._load_theme_colors(slide.part)
    checkpoint()


def test_parallel_imports_do_not_share_cancellation(pptx_source):
    barrier = Barrier(2)

    def run(cancel):
        polls = 0

        def signal():
            nonlocal polls
            polls += 1
            if polls == 12:
                barrier.wait(timeout=30)
            return cancel and polls >= 12

        return read_document(pptx_source, options=ImportOptions(cancelled=signal))

    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(run, True)
        second = pool.submit(run, False)
        cancelled, completed = first.result(timeout=60), second.result(timeout=60)
    assert cancelled.document is None
    assert any(issue.code == 'import.cancelled' for issue in cancelled.issues)
    assert completed.document is not None and len(completed.document.sections) == 3
    assert document_to_json(completed.document) == document_to_json(read_document(pptx_source).document)
