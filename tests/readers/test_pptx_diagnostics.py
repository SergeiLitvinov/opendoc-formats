"""Own source-object fixtures for partial PPTX preservation diagnostics."""

import io
import zipfile

import pytest
from lxml import etree
from opendoc_model import document_from_json, document_to_json, get_integration
from PIL import Image as PillowImage
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from opendoc_formats import read_document


def _png():
    output = io.BytesIO()
    PillowImage.new("RGB", (8, 8), (12, 34, 56)).save(output, format="PNG")
    output.seek(0)
    return output


def _verify_locations(source, document):
    ledger = get_integration(document)
    assert ledger is not None and not ledger.assessment_complete
    with zipfile.ZipFile(source) as archive:
        for record in ledger.preservation:
            provenance = record.provenance
            assert provenance.source_path == str(source)
            root = etree.fromstring(archive.read(provenance.package_part.lstrip("/")))
            if "#xpath=" in record.issue.location:
                xpath = record.issue.location.split("#xpath=", 1)[1]
                assert root.xpath(xpath, namespaces={key: value for key, value in root.nsmap.items() if key})
    assert get_integration(document_from_json(document_to_json(document))) == ledger
    transport = source.with_suffix(".json")
    transport.write_text(document_to_json(document), encoding="utf-8")
    assert read_document(transport).issues == read_document(source).issues


def test_transition_timing_group_media_and_3d_have_source_locations(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    text = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    text.text = "Retained group text"
    transition = OxmlElement("p:transition")
    transition.append(OxmlElement("p:fade"))
    slide._element.append(transition)
    slide._element.append(OxmlElement("p:timing"))
    properties = text._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
    properties.append(OxmlElement("a:scene3d"))
    properties.append(OxmlElement("a:sp3d"))
    media = text._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}nvSpPr")
    media.append(OxmlElement("a:audioFile"))
    media.append(OxmlElement("a:videoFile"))
    source = tmp_path / "objects.pptx"
    presentation.save(source)
    result = read_document(source)
    assert result.success and not result.lossless
    assert {issue.reason for issue in result.issues} == {
        "original-package-not-retained", "flattened-group", "unsupported-transition", "unsupported-timing",
        "unsupported-audio", "unsupported-video", "unsupported-3d",
    }
    assert "Retained group text" in str(result.document)
    assert any(record.provenance.object_id == str(text.shape_id) for record in get_integration(result.document).preservation)
    _verify_locations(source, result.document)


def test_ole_has_actual_visual_preview_and_opaque_payload_is_not_claimed(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_ole_object(io.BytesIO(b"own inert embedded bytes"), "Package",
                                      Inches(1), Inches(1), Inches(2), Inches(2), icon_file=_png())
    source = tmp_path / "ole.pptx"
    presentation.save(source)
    result = read_document(source)
    assert result.success
    records = [record for record in get_integration(result.document).preservation if record.issue.code == "pptx.ole"]
    assert len(records) == 1 and records[0].state.value == "visual"
    assert records[0].provenance.object_id == str(shape.shape_id)
    resource = result.document.resources[records[0].issue.measurement["resource_id"]]
    assert resource.data == _png().getvalue()
    assert not any(item.data == b"own inert embedded bytes" for item in result.document.resources.values())
    _verify_locations(source, result.document)


@pytest.mark.parametrize("uri,feature", [
    ("urn:own:unsupported-graphic", "pptx.graphic-object"),
    ("http://schemas.openxmlformats.org/drawingml/2006/diagram", "pptx.diagram"),
])
def test_unknown_graphic_and_missing_picture_are_lost(tmp_path, uri, feature):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(1), Inches(1))
    data = frame._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}graphicData")
    for child in list(data):
        data.remove(child)
    data.set("uri", uri)
    picture = slide.shapes.add_picture(_png(), Inches(3), Inches(1))
    blip = picture._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    blip.attrib.clear()
    source = tmp_path / "lost.pptx"
    presentation.save(source)
    result = read_document(source)
    records = get_integration(result.document).preservation
    assert any(record.issue.code == feature and record.issue.reason == "unsupported-graphic"
               and record.state.value == "lost" for record in records)
    assert any(record.issue.reason == "missing-image-resource" for record in records)
    _verify_locations(source, result.document)


def test_missing_geometry_and_zero_extent_are_located_and_slide_order_is_independent(tmp_path):
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    missing = first.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
    missing.text = "Missing geometry"
    properties = missing._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
    properties.remove(properties.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm"))
    first.shapes.add_textbox(Inches(1), Inches(1), 0, Inches(1))
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text = "Retained"
    slide_ids = presentation.slides._sldIdLst
    slide_ids.insert(0, slide_ids[-1])
    source = tmp_path / "geometry.pptx"
    presentation.save(source)
    result = read_document(source)
    records = [record for record in get_integration(result.document).preservation if record.issue.code == "pptx.shape"]
    assert {record.issue.reason for record in records} == {"missing-geometry", "nonpositive-geometry"}
    assert all(record.provenance.page == 2 and record.provenance.package_part == "/ppt/slides/slide1.xml" for record in records)
    assert all(record.provenance.object_id for record in records)
    assert result.document.sections[1].provenance.package_part == "/ppt/slides/slide1.xml"
    _verify_locations(source, result.document)
