"""Physical source bytes and unproved edges survive PPTX import and JSON transport."""

import io
import zipfile
from copy import deepcopy
from dataclasses import replace

import pytest
from lxml import etree
from opendoc_model import (
    ArtifactLimitError,
    DocumentLimits,
    document_from_json,
    document_to_json,
    get_integration,
)
from pptx import Presentation
from pptx.util import Inches

from opendoc_formats import read_document, write_document
from opendoc_formats.anatomy import inspect_binary
from opendoc_formats.errors import InvalidDocumentError, OperationCancelledError, ResourceLimitError, UnsupportedDocumentError
from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.readers.pptx_package import capture_pptx_package
from tests.support.test_ooxml_pptx_package import source as source_with_dependencies

REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'


def rewrite(path, updates):
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as original, zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in original.infolist():
            target.writestr(item, updates.pop(item.filename, original.read(item)))
        for name, data in updates.items():
            target.writestr(name, data)
    path.write_bytes(buffer.getvalue())


def own_pptx(tmp_path):
    path = source_with_dependencies(tmp_path)
    with zipfile.ZipFile(path) as archive:
        content_types = etree.fromstring(archive.read('[Content_Types].xml'))
    for name, kind in (('/own/unknown.bin', 'application/octet-stream'), ('/own/no-decode.xml', 'application/x-own-binary')):
        etree.SubElement(content_types, '{'+CT+'}Override', PartName=name, ContentType=kind)
    rewrite(path, {
        '[Content_Types].xml': etree.tostring(content_types),
        'own/unknown.bin': b'own inert bytes\x00\xff',
        'own/no-decode.xml': b'\x00binary despite suffix: declared media type controls inspection',
        'own/_rels/unknown.bin.rels': (
            f'<Relationships xmlns="{REL}"><Relationship Id="rIdOwn" Type="urn:own:attachment" '
            'Target="no-decode.xml"/></Relationships>'
        ).encode(),
    })
    return path


def test_every_original_payload_and_proved_edge_survives_two_json_cycles(tmp_path):
    source = own_pptx(tmp_path)
    original_bytes = source.read_bytes()
    physical = inspect_binary(source, container='zip', capture=True)
    result = read_document(source)
    assert result.success and not result.lossless, result.issues
    model = result.document
    graph = deepcopy(model.package)
    assert graph.format == 'pptx' and graph.root == '/ppt/presentation.xml'
    assert graph.validate() == []
    with zipfile.ZipFile(source) as archive:
        assert set(graph.parts) == {'/'+name for name in archive.namelist()}
        for name, part in graph.parts.items():
            assert part.data == archive.read(name[1:])
    assert graph.parts['/own/no-decode.xml'].data.startswith(b'\x00binary')
    assert graph.parts['/own/no-decode.xml'].media_type == 'application/x-own-binary'
    proved = {(edge.owner, dict(edge.attributes)['Id']) for edge in physical.opc.relationships
              if edge.status in {'internal', 'external'}}
    assert {(edge.source, edge.id) for edge in graph.relationships} == proved
    assert any(edge.external and edge.target == 'https://example.invalid/never-fetch' for edge in graph.relationships)
    assert any(edge.source == '/own/unknown.bin' and edge.target == '/own/no-decode.xml' for edge in graph.relationships)
    for _ in range(2):
        model = document_from_json(document_to_json(model))
        assert model.package == graph
        records = [record for record in get_integration(model).preservation if record.issue.code == 'pptx.package-part']
        assert {record.provenance.package_part for record in records} == set(graph.parts)
        for record in records:
            item = next(item for item in physical.items if item.locator == 'zip:'+record.provenance.package_part)
            assert record.state.value == 'opaque'
            assert record.issue.measurement['sha256'] == item.sha256
            assert record.issue.measurement['bytes'] == len(item.payload)
    assert source.read_bytes() == original_bytes


@pytest.mark.parametrize('bad', ['missing-target', 'missing-owner', 'duplicate-id', 'invalid-type', 'escape'])
def test_backend_repair_does_not_erase_rejected_physical_relationship(tmp_path, bad):
    source = own_pptx(tmp_path)
    name = 'own/_rels/unknown.bin.rels'
    attributes = 'Id="rIdOwn" Type="urn:own:attachment" Target="no-decode.xml"'
    if bad == 'missing-target':
        attributes = attributes.replace('no-decode.xml', 'missing.bin')
    elif bad == 'missing-owner':
        name = 'absent/_rels/owner.xml.rels'
    elif bad == 'invalid-type':
        attributes = attributes.replace('urn:own:attachment', 'not a URI')
    elif bad == 'escape':
        attributes = attributes.replace('no-decode.xml', '../../outside.bin')
    edge = f'<Relationship {attributes}/>'
    payload = f'<Relationships xmlns="{REL}">{edge}{edge if bad == "duplicate-id" else ""}</Relationships>'.encode()
    rewrite(source, {name: payload})
    original = source.read_bytes()
    model = read_pptx_model(source)
    assert model.package.parts['/'+name].data == payload
    assert model.package.validate() == []
    records = [r for r in get_integration(model).preservation if r.issue.code == 'pptx.package-audit']
    assert records and all(r.issue.location.startswith(str(source)+'!/') for r in records)
    assert all(r.provenance.package_part in model.package.parts for r in records)
    assert not any(edge.source == ('/absent/owner.xml' if bad == 'missing-owner' else '/own/unknown.bin')
                   and edge.id == 'rIdOwn' for edge in model.package.relationships)
    assert document_from_json(document_to_json(model)).package == model.package
    assert source.read_bytes() == original


def test_original_slide_part_names_survive_engine_display_order_renaming(tmp_path):
    deck = Presentation()
    for text in ('First', 'Second'):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = text
    ids = deck.slides._sldIdLst
    ids.insert(0, ids[-1])
    source = tmp_path/'reordered.pptx'
    deck.save(source)
    model = read_pptx_model(source)
    assert model.sections[0].provenance.package_part == '/ppt/slides/slide2.xml'
    assert model.sections[1].provenance.package_part == '/ppt/slides/slide1.xml'
    with zipfile.ZipFile(source) as archive:
        for name in ('/ppt/slides/slide1.xml', '/ppt/slides/slide2.xml', '/ppt/_rels/presentation.xml.rels'):
            assert model.package.parts[name].data == archive.read(name[1:])


@pytest.mark.parametrize('field,value', [('max_bytes', 8), ('max_embedded_bytes', 0), ('max_embedded_bytes', 512),
                                         ('max_nodes', 2), ('max_depth', 1)])
def test_source_capture_is_bounded_before_engine_open(tmp_path, monkeypatch, field, value):
    import pptx
    source = own_pptx(tmp_path)
    before = source.read_bytes()
    def unexpected(*args, **kwargs):
        pytest.fail('Native engine opened before physical package was accepted')
    monkeypatch.setattr(pptx, 'Presentation', unexpected)
    with pytest.raises((ResourceLimitError, ArtifactLimitError)):
        read_pptx_model(source, limits=replace(DocumentLimits(), **{field: value}))
    assert source.read_bytes() == before


def test_source_xml_policy_rejects_dtd_in_unreachable_component(tmp_path):
    source = own_pptx(tmp_path)
    rewrite(source, {'own/unreachable.xml': b'<!DOCTYPE own [<!ENTITY x "not allowed">]><own>&x;</own>'})
    with pytest.raises(UnsupportedDocumentError, match='DTD'):
        read_pptx_model(source)


def test_package_capture_cancelled_without_changing_source(tmp_path, monkeypatch):
    from opendoc_formats.readers import pptx_package
    source = own_pptx(tmp_path)
    before = source.read_bytes()
    invoked = False
    original = pptx_package.part_name
    def stop(name):
        nonlocal invoked
        invoked = True
        return original(name)
    monkeypatch.setattr(pptx_package, 'part_name', stop)
    with pytest.raises(OperationCancelledError):
        read_pptx_model(source, cancelled=lambda: invoked)
    assert invoked and source.read_bytes() == before


def test_regenerated_pptx_does_not_claim_native_package_restoration(tmp_path):
    model = read_pptx_model(own_pptx(tmp_path))
    before = document_to_json(model)
    target = tmp_path/'generated.pptx'
    result = write_document(model, target)
    assert result.success and not result.lossless
    assert any(issue.feature == 'package' and issue.severity.value == 'loss' for issue in result.issues)
    assert document_to_json(model) == before


def test_missing_entry_point_does_not_invent_a_model_root(tmp_path):
    source = own_pptx(tmp_path)
    rewrite(source, {'_rels/.rels': f'<Relationships xmlns="{REL}"/>'.encode()})
    with pytest.raises(InvalidDocumentError, match='one proved physical'):
        capture_pptx_package(source.read_bytes())


def test_foreign_source_package_never_leaks_into_docx(tmp_path):
    model = read_pptx_model(own_pptx(tmp_path))
    before = document_to_json(model)
    target = tmp_path/'projection.docx'
    report = write_document(model, target)
    assert report.success and not report.lossless, report.issues
    assert any(issue.feature == 'package-graph' for issue in report.issues)
    with zipfile.ZipFile(target) as archive:
        assert not any(name.startswith(('ppt/', 'own/')) for name in archive.namelist())
    assert document_to_json(model) == before


def test_full_physical_inventory_requires_editable_bindings_before_restore(tmp_path):
    from opendoc_formats.ooxml.package import restore_package_graph
    model = read_pptx_model(own_pptx(tmp_path))
    deck = Presentation()
    before = tuple(deck.part.rels), deck.part.blob
    with pytest.raises(ValueError, match='explicit editable binding'):
        restore_package_graph(deck.part, model.package)
    assert (tuple(deck.part.rels), deck.part.blob) == before


def test_projection_and_inventory_use_same_captured_bytes(tmp_path, monkeypatch):
    import pptx
    source = own_pptx(tmp_path)
    original = source.read_bytes()
    native = pptx.Presentation
    calls = []

    def open_captured(stream):
        assert isinstance(stream, io.BytesIO)
        assert stream.getvalue() == original
        calls.append(True)
        source.unlink()  # prove the engine cannot reopen another generation of the source
        return native(stream)

    monkeypatch.setattr(pptx, 'Presentation', open_captured)
    model = read_pptx_model(source)
    assert calls == [True] and 'Own source text' in str(model)
    with zipfile.ZipFile(io.BytesIO(original)) as archive:
        assert model.package.parts['/ppt/slides/slide1.xml'].data == archive.read('ppt/slides/slide1.xml')
