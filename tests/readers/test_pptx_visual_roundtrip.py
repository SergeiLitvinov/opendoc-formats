"""Own slides for effective backgrounds and zero-axis editable freeforms."""

import pytest
from opendoc_model import document_from_json, document_to_json
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model


@pytest.mark.parametrize("owner", ["master", "layout", "slide"])
def test_effective_solid_background_two_json_cycles(tmp_path, owner):
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[6])
    slide.shapes.add_textbox(Pt(30), Pt(30), Pt(200), Pt(40)).text = "Own background"
    parent = {"master": source.slide_master, "layout": slide.slide_layout, "slide": slide}[owner]
    parent.background.fill.solid()
    parent.background.fill.fore_color.rgb = RGBColor(24, 72, 96)
    path = tmp_path / "source.pptx"
    source.save(path)
    original = path.read_bytes()
    for cycle in range(2):
        model = read_pptx_model(path)
        assert model.sections[0].properties.get("background_fill") == "#184860"
        model = document_from_json(document_to_json(model))
        path = tmp_path / f"cycle-{cycle}.pptx"
        assert write_pptx_model(model, path).success
        assert str(Presentation(path).slides[0].background.fill.fore_color.rgb) == "184860"
    assert (tmp_path / "source.pptx").read_bytes() == original


def test_freeform_axis_lines_two_json_cycles(tmp_path):
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[6])
    for x, y, origin_y in ((4000, 0, 1), (0, 2000, 2)):
        builder = slide.shapes.build_freeform(0, 0, scale=Inches(1) / 1000)
        builder.add_line_segments([(x, y)], close=False)
        shape = builder.convert_to_shape(Inches(1), Inches(origin_y))
        shape.fill.background()
        shape.line.color.rgb = RGBColor(24, 72, 96)
        shape.line.width = Pt(6)
    path = tmp_path / "source.pptx"
    source.save(path)
    expected = [(shape.left, shape.top, shape.width, shape.height) for shape in slide.shapes]
    for cycle in range(2):
        model = document_from_json(document_to_json(read_pptx_model(path)))
        assert len(model.sections[0].blocks) == 2
        path = tmp_path / f"cycle-{cycle}.pptx"
        assert write_pptx_model(model, path).success
        shapes = Presentation(path).slides[0].shapes
        assert [(shape.left, shape.top, shape.width, shape.height) for shape in shapes] == expected
        assert all(shape._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}custGeom')
                   is not None for shape in shapes)


def test_table_theme_audit(tmp_path):
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Pt(60), Pt(60), Pt(400), Pt(140)).table
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"Own cell {r}:{c}"
    path = tmp_path / "source.pptx"
    source.save(path)
    for cycle in range(2):
        model = document_from_json(document_to_json(read_pptx_model(path)))
        path = tmp_path / f"cycle-{cycle}.pptx"
        report = write_pptx_model(model, path)
        assert report.success
        assert any(issue.feature == "styles" for issue in report.issues)
        restored = Presentation(path).slides[0].shapes[0].table
        assert restored._tbl.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill') == []


def test_nearest_background_overrides_master_without_mutating_source(tmp_path):
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[6])
    for owner, color in ((source.slide_master, "FF0000"), (slide.slide_layout, "00FF00"), (slide, "0000FF")):
        owner.background.fill.solid()
        owner.background.fill.fore_color.rgb = RGBColor.from_string(color)
    path = tmp_path / "source.pptx"
    source.save(path)
    before = path.read_bytes()
    assert read_pptx_model(path).sections[0].properties["background_fill"] == "#0000FF"
    assert path.read_bytes() == before
    slide.background.fill.background()
    source.save(path)
    assert "background_fill" not in read_pptx_model(path).sections[0].properties
