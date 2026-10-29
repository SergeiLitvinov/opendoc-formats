"""Real python-pptx packages retain opaque bytes, source rIds and generated neighbors."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree
from opendoc_model import DocumentModel, document_from_json, document_to_json
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TARGET_MODE as MODE
from pptx.opc.package import Part, _Relationship
from pptx.opc.packuri import PackURI
from pptx.util import Inches

from opendoc_formats.ooxml.package import load_package_graph, restore_package_graph

R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
OWN = 'urn:opendoc-formats:own-package'
LINK = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink'
IMAGE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'
PNG = Path(__file__).resolve().parents[1]/'native/data/table-image.png'


def native_edge(owner, rid, kind, target, *, external=False):
    """Construct source through the engine API, independently of the helper under test."""
    owner.rels._rels[rid] = _Relationship(
        owner.partname.baseURI, rid, kind, MODE.EXTERNAL if external else MODE.INTERNAL, target,
    )


def source(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = 'Own source text'
    package = slide.part.package
    custom = Part(PackURI('/ppt/own/data.xml'), 'application/xml', package, b'<own>exact payload</own>')
    binary = Part(PackURI('/ppt/media/image1.png'), 'image/png', package, PNG.read_bytes())
    native_edge(slide.part, 'rIdOwn', OWN, custom)
    native_edge(custom, 'rIdPixels', IMAGE, binary)
    native_edge(custom, 'rIdPixelsAgain', IMAGE, binary)
    native_edge(binary, 'rIdCycle', OWN, custom)
    native_edge(custom, 'rIdExternal', LINK, 'https://example.invalid/never-fetch', external=True)
    output = tmp_path/'independent.pptx'
    deck.save(output)
    return output


def graph_at(slide):
    graph = load_package_graph(
        slide.part, format_name='ooxml', supported_relationships={OWN}, recursive_relationships={OWN},
    )
    assert graph is not None and graph.validate() == []
    return graph


def test_pptx_recursive_graph_keeps_bytes_cycles_and_distinct_ids_after_json(tmp_path):
    original = source(tmp_path)
    original_bytes = original.read_bytes()
    graph = graph_at(Presentation(original).slides[0])
    expected = deepcopy(graph)
    for cycle in range(2):
        model = document_from_json(document_to_json(DocumentModel(package=graph)))
        before = document_to_json(model)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = 'Generated neighbor'
        picture = slide.shapes.add_picture(BytesIO(PNG.read_bytes()), 0, Inches(2))
        generated_id = picture._pic.blipFill.blip.rEmbed
        generated = slide.part.rels[generated_id].target_part
        assert str(generated.partname) == '/ppt/media/image1.png'
        restored = restore_package_graph(slide.part, model.package)
        assert all(part.package is slide.part.package for part in restored.values())
        assert generated.partname != '/ppt/media/image1.png'
        assert slide.part.rels[generated_id].target_part is generated
        output = tmp_path/f'cycle-{cycle}.pptx'
        deck.save(output)
        with ZipFile(output) as archive:
            assert archive.read('ppt/own/data.xml') == expected.parts['/ppt/own/data.xml'].data
            assert archive.read('ppt/media/image1.png') == PNG.read_bytes()
            assert archive.read(str(generated.partname).lstrip('/')) == PNG.read_bytes()
            assert len(archive.namelist()) == len(set(archive.namelist()))
        reloaded = Presentation(output)
        assert reloaded.slides[0].shapes[0].text == 'Generated neighbor'
        assert reloaded.slides[0].shapes[1].image.blob == PNG.read_bytes()
        graph = graph_at(reloaded.slides[0])
        assert graph == expected
        assert document_to_json(model) == before
    assert original.read_bytes() == original_bytes


@pytest.mark.parametrize('bound', [False, True], ids=['root', 'bound-part'])
def test_pptx_rid_collision_rewrites_every_generated_reference_and_keeps_opaque(tmp_path, bound):
    graph = graph_at(Presentation(source(tmp_path)).slides[0])
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    owner = slide.part
    if bound:
        from pptx.opc.package import XmlPart
        owner = XmlPart(PackURI('/ppt/own/data.xml'), 'application/xml', slide.part.package, etree.Element('own'))
        slide.part.relate_to(owner, OWN)
        rid = 'rIdExternal'
        container = owner._element
    else:
        rid = 'rIdOwn'
        container = slide._element
    native_edge(owner, rid, LINK, 'https://example.invalid/generated', external=True)
    references = [etree.SubElement(container, 'own') for _ in range(3)]
    for element in references:
        element.set(R+'id', rid)
        element.set('unrelated', rid)
    opaque = references[-1]
    if bound:
        owner._opendoc_opaque_roots = (opaque,)
    restored = restore_package_graph(
        slide.part, graph, protected_roots=() if bound else (opaque,),
        source_parts={'/ppt/own/data.xml': owner} if bound else None,
    )
    replacement = references[0].get(R+'id')
    assert replacement != rid
    assert references[1].get(R+'id') == replacement
    assert opaque.get(R+'id') == rid
    assert [node.get('unrelated') for node in references] == [rid]*3
    assert owner.rels[replacement].target_ref == 'https://example.invalid/generated'
    if bound:
        assert restored['/ppt/own/data.xml'] is owner
        assert owner.rels[rid].target_ref == 'https://example.invalid/never-fetch'
    else:
        assert owner.rels[rid].target_part is restored['/ppt/own/data.xml']
    output = tmp_path/f'collision-{bound}.pptx'
    deck.save(output)
    reopened = Presentation(output).slides[0].part
    if bound:
        reopened = reopened.rels['rIdOwn'].target_part
    assert reopened.rels[replacement].target_ref == 'https://example.invalid/generated'
    assert reopened.rels[rid].reltype == (LINK if bound else OWN)


def test_pptx_unretained_opaque_relationship_refuses_before_mutation(tmp_path):
    graph = graph_at(Presentation(source(tmp_path)).slides[0])
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    opaque = etree.SubElement(slide._element, 'own')
    opaque.set(R+'embed', 'rIdMissing')
    before = slide.part.blob, tuple(slide.part.rels)
    with pytest.raises(ValueError, match='unretained source relationship'):
        restore_package_graph(slide.part, graph, protected_roots=(opaque,))
    assert (slide.part.blob, tuple(slide.part.rels)) == before
