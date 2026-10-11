"""Own native sources exercise caller-selected budgets across import and JSON."""

import hashlib

import pytest
from opendoc_model import DocumentLimits, DocumentModel, Paragraph, Section, TextRun, document_to_json
from pptx import Presentation
from pptx.util import Inches

from opendoc_formats import ExportOptions, ImportOptions, read_document, write_document


@pytest.mark.parametrize('path,quota,maximum', [
    ('$.sections[0]', 'nodes', 125), (None, 'work', 0), ('$.resources', None, None),
])
def test_structured_sdk_quota_overrides_message_without_inventing_fields(tmp_path, path, quota, maximum):
    from opendoc_model import ArtifactLimitError, ConversionReport

    from opendoc_formats.export import ExporterRegistry, ExporterSpec

    registry = ExporterRegistry()

    def writer(document, destination):
        destination.write_bytes(b'partial output')
        raise ArtifactLimitError('misleading: bytes quota exceeded (999)', path=path, quota=quota, maximum=maximum)

    registry.register(ExporterSpec('own', ('.own',), writer))
    destination = tmp_path / 'structured.own'
    destination.write_bytes(b'previous destination')
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun('Own quota')])])])
    report = registry.write(model, destination)
    assert isinstance(report, ConversionReport) and not report.success
    assert destination.read_bytes() == b'previous destination'
    issue = next(i for i in report.issues if i.feature == 'export.document-limit')
    assert issue.message == 'misleading: bytes quota exceeded (999)'
    assert issue.location == str(destination) + ('#' + path if path is not None else '')
    measured = report.metrics['document_limit']
    assert measured.get('quota') == quota and measured.get('maximum') == maximum
    assert ('quota' in measured) == (quota is not None)
    assert ('maximum' in measured) == (maximum is not None)


def test_legacy_sdk_quota_preserves_previous_location_and_measurement():
    from opendoc_model import ArtifactLimitError

    from opendoc_formats.support.document_limits import limit_diagnostic

    issue = limit_diagnostic(ArtifactLimitError('$.sections[0]: nodes quota exceeded (123)'),
                             DocumentLimits(), 'own.docx', 'import.document-limit')
    assert issue.location == 'own.docx#$.sections[0]'
    assert issue.measurement['quota'] == 'nodes' and issue.measurement['maximum'] == 123


@pytest.mark.parametrize('rule,value', [('minimum', 14), ('exact', 18), ('relative', 1.15)])
def test_model_line_spacing_two_public_json_cycles_and_native_application(tmp_path, rule, value):
    from opendoc_model import IssueSeverity, LineSpacing, get_line_spacing, set_line_spacing

    paragraph = Paragraph([TextRun('Own declared spacing')])
    set_line_spacing(paragraph, LineSpacing(rule, value))
    model = DocumentModel(sections=[Section(blocks=[paragraph])])
    expected = document_to_json(model)
    for cycle in range(2):
        path = tmp_path / f'spacing-{cycle}.json'
        assert write_document(model, path).success
        result = read_document(path)
        assert result.success
        model = result.document
        assert document_to_json(model) == expected
        assert get_line_spacing(model.sections[0].blocks[0]) == LineSpacing(rule, value)
    target = tmp_path / 'spacing.docx'
    target.write_bytes(b'previous destination')
    report = write_document(model, target, options=ExportOptions(
        quality_gate=lambda report, _: not any(issue.severity is IssueSeverity.LOSS for issue in report.issues),
    ))
    assert report.success, report.issues
    for _ in range(2):
        restored = read_document(target).document
        assert get_line_spacing(restored.sections[0].blocks[0]) == LineSpacing(rule, value)
        assert write_document(restored, target).success
    assert document_to_json(model) == expected


@pytest.fixture(scope="module")
def large_pptx(tmp_path_factory):
    source = tmp_path_factory.mktemp("caller-limits") / "own.pptx"
    deck = Presentation()
    for slide_number in range(35):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for shape_number in range(6):
            frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame
            for run_number in range(32):
                frame.paragraphs[0].add_run().text = f"Own {slide_number}/{shape_number}/{run_number} "
    deck.save(source)
    return source


def test_native_pptx_selected_budget_and_two_public_json_cycles(large_pptx, tmp_path):
    before = hashlib.sha256(large_pptx.read_bytes()).digest()
    refused = read_document(large_pptx)
    assert not refused.success and refused.document is None
    issue, = refused.issues
    assert issue.code == "import.document-limit" and "#$.sections[" in issue.location
    assert issue.measurement["quota"] == "nodes" and issue.measurement["maximum"] == 100000
    limits = DocumentLimits(max_nodes=200000)
    result = read_document(large_pptx, options=ImportOptions(document_limits=limits))
    assert result.success, result.issues
    assert not result.assessment_complete and not result.lossless
    document = result.document
    texts = [run.text for section in document.sections for block in section.blocks for run in block.content]
    assert len(document.sections) == 35 and len(texts) == 6720
    snapshot = document_to_json(document, limits=limits)
    for cycle in range(2):
        path = tmp_path / f"cycle-{cycle}.json"
        report = write_document(document, path, options=ExportOptions(document_limits=limits))
        assert report.success, report.issues
        result = read_document(path, options=ImportOptions(document_limits=limits))
        assert result.success and not result.lossless and not result.assessment_complete
        document = result.document
        assert document_to_json(document, limits=limits) == snapshot
    assert hashlib.sha256(large_pptx.read_bytes()).digest() == before


def test_native_pptx_export_respects_selected_budget_and_preserves_refused_target(large_pptx, tmp_path):
    limits = DocumentLimits(max_nodes=200000)
    before = large_pptx.read_bytes()
    document = read_document(large_pptx, options=ImportOptions(document_limits=limits)).document
    expected_texts = [run.text for section in document.sections for block in section.blocks for run in block.content]
    target = tmp_path / 'own-native.pptx'
    sentinel = b'previous destination'
    target.write_bytes(sentinel)
    refused = write_document(document, target)
    assert not refused.success and target.read_bytes() == sentinel
    assert any(issue.feature == 'export.document-limit' for issue in refused.issues)
    for cycle in range(2):
        report = write_document(document, target, options=ExportOptions(document_limits=limits))
        assert report.success, report.issues
        document = read_document(target, options=ImportOptions(document_limits=limits)).document
        assert len(document.sections) == 35
        texts = [run.text for section in document.sections for block in section.blocks for run in block.content]
        assert len(texts) == 6720 and texts == expected_texts
    assert large_pptx.read_bytes() == before


@pytest.mark.parametrize('format_id,suffix', [('docx', '.docx'), ('pptx', '.pptx'), ('html', '.html'),
                                             ('epub', '.epub'), ('latex', '.tex'), ('pdf', '.pdf')])
def test_every_native_model_writer_uses_selected_budget(tmp_path, format_id, suffix):
    from opendoc_model import ArtifactLimitError

    document = DocumentModel(sections=[Section(blocks=[
        Paragraph([TextRun('Own bounded run ') for _ in range(32)]) for _ in range(210)
    ])])
    limits = DocumentLimits(max_nodes=300000)
    with pytest.raises(ArtifactLimitError):
        document.validate()
    assert not document.validate(limits=limits)
    target = tmp_path / ('large' + suffix)
    target.write_bytes(b'previous destination')
    refused = write_document(document, target, format_id=format_id)
    assert not refused.success and target.read_bytes() == b'previous destination'
    report = write_document(document, target, format_id=format_id, options=ExportOptions(document_limits=limits))
    assert report.success, report.issues
    assert target.stat().st_size > len(b'previous destination')
    if format_id == 'pdf':
        import fitz

        with fitz.open(target) as native:
            text = ' '.join(' '.join(page.get_text().split()) for page in native)
            assert text.count('Own bounded run') == 6720
    elif format_id == 'docx':
        from docx import Document

        assert sum(p.text.count('Own bounded run') for p in Document(target).paragraphs) == 6720
    elif format_id == 'pptx':
        deck = Presentation(target)
        assert sum(shape.text.count('Own bounded run') for slide in deck.slides
                   for shape in slide.shapes if shape.has_text_frame) == 6720
    elif format_id == 'epub':
        from zipfile import ZipFile

        with ZipFile(target) as archive:
            assert sum(archive.read(name).decode().count('Own bounded run') for name in archive.namelist()
                       if name.endswith('.xhtml') and 'nav' not in name) == 6720
    else:
        assert target.read_text(encoding='utf-8').count('Own bounded run') == 6720


def test_prepared_model_quota_refuses_atomically_with_selected_budget(tmp_path):
    from opendoc_model import ArtifactLimitError

    document = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun('Own late quota')])])])
    low, high = 1, 1000
    while low < high:
        middle = (low + high) // 2
        try:
            document.validate(limits=DocumentLimits(max_nodes=middle))
        except ArtifactLimitError:
            low = middle + 1
        else:
            high = middle
    limits = DocumentLimits(max_nodes=low)
    assert not document.validate(limits=limits)
    before = document_to_json(document)
    target = tmp_path / 'late.html'
    target.write_bytes(b'previous destination')
    report = write_document(document, target, options=ExportOptions(document_limits=limits))
    assert not report.success and target.read_bytes() == b'previous destination'
    issue, = [issue for issue in report.issues if issue.feature == 'export.document-limit']
    assert issue.location.startswith(str(target) + '#$')
    assert report.metrics['document_limit']['quota'] == 'nodes'
    assert report.metrics['document_limit']['maximum'] == low
    assert report.metrics['document_limit']['limits']['max_nodes'] == low
    assert document_to_json(document) == before


def test_unknown_writer_quota_keeps_original_text_and_destination(tmp_path):
    from opendoc_model import ArtifactLimitError

    from opendoc_formats import ExporterRegistry, ExporterSpec

    def custom_writer(document, path):
        path.write_bytes(b'partial staged bytes')
        raise ArtifactLimitError('Unknown native quota representation')

    registry = ExporterRegistry()
    registry.register(ExporterSpec('custom', ('.own',), custom_writer))
    target = tmp_path / 'custom.own'
    target.write_bytes(b'previous destination')
    report = registry.write(DocumentModel(), target)
    assert not report.success and target.read_bytes() == b'previous destination'
    issue, = report.issues
    assert issue.feature == 'export.document-limit'
    assert issue.location == str(target) and issue.message == 'Unknown native quota representation'
    assert 'quota' not in report.metrics['document_limit']
    assert 'maximum' not in report.metrics['document_limit']


def test_public_pdf_passes_cancellation_inside_writer(tmp_path):
    calls = []

    def cancelled():
        calls.append(True)
        return len(calls) > 1

    target = tmp_path / 'cancelled.pdf'
    target.write_bytes(b'previous destination')
    document = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun('Own cancelled PDF')])])])
    report = write_document(document, target, options=ExportOptions(cancelled=cancelled))
    assert not report.success and target.read_bytes() == b'previous destination'
    assert len(calls) >= 2
    assert any('cancel' in (issue.feature + issue.message).lower() for issue in report.issues)


@pytest.mark.parametrize('format_id', ['html', 'pdf', 'pptx'])
def test_docx_private_renderer_view_uses_selected_budget(tmp_path, format_id):
    from docx import Document

    source = tmp_path / 'own-large.docx'
    native = Document()
    for paragraph_number in range(210):
        paragraph = native.add_paragraph()
        for run_number in range(32):
            paragraph.add_run(f'Own {paragraph_number}/{run_number} ')
    native.save(source)
    before = source.read_bytes()
    limits = DocumentLimits(max_nodes=500000)
    imported = read_document(source, options=ImportOptions(document_limits=limits))
    assert imported.success, imported.issues
    document = imported.document
    snapshot = document_to_json(document, limits=limits)
    target = tmp_path / ('own-rendered.' + format_id)
    report = write_document(document, target, options=ExportOptions(document_limits=limits))
    assert report.success, report.issues
    assert document_to_json(document, limits=limits) == snapshot
    assert source.read_bytes() == before


def test_inherited_color_view_growth_refuses_before_native_publication(tmp_path):
    from opendoc_model import ArtifactLimitError, TextStyle
    from opendoc_model.color import ColorValue

    document = DocumentModel(source_format='docx', styles={'own': TextStyle(color=ColorValue.from_hex('#123456'))},
                             sections=[Section(blocks=[Paragraph([TextRun('Own inherited')], style_id='own')])])
    maximum = 1
    while True:
        try:
            document.validate(limits=DocumentLimits(max_nodes=maximum))
            break
        except ArtifactLimitError:
            maximum += 1
    before = document_to_json(document)
    target = tmp_path / 'color-growth.pptx'
    target.write_bytes(b'previous destination')
    report = write_document(document, target, options=ExportOptions(document_limits=DocumentLimits(max_nodes=maximum)))
    assert not report.success and target.read_bytes() == b'previous destination'
    assert any(issue.feature == 'export.document-limit' for issue in report.issues)
    assert report.metrics['document_limit']['maximum'] == maximum
    assert document_to_json(document) == before


@pytest.mark.parametrize("limits,quota,maximum", [
    (DocumentLimits(max_nodes=1), "nodes", 1),
    (DocumentLimits(max_depth=1), "depth", 1),
    (DocumentLimits(max_bytes=64), "bytes", 64),
    (DocumentLimits(max_embedded_bytes=0), "embedded bytes", 0),
])
def test_docx_quota_refusal_has_location_and_selected_limits(tmp_path, limits, quota, maximum):
    from docx import Document

    source = tmp_path / "own.docx"
    native = Document()
    native.add_paragraph("Own bounded source")
    native.save(source)
    before = source.read_bytes()
    result = read_document(source, options=ImportOptions(document_limits=limits))
    assert not result.success and result.document is None
    issue, = result.issues
    assert issue.code == "import.document-limit" and "#" in issue.location
    assert issue.measurement["quota"] == quota and issue.measurement["maximum"] == maximum
    assert source.read_bytes() == before


def test_cancel_and_input_limit_precede_native_read(large_pptx):
    for options, code in [
        (ImportOptions(cancelled=lambda: True), "import.cancelled"),
        (ImportOptions(max_input_bytes=1), "import.input-limit"),
    ]:
        result = read_document(large_pptx, options=options)
        assert not result.success and result.document is None and result.issues[0].code == code


@pytest.mark.parametrize("format_id", ["html", "epub", "latex", "pdf", "txt", "json"])
def test_registered_readers_return_small_budget_refusals(tmp_path, format_id):
    suffix = {"latex": "tex"}.get(format_id, format_id)
    source = tmp_path / f"own.{suffix}"
    model = DocumentModel(sections=[Section(blocks=[Paragraph(content=[TextRun("Own source")])])])
    if format_id == "html":
        source.write_text("<p>Own source</p>", encoding="utf-8")
    elif format_id == "latex":
        source.write_text("Own source", encoding="utf-8")
    elif format_id == "txt":
        source.write_text("Own source", encoding="utf-8")
    elif format_id == "json":
        source.write_text(document_to_json(model), encoding="utf-8")
    elif format_id == "pdf":
        import pymupdf

        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((72, 72), "Own source")
            pdf.save(source)
    else:
        assert write_document(model, source).success
    result = read_document(source, options=ImportOptions(document_limits=DocumentLimits(max_nodes=1)))
    assert not result.success and result.document is None
    assert result.issues[0].code == "import.document-limit"
