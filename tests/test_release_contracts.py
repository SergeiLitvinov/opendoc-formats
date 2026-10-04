"""Regressions for pre-publication validation, options and resource boundaries."""

from pathlib import Path

import pytest
from opendoc import ConversionReport, DocumentModel

from opendoc_formats import (
    AdapterRegistry,
    AdapterSpec,
    ExporterRegistry,
    ExporterSpec,
    ExportOptions,
    ImportOptions,
    read_document,
    write_document,
)
from opendoc_formats.support.output_validation import validate_output


@pytest.mark.parametrize("format_id", ["json", "html", "docx", "pptx", "pdf", "latex"])
def test_corrupt_export_is_never_published(tmp_path, format_id):
    target = tmp_path / "result.custom"
    target.write_bytes(b"previous")

    def writer(document, path):
        path.write_bytes(b"corrupt content")
        return ConversionReport(path)

    registry = ExporterRegistry()
    registry.register(ExporterSpec("custom", (".custom",), writer, validator=lambda path: validate_output(path, format_id)))
    report = registry.write(DocumentModel(), target)
    assert not report.success
    assert report.issues[-1].feature == "export.invalid-output"
    assert target.read_bytes() == b"previous"
    assert not list(tmp_path.glob(".opendoc-formats-export-*"))


def test_cancellation_during_validation_preserves_previous_result(tmp_path):
    target = tmp_path / "result.custom"
    target.write_text("previous", encoding="utf-8")
    cancelled = False

    def validator(path):
        nonlocal cancelled
        assert path.read_text(encoding="utf-8") == "replacement"
        cancelled = True

    def writer(document, path):
        path.write_text("replacement", encoding="utf-8")
        return ConversionReport(path)

    registry = ExporterRegistry()
    registry.register(ExporterSpec("custom", (".custom",), writer, validator=validator))
    report = registry.write(DocumentModel(), target, options=ExportOptions(cancelled=lambda: cancelled))
    assert not report.success and report.issues[-1].feature == "export.cancelled"
    assert target.read_text(encoding="utf-8") == "previous"


def test_verification_can_be_disabled_explicitly(tmp_path):
    def writer(document, path):
        path.write_text("custom", encoding="utf-8")
        return ConversionReport(path)

    def validator(path):
        raise ValueError("custom validation")

    registry = ExporterRegistry()
    registry.register(ExporterSpec("custom", (".custom",), writer, validator=validator))
    target = tmp_path / "result.custom"
    assert registry.write(DocumentModel(), target, options=ExportOptions(verify_output=False)).success


def test_wrong_writer_return_keeps_destination(tmp_path):
    registry = ExporterRegistry()
    registry.register(ExporterSpec("custom", (".custom",), lambda document, path: None))
    output = tmp_path / "result.custom"
    output.write_bytes(b"previous")
    report = registry.write(DocumentModel(), output)
    assert not report.success and report.issues[-1].feature == "export.invalid-result"
    assert output.read_bytes() == b"previous"


@pytest.mark.parametrize("kwargs", [{"pdf_mode": "unknown"}, {"ocr_engine_factory": 0}, {"max_input_bytes": True}])
def test_bad_import_options_fail_at_construction(kwargs):
    with pytest.raises(ValueError):
        ImportOptions(**kwargs)


@pytest.mark.parametrize("options", [False, 0, {}])
def test_falsey_invalid_options_are_not_replaced_with_defaults(tmp_path, options):
    with pytest.raises(ValueError, match="options"):
        read_document(tmp_path / "input.txt", options=options)
    with pytest.raises(ValueError, match="options"):
        write_document(DocumentModel(), tmp_path / "output.txt", options=options)


@pytest.mark.parametrize("extensions", [(".",), (".UPPER",), (".same", ".same"), (None,)])
def test_invalid_extensions_are_rejected(extensions):
    with pytest.raises(ValueError):
        AdapterSpec("custom", extensions, lambda path, options: DocumentModel())
    with pytest.raises(ValueError):
        ExporterSpec("custom", extensions, lambda document, path: ConversionReport(path))


def test_import_validation_rejects_non_json_metadata(tmp_path):
    path = tmp_path / "source.custom"
    path.write_bytes(b"input")
    registry = AdapterRegistry()
    registry.register(
        AdapterSpec("custom", (".custom",), lambda path, options: DocumentModel(metadata={"invalid": Path("file")}))
    )
    result = registry.read(path)
    assert not result.success and result.issues[-1].code == "import.invalid-model"


def test_pptx_viewer_does_not_fetch_scripts_from_network(tmp_path):
    from pptx import Presentation

    from opendoc_formats.writers.pptx_to_html import convert

    source = tmp_path / "input.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)
    result = convert(source, tmp_path / "viewer")
    assert result.success, result.to_dict()
    from html.parser import HTMLParser

    class ScriptSources(HTMLParser):
        sources = []

        def handle_starttag(self, tag, attrs):
            if tag == "script" and "src" in dict(attrs):
                self.sources.append(dict(attrs)["src"])

    parser = ScriptSources()
    parser.feed((tmp_path / "viewer/index.html").read_text(encoding="utf-8"))
    assert parser.sources and all(not src.startswith(("http:", "https:", "//")) for src in parser.sources)


def test_missing_dotted_backend_is_diagnostic(tmp_path):
    source = tmp_path / "source.custom"
    source.write_bytes(b"input")
    imports = AdapterRegistry()
    imports.register(AdapterSpec("custom", (".custom",), lambda path, options: DocumentModel(), ("absent_backend.child",)))
    assert imports.read(source).issues[-1].code == "import.backend-unavailable"
    exports = ExporterRegistry()
    exports.register(
        ExporterSpec("custom", (".custom",), lambda document, path: ConversionReport(path), ("absent_backend.child",))
    )
    assert exports.write(DocumentModel(), tmp_path / "result.custom").issues[-1].feature == "export.backend-unavailable"
