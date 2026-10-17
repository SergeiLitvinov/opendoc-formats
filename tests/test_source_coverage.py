"""Reproducible limits of source import, linked to finite coverage tasks."""

from opendoc_model import Formula, Table

from opendoc_formats import read_document, write_document
from opendoc_formats.readers.epub_model import read_epub_model
from opendoc_formats.readers.pptx import read_pptx_model


def test_tex_source_has_finite_native_reader(tmp_path):
    source = tmp_path / "source.tex"
    source.write_text(r"\documentclass{article}\begin{document}Hello\end{document}", encoding="utf-8")
    for identifier in (None, "latex"):
        result = read_document(source, format_id=identifier)
        assert result.success and not result.lossless and not result.assessment_complete
        assert result.document.sections[0].blocks[0].plain_text == "Hello"
        assert result.document.resources["latex-original-source"].data == source.read_bytes()


def test_epub_table_math_are_semantic_but_nonlinear_content_is_lost(tmp_path):
    import zipfile

    source = tmp_path / "source.epub"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", '''<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
          <rootfiles><rootfile full-path="book.opf" media-type="application/oebps-package+xml"/></rootfiles></container>''')
        archive.writestr("book.opf", '''<package xmlns="http://www.idpf.org/2007/opf" version="3.0">
          <metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier>coverage-fixture</dc:identifier>
          <dc:title>Own generated source fixture</dc:title><dc:language>en</dc:language></metadata>
          <manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/>
          <item id="appendix" href="appendix.xhtml" media-type="application/xhtml+xml"/></manifest>
          <spine><itemref idref="chapter"/><itemref idref="appendix" linear="no"/></spine></package>''')
        archive.writestr("chapter.xhtml", '''<html xmlns="http://www.w3.org/1999/xhtml"><body>
          <p>Visible <math xmlns="http://www.w3.org/1998/Math/MathML"><mi>x</mi></math></p>
          <table><tr><td>TableOnly</td></tr></table></body></html>''')
        archive.writestr("appendix.xhtml", "<html><body><p>NonlinearOnly</p></body></html>")
    document = read_epub_model(source)
    blocks = [block for section in document.sections for block in section.blocks]
    text = str(document)
    assert "Visible" in text and "TableOnly" in text and "NonlinearOnly" not in text
    table = next(block for block in blocks if isinstance(block, Table))
    assert table.rows[0].cells[0].blocks[0].plain_text == "TableOnly"
    assert isinstance(blocks[0].content[1], Formula)
    result = read_document(source)
    assert result.success and not result.lossless and not result.assessment_complete
    assert {issue.code for issue in result.issues} == {"epub.spine", "epub.package"}
    assert all(issue.severity.value == "loss" and str(source) in issue.location for issue in result.issues)
    assert any("appendix.xhtml" in issue.location and issue.reason == "nonlinear-spine" for issue in result.issues)
    from opendoc_formats import ImportOptions

    inclusive = read_document(source, options=ImportOptions(epub_include_nonlinear=True))
    assert inclusive.success and "NonlinearOnly" in str(inclusive.document)
    assert not any(issue.reason == "nonlinear-spine" for issue in inclusive.issues)
    assert inclusive.document.sections[1].properties["epub"]["linear"] == "no"
    assert inclusive.document.metadata["epub"]["source_spine"] == document.metadata["epub"]["source_spine"]
    from opendoc_model import document_from_json, document_to_json, get_integration

    inclusive_restored = document_from_json(document_to_json(inclusive.document))
    assert inclusive_restored.metadata["epub"]["include_nonlinear"] is True
    assert inclusive_restored.sections[1].properties["epub"]["linear"] == "no"
    assert inclusive_restored.sections[1].blocks[0].plain_text == "NonlinearOnly"
    restored = document_from_json(document_to_json(document))
    ledger = get_integration(restored)
    assert ledger == get_integration(document)
    assert all(record.provenance.source_path == str(source) for record in ledger.preservation)
    assert next(
        record for record in ledger.preservation if record.issue.reason == "nonlinear-spine"
    ).provenance.package_part.endswith(".xhtml")
    with zipfile.ZipFile(source) as archive:
        assert all(record.provenance.package_part.lstrip("/") in archive.namelist() for record in ledger.preservation)
    report = write_document(document, tmp_path / "output.epub")
    assert report.success and not report.lossless and report.metrics["output_verified"]
    assert any(issue.feature == "epub.spine" for issue in report.issues)
    copied = read_document(tmp_path / "output.epub")
    assert copied.success and "TableOnly" in str(copied.document) and "NonlinearOnly" not in str(copied.document)


def test_pptx_import_does_not_preserve_original_package(tmp_path):
    from pptx import Presentation
    from pptx.oxml.xmlchemy import OxmlElement

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    transition = OxmlElement("p:transition")
    transition.append(OxmlElement("p:fade"))
    slide._element.append(transition)
    source = tmp_path / "transition.pptx"
    presentation.save(source)
    document = read_pptx_model(source)
    assert document.package is None
    assert "transition" not in document.sections[0].properties
    result = read_document(source)
    assert result.success and not result.lossless and not result.assessment_complete
    assert any(issue.code == "pptx.transition" and issue.reason == "unsupported-transition" for issue in result.issues)
    target = tmp_path / "output.pptx"
    assert write_document(document, target).success
    restored = Presentation(target)
    assert not restored.slides[0]._element.xpath("./p:transition")
