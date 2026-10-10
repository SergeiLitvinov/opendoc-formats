"""Own direct table/cell declarations through two public DOCX/JSON cycles."""

from dataclasses import replace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from opendoc_model import Border, ColorValue, get_borders, set_border, set_borders

from opendoc_formats import read_document, write_document

TABLE_SIDES = ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')


def _source(path, style='Normal Table', special=False):
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.style = style
    for index, cell in enumerate(table._cells):
        cell.text = f'Own cell {index}'
    container = OxmlElement('w:tblBorders')
    for side in TABLE_SIDES:
        child = OxmlElement('w:' + side)
        for key, value in {'val': 'single', 'sz': '8', 'color': '000000'}.items():
            child.set(qn('w:' + key), value)
        container.append(child)
    table._tbl.tblPr.append(container)
    cells = OxmlElement('w:tcBorders')
    for side, val in [('top', 'none'), ('bottom', 'nil'), ('tl2br', 'double'), ('tr2bl', 'dotted')]:
        child = OxmlElement('w:' + side)
        child.set(qn('w:val'), val)
        if val not in ('none', 'nil'):
            child.set(qn('w:sz'), '16')
            child.set(qn('w:color'), '112233')
            child.set(qn('w:space'), '3')
        cells.append(child)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(cells)
    if special:
        child = container[0]
        child.set(qn('w:val'), 'triple')
        child.set(qn('w:themeColor'), 'accent1')
        child.set(qn('w:themeTint'), '80')
        child.set(qn('w:shadow'), '1')
    doc.save(path)


def _native(path):
    doc = Document(path)
    table = doc.tables[0]
    return ([(child.tag, dict(child.attrib)) for child in table._tbl.tblPr.find(qn('w:tblBorders'))],
            [(child.tag, dict(child.attrib)) for child in table.cell(0, 0)._tc.tcPr.find(qn('w:tcBorders'))])


def test_six_direct_table_borders_and_cell_suppression_diagonals_in_two_cycles(tmp_path):
    source = tmp_path / 'own.docx'
    _source(source)
    original, expected = source.read_bytes(), _native(source)
    path = source
    for cycle in range(2):
        imported = read_document(path)
        assert imported.success
        table = imported.document.sections[0].blocks[0]
        borders = get_borders(table)
        assert set(borders.sides) == {'top', 'right', 'bottom', 'left', 'inside-horizontal', 'inside-vertical'}
        assert borders.sides['top'].style == 'solid' and borders.sides['top'].width_pt == 1
        cell = get_borders(table.rows[0].cells[0])
        assert cell.sides['top'].kind == 'none' and cell.sides['bottom'].kind == 'nil'
        assert cell.sides['diagonal-down'].width_pt == 2 and cell.sides['diagonal-up'].offset_pt == 3
        transport = tmp_path / f'cycle-{cycle}.json'
        assert write_document(imported.document, transport).success
        path = transport.with_suffix('.docx')
        report = write_document(read_document(transport).document, path)
        assert report.success and not any('border' in issue.feature for issue in report.issues)
        actual = _native(path)
        assert [sorted(part) for part in actual] == [sorted(part) for part in expected]
    assert source.read_bytes() == original


def test_common_border_edit_and_removal_override_native_attributes(tmp_path):
    source = tmp_path / 'own.docx'
    _source(source, style='Table Grid')
    model = read_document(source).document
    table = model.sections[0].blocks[0]
    old = get_borders(table).sides['top']
    set_border(table, 'top', replace(old, style='double', width_pt=2, color=ColorValue.from_hex('AA5500'), offset_pt=1))
    set_border(table.rows[0].cells[0], 'top', Border(kind='nil'))
    output = tmp_path / 'edited.docx'
    assert write_document(model, output).success
    native, cells = _native(output)
    assert dict(native)[qn('w:top')] == {qn('w:val'): 'double', qn('w:sz'): '16', qn('w:color'): 'AA5500', qn('w:space'): '1'}
    assert dict(cells)[qn('w:top')] == {qn('w:val'): 'nil'}
    set_borders(table, None)
    set_borders(table.rows[0].cells[0], None)
    transport = tmp_path / 'removed.json'
    assert write_document(model, transport).success
    assert write_document(read_document(transport).document, output).success
    native_table = Document(output).tables[0]
    assert native_table._tbl.tblPr.find(qn('w:tblBorders')) is None
    assert native_table.cell(0, 0)._tc.tcPr.find(qn('w:tcBorders')) is None
    assert native_table.style.name == 'Table Grid'


def test_theme_and_unsupported_style_are_retained_with_located_loss(tmp_path):
    source = tmp_path / 'native.docx'
    _source(source, special=True)
    imported = read_document(source)
    assert any(issue.code == 'docx.table-borders' and issue.location for issue in imported.issues)
    top = get_borders(imported.document.sections[0].blocks[0]).sides['top']
    assert top.style is None and top.color is None
    transport = tmp_path / 'native.json'
    assert write_document(imported.document, transport).success
    output = tmp_path / 'native-out.docx'
    report = write_document(read_document(transport).document, output)
    assert report.success
    assert any(issue.feature == 'docx.table-borders' and issue.location for issue in report.issues)
    assert [sorted(part) for part in _native(output)] == [sorted(part) for part in _native(source)]


@pytest.mark.parametrize('border', [Border(style='solid', width_pt=.3), Border(style='solid', offset_pt=32), Border()])
def test_unrepresentable_border_refuses_atomically(tmp_path, border):
    source = tmp_path / 'own.docx'
    _source(source)
    model = read_document(source).document
    set_border(model.sections[0].blocks[0], 'top', border)
    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous artifact')
    report = write_document(model, output)
    assert not report.success
    assert any(issue.feature == 'docx-table-borders' and issue.location for issue in report.issues)
    assert output.read_bytes() == b'previous artifact'


def test_cross_format_border_loss_is_located_and_json_retains_declaration(tmp_path):
    source = tmp_path / 'own.docx'
    _source(source)
    model = read_document(source).document
    output = tmp_path / 'other.html'
    report = write_document(model, output)
    assert report.success
    assert any(issue.feature == 'table-borders' and issue.location == 'sections[0].blocks[0]' for issue in report.issues)
    transport = tmp_path / 'same.json'
    assert write_document(model, transport).success
    assert get_borders(read_document(transport).document.sections[0].blocks[0]) == get_borders(model.sections[0].blocks[0])


def test_common_borders_do_not_block_unrelated_legacy_html_migration(tmp_path):
    from opendoc_model import (
        BorderSet,
        DocumentModel,
        Paragraph,
        Section,
        Table,
        TableCell,
        TableRow,
        TextRun,
        get_table_cell_semantics,
    )

    cell = TableCell([Paragraph([TextRun('Own header')])], properties={'html_cell_tag': 'th', 'anchor_id': 'heading'})
    table = Table([TableRow([cell])])
    set_borders(table, BorderSet({'top': Border(style='solid', width_pt=1)}))
    model = DocumentModel(sections=[Section(blocks=[table])])
    transport = tmp_path / 'legacy.json'
    assert write_document(model, transport).success
    restored = read_document(transport).document.sections[0].blocks[0]
    assert get_table_cell_semantics(restored.rows[0].cells[0]).role == 'header'
    assert get_borders(restored) == get_borders(table)
