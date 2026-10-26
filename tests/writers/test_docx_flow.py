"""Own ruled PDF: explicit editable reflow, honest grid mapping and atomic refusal."""

import hashlib
import io
from copy import deepcopy
from threading import Event
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from opendoc_model import (
    ArtifactLimitError,
    Border,
    BorderSet,
    Box,
    DocumentLimits,
    DocumentModel,
    Heading,
    Image,
    IssueSeverity,
    Paragraph,
    PreservationState,
    Table,
    TableCell,
    TableRow,
    TextRun,
    TextStyle,
    document_from_json,
    document_to_json,
    get_borders,
    set_borders,
    set_heading,
)

from opendoc_formats import DocxProfile, ExportOptions, get_object_mapping, read_document, write_document
from opendoc_formats.errors import OperationCancelledError
from opendoc_formats.support.cancellation import cancellation_scope
from opendoc_formats.writers.docx_writer import write_docx_model


def grid_source_model(source, *, fragmented=False, styled=False, cell_styled=False):
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=600, height=500)
        page.insert_text((70, 100), 'Own outside table')
        for x in (70, 170, 270, 370):
            width, color = ((.5, (1, 0, 0)) if x == 70 else (1.5, (0, 1, 0)) if x == 370 else
                            (2.25, (0, 1, 1))) if styled else (1, (0, 0, 0))
            for start, end in (((150, 190), (190, 230), (230, 270)) if fragmented else ((150, 270),)):
                if cell_styled:
                    width, color = cell_edge_style('v', x, start)
                page.draw_line((x, start), (x, end), width=width, color=color)
        for y in (150, 190, 230, 270):
            width, color = ((1, (0, 0, 0)) if y == 150 else (3, (0, 0, 1)) if y == 270 else
                            (.75, (1, 0, 1))) if styled else (1, (0, 0, 0))
            for start, end in (((70, 170), (170, 270), (270, 370)) if fragmented else ((70, 370),)):
                if cell_styled:
                    width, color = cell_edge_style('h', y, start)
                page.draw_line((start, y), (end, y), width=width, color=color)
        for row in range(3):
            for column in range(3):
                page.insert_text((80+100*column, 175+40*row), f'R{row+1}C{column+1}')
        pdf.new_page(width=600, height=500).insert_text((70, 100), 'Own other page')
        pdf.save(source)
    imported = read_document(source)
    assert imported.success
    return imported.document


def cell_edge_style(axis, fixed, start):
    palette = ((.5, (1, 0, 0)), (.75, (1, 0, 1)), (1, (0, 0, 0)),
               (1.5, (0, 1, 0)), (2.25, (0, 1, 1)), (3, (0, 0, 1)))
    index = (round((fixed-70)/100)+round((start-150)/40) if axis == 'v' else
             round((fixed-150)/40)+2*round((start-70)/100))
    return palette[index % len(palette)]


@pytest.fixture
def source_model(tmp_path):
    return grid_source_model(tmp_path/'own-grid.pdf')


def merged_source_model(path, *, styled=False):
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=600, height=500)
        for x, top, bottom in ((70, 150, 270), (170, 190, 270), (270, 150, 270), (370, 150, 270)):
            width, color = ((.5, (1, 0, 0)) if x == 70 else (.75, (1, 0, 1)) if x == 170 else
                            (2.25, (0, 1, 1)) if x == 270 else (1.5, (0, 1, 0))) if styled else (1, (0, 0, 0))
            page.draw_line((x, top), (x, bottom), width=width, color=color)
        for y, left, right in ((150, 70, 370), (190, 70, 370), (230, 170, 370), (270, 70, 370)):
            width, color = ((1, (0, 0, 0)) if y == 150 else (.5, (1, 0, 0)) if y == 190 else
                            (1.5, (0, 1, 0)) if y == 230 else (3, (0, 0, 1))) if styled else (1, (0, 0, 0))
            page.draw_line((left, y), (right, y), width=width, color=color)
        for position, text in (((80, 175), 'Top merge'), ((280, 175), 'R1C3'),
                               ((80, 215), 'Tall merge'), ((180, 215), 'R2C2'), ((280, 215), 'R2C3'),
                               ((180, 255), 'R3C2'), ((280, 255), 'R3C3')):
            page.insert_text(position, text)
        pdf.save(path)
    result = read_document(path)
    assert result.success, result.issues
    return result.document


def crop_merged_model(model, mode):
    table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
    if mode == 'top':
        table.rows = table.rows[:1]
        table.box = Box(70, 150, 300, 40)
    else:
        table.rows = table.rows[1:]
        if mode == 'plain':
            table.rows[0].cells = table.rows[0].cells[1:]
            table.box = Box(170, 190, 200, 80)
        else:
            table.box = Box(70, 190, 300, 80)
    return table


@pytest.mark.parametrize('mode,excluded', [('top', 3), ('bottom', 1), ('plain', 2)])
def test_cropped_merged_source_keeps_original_topology_through_json_and_native_cycles(tmp_path, mode, excluded):
    model = merged_source_model(tmp_path/'source.pdf')
    source_table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
    original_bounds = source_table.properties['pdf_source_cell_bounds']
    assert len(original_bounds) == 7
    model = document_from_json(document_to_json(model))
    table = crop_merged_model(model, mode)
    assert table.properties['pdf_source_cell_bounds'] == original_bounds
    expected_spans = [[(cell.row_span, cell.column_span) for cell in row.cells] for row in table.rows]
    expected_text = [['\n'.join(block.plain_text for block in cell.blocks) for cell in row.cells] for row in table.rows]
    before = document_to_json(model)
    target = tmp_path/'crop.docx'
    report = write_document(model, target, options=ExportOptions(docx_profile=DocxProfile('flow')))
    assert report.success, report.issues
    mapping = get_object_mapping(report)
    assert len(mapping.sources) == sum(len(entry.source_ids) for entry in mapping.entries) == 8
    assert sum(len(entry.source_ids) for entry in mapping.entries if entry.state is PreservationState.LOST) == excluded
    assert document_to_json(model) == before
    for cycle in range(3):
        imported = read_document(target)
        assert imported.success, imported.issues
        restored = next(block for block in imported.document.sections[0].blocks if isinstance(block, Table))
        assert [[(cell.row_span, cell.column_span) for cell in row.cells] for row in restored.rows] == expected_spans
        assert [['\n'.join(block.plain_text for block in cell.blocks) for cell in row.cells]
                for row in restored.rows] == expected_text
        if cycle < 2:
            target = tmp_path/f'crop-cycle-{cycle}.docx'
            assert write_document(document_from_json(document_to_json(imported.document)), target).success


@pytest.mark.parametrize('fault', ['overlap', 'incomplete', 'off-grid', 'cut-source-cell'])
def test_original_merge_inventory_is_validated_before_consuming_vectors(tmp_path, fault):
    model = merged_source_model(tmp_path/'source.pdf')
    table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
    basis = table.properties['pdf_source_cell_bounds']
    if fault == 'overlap':
        basis.append(list(basis[0]))
    elif fault == 'incomplete':
        basis.pop()
    elif fault == 'off-grid':
        basis[0][0] += .5
    else:
        first = table.rows[0].cells[0]
        first.column_span = 1
        first.properties['pdf_cell_bbox'][2] = 170
        table.rows[0].cells.insert(1, TableCell([Paragraph([TextRun('Unmeasured split')])],
                                              properties={'pdf_cell_bbox': [170, 150, 270, 190]}))
    before = document_to_json(model)
    target = tmp_path/'previous.docx'
    target.write_bytes(b'previous')
    report = write_document(model, target, options=ExportOptions(docx_profile=DocxProfile('flow')))
    assert not report.success and target.read_bytes() == b'previous'
    assert document_to_json(model) == before


def test_source_inventory_grid_budget_and_cancellation_are_native_limits(monkeypatch):
    from opendoc_formats.writers import docx_flow

    bounds = [[0, 0, 3, 3]]
    with pytest.raises(ArtifactLimitError, match='Source table grid'):
        docx_flow._source_segments(bounds, [0, 1, 2, 3], [0, 1, 2, 3], DocumentLimits(max_nodes=8))
    stop, calls, original = Event(), [], docx_flow.checkpoint

    def checkpoint():
        calls.append(True)
        if len(calls) == 3:
            stop.set()
        original()

    monkeypatch.setattr(docx_flow, 'checkpoint', checkpoint)
    with cancellation_scope(stop.is_set, 'Source cell inventory'):
        with pytest.raises(OperationCancelledError):
            docx_flow._source_segments(bounds, [0, 1, 2, 3], [0, 1, 2, 3], None)
        stop.clear()


@pytest.mark.parametrize('public', [False, True])
def test_measured_merged_pdf_grid_retains_native_spans_and_two_cycles(tmp_path, public):
    model = merged_source_model(tmp_path/'merges.pdf')
    table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
    spans = [[(cell.row_span, cell.column_span) for cell in row.cells] for row in table.rows]
    assert spans == [[(1, 2), (1, 1)], [(2, 1), (1, 1), (1, 1)], [(1, 1), (1, 1)]]
    before = document_to_json(model)
    target = tmp_path/'merged-flow.docx'
    report = (write_document(model, target, options=ExportOptions(docx_profile=DocxProfile('flow'))) if public else
              write_docx_model(model, target, profile=DocxProfile('flow')))
    assert report.success, report.issues
    mapping = get_object_mapping(report)
    assert mapping.publication == 'published' and len(mapping.sources) == 8
    assert sum(len(entry.source_ids) for entry in mapping.entries) == 8
    assert all(entry.state is PreservationState.VISUAL for entry in mapping.entries)
    assert len(mapping.targets[0].extra['cell_topology']) == 7
    assert document_to_json(model) == before
    for index in range(3):
        with ZipFile(target) as package:
            assert not any(name.startswith('word/media/') for name in package.namelist())
        native = Document(target).tables[0]
        assert native.cell(0, 0)._tc is native.cell(0, 1)._tc
        assert native.cell(1, 0)._tc is native.cell(2, 0)._tc
        assert native.cell(0, 0).text == 'Top merge' and native.cell(1, 0).text == 'Tall merge'
        assert abs(native.cell(0, 0).width.pt-200) < .1
        restored = read_document(target)
        assert restored.success, restored.issues
        restored_table = next(block for block in restored.document.sections[0].blocks if isinstance(block, Table))
        assert [[(cell.row_span, cell.column_span) for cell in row.cells] for row in restored_table.rows] == spans
        if index < 2:
            target = tmp_path/f'merged-cycle-{index}.docx'
            assert write_document(document_from_json(document_to_json(restored.document)), target).success


@pytest.mark.parametrize('fault', ['span', 'reordered-geometry', 'crop', 'missing-line'])
def test_merged_grid_refuses_unproved_topology_atomically(tmp_path, fault):
    model = merged_source_model(tmp_path/'merges.pdf')
    table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
    if fault == 'span':
        table.rows[0].cells[0].column_span = 1
    elif fault == 'reordered-geometry':
        table.rows[1].cells.reverse()
    elif fault == 'crop':
        table.properties.pop('pdf_source_cell_bounds', None)  # Legacy snapshot has no independent original inventory.
        table.rows = table.rows[:1]
    else:
        vector = next(block for block in model.sections[0].blocks if isinstance(block, Paragraph)
                      and len(block.content) == 1 and isinstance(block.content[0], Image))
        model.sections[0].blocks.remove(vector)
    before = document_to_json(model)
    target = tmp_path/'previous.docx'
    target.write_bytes(b'previous')
    report = write_document(model, target, options=ExportOptions(docx_profile=DocxProfile('flow')))
    assert not report.success and target.read_bytes() == b'previous'
    assert document_to_json(model) == before


def test_changed_native_merge_is_refused_before_publication(tmp_path, monkeypatch):
    from opendoc_formats.writers import docx_flow

    model = merged_source_model(tmp_path/'merges.pdf')
    before = document_to_json(model)
    original = docx_flow.finalize_flow_mapping

    def tamper(report, path, *, limits):
        output = io.BytesIO()
        with ZipFile(path) as source, ZipFile(output, 'w') as target:
            for info in source.infolist():
                value = source.read(info.filename)
                if info.filename == 'word/document.xml':
                    root = ET.fromstring(value)
                    root.find('.//'+qn('w:gridSpan')).set(qn('w:val'), '1')
                    value = ET.tostring(root)
                target.writestr(info, value)
        path.write_bytes(output.getvalue())
        original(report, path, limits=limits)

    monkeypatch.setattr(docx_flow, 'finalize_flow_mapping', tamper)
    target = tmp_path/'previous.docx'
    target.write_bytes(b'previous')
    report = write_document(model, target, options=ExportOptions(docx_profile=DocxProfile('flow')))
    assert not report.success and target.read_bytes() == b'previous'
    assert any('merged-cell topology changed' in issue.message for issue in report.issues)
    assert document_to_json(model) == before


@pytest.mark.parametrize('crop', [False, True])
@pytest.mark.parametrize('public', [False, True])
def test_fragmented_grid_matches_whole_stroke_native_table(source_model, tmp_path, crop, public):
    fragmented = grid_source_model(tmp_path/'fragmented-grid.pdf', fragmented=True)
    with pymupdf.open(tmp_path/'own-grid.pdf') as whole, pymupdf.open(tmp_path/'fragmented-grid.pdf') as split:
        assert all(first.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).samples ==
                   second.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).samples
                   for first, second in zip(whole, split, strict=True))
    if crop:
        _crop(source_model)
        _crop(fragmented)
    before = document_to_json(fragmented)
    outputs = [tmp_path/'whole.docx', tmp_path/'fragments.docx']
    reports = []
    for model, path in zip((source_model, fragmented), outputs, strict=True):
        report = (write_document(model, path, options=ExportOptions(docx_profile=DocxProfile('flow'))) if public
                  else write_docx_model(model, path, profile=DocxProfile('flow')))
        assert report.success, report.issues
        reports.append(report)
    tables = []
    for path in outputs:
        with ZipFile(path) as package:
            root = ET.fromstring(package.read('word/document.xml'))
            tables.append(ET.tostring(root.find('.//'+qn('w:tbl'))))
            assert not any(name.startswith('word/media/') for name in package.namelist())
    assert tables[0] == tables[1]
    mapping = get_object_mapping(reports[1])
    assert mapping.publication == 'published' and len(mapping.sources) == 24
    assert sum(len(entry.source_ids) for entry in mapping.entries) == 24
    assert sum(len(entry.source_ids) for entry in mapping.entries if entry.state is PreservationState.LOST) == (12 if crop else 0)
    assert document_to_json(fragmented) == before
    restored = read_document(outputs[1])
    assert restored.success
    target = tmp_path/'fragment-cycle.docx'
    assert write_document(document_from_json(document_to_json(restored.document)), target).success
    native = Document(target)
    assert [[cell.text for cell in row.cells] for row in native.tables[0].rows] == (
        [['R2C1', 'R2C2'], ['R3C1', 'R3C2']] if crop else
        [[f'R{row}C{column}' for column in range(1, 4)] for row in range(1, 4)])


@pytest.mark.parametrize('fault', ['gap', 'overlap', 'cumulative-gap', 'cumulative-overlap'])
def test_fragmented_grid_refuses_unproven_coverage_atomically(tmp_path, fault):
    model = grid_source_model(tmp_path/'fragmented.pdf', fragmented=True)
    vectors = [block.content[0] for block in model.sections[0].blocks if isinstance(block, Paragraph)
               and len(block.content) == 1 and isinstance(block.content[0], Image)]
    if fault.startswith('cumulative-'):
        # Each individual discrepancy is below the coordinate tolerance, but
        # their sum exceeds it on a single source border.
        vertical = sorted((image for image in vectors if image.box.x == 70 and image.box.width == 0),
                          key=lambda image: image.box.y)
        assert len(vertical) == 3
        amount = -.00006 if fault == 'cumulative-gap' else .00006
        for image in vertical[:2]:
            model.resources[image.resource_id].properties['items'][0][2][1] += amount
            image.box.height += amount
    else:
        image = vectors[0]
        points = model.resources[image.resource_id].properties['items'][0]
        vertical = points[1][0] == points[2][0]
        axis = 1 if vertical else 0
        amount = -1 if fault == 'gap' else 1
        points[2][axis] += amount
        if vertical:
            image.box.height += amount
        else:
            image.box.width += amount
    before = document_to_json(model)
    target = tmp_path/'previous.docx'
    target.write_bytes(b'previous')
    report = write_document(model, target, options=ExportOptions(docx_profile=DocxProfile('flow')))
    assert not report.success and any(issue.feature == 'docx.flow-profile' for issue in report.issues)
    assert target.read_bytes() == b'previous' and document_to_json(model) == before


@pytest.mark.parametrize('adjustment,accepted', [(-.00004, True), (.00004, True),
                                              (-.00006, False), (.00006, False)])
def test_fragment_tolerance_is_total_per_border_including_both_endpoints(adjustment, accepted):
    from opendoc_formats.writers import docx_flow

    expected = (0, 0, 0, 3)
    lines = [(0, 0, 0, 1+adjustment), (0, 1, 0, 2+adjustment), (0, 2, 0, 3)]
    matches = [(line, None, 'owned', None, expected) for line in lines]
    assert docx_flow._complete_strokes(matches, {expected}) is accepted


@pytest.mark.parametrize('mode', ['endpoints', 'mixed-sign-joints'])
@pytest.mark.parametrize('amount,accepted', [(.00004, True), (.00006, False)])
def test_fragment_error_budget_charges_endpoints_and_cannot_cancel_opposite_errors(mode, amount, accepted):
    from opendoc_formats.writers import docx_flow

    expected = (0, 0, 0, 3)
    lines = ([(0, amount, 0, 1), (0, 1, 0, 2), (0, 2, 0, 3-amount)] if mode == 'endpoints' else
             [(0, 0, 0, 1+amount), (0, 1, 0, 2-amount), (0, 2, 0, 3)])
    matches = [(line, None, 'owned', None, expected) for line in lines]
    assert docx_flow._complete_strokes(matches, {expected}) is accepted


def test_cancellation_inside_fragment_coverage_check(monkeypatch):
    from opendoc_formats.writers import docx_flow

    stop, calls, original = Event(), [], docx_flow.checkpoint

    def checkpoint():
        calls.append(True)
        if len(calls) == 6:
            stop.set()
        original()

    monkeypatch.setattr(docx_flow, 'checkpoint', checkpoint)
    matches = [((0, index, 0, index+1), None, 'owned', None, (0, 0, 0, 3)) for index in range(3)]
    with cancellation_scope(stop.is_set, 'Grid coverage'):
        with pytest.raises(OperationCancelledError):
            docx_flow._complete_strokes(matches, {(0, 0, 0, 3)})
        stop.clear()
    assert len(calls) == 6


@pytest.mark.parametrize('segment,visible', [
    ((69.99995, 190, 69.99995, 230), True),
    ((70, 189.99995, 170, 189.99995), True),
    ((70, 150, 70, 190), False),
    ((270, 230, 370, 230), False),
])
def test_fragment_mapping_respects_coordinate_tolerance_and_positive_overlap(segment, visible):
    from opendoc_formats.writers import docx_flow

    assert docx_flow._visible_fragment(segment, (70, 190, 270, 270)) is visible


def _crop(model, *, box=True):
    table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
    table.rows = table.rows[1:]
    for row in table.rows:
        row.cells = row.cells[:2]
    table.box = Box(70, 190, 200, 80) if box else None
    # Public semantic edits leave original resources/native hints untouched.
    outside = next(block for block in model.sections[0].blocks if isinstance(block, Paragraph)
                   and any(getattr(run, 'text', '') == 'Own outside table' for run in block.content))
    set_heading(outside, Heading(2))
    remaining = [block for block in model.sections[0].blocks if block is not table and block is not outside]
    model.sections[0].blocks = [table, outside, *remaining]


@pytest.mark.parametrize('box', [False, True])
@pytest.mark.parametrize('public', [False, True])
def test_crop_reorder_and_two_native_json_cycles(source_model, tmp_path, box, public):
    _crop(source_model, box=box)
    before = document_to_json(source_model)
    output = tmp_path / 'flow.docx'
    model = document_from_json(before)
    if public:
        report = write_document(model, output, options=ExportOptions(docx_profile=DocxProfile('flow')))
    else:
        report = write_docx_model(model, output, profile=DocxProfile('flow'))
    assert report.success, report.issues
    assert [(issue.feature, issue.severity) for issue in report.issues if issue.severity is IssueSeverity.LOSS] == [
        ('docx.flow-layout', IssueSeverity.LOSS),
    ]
    mapping = get_object_mapping(report)
    assert mapping.publication == 'published' and not mapping.assessment_complete
    assert mapping.result_id == 'sha256:'+hashlib.sha256(output.read_bytes()).hexdigest()
    assert len(mapping.sources) == 8 and len(mapping.targets) == 1
    assert len([source for entry in mapping.entries for source in entry.source_ids]) == 8
    assert len(next(entry for entry in mapping.entries if entry.state is PreservationState.LOST).source_ids) == 2
    assert mapping.targets[0].native_locator.format == 'docx'
    assert document_to_json(source_model) == before == document_to_json(model)
    with ZipFile(output) as package:
        assert not any(name.startswith('word/media/') for name in package.namelist())
    for cycle in range(3):
        native = Document(output)
        assert [[cell.text for cell in row.cells] for row in native.tables[0].rows] == [['R2C1', 'R2C2'], ['R3C1', 'R3C2']]
        assert native.paragraphs[0].text == 'Own outside table'
        assert native.paragraphs[0]._p.find(qn('w:pPr')).find(qn('w:outlineLvl')).get(qn('w:val')) == '1'
        body = list(native._element.body)
        assert body[0].tag.endswith('}tbl') and body[1].tag.endswith('}p')
        assert any(paragraph.text == 'Own other page' for paragraph in native.paragraphs)
        result = read_document(output)
        assert result.success
        table = next(block for block in result.document.sections[0].blocks if isinstance(block, Table))
        borders = get_borders(table)
        assert borders is not None and len(borders.sides) == 6
        if cycle < 2:
            output = tmp_path / f'native-{cycle}.docx'
            assert write_document(document_from_json(document_to_json(result.document)), output).success


@pytest.mark.parametrize('profile', [None, DocxProfile()])
def test_default_still_rejects_unmeasured_pdf_grid(source_model, tmp_path, profile):
    _crop(source_model, box=False)
    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous')
    report = write_docx_model(source_model, output, profile=profile)
    assert not report.success and output.read_bytes() == b'previous'


@pytest.mark.parametrize('fault', ['missing', 'duplicate', 'dash', 'unrelated', 'ambiguous', 'hole', 'edited-border'])
def test_graphics_are_not_dropped_silently(source_model, tmp_path, fault):
    section = source_model.sections[0]
    vectors = [block for block in section.blocks if isinstance(block, Paragraph)
               and len(block.content) == 1 and isinstance(block.content[0], Image)]
    table = next(block for block in section.blocks if isinstance(block, Table))
    if fault == 'missing':
        section.blocks.remove(vectors[0])
    elif fault == 'duplicate':
        section.blocks.append(deepcopy(vectors[0]))
    elif fault == 'dash':
        source_model.resources[vectors[0].content[0].resource_id].properties['dashes'] = '[2 2] 0'
    elif fault == 'unrelated':
        resource = source_model.resources[vectors[0].content[0].resource_id]
        resource.properties['items'][0][1][0] = 71
        resource.properties['items'][0][2][0] = 71
        vectors[0].content[0].box.x = 71
    elif fault == 'ambiguous':
        section.blocks.append(deepcopy(table))
    elif fault == 'edited-border':
        set_borders(table, BorderSet({'top': Border(kind='line', style='solid', width_pt=.5)}))
    else:
        table.rows[0].cells = table.rows[0].cells[:2]
    before = document_to_json(source_model)
    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous')
    report = write_document(source_model, output, options=ExportOptions(docx_profile=DocxProfile('flow')))
    assert not report.success and any(issue.feature == 'docx.flow-profile' for issue in report.issues)
    assert output.read_bytes() == b'previous' and document_to_json(source_model) == before


def test_quality_gate_and_limits_keep_previous_destination(source_model, tmp_path):
    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous')
    report = write_document(source_model, output, options=ExportOptions(
        docx_profile=DocxProfile('flow'), quality_gate=lambda report, path: report.lossless,
    ))
    assert not report.success and any(issue.feature == 'export.quality-rejected' for issue in report.issues)
    assert output.read_bytes() == b'previous'
    report = write_document(source_model, output, options=ExportOptions(
        docx_profile=DocxProfile('flow'), document_limits=DocumentLimits(max_nodes=2),
    ))
    assert not report.success and output.read_bytes() == b'previous'


@pytest.mark.parametrize('public', [False, True])
def test_cancel_during_grid_matching(source_model, tmp_path, monkeypatch, public):
    from opendoc_formats.writers import docx_flow

    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous')
    original = document_to_json(source_model)
    stop = Event()
    calls = []
    native = docx_flow._line

    def first(*args):
        assert not calls, 'Matching continued after cancellation'
        result = native(*args)
        calls.append(True)
        stop.set()
        return result

    monkeypatch.setattr(docx_flow, '_line', first)
    if public:
        report = write_document(source_model, output, options=ExportOptions(
            docx_profile=DocxProfile('flow'), cancelled=stop.is_set,
        ))
        assert not report.success and any(issue.feature == 'export.cancelled' for issue in report.issues)
    else:
        with pytest.raises(OperationCancelledError):
            write_docx_model(source_model, output, profile=DocxProfile('flow'), cancelled=stop.is_set)
    assert document_to_json(source_model) == original and output.read_bytes() == b'previous'
    assert calls == [True]


def test_invalid_profile_is_explicit(tmp_path):
    with pytest.raises(ValueError):
        DocxProfile('unknown')
    with pytest.raises(ValueError):
        ExportOptions(docx_profile='flow')
    report = write_docx_model(DocumentModel(), tmp_path / 'invalid.docx', profile=DocxProfile('flow'))
    assert not report.success


def test_profile_does_not_shift_existing_positional_options():
    gate = lambda report, path: True  # noqa: E731
    options = ExportOptions(DocumentLimits(), None, True, None, gate)
    assert options.quality_gate is gate and options.docx_profile is None


@pytest.mark.parametrize(('key', 'value'), [
    ('pdf_paint_type', 3), ('pdf_paint_opacity', .5), ('pdf_paint_unknown', True),
    ('pdf_graphics_state_unknown', True),
])
def test_unknown_or_nonopaque_text_is_refused_atomically(source_model, tmp_path, key, value):
    from opendoc_model import walk_model

    run = next(ref for ref in walk_model(source_model) if isinstance(ref.node, TextRun))
    run.node.properties[key] = value
    before = document_to_json(source_model)
    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous')
    report = write_document(source_model, output, options=ExportOptions(docx_profile=DocxProfile('flow')))
    assert not report.success
    assert any(issue.feature == 'docx.flow-profile' and issue.location == run.path for issue in report.issues)
    assert document_to_json(source_model) == before and output.read_bytes() == b'previous'


def test_border_width_outside_native_word_enum_is_refused(source_model, tmp_path):
    image = next(block.content[0] for block in source_model.sections[0].blocks if isinstance(block, Paragraph)
                 and len(block.content) == 1 and isinstance(block.content[0], Image))
    source_model.resources[image.resource_id].properties['width'] = .125
    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous')
    before = document_to_json(source_model)
    report = write_document(source_model, output, options=ExportOptions(docx_profile=DocxProfile('flow')))
    assert not report.success and any(issue.feature == 'docx.flow-profile' for issue in report.issues)
    assert document_to_json(source_model) == before and output.read_bytes() == b'previous'


def test_font_and_native_diagnostics_keep_canonical_source_locations(source_model, tmp_path):
    vectors = [block for block in source_model.sections[0].blocks if isinstance(block, Paragraph)
               and len(block.content) == 1 and isinstance(block.content[0], Image)]
    rest = [block for block in source_model.sections[0].blocks if block not in vectors]
    after = Paragraph([TextRun('After grid', TextStyle(font_family='OwnMissingMonoSource'))])
    source_model.sections[0].blocks = [*vectors, *rest, after]
    report = write_document(source_model, tmp_path / 'located.docx', options=ExportOptions(docx_profile=DocxProfile('flow')))
    assert report.success, report.issues
    assert any(issue.feature == 'font-substitution' and issue.location == 'sections[0].blocks[10].content[0]'
               for issue in report.issues)
    invalid = Table([TableRow([TableCell([Paragraph([TextRun('Bad native border')])])])])
    set_borders(invalid, BorderSet({'top': Border(kind='line', style='solid', width_pt=.3)}))
    source_model.sections[0].blocks[-1] = invalid
    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous')
    report = write_document(source_model, output, options=ExportOptions(docx_profile=DocxProfile('flow')))
    assert not report.success
    assert any(issue.severity is IssueSeverity.ERROR and issue.location == 'sections[0].blocks[10]' for issue in report.issues)
    assert output.read_bytes() == b'previous'


def test_signal_after_direct_commit_does_not_claim_rollback(source_model, tmp_path, monkeypatch):
    from pathlib import Path

    from opendoc_formats.support import io

    output = tmp_path / 'committed.docx'
    stop = Event()
    native = io.os.replace

    def commit(source, destination):
        result = native(source, destination)
        if Path(destination) == output:
            stop.set()
        return result

    monkeypatch.setattr(io.os, 'replace', commit)
    report = write_docx_model(source_model, output, profile=DocxProfile('flow'), cancelled=stop.is_set)
    assert stop.is_set() and report.success
    assert get_object_mapping(report).publication == 'published'
    assert Document(output).tables[0].cell(0, 0).text == 'R1C1'
