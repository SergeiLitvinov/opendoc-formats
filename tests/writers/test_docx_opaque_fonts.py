"""Retained OOXML must not lose glyphs when native font programs are embedded."""

from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fontTools.ttLib import TTFont
from opendoc_model import ResourceKind, document_from_json, document_to_json, iter_font_references

from opendoc_formats import read_document, write_document


def test_smart_tag_keeps_font_coverage_and_known_loss_in_two_json_cycles(tmp_path):
    source = tmp_path / 'own.docx'
    native = Document()
    paragraph = native.add_paragraph('Before ')
    wrapper = OxmlElement('w:smartTag')
    wrapper.set(qn('w:uri'), 'urn:opendoc-formats:own-acceptance')
    wrapper.set(qn('w:element'), 'OwnToken')
    run, text = OxmlElement('w:r'), OxmlElement('w:t')
    text.text = 'OwnSmartToken'
    run.append(text)
    wrapper.append(run)
    paragraph._p.append(wrapper)
    paragraph.add_run(' After')
    native.save(source)
    original = source.read_bytes()
    current = source
    for cycle in range(2):
        imported = read_document(current)
        assert imported.success
        assert sum(issue.code == 'docx.smart-tags' for issue in imported.issues) == 1
        model = document_from_json(document_to_json(imported.document))
        before = document_to_json(model)
        current = tmp_path / f'cycle-{cycle}.docx'
        report = write_document(model, current)
        assert report.success, report.issues
        assert not any(issue.feature == 'font-glyphs' for issue in report.issues)
        assert report.metrics['font_embedding']['subset_policy'] == 'full-for-retained-ooxml'
        assert report.metrics['font_embedding']['subset_bytes'] == report.metrics['font_embedding']['source_bytes']
        assert document_to_json(model) == before
        restored = read_document(current).document
        fonts = [resource for resource in restored.resources.values() if resource.kind is ResourceKind.FONT]
        assert fonts
        assert list(iter_font_references(restored))
        for resource in fonts:
            with TTFont(BytesIO(resource.data)) as font:
                assert set(map(ord, 'Before OwnSmartToken After')) <= set(font.getBestCmap())
        raw = Document(current).paragraphs[0]._p
        retained = raw.find(qn('w:smartTag'))
        assert retained is not None and ''.join(node.text or '' for node in retained.iter(qn('w:t'))) == 'OwnSmartToken'
    assert source.read_bytes() == original


def test_plain_document_keeps_semantic_subsetting_policy(tmp_path):
    source = tmp_path / 'plain.docx'
    native = Document()
    native.add_paragraph('Own ordinary text')
    native.save(source)
    model = read_document(source).document
    report = write_document(model, tmp_path / 'output.docx')
    assert report.success
    assert report.metrics['font_embedding']['subset_policy'] == 'semantic-usage'
