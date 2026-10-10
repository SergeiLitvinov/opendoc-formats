"""Own finite opaque Office objects keep source relationships through two cycles."""

import zipfile

import pytest
from docx import Document
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from lxml import etree

from opendoc_formats import read_document, write_document

R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


@pytest.mark.parametrize('kind', ['ole', 'smartart'])
def test_opaque_object_parts_and_reference_targets_survive_two_cycles(tmp_path, kind):
    document = Document()
    package = document.part.package
    expected = {}

    def part(name, media_type, data, relationship):
        expected[name.lstrip('/')] = data
        item = Part(PackURI(name), media_type, data, package)
        return document.part.relate_to(item, relationship)

    if kind == 'ole':
        preview = part('/word/media/own.emf', 'image/x-emf', b'own inert preview', R + '/image')
        payload = part('/word/embeddings/own.xlsx', 'application/octet-stream', b'own inert payload', R + '/package')
        content = (f'<w:object xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
                   f'<v:shape><v:imagedata r:id="{preview}"/></v:shape>'
                   f'<o:OLEObject Type="Embed" r:id="{payload}"/></w:object>')
    else:
        ids = [part('/word/diagrams/' + name + '.xml', 'application/xml', b'<own/>', R + '/diagram' + rel)
               for name, rel in [('data', 'Data'), ('layout', 'Layout'), ('quickStyle', 'QuickStyle'), ('colors', 'Colors')]]
        part('/word/diagrams/drawing.xml', 'application/xml', b'<ownDrawing/>',
             'http://schemas.microsoft.com/office/2007/relationships/diagramDrawing')
        content = ('<w:drawing><d:relIds xmlns:d="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
                   + ' '.join(f'r:{key}="{value}"' for key, value in zip(('dm', 'lo', 'qs', 'cs'), ids))
                   + '/></w:drawing>')
    document.add_paragraph()._p.append(parse_xml(
        f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="{R}">{content}</w:r>'
    ))
    source = tmp_path / 'source.docx'
    document.save(source)
    original = source.read_bytes()
    original_doc = Document(source)
    targets = {rel.rId: rel.target_ref for rel in original_doc.part.rels.values()}
    path = source
    for cycle in range(2):
        result = read_document(path)
        assert result.success
        transport = tmp_path / f'cycle-{cycle}.json'
        assert write_document(result.document, transport).success
        path = transport.with_suffix('.docx')
        assert write_document(read_document(transport).document, path).success
        reopened = Document(path)
        for node in reopened.element.iter():
            for key, value in node.attrib.items():
                if key.startswith('{' + R + '}'):
                    assert reopened.part.rels[value].target_ref == targets[value]
        with zipfile.ZipFile(path) as archive:
            assert all(archive.read(name) == data for name, data in expected.items())
            etree.fromstring(archive.read('word/document.xml'))
    assert source.read_bytes() == original


@pytest.mark.parametrize('missing_graph', [False, True])
def test_unretained_opaque_reference_refuses_without_replacing_file(tmp_path, missing_graph):
    from opendoc_model import DocumentModel, PackageGraph, Paragraph, Section, TextRun

    raw = (f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="{R}" '
           'xmlns:o="urn:schemas-microsoft-com:office:office"><w:object>'
           '<o:OLEObject r:id="rIdMissing"/></w:object></w:r>')
    model = DocumentModel(sections=[Section(blocks=[Paragraph(content=[
        TextRun(text='', properties={'docx_raw_inline_xml': raw}),
    ])])], package=PackageGraph(format='ooxml', root='/word/document.xml'))
    if missing_graph:
        model.package = None
    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous file')
    result = write_document(model, output)
    assert not result.success
    assert any(issue.feature == 'package-graph' for issue in result.issues)
    assert output.read_bytes() == b'previous file'
