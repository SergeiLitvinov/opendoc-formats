"""Editable native group structure, authoritative SDK edits and atomic refusal."""

from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree
from opendoc_model import (
    AffineTransform,
    Anchor,
    Box,
    ColorValue,
    DocumentModel,
    GroupColorSpace,
    GroupCompositing,
    Image,
    Length,
    PageSettings,
    Paint,
    Paragraph,
    PathCommand,
    Resource,
    ResourceKind,
    Section,
    VectorGroup,
    VectorPath,
    VectorScene,
    document_from_json,
    document_to_json,
    get_integration,
    get_vector_scene,
    remove_vector_scene,
    set_anchor,
    set_scene_clip,
    set_scene_group,
    set_scope_color_space,
    set_vector_scene,
)

from opendoc_formats import read_document, write_document
from tests.writers.test_pdf_docx_graphics import grouped_form


def rectangle(key, x, y, width, height, color, opacity=1):
    commands = tuple(PathCommand(op, values) for op, values in [('move', (x, y)), ('line', (x+width, y)),
                     ('line', (x+width, y+height)), ('line', (x, y+height)), ('close', ())])
    return VectorPath(key, commands, fill=Paint(ColorValue.from_hex(color), opacity))


def scene_document():
    neutral = GroupCompositing(1, False, False, 'normal', color_space=GroupColorSpace('rgb'))
    image = Image('preview', box=Box(0, 0, 400, 300), properties={
        'placement': 'anchor', 'wrap': 'none', 'horizontal_relative_from': 'page', 'vertical_relative_from': 'page',
    })
    set_anchor(image, Anchor('native-group-node'))
    page = PageSettings(Length(400), Length(300), Length(0), Length(0), Length(0), Length(0))
    model = DocumentModel(source_format='docx', sections=[Section(blocks=[Paragraph(content=[image])], page=page)],
                          resources={'preview': Resource('preview', ResourceKind.VECTOR_IMAGE, 'image/svg+xml',
                                     data=b'<svg xmlns="http://www.w3.org/2000/svg"><rect fill="black"/></svg>')})
    scene = VectorScene('native-proof', 400, 300,
                        paths=(rectangle('Backdrop', 0, 0, 400, 300, 'E6D999'),
                               rectangle('Red', 80, 100, 110, 100, 'FF0000'),
                               rectangle('Blue', 130, 100, 110, 100, '0000FF', .5)),
                        groups=(VectorGroup('Pair', ('Red', 'Blue'), compositing=neutral),),
                        roots=('Backdrop', 'Pair'), node_id='native-group-node', scope_version=1)
    set_vector_scene(model, scene)
    return model


def scene_summary(scene):
    return ([(p.extra.get('docx_name', p.id), p.commands, p.fill) for p in scene.paths],
            [(g.extra.get('docx_name', g.id), g.transform, g.compositing,
              [next(p.extra.get('docx_name', p.id) for p in scene.paths if p.id == child) for child in g.children])
             for g in scene.groups])


def test_native_groups_and_individual_paints_survive_two_public_cycles(tmp_path):
    model = scene_document()
    original = document_to_json(model)
    destination = tmp_path / 'group.docx'
    assert write_document(model, destination).success
    assert document_to_json(model) == original
    previous = None
    for index in range(2):
        with ZipFile(destination) as package:
            xml = etree.fromstring(package.read('word/document.xml'))
            assert len(xml.xpath('//*[local-name()="wgp"]')) == 1
            assert len(xml.xpath('//*[local-name()="wpc"]')) == 1
            assert not xml.xpath('//*[local-name()="grpSp"]')
            assert len(xml.xpath('//*[local-name()="wsp"]')) == 3
            assert not xml.xpath('//*[local-name()="blip" or local-name()="effectDag"]')
            assert not any(n.startswith('word/media/') for n in package.namelist())
        model = document_from_json(document_to_json(read_document(destination).document))
        scene = get_integration(model).scenes[0]
        assert scene.id == 'native-proof' and scene.node_id
        summary = scene_summary(scene)
        assert previous is None or summary == previous
        previous = summary
        assert [p.fill.opacity for p in scene.paths] == [1, 1, .5]
        assert scene.groups[0].compositing.color_space == GroupColorSpace('rgb')
        assert 'docx_group_pending' not in document_to_json(model)
        destination = tmp_path / f'cycle-{index}.docx'
        assert write_document(model, destination).success


@pytest.mark.parametrize('edit', ['opacity', 'blend', 'knockout', 'isolation', 'unknown', 'gray', 'remove-space', 'icc'])
def test_unsupported_group_edits_never_fall_back_to_svg_preview(tmp_path, edit):
    model = scene_document()
    scene = get_vector_scene(model, 'native-proof')
    group = scene.groups[0]
    if edit in {'gray', 'remove-space'}:
        set_scope_color_space(model, scene.id, group.id, GroupColorSpace('gray') if edit == 'gray' else None)
    elif edit == 'icc':
        model.add_resource(Resource('profile', ResourceKind.ATTACHMENT, 'application/vnd.iccprofile', data=b'own-test'))
        set_scope_color_space(model, scene.id, group.id, GroupColorSpace('icc', 'profile'))
    else:
        changes = {'opacity': {'opacity': .5}, 'blend': {'blend_mode': 'multiply'}, 'knockout': {'knockout': True},
                   'isolation': {'isolated': True}, 'unknown': {'opacity': None}}[edit]
        set_scene_group(model, scene.id, group.id, replace(group, compositing=replace(group.compositing, **changes)))
    before = document_to_json(model)
    destination = tmp_path / 'previous.docx'
    destination.write_bytes(b'previous')
    result = write_document(model, destination)
    assert not result.success and destination.read_bytes() == b'previous'
    assert any(i.feature == 'docx.graphics-group' for i in result.issues)
    assert document_to_json(model) == before


def test_group_removal_releases_children_without_resurrecting_old_group(tmp_path):
    model = scene_document()
    set_scene_group(model, 'native-proof', 'Pair', None)
    destination = tmp_path / 'unwrapped.docx'
    assert write_document(model, destination).success
    restored = read_document(destination).document
    scene = get_integration(restored).scenes[0]
    assert not scene.groups and len(scene.paths) == 3 and len(scene.roots) == 3


def test_removed_scene_does_not_resurrect_derived_preview(tmp_path):
    destination = tmp_path / 'group.docx'
    assert write_document(scene_document(), destination).success
    model = read_document(destination).document
    remove_vector_scene(model, 'native-proof')
    before = destination.read_bytes()
    assert not write_document(model, destination).success
    assert destination.read_bytes() == before


def test_native_child_order_color_and_group_translation_are_authoritative(tmp_path):
    model = scene_document()
    scene = get_vector_scene(model, 'native-proof')
    paths = tuple(replace(p, fill=Paint(ColorValue.from_hex('00FF00'), .5)) if p.id == 'Blue' else p for p in scene.paths)
    group = replace(scene.groups[0], children=('Blue', 'Red'), transform=AffineTransform(e=10, f=5))
    set_vector_scene(model, replace(scene, paths=paths, groups=(group,)))
    destination = tmp_path / 'edited.docx'
    assert write_document(model, destination).success
    restored = get_integration(read_document(destination).document).scenes[0]
    pair = restored.groups[0]
    assert pair.transform == AffineTransform(e=10, f=5)
    names = {p.id: p.extra['docx_name'] for p in restored.paths}
    assert [names[c] for c in pair.children] == ['Blue', 'Red']
    assert next(p.fill.color for p in restored.paths if p.extra['docx_name'] == 'Blue') == ColorValue.from_hex('00FF00')


@pytest.mark.parametrize('extension', ['html', 'pdf', 'txt'])
def test_other_targets_refuse_instead_of_using_stale_preview(tmp_path, extension):
    destination = tmp_path / ('previous.'+extension)
    destination.write_bytes(b'previous')
    result = write_document(scene_document(), destination)
    assert not result.success and destination.read_bytes() == b'previous'
    assert any(i.feature == 'export.graphics-scopes' for i in result.issues)


def test_pdf_group_explicitly_edited_into_native_profile_keeps_structure(tmp_path):
    source, destination = tmp_path / 'source.pdf', tmp_path / 'native.docx'
    grouped_form(source, opacity=.5, isolated=False)
    model = read_document(source).document
    scene = get_vector_scene(model, model.sections[0].properties['pdf_scene_id'])
    group = next(g for g in scene.groups if g.compositing.opacity == .5)
    set_scene_group(model, scene.id, group.id, replace(group, compositing=replace(group.compositing, opacity=1)))
    # Explicit user edits, not silent dropping of a Form's persistent clip.
    for clip in scene.clips:
        set_scene_clip(model, scene.id, clip.id, None)
    original = document_to_json(model)
    result = write_document(model, destination)
    assert result.success, result.issues
    assert document_to_json(model) == original
    for index in range(2):
        with ZipFile(destination) as package:
            xml = etree.fromstring(package.read('word/document.xml'))
            assert len(xml.xpath('//*[local-name()="wgp"]//*[local-name()="wsp"]')) == 2
        restored = document_from_json(document_to_json(read_document(destination).document))
        scene = get_integration(restored).scenes[0]
        assert len(scene.paths) == 2 and len(scene.groups) == 1
        assert [p.fill.opacity for p in scene.paths] == [1, .5]
        destination = tmp_path / f'native-{index}.docx'
        assert write_document(restored, destination).success


def native_word_document(destination, mutate=None):
    from docx import Document
    from docx.shared import Pt

    drawing = etree.parse(str(Path(__file__).parents[1] / 'corpus/native/word-canvas-group.xml')).getroot()
    if mutate:
        mutate(drawing)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Pt(400), Pt(300)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(0)
    wrapper = etree.Element('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
    wrapper.append(drawing)
    document.add_paragraph().add_run()._r.append(wrapper)
    document.save(destination)


def test_independent_word_canvas_retains_background_children_and_native_group(tmp_path):
    destination = tmp_path / 'word.docx'
    native_word_document(destination)
    before = None
    for index in range(3):
        model = document_from_json(document_to_json(read_document(destination).document))
        scene = get_integration(model).scenes[0]
        assert len(scene.groups) == 1 and scene.groups[0].extra['docx_name'] == 'OwnPair'
        assert [p.fill.opacity for p in scene.paths] == [1, 1, .5]
        assert scene.paths[0].extra['docx_canvas_background']
        summary = scene_summary(scene)
        assert before is None or summary == before
        before = summary
        destination = tmp_path / f'word-cycle-{index}.docx'
        assert write_document(model, destination).success
        with ZipFile(destination) as package:
            assert not any(n.startswith('word/media/') for n in package.namelist())


@pytest.mark.parametrize('edit', ['extension', 'adjustment', 'effect'])
def test_unverified_native_word_properties_are_not_silently_lost(tmp_path, edit):
    def mutate(root):
        ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        if edit == 'extension':
            root.xpath('//*[local-name()="ext" and @uri]')[0].set('uri', 'unverified')
        elif edit == 'adjustment':
            etree.SubElement(root.xpath('//*[local-name()="avLst"]')[0], f'{{{ns}}}gd', name='own', fmla='val 50')
        else:
            etree.SubElement(root.xpath('//*[local-name()="effectLst"]')[0], f'{{{ns}}}blur', rad='100')
    destination = tmp_path / 'unverified.docx'
    native_word_document(destination, mutate)
    with pytest.raises(ValueError, match='Unverified|Adjusted|effect'):
        read_document(destination)


def test_nested_group_xml_is_not_claimed_as_independently_editable_in_word(tmp_path):
    model = scene_document()
    scene = get_vector_scene(model, 'native-proof')
    outer = replace(scene.groups[0], id='outer', children=('Pair',))
    set_vector_scene(model, replace(scene, groups=(*scene.groups, outer), roots=('Backdrop', 'outer')))
    destination = tmp_path / 'nested.docx'
    destination.write_bytes(b'previous')
    result = write_document(model, destination)
    assert not result.success and destination.read_bytes() == b'previous'
    assert any('Nested native groups' in issue.message for issue in result.issues)
