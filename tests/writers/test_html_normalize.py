"""Нативные закладки Word становятся внутренней навигацией HTML."""

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from opendoc_formats.readers.docx import read_docx_model
from opendoc_formats.writers.html_writer import write_html_model


def build_bookmark_docx(path):
    document = Document()
    paragraph = document.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "target_section")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Jump to target"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    for index in range(35):
        document.add_paragraph(f"Body paragraph {index}")
    paragraph = document.add_paragraph("Before bookmark. ")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "7")
    start.set(qn("w:name"), "target_section")
    paragraph._p.append(start)
    paragraph.add_run("Target text")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "7")
    paragraph._p.append(end)
    document.save(path)
    return path


def test_native_docx_bookmark_and_link_reach_html(tmp_path):
    from opendoc_model.document_codec import document_to_dict

    source = build_bookmark_docx(tmp_path / "source.docx")
    model = read_docx_model(source)
    before = document_to_dict(model)
    output = tmp_path / "result.html"
    report = write_html_model(model, output)
    assert report.success
    assert report.metrics["html_navigation"]["internal_links"] == 1
    assert report.metrics["html_navigation"]["verified"] is True
    html = output.read_text(encoding="utf-8")
    assert 'href="#target_section"' in html
    assert '<span id="target_section"></span>' in html
    assert document_to_dict(model) == before


def test_navigation_and_inline_position_survive_two_serialized_cycles(tmp_path):
    from opendoc_model.document_codec import document_from_json, document_to_json

    from opendoc_formats.readers.html_text import read_html_model
    from opendoc_formats.writers.docx_writer import write_docx_model

    model = read_docx_model(build_bookmark_docx(tmp_path / "source.docx"))
    for index in range(2):
        output = tmp_path / f"cycle-{index}.html"
        report = write_html_model(model, output)
        assert report.success and report.metrics["html_navigation"]["verified"]
        assert report.metrics["html_navigation"]["internal_links"] == 1
        imported = document_from_json(document_to_json(read_html_model(output)))
        target_paragraph = next(block for block in imported.sections[0].blocks if "Before bookmark." in block.plain_text)
        marker = next(i for i, item in enumerate(target_paragraph.content) if item.properties.get("anchor_id"))
        assert target_paragraph.content[marker - 1].text == "Before bookmark. "
        assert target_paragraph.content[marker + 1].text == "Target text"
        docx = tmp_path / f"cycle-{index}.docx"
        assert write_docx_model(imported, docx).success
        saved = Document(docx)
        starts = saved.element.xpath(".//w:bookmarkStart")
        link = saved.element.xpath(".//w:hyperlink")[0]
        assert len(starts) == 1 and starts[0].get(qn("w:name")) == link.get(qn("w:anchor"))
        children = list(starts[0].getparent())
        position = children.index(starts[0])
        assert "Before bookmark." in "".join(node.text or "" for child in children[:position] for node in child.iter(qn("w:t")))
        assert "Target text" in "".join(node.text or "" for child in children[position:] for node in child.iter(qn("w:t")))
        model = read_docx_model(docx)




def test_external_link_and_native_word_data_remain_unchanged(tmp_path):
    from opendoc_model.document_model import DocumentModel, Paragraph, Section, TextRun

    from opendoc_formats.writers.html_normalize import HtmlNormalizeStage
    from opendoc_formats.writers.stages import StageContext

    run = TextRun("External", link="https://example.com/path#fragment", properties={"hyperlink_anchor": "fragment"})
    model = DocumentModel(sections=[Section(blocks=[Paragraph([run])])])
    result = HtmlNormalizeStage().execute(model, StageContext(tmp_path / "out.html"))
    prepared = result.value.sections[0].blocks[0].content[0]
    assert prepared.link == run.link == prepared.properties["html_normalize"]["link"]
    assert prepared.properties["hyperlink_anchor"] == "fragment"
    assert result.report.metrics["html_normalization"]["word_internal_links"] == 0
    assert "html_normalize" not in run.properties


@pytest.mark.parametrize(
    "content, text",
    [
        ('<span id="point"></span> Text ', "Text"),
        ('Before <span id="point"></span> after', "Before after"),
        ('<span id="point"></span>', ""),
    ],
)
def test_empty_inline_target_survives_html_import_without_whitespace_changes(tmp_path, content, text):
    from opendoc_formats.readers.html_text import read_html_model

    source = tmp_path / "source.html"
    source.write_text('<p><a href="#point">Go</a></p><p>' + content + "</p>", encoding="utf-8")
    model = read_html_model(source)
    assert len(model.sections[0].blocks) == 2
    assert model.sections[0].blocks[1].plain_text == text
    report = write_html_model(model, tmp_path / "result.html")
    assert report.success and report.metrics["html_navigation"]["verified"]


@pytest.mark.parametrize("existing", [False, True])
def test_cancelled_normalization_preserves_output(tmp_path, monkeypatch, existing):
    from opendoc_model.document_model import DocumentModel

    from opendoc_formats.writers.html_normalize import HtmlNormalizeStage
    from opendoc_formats.writers.stages import StageContext

    execute = HtmlNormalizeStage.execute

    def cancel(self, value, context):
        return execute(self, value, StageContext(context.output_path, cancelled=lambda: True))

    monkeypatch.setattr(HtmlNormalizeStage, "execute", cancel)
    output = tmp_path / "result.html"
    if existing:
        output.write_bytes(b"previous")
    report = write_html_model(DocumentModel(), output)
    assert not report.success and report.issues[0].feature == "cancelled"
    if existing:
        assert output.read_bytes() == b"previous"
    else:
        assert not output.exists()


def test_normalization_traverses_tables_and_rendered_headers(tmp_path):
    from opendoc_model.document_model import DocumentModel, Paragraph, Section, Table, TableCell, TableRow, TextRun

    from opendoc_formats.writers.html_normalize import HtmlNormalizeStage
    from opendoc_formats.writers.stages import StageContext

    marker = TextRun("", properties={"bookmark_start": {"id": "1", "name": "inside_table"}})
    link = TextRun("Go", properties={"hyperlink_anchor": "inside_table"})
    table = Table([TableRow([TableCell([Paragraph([marker, TextRun("Target")])])])])
    model = DocumentModel(sections=[Section(blocks=[table], headers=[Paragraph([link])])])
    result = HtmlNormalizeStage().execute(model, StageContext(tmp_path / "out.html"))
    assert result.report.metrics["html_normalization"] == {
        "stage": "html.normalize",
        "word_bookmarks": 1,
        "word_internal_links": 1,
        "inline_anchors": 1,
    }
    prepared_marker = result.value.sections[0].blocks[0].rows[0].cells[0].blocks[0].content[0]
    assert prepared_marker.properties["html_normalize"]["anchor_id"] == "inside_table"
    assert result.value.sections[0].headers[0].content[0].properties["html_normalize"]["link"] == "#inside_table"
    assert "html_normalize" not in marker.properties and link.link is None


def test_model_anchor_is_encoded_once_and_escaped_in_html(tmp_path):
    from opendoc_model.document_model import DocumentModel, Paragraph, Section, TextRun

    name = 'Раздел%20&"'
    model = DocumentModel(
        sections=[
            Section(
                blocks=[
                    Paragraph(
                        [
                            TextRun("Go", properties={"hyperlink_anchor": name}),
                            TextRun("", properties={"bookmark_start": {"id": "1", "name": name}}),
                            TextRun("Target"),
                        ]
                    )
                ]
            )
        ]
    )
    output = tmp_path / "out.html"
    report = write_html_model(model, output)
    assert report.success and report.metrics["html_navigation"]["verified"]
    html = output.read_text(encoding="utf-8")
    assert "%2520%26%22" in html
    assert 'id="Раздел%20&amp;&quot;"' in html


def test_unnamed_word_bookmark_reports_loss(tmp_path):
    from opendoc_model.document_model import DocumentModel, Paragraph, Section, TextRun

    model = DocumentModel(
        sections=[
            Section(
                blocks=[
                    Paragraph(
                        [
                            TextRun("", properties={"bookmark_start": {"id": "1"}}),
                            TextRun("Text"),
                        ]
                    )
                ]
            )
        ]
    )
    report = write_html_model(model, tmp_path / "out.html")
    assert report.success and not report.lossless
    issue = next(issue for issue in report.issues if issue.feature == "bookmark")
    assert issue.location == "sections[0].blocks[0].content[0]"
