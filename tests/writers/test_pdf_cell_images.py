"""Own DOCX cell pictures: bounded native PDF placement, captions and full pixels."""

import math
from io import BytesIO
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from opendoc_model import (
    ConversionReport,
    DocumentModel,
    Image,
    ImageCrop,
    Paragraph,
    Table,
    TableCell,
    TableRow,
    document_from_json,
    document_to_json,
)

from opendoc_formats import read_document, write_document
from tests.writers.test_pdf_docx_rasters import own_pixels


def test_owned_rgba_fixture_is_reproducible():
    assert (Path(__file__).resolve().parents[1]/'native/data/table-image.png').read_bytes() == own_pixels()


def own_cell_document(path, form='plain', alignment='left', lines=8):
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Pt(240), Pt(300)
    page.left_margin = page.right_margin = page.top_margin = page.bottom_margin = Pt(20)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for index in range(lines):
        document.add_paragraph(f'Own preceding line {index}')
    table = document.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    if form == 'merged':
        cell = table.cell(0, 0).merge(table.cell(0, 1))
    else:
        table.cell(0, 1).text = 'Own neighbor text'
        cell = table.cell(0, 0)
        if form == 'nested':
            cell = cell.add_table(rows=1, cols=1).cell(0, 0)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = getattr(WD_ALIGN_PARAGRAPH, alignment.upper())
    paragraph.add_run().add_picture(BytesIO(own_pixels()), width=Pt(72), height=Pt(72))
    cell.add_paragraph('Own image caption')
    document.add_paragraph('Following table text must remain')
    document.save(path)
    result = read_document(path)
    assert result.success, result.issues
    return result.document


def cell_images(table):
    for row in table.rows:
        for cell in row.cells:
            for block in cell.blocks:
                if isinstance(block, Paragraph):
                    yield from (image for image in block.content if isinstance(image, Image))
                elif isinstance(block, Table):
                    yield from cell_images(block)


@pytest.mark.parametrize('form', ['plain', 'merged', 'nested'])
@pytest.mark.parametrize('alignment', ['left', 'center', 'right'])
def test_cell_image_crop_rotation_pixels_and_text_survive_two_json_cycles(tmp_path, form, alignment):
    model = own_cell_document(tmp_path/'source.docx', form, alignment)
    table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
    image = next(cell_images(table))
    image.box.rotation = 3
    image.crop = ImageCrop(left=.1, top=.2, right=.1, bottom=.1)
    before = document_to_json(model)
    reference = None
    for cycle in range(2):
        current = document_from_json(document_to_json(model))
        path = tmp_path/f'cell-{cycle}.pdf'
        report = write_document(current, path)
        assert report.success, report.issues
        assert any(issue.feature == 'pdf-table-flow-layout' for issue in report.issues)
        assert document_to_json(current) == before
        with pymupdf.open(path) as pdf:
            text = ' '.join(''.join(page.get_text() for page in pdf).split())
            assert text.count('Own image caption') == 1
            assert text.count('Following table text must remain') == 1
            if form != 'merged':
                assert text.count('Own neighbor text') == 1
            rasters = [(number, page, picture) for number, page in enumerate(pdf) for picture in page.get_image_info()]
            assert len(rasters) == 1
            _, page, picture = rasters[0]
            matrix = pymupdf.Matrix(picture['transform'])
            corners = [pymupdf.Point(x, y)*matrix for x, y in ((.1, .2), (.9, .2), (.9, .9), (.1, .9))]
            visible = pymupdf.Rect(min(p.x for p in corners), min(p.y for p in corners),
                                  max(p.x for p in corners), max(p.y for p in corners))
            expected = 72*(math.cos(math.radians(3))+math.sin(math.radians(3)))
            assert visible.width == pytest.approx(expected, abs=.01)
            assert visible.height == pytest.approx(expected, abs=.01)
            assert visible.x0 >= 20 and visible.y0 >= 20
            assert visible.x1 <= 220 and visible.y1 <= 280
            assert 'Own image caption' in ' '.join(page.get_text().split())
            for word in page.get_text('words'):
                assert not (visible & pymupdf.Rect(word[:4])).get_area()
            native = page.get_images(full=True)[0]
            pixels = pymupdf.Pixmap(pdf, native[0])
            if native[1]:
                pixels = pymupdf.Pixmap(pixels, pymupdf.Pixmap(pdf, native[1]))
            original = pymupdf.Pixmap(own_pixels())
            assert (pixels.width, pixels.height, pixels.samples) == (original.width, original.height, original.samples)
            rendered = [page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).samples for page in pdf]
            assert reference is None or rendered == reference
            reference = rendered


@pytest.mark.parametrize('form', ['plain', 'merged', 'nested'])
def test_too_tall_cell_image_never_overlaps_page_margins_without_text_guard(tmp_path, form):
    model = own_cell_document(tmp_path/'source.docx', form, lines=0)
    table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
    image = next(cell_images(table))
    image.box.height = 260
    # No text occurrence can reveal a clipped/overflowing raster in this case.
    def remove_text(value):
        for row in value.rows:
            for cell in row.cells:
                for index, block in enumerate(cell.blocks):
                    if isinstance(block, Table):
                        remove_text(block)
                    else:
                        cell.blocks[index] = Paragraph([part for part in block.content if isinstance(part, Image)])

    remove_text(table)
    model.sections[0].blocks = [table]
    before = document_to_json(model)
    path = tmp_path/'previous.pdf'
    path.write_bytes(b'previous artifact')
    report = write_document(model, path)
    assert not report.success and path.read_bytes() == b'previous artifact'
    assert any('overflows' in issue.message or 'content area' in issue.message for issue in report.issues)
    assert document_to_json(model) == before


@pytest.mark.parametrize('form', ['plain', 'horizontal', 'vertical'])
def test_measured_column_ratios_follow_native_spans_and_do_not_change_html(form):
    from opendoc_formats.writers.html_writer import _HtmlRenderer
    from opendoc_formats.writers.pdf_flow import PdfFlowRenderer

    rows = [TableRow([TableCell(), TableCell()])]
    if form == 'horizontal':
        rows.append(TableRow([TableCell(column_span=2)]))
    elif form == 'vertical':
        rows[0].cells[0].row_span = 2
        rows.append(TableRow([TableCell()]))
    table = Table(rows, properties={'grid_widths_twips': [1000, 3000]})
    renderer = PdfFlowRenderer(DocumentModel(), ConversionReport(None), {})
    assert renderer._table_cell_styles(table, 0, 0) == ['width:25.0000000000%']
    assert renderer._table_cell_styles(table, 0, 1) == ['width:75.0000000000%']
    if form != 'plain':
        assert renderer._table_cell_styles(table, 1, 0) == [
            'width:100.0000000000%' if form == 'horizontal' else 'width:75.0000000000%']
    assert _HtmlRenderer(DocumentModel(), ConversionReport(None))._table_cell_styles(table, 0, 0) == []
