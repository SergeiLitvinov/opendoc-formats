"""Independent source tables and painted occurrences, never converter snapshots."""

from io import BytesIO
from zipfile import ZipFile

import pymupdf as fitz
import pytest
from docx import Document
from opendoc_model import Image, Table, document_from_json, document_to_json, walk_model
from PIL import Image as PillowImage

from opendoc_formats import read_document, write_document
from opendoc_formats.writers.pdf_to_docx import PdfModelConverter


def merged_table(path):
    with fitz.open() as native:
        page = native.new_page(width=500, height=400)
        for start, end in [((60, 100), (420, 100)), ((60, 140), (420, 140)),
                           ((180, 190), (420, 190)), ((60, 240), (420, 240)),
                           ((60, 100), (60, 240)), ((180, 140), (180, 240)),
                           ((300, 100), (300, 240)), ((420, 100), (420, 240))]:
            page.draw_line(start, end, width=.75)
        for point, text in [((67, 122), 'Merged heading'), ((307, 122), 'Right heading'),
                            ((67, 162), 'Tall cell'), ((187, 162), 'Center A'), ((307, 162), 'Right A'),
                            ((187, 212), 'Center B'), ((307, 212), 'Right B')]:
            page.insert_text(point, text, fontsize=11, fontname='Times-Roman')
        native.save(path)


def raster_bytes():
    image = PillowImage.new('RGBA', (80, 40), (220, 20, 30, 255))
    for x in range(40, 80):
        for y in range(40):
            image.putpixel((x, y), (10, 160, 80, 128 if y < 20 else 255))
    stream = BytesIO()
    image.save(stream, format='PNG')
    return stream.getvalue()


def layered_page(path):
    with fitz.open() as native:
        page = native.new_page(width=400, height=400)
        page.draw_rect((40, 40, 360, 180), color=None, fill=(1, .9, .6))
        page.insert_text((60, 95), 'Layers and raster', fontsize=14, fontname='Times-Roman')
        image = raster_bytes()
        page.insert_image((220, 70, 320, 130), stream=image, keep_proportion=False)
        page.draw_rect((100, 80, 140, 105), color=None, fill=(.1, .2, .8))
        page.insert_image((50, 220, 130, 300), stream=image, rotate=90, keep_proportion=False)
        page.insert_image((220, 70, 320, 130), stream=image, keep_proportion=False)
        native.save(path)


def multiline_table(path):
    with fitz.open() as native:
        page = native.new_page(width=500, height=400)
        for y in (100, 180, 260):
            page.draw_line((60, y), (420, y), width=.75)
        for x in (60, 240, 420):
            page.draw_line((x, 100), (x, 260), width=.75)
        for row in range(2):
            for column in range(2):
                page.insert_text((67 + column * 180, 122 + row * 80), f'First {row}-{column}',
                                 fontsize=11, fontname='Times-Roman')
                page.insert_text((72 + column * 180, 148 + row * 80), f'Second {row}-{column}',
                                 fontsize=11, fontname='Times-Roman')
        native.save(path)


def borderless_table(path):
    with fitz.open() as native:
        page = native.new_page(width=400, height=300)
        for y, values in ((90, ('Name', 'Value')), (120, ('Alpha', 'Ten')), (150, ('Beta', 'Twenty'))):
            for x, text in zip((60, 250), values):
                page.insert_text((x, y), text, fontsize=11, fontname='Times-Roman')
        native.save(path)


def table_with_images(path, *, impossible=False):
    with fitz.open() as native:
        page = native.new_page(width=500, height=400)
        for y in (100, 160, 220):
            page.draw_line((60, y), (420, y), width=.75)
        for x in (60, 240, 420):
            page.draw_line((x, 100), (x, 220), width=.75)
        if not impossible:
            page.insert_image((63, 110, 180, 150), stream=raster_bytes(), keep_proportion=False)
        page.insert_text((67, 125), 'Before picture', fontsize=11, fontname='Times-Roman')
        page.insert_image((65, 112, 190, 150) if impossible else (70, 118, 95, 135),
                          stream=raster_bytes(), keep_proportion=False)
        if impossible:
            page.insert_text((67, 143), 'After picture', fontsize=11, fontname='Times-Roman')
        for point, text in [((247, 125), 'Right top'), ((67, 190), 'Left bottom'), ((247, 190), 'Right bottom')]:
            page.insert_text(point, text, fontsize=11, fontname='Times-Roman')
        native.save(path)


def test_merged_pdf_cells_are_native_editable_spans_through_two_cycles(tmp_path):
    source, output = tmp_path / 'merged.pdf', tmp_path / 'merged.docx'
    merged_table(source)
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    for index in range(3):
        model = read_document(output).document
        table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
        assert [len(row.cells) for row in table.rows] == [2, 3, 2]
        assert table.rows[0].cells[0].column_span == 2
        assert table.rows[1].cells[0].row_span == 2
        assert table.rows[2].cells[0].blocks[0].plain_text == 'Center B'
        assert table.properties.grid_widths_twips == [2400, 2400, 2400]
        assert [row.properties['docx_height_twips'] for row in table.rows] == [800, 1000, 1000]
        if index < 2:
            output = tmp_path / f'merged-cycle{index}.docx'
            report = write_document(document_from_json(document_to_json(model)), output)
            assert report.success, report.issues
    native = Document(output)
    assert native.tables[0].cell(1, 0)._tc is native.tables[0].cell(2, 0)._tc
    assert native.tables[0].cell(0, 0)._tc is native.tables[0].cell(0, 1)._tc
    continuation = native.tables[0].rows[2]._tr.tc_lst[0]
    assert continuation.xpath('./w:tcPr/w:tcBorders/w:left/@w:val') == ['single']
    assert continuation.xpath('./w:tcPr/w:tcBorders/w:bottom/@w:val') == ['single']


def test_paint_order_rotated_rasters_and_repeated_alpha_occurrences(tmp_path):
    source, output = tmp_path / 'layers.pdf', tmp_path / 'layers.docx'
    layered_page(source)
    source_model = read_document(source).document
    source_images = [ref.node for ref in walk_model(source_model) if isinstance(ref.node, Image)]
    assert len(source_images) == 5
    assert sorted(image.properties['pdf_paint_seq'] for image in source_images) == [0, 2, 3, 4, 5]
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    native = Document(output)
    levels = [int(anchor.get('relativeHeight')) for anchor in native.element.xpath('.//wp:anchor')]
    assert sorted(levels) == [2, 6, 10, 14, 18, 22]
    for index in range(2):
        model = document_from_json(document_to_json(read_document(output).document))
        images = [ref.node for ref in walk_model(model) if isinstance(ref.node, Image)]
        assert len(images) == 5
        assert sorted(image.properties.relative_height for image in images) == [2, 10, 14, 18, 22]
        assert any(image.box.rotation % 360 == 270 for image in images)
        output = tmp_path / f'layers-cycle{index}.docx'
        report = write_document(model, output)
        assert report.success, report.issues
    with ZipFile(output) as archive:
        rasters = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/') and name.endswith('.png')]
    source_raster = next(resource for resource in source_model.resources.values() if resource.media_type == 'image/png')
    pixels = PillowImage.open(BytesIO(source_raster.data)).convert('RGBA').tobytes()
    assert any(PillowImage.open(BytesIO(data)).convert('RGBA').tobytes() == pixels for data in rasters)


def test_soft_mask_preserves_stored_straight_samples_without_roundoff():
    from opendoc_formats.readers.pdf_image_alpha import straight_alpha_png

    base = PillowImage.new('RGB', (2, 2), (9, 159, 79))
    mask = PillowImage.new('L', (2, 2))
    mask.putdata([0, 1, 128, 255])
    rgb, opacity = BytesIO(), BytesIO()
    base.save(rgb, format='PNG')
    mask.save(opacity, format='PNG')
    restored = PillowImage.open(BytesIO(straight_alpha_png(rgb.getvalue(), opacity.getvalue())))
    assert [restored.getpixel((x, y)) for y in range(2) for x in range(2)] == [
        (9, 159, 79, alpha) for alpha in (0, 1, 128, 255)]


def test_multiline_cell_keeps_distinct_baselines_and_line_indents(tmp_path):
    source, output = tmp_path / 'multiline.pdf', tmp_path / 'multiline.docx'
    multiline_table(source)
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    for index in range(3):
        model = read_document(output).document
        table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
        paragraphs = table.rows[0].cells[0].blocks
        assert [p.plain_text for p in paragraphs] == ['First 0-0', 'Second 0-0']
        assert paragraphs[1].properties.left_indent_pt == 5
        assert paragraphs[0].properties.space_after_pt > 13
        if index < 2:
            output = tmp_path / f'multiline-cycle{index}.docx'
            report = write_document(document_from_json(document_to_json(model)), output)
            assert report.success, report.issues


def test_shear_is_not_silently_changed_to_a_rectangular_picture():
    from opendoc_model import Box

    from opendoc_formats.writers.pdf_docx_images import prepare_image_transform

    image = Image('own', box=Box(0, 0, 100, 100), properties={'pdf_image_transform': [100, 0, 20, 100, 0, 0]})
    with pytest.raises(ValueError, match='Sheared'):
        prepare_image_transform(image)


def test_virtual_borderless_bands_fit_native_font_without_clipping(tmp_path):
    source, output = tmp_path / 'borderless.pdf', tmp_path / 'borderless.docx'
    borderless_table(source)
    before = document_to_json(read_document(source).document)
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    for index in range(3):
        model = read_document(output).document
        table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
        assert [cell.blocks[0].plain_text for row in table.rows for cell in row.cells] == [
            'Name', 'Value', 'Alpha', 'Ten', 'Beta', 'Twenty']
        assert table.rows[-1].properties['docx_height_twips'] >= 250
        if index < 2:
            output = tmp_path / f'borderless-cycle{index}.docx'
            report = write_document(document_from_json(document_to_json(model)), output)
            assert report.success, report.issues
    assert document_to_json(read_document(source).document) == before


def test_table_rasters_can_be_behind_and_above_distinct_body_text(tmp_path):
    source, output = tmp_path / 'mixed.pdf', tmp_path / 'mixed.docx'
    table_with_images(source)
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    model = read_document(output).document
    rasters = [ref.node for ref in walk_model(model) if isinstance(ref.node, Image)
               and model.resources[ref.node.resource_id].media_type == 'image/png']
    assert len(rasters) == 2
    assert sorted(image.properties.behind_doc for image in rasters) == [False, True]


def test_impossible_body_plane_interleaving_preserves_previous_destination(tmp_path):
    source, output = tmp_path / 'interleaved.pdf', tmp_path / 'previous.docx'
    table_with_images(source, impossible=True)
    output.write_bytes(b'previous destination')
    report = PdfModelConverter().convert(source, output)
    assert not report.success
    assert output.read_bytes() == b'previous destination'
    assert any('both above and below' in issue.message for issue in report.issues)


def test_scan_preserves_native_pixels_without_text_threshold_rasterization(tmp_path):
    source, output = tmp_path / 'scan.pdf', tmp_path / 'scan.docx'
    with fitz.open() as native:
        page = native.new_page(width=400, height=300)
        page.insert_image(page.rect, stream=raster_bytes(), keep_proportion=False)
        native.save(source)
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    model = read_document(output).document
    images = [ref.node for ref in walk_model(model) if isinstance(ref.node, Image)]
    assert len(images) == 1
    assert (images[0].box.x, images[0].box.y, images[0].box.width, images[0].box.height) == (0, 0, 400, 300)


def test_unmeasured_form_clipping_is_rejected_without_showing_hidden_pixels(tmp_path):
    source, output = tmp_path / 'form-crop.pdf', tmp_path / 'previous.docx'
    with fitz.open(stream=raster_bytes(), filetype='png') as image:
        image_pdf = fitz.open(stream=image.convert_to_pdf(), filetype='pdf')
    with image_pdf, fitz.open() as native:
        page = native.new_page(width=400, height=300)
        page.show_pdf_page(fitz.Rect(60, 80, 280, 180), image_pdf, clip=fitz.Rect(10, 5, 65, 35),
                           keep_proportion=False)
        native.save(source)
    output.write_bytes(b'previous destination')
    report = PdfModelConverter().convert(source, output)
    assert not report.success
    assert any('clipping or graphics group' in issue.message for issue in report.issues)
    assert output.read_bytes() == b'previous destination'


def test_native_jpeg_bytes_are_embedded_without_reencoding(tmp_path):
    image = PillowImage.new('RGB', (97, 53), (30, 150, 210))
    encoded = BytesIO()
    image.save(encoded, format='JPEG', quality=91)
    source, output = tmp_path / 'photo.pdf', tmp_path / 'photo.docx'
    with fitz.open() as native:
        page = native.new_page(width=400, height=300)
        page.insert_image((60, 80, 254, 186), stream=encoded.getvalue())
        native.save(source)
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    with ZipFile(output) as archive:
        assert any(archive.read(name) == encoded.getvalue() for name in archive.namelist() if name.startswith('word/media/'))


def test_public_model_route_has_the_same_layout_and_keeps_source_immutable(tmp_path):
    source, output = tmp_path / 'merged.pdf', tmp_path / 'public.docx'
    merged_table(source)
    model = read_document(source).document
    before = document_to_json(model)
    report = write_document(model, output)
    assert report.success, report.issues
    assert report.metrics['pdf_docx_layout_profile'] == 'opendoc-v1'
    assert report.metrics['assessment_complete'] is False
    assert document_to_json(model) == before
    restored = read_document(output).document
    table = next(block for block in restored.sections[0].blocks if isinstance(block, Table))
    assert table.rows[1].cells[0].row_span == 2
    assert len(Document(output).element.xpath('.//wp:anchor')) == 0


def test_native_grid_expansion_respects_selected_budget_before_allocation(tmp_path):
    from opendoc_model import DocumentLimits, DocumentModel, Paragraph, Section, TableCell, TableRow, TextRun

    from opendoc_formats import ExportOptions

    model = DocumentModel(sections=[Section([Table([TableRow([TableCell(
        blocks=[Paragraph([TextRun('Sparse native grid')])], column_span=2000)])])])])
    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous destination')
    report = write_document(model, output, options=ExportOptions(document_limits=DocumentLimits(max_nodes=1000)))
    assert not report.success
    assert output.read_bytes() == b'previous destination'
    assert report.metrics['document_limit']['quota'] == 'native table grid'
    assert report.metrics['document_limit']['maximum'] == 1000


def test_reflected_native_image_uses_word_transform_without_resampling(tmp_path):
    source, output = tmp_path / 'mirror.pdf', tmp_path / 'mirror.docx'
    with fitz.open() as native:
        page = native.new_page(width=400, height=300)
        page.insert_image((80, 80, 240, 160), stream=raster_bytes(), keep_proportion=False)
        name = page.get_images(full=True)[0][7]
        native.update_stream(page.get_contents()[-1], f'q -160 0 0 80 240 140 cm /{name} Do Q'.encode('ascii'))
        native.save(source)
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    model = read_document(output).document
    image = next(ref.node for ref in walk_model(model) if isinstance(ref.node, Image))
    assert image.box.rotation == 180
    assert image.properties.flip_vertical is True
    assert image.box.width == 160 and image.box.height == 80


def test_embedded_font_program_precedes_base14_alias(tmp_path):
    from fontTools.ttLib import TTFont

    from tests.writers.test_pdf_unicode import _font

    font = tmp_path / 'own-base14-name.ttf'
    _font(font)
    with TTFont(font) as native:
        for record in list(native['name'].names):
            if record.nameID in (1, 4, 6):
                native['name'].setName('Helvetica', record.nameID, record.platformID, record.platEncID, record.langID)
        native.save(font)
    source, output = tmp_path / 'own-font.pdf', tmp_path / 'own-font.docx'
    with fitz.open() as native:
        page = native.new_page(width=400, height=300)
        page.insert_font(fontname='OwnNamed', fontfile=str(font))
        page.insert_text((60, 90), 'f i t', fontsize=12, fontname='OwnNamed')
        native.save(source)
    assert read_document(source).document.sections[0].blocks[0].content[0].style.font_family == 'Helvetica'
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    assert not any(issue.feature == 'font-substitution' for issue in report.issues)
    restored = read_document(output).document.sections[0].blocks[0].content[0]
    assert restored.style.font_family.startswith('ODF Embedded ')
    assert restored.text == 'f i t'


def test_same_family_different_programs_keep_explicit_model_binding(tmp_path):
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    from fontTools.ttLib import TTFont
    from opendoc_model import ResourceKind, TextRun, get_font_reference, set_font_reference

    from tests.writers.test_pdf_unicode import _font

    fonts = [tmp_path / 'rectangle.ttf', tmp_path / 'triangle.ttf']
    for font in fonts:
        _font(font)
    with TTFont(fonts[1]) as native:
        pen = TTGlyphPen(None)
        pen.moveTo((40, 0))
        pen.lineTo((400, 0))
        pen.lineTo((220, 650))
        pen.closePath()
        native['glyf']['f'] = pen.glyph()
        native.save(fonts[1])
    source, output = tmp_path / 'twins.pdf', tmp_path / 'previous.docx'
    with fitz.open() as native:
        page = native.new_page(width=400, height=300)
        for index, font in enumerate(fonts):
            name = f'Own{index}'
            page.insert_font(fontname=name, fontfile=str(font))
            page.insert_text((60, 90 + 70 * index), 'f', fontsize=20, fontname=name)
        native.save(source)
    model = read_document(source).document
    resources = [resource for resource in model.resources.values() if resource.kind is ResourceKind.FONT]
    assert len(resources) == 2 and resources[0].data != resources[1].data
    families = {ref.node.style.font_family for ref in walk_model(model) if isinstance(ref.node, TextRun)}
    assert len(families) == 1 and next(iter(families)).startswith('OwnLigature')
    runs = [ref.node for ref in walk_model(model) if isinstance(ref.node, TextRun)]
    programs = [model.resources[get_font_reference(run.style).resource_id].data for run in runs]
    assert programs == [font.read_bytes() for font in fonts]
    assert all(get_font_reference(run.style).face_index is None for run in runs)
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        report = write_document(model, output)
        assert report.success, report.issues
        restored = read_document(output).document
        restored_runs = [ref.node for ref in walk_model(restored) if isinstance(ref.node, TextRun)]
        assert len({get_font_reference(run.style).resource_id for run in restored_runs}) == 2
        assert len({run.style.font_family for run in restored_runs}) == 2
        # Compare actual glyph outlines, independent of Word aliases and subsetting bytes.
        for original, run in zip(fonts, restored_runs, strict=True):
            data = restored.resources[get_font_reference(run.style).resource_id].data
            with TTFont(original) as expected, TTFont(BytesIO(data)) as actual:
                glyph = actual.getBestCmap()[ord('f')]
                assert list(actual['glyf'][glyph].coordinates) == list(expected['glyf']['f'].coordinates)
        model = restored
    # A model edit swaps actual programs while keeping the same visible family label.
    model = read_document(source).document
    runs = [ref.node for ref in walk_model(model) if isinstance(ref.node, TextRun)]
    references = [get_font_reference(run.style) for run in runs]
    for run, reference in zip(runs, reversed(references), strict=True):
        set_font_reference(run.style, reference)
    # Inherit the first binding through a named style using the SDK cascade.
    from opendoc_model import TextStyle

    model.styles['bound'] = TextStyle(font_family=runs[0].style.font_family)
    set_font_reference(model.styles['bound'], get_font_reference(runs[0].style))
    model.styles['bound-child'] = TextStyle(properties={'base_style_id': 'bound'})
    set_font_reference(runs[0].style, None)
    model.sections[0].blocks[0].style_id = 'bound-child'
    before = document_to_json(model)
    report = write_document(model, output)
    assert report.success, report.issues
    assert document_to_json(model) == before
    restored = read_document(output).document
    restored_runs = [ref.node for ref in walk_model(restored) if isinstance(ref.node, TextRun)]
    for original, run in zip(reversed(fonts), restored_runs, strict=True):
        data = restored.resources[get_font_reference(run.style).resource_id].data
        with TTFont(original) as expected, TTFont(BytesIO(data)) as actual:
            assert list(actual['glyf'][actual.getBestCmap()[ord('f')]].coordinates) == list(expected['glyf']['f'].coordinates)
    # Without native binding evidence the previous destination remains protected.
    model = read_document(source).document
    for location in walk_model(model):
        if isinstance(location.node, TextRun):
            set_font_reference(location.node.style, None)
    output.write_bytes(b'previous destination')
    report = write_document(model, output)
    assert not report.success
    assert any(issue.feature == 'font-resource-binding' for issue in report.issues)
    assert output.read_bytes() == b'previous destination'


def test_pdf_font_binding_unknown_collection_face_is_not_defaulted(tmp_path):
    from opendoc_model import FontReference, TextRun, get_font_reference, set_font_reference

    from tests.writers.test_pdf_unicode import _font

    font, source, output = tmp_path / 'own.ttf', tmp_path / 'own.pdf', tmp_path / 'previous.docx'
    _font(font)
    with fitz.open() as native:
        page = native.new_page()
        page.insert_font(fontname='Own', fontfile=str(font))
        page.insert_text((54, 90), 'f', fontname='Own')
        native.save(source)
    model = read_document(source).document
    run = next(location.node for location in walk_model(model) if isinstance(location.node, TextRun))
    reference = get_font_reference(run.style)
    set_font_reference(run.style, FontReference(reference.resource_id, face_index=1))
    output.write_bytes(b'previous destination')
    report = write_document(model, output)
    assert not report.success
    assert any(issue.feature == 'font-unavailable' for issue in report.issues)
    assert output.read_bytes() == b'previous destination'


def test_native_font_buffer_survives_repeated_content_view_close(tmp_path):
    import subprocess
    import sys

    from tests.writers.test_pdf_unicode import _font

    font, source = tmp_path / 'own.ttf', tmp_path / 'annotated.pdf'
    _font(font)
    with fitz.open() as native:
        page = native.new_page()
        page.insert_font(fontname='Own', fontfile=str(font))
        page.insert_text((54, 90), 'f', fontname='Own')
        page.add_text_annot((180, 180), 'Independent appearance')
        native.save(source)
    # A subprocess turns native allocator aborts into an ordinary, located regression failure.
    script = '''
import gc, sys
import pymupdf
from opendoc_model import TextRun, get_font_reference, walk_model
from opendoc_formats import read_document
for index in range(12):
    model = read_document(sys.argv[1]).document
    runs = [item.node for item in walk_model(model) if isinstance(item.node, TextRun) and item.node.text == 'f']
    assert runs and all(get_font_reference(run.style) is not None for run in runs)
    del model, runs
    gc.collect()
    with pymupdf.open(sys.argv[1]) as native:
        assert native[0].get_pixmap().width > 0
'''
    result = subprocess.run([sys.executable, '-c', script, str(source)], capture_output=True, text=True, timeout=60)
    assert result.returncode == 0, result.stdout + result.stderr


def test_font_subsetting_keeps_license_and_localized_notices(tmp_path):
    from fontTools.ttLib import TTFont

    from opendoc_formats.fonts.embedding import FontUsage, subset_font
    from tests.writers.test_pdf_unicode import _font

    font = tmp_path / 'notices.ttf'
    _font(font)
    with TTFont(font) as native:
        native['name'].setName('Own copyright fixture', 0, 3, 1, 0x409)
        native['name'].setName('Own license fixture', 13, 3, 1, 0x409)
        native['name'].setName('Own localized license fixture', 13, 3, 1, 0x419)
        native['name'].setName('https://example.invalid/own-font-license', 14, 3, 1, 0x409)
        native.save(font)
    subset = subset_font(FontUsage('OwnLigature', font, False, False, frozenset('f'), True, True))
    with TTFont(BytesIO(subset)) as native:
        assert native['name'].getName(0, 3, 1, 0x409).toUnicode() == 'Own copyright fixture'
        assert native['name'].getName(13, 3, 1, 0x409).toUnicode() == 'Own license fixture'
        assert native['name'].getName(13, 3, 1, 0x419).toUnicode() == 'Own localized license fixture'
        assert native['name'].getName(14, 3, 1, 0x409).toUnicode() == 'https://example.invalid/own-font-license'


def test_unmeasured_thin_svg_stroke_never_replaces_previous_document(tmp_path):
    source, output = tmp_path / 'stroke.pdf', tmp_path / 'previous.docx'
    with fitz.open() as native:
        page = native.new_page(width=400, height=300)
        page.draw_line((60, 80), (320, 80), width=.75)
        native.save(source)
    output.write_bytes(b'previous destination')
    report = PdfModelConverter().convert(source, output)
    assert not report.success
    assert any('zero-area PDF strokes' in issue.message for issue in report.issues)
    assert output.read_bytes() == b'previous destination'


@pytest.mark.parametrize('interaction', ['highlight', 'uri'])
def test_public_model_route_does_not_silently_drop_pdf_interactions(tmp_path, interaction):
    source, output = tmp_path / 'interactive.pdf', tmp_path / 'previous.docx'
    with fitz.open() as native:
        page = native.new_page(width=400, height=300)
        page.insert_text((60, 90), 'Interactive source text', fontsize=12, fontname='Times-Roman')
        if interaction == 'highlight':
            page.add_highlight_annot(fitz.Rect(60, 75, 150, 95))
        else:
            page.insert_link({'kind': fitz.LINK_URI, 'from': fitz.Rect(60, 75, 150, 95),
                              'uri': 'https://example.invalid/own'})
        native.save(source)
    output.write_bytes(b'previous destination')
    model = read_document(source).document
    report = write_document(document_from_json(document_to_json(model)), output)
    assert not report.success
    assert any(issue.feature == 'pdf-docx.interactions' and issue.location == 'integration.annotations'
               for issue in report.issues)
    assert output.read_bytes() == b'previous destination'


@pytest.mark.parametrize('case', ['font-size', 'baseline'])
def test_native_word_quantization_is_not_a_silent_typography_loss(tmp_path, case):
    source, output = tmp_path / 'precision.pdf', tmp_path / 'previous.docx'
    with fitz.open() as native:
        page = native.new_page(width=400, height=300)
        if case == 'font-size':
            page.insert_text((60, 90), 'Measured size', fontsize=10.8, fontname='Times-Roman')
        else:
            page.insert_text((60, 120), 'E = mc', fontsize=14, fontname='Courier')
            page.insert_text((110.4, 114.25), '2', fontsize=8, fontname='Courier')
        native.save(source)
    output.write_bytes(b'previous destination')
    report = PdfModelConverter().convert(source, output)
    assert not report.success
    assert any('without rounding' in issue.message for issue in report.issues)
    assert output.read_bytes() == b'previous destination'


def test_pdf_font_without_unicode_cmap_is_not_replaced_by_system_glyphs(tmp_path):
    import re

    from fontTools.ttLib import TTFont

    from tests.writers.test_pdf_unicode import _font

    font = tmp_path / 'own.ttf'
    _font(font)
    source, output = tmp_path / 'missing-cmap.pdf', tmp_path / 'previous.docx'
    with fitz.open() as native:
        page = native.new_page(width=400, height=300)
        font_xref = page.insert_font(fontname='Own', fontfile=str(font))
        page.insert_text((60, 90), 'f', fontsize=12, fontname='Own')
        descendant = int(re.search(r'(\d+) 0 R', native.xref_get_key(font_xref, 'DescendantFonts')[1])[1])
        descriptor = int(native.xref_get_key(descendant, 'FontDescriptor')[1].split()[0])
        program = int(native.xref_get_key(descriptor, 'FontFile2')[1].split()[0])
        _, _, _, data = native.extract_font(font_xref)
        with TTFont(BytesIO(data)) as original:
            for table in original['cmap'].tables:
                if table.isUnicode():
                    table.cmap.pop(ord('f'), None)
            encoded = BytesIO()
            original.save(encoded)
        native.update_stream(program, encoded.getvalue())
        native.save(source)
    with fitz.open(source) as native:
        assert native[0].get_text().strip() == 'f'  # PDF ToUnicode still addresses the original glyph.
    output.write_bytes(b'previous destination')
    report = PdfModelConverter().convert(source, output)
    assert not report.success
    assert any(issue.feature == 'font-glyphs' for issue in report.issues)
    assert output.read_bytes() == b'previous destination'
