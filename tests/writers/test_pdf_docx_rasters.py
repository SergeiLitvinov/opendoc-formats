"""Native PDF clipping reference and real DOCX inline flow, without image resampling."""

import io
import math

import pymupdf as fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from opendoc_model import (
    Box,
    DocumentModel,
    Image,
    ImageCrop,
    Length,
    PageSettings,
    Paragraph,
    Resource,
    ResourceKind,
    Section,
    TextRun,
    document_from_json,
    document_to_json,
)
from PIL import Image as Pixels
from PIL import ImageDraw

from opendoc_formats import read_document, write_document


def own_pixels():
    image = Pixels.new("RGBA", (120, 180), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 119, 179), outline="red", width=6)
    draw.rectangle((12, 36, 65, 100), fill="blue")
    draw.rectangle((66, 36, 107, 125), fill="green")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.mark.parametrize("angle", [3, 12.5, -17, 90])
def test_rotated_crop_against_independent_native_form_reference(tmp_path, angle):
    raw = own_pixels()
    crop = ImageCrop(left=0.1, top=0.2, right=0.1, bottom=0.1)
    model = DocumentModel(
        sections=[Section(blocks=[Image("own", box=Box(70, 80, 108, 72, rotation=angle), crop=crop)])],
        resources={"own": Resource("own", ResourceKind.RASTER_IMAGE, "image/png", data=raw)},
    )
    radians = math.radians(angle)
    w = abs(108 * math.cos(radians)) + abs(72 * math.sin(radians))
    h = abs(108 * math.sin(radians)) + abs(72 * math.cos(radians))
    # Reference uses the documented Form XObject API with source clipping and
    # arbitrary rotation; production writes a single native image placement.
    with fitz.open(stream=raw, filetype="png") as image:
        image_pdf = fitz.open(stream=image.convert_to_pdf(), filetype="pdf")
    with image_pdf, fitz.open() as cropped, fitz.open() as reference:
        source_rect = image_pdf[0].rect
        clip = fitz.Rect(source_rect.width * 0.1, source_rect.height * 0.2,
                         source_rect.width * 0.9, source_rect.height * 0.9)
        frame = cropped.new_page(width=108, height=72)
        frame.show_pdf_page(frame.rect, image_pdf, clip=clip, keep_proportion=False)
        page = reference.new_page(width=595.27559, height=841.88976)
        page.show_pdf_page(fitz.Rect(124-w/2, 116-h/2, 124+w/2, 116+h/2), cropped, rotate=-angle)
        reference.save(tmp_path / "reference.pdf")
        expected = page.get_pixmap(alpha=False)
        for cycle in range(2):
            model = document_from_json(document_to_json(model))
            before = document_to_json(model)
            target = tmp_path / f"crop-{cycle}.pdf"
            result = write_document(model, target)
            assert result.success, result.issues
            assert document_to_json(model) == before
            with fitz.open(target) as pdf:
                actual = pdf[0].get_pixmap(alpha=False)
                assert (actual.width, actual.height) == (expected.width, expected.height)
                differences = [abs(a-b) for a, b in zip(actual.samples, expected.samples, strict=True)]
                assert pdf[0].get_image_info()[0]["transform"] == pytest.approx(
                    page.get_image_info()[0]["transform"], abs=0.002,
                )
                corners = [fitz.Point(x-124, y-116) * fitz.Matrix(angle) + fitz.Point(124, 116)
                           for x, y in ((70, 80), (178, 80), (178, 152), (70, 152))]
                # Nested reference Forms clip the same edge twice: allow edge AA,
                # require interior pixels to agree and bound error by image area.
                assert sum(differences) / (108 * 72 * 3) < 1
                for index, delta in enumerate(differences):
                    if delta <= 2:
                        continue
                    point = fitz.Point((index // 3) % actual.width + 0.5, (index // 3) // actual.width + 0.5)
                    distances = []
                    for start, end in zip(corners, corners[1:] + corners[:1], strict=True):
                        vector = end-start
                        factor = max(0, min(1, ((point-start).x * vector.x + (point-start).y * vector.y)
                                            / (vector.x**2 + vector.y**2)))
                        distances.append(abs(point - (start + vector * factor)))
                    assert min(distances) <= 1.5


def test_docx_inline_crop_rotation_reserves_real_flow_position(tmp_path):
    source = Document()
    source.add_paragraph("Before own image")
    shape = source.add_paragraph().add_run().add_picture(io.BytesIO(own_pixels()), width=Inches(1.5))
    fill = shape._inline.xpath(".//pic:blipFill")[0]
    crop = OxmlElement("a:srcRect")
    for key, value in (("l", "10000"), ("t", "20000"), ("r", "10000"), ("b", "10000")):
        crop.set(key, value)
    fill.insert(1, crop)
    shape._inline.xpath(".//pic:spPr/a:xfrm")[0].set("rot", "180000")
    source.add_paragraph("After own image")
    path = tmp_path / "own.docx"
    source.save(path)
    original = path.read_bytes()
    result = read_document(path)
    assert result.success
    model = result.document
    expected = None
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        before = document_to_json(model)
        target = tmp_path / f"inline-{cycle}.pdf"
        exported = write_document(model, target)
        assert exported.success, exported.issues
        assert not any(issue.feature == "pdf-raster-unsupported" for issue in exported.issues)
        assert document_to_json(model) == before
        with fitz.open(target) as pdf:
            assert len(pdf) == 1
            images = pdf[0].get_image_info()
            assert len(images) == 1
            a, b, c, d, e, f = images[0]["transform"]
            assert a > 0 and b > 0 and c < 0 and d > 0  # DrawingML clockwise rotation
            assert b / a == pytest.approx(math.tan(math.radians(3)), abs=1e-6)
            before_box = pdf[0].search_for("Before own image")[0]
            after_box = pdf[0].search_for("After own image")[0]
            assert e + a * 0.1 + c * 0.2 > 20
            assert f + b * 0.1 + d * 0.2 > before_box.y1
            assert after_box.y0 - before_box.y1 > 140
            pixels = pdf[0].get_pixmap(alpha=False).samples
            if expected is None:
                expected = pixels
            else:
                assert pixels == expected
    assert path.read_bytes() == original


@pytest.mark.parametrize("placement", ["inline", "anchor"])
def test_centered_flow_raster_on_later_page_and_anchor_diagnostic(tmp_path, placement):
    props = {"placement": placement}
    if placement == "anchor":
        props.update(horizontal_relative_from="column", horizontal_align="center",
                     vertical_relative_from="paragraph", wrap="tight")
    picture = Image("own", box=Box(0, 0, 60, 45, rotation=3), properties=props)
    model = DocumentModel(
        sections=[Section(
            blocks=[*[Paragraph([TextRun(f"Before row {i}")]) for i in range(15)],
                    Paragraph([picture], alignment="center"), Paragraph([TextRun("After own image")])],
            page=PageSettings(width=Length(240), height=Length(200), margin_left=Length(20), margin_right=Length(20),
                              margin_top=Length(20), margin_bottom=Length(20)),
        )],
        resources={"own": Resource("own", ResourceKind.RASTER_IMAGE, "image/png", data=own_pixels())},
    )
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        target = tmp_path / f"later-{cycle}.pdf"
        report = write_document(model, target)
        assert report.success, report.issues
        assert any(issue.feature == "pdf-raster-anchor-wrap" for issue in report.issues) == (placement == "anchor")
        with fitz.open(target) as pdf:
            assert len(pdf) > 1
            containing = [i for i, page in enumerate(pdf) if page.get_image_info()]
            assert len(containing) == 1 and containing[0] > 0
            info = pdf[containing[0]].get_image_info()[0]
            assert (info["bbox"][0] + info["bbox"][2]) / 2 == pytest.approx(120, abs=0.01)
            assert len(pdf[containing[0]].search_for("After own image")) == 1


def test_flow_header_repeats_on_each_page(tmp_path):
    picture = Image("own", box=Box(0, 0, 24, 12), properties={"placement": "inline"})
    model = DocumentModel(
        sections=[Section(
            headers=[Paragraph([picture], alignment="right")],
            blocks=[Paragraph([TextRun(f"Own body row {i}")]) for i in range(20)],
            page=PageSettings(width=Length(240), height=Length(200), margin_left=Length(20), margin_right=Length(20),
                              margin_top=Length(40), margin_bottom=Length(20)),
        )],
        resources={"own": Resource("own", ResourceKind.RASTER_IMAGE, "image/png", data=own_pixels())},
    )
    target = tmp_path / "header.pdf"
    report = write_document(model, target)
    assert report.success, report.issues
    with fitz.open(target) as pdf:
        assert len(pdf) > 1
        for page in pdf:
            images = page.get_image_info()
            assert len(images) == 1
            assert images[0]["bbox"][3] < 40
            assert images[0]["bbox"][2] == pytest.approx(220, abs=0.01)


@pytest.mark.parametrize("fault", ["crop", "oversize", "png-limit"])
def test_unsafe_raster_rejects_before_publication(tmp_path, fault):
    picture = Image("own", box=Box(0, 0, 60, 45, rotation=3), properties={"placement": "inline"})
    raw = own_pixels()
    if fault == "crop":
        picture.crop = ImageCrop(left=0.9999999)
    elif fault == "oversize":
        picture.box.height = 10000
    else:
        import struct

        raw = raw[:16] + struct.pack(">II", 100000, 100000) + raw[24:]
    model = DocumentModel(
        sections=[Section(blocks=[Paragraph([picture])])],
        resources={"own": Resource("own", ResourceKind.RASTER_IMAGE, "image/png", data=raw)},
    )
    target = tmp_path / "previous.pdf"
    target.write_bytes(b"previous")
    report = write_document(model, target)
    assert not report.success
    assert any(issue.feature == "pdf-raster-unsupported" and issue.location == "sections[0].blocks[0].content[0]"
               for issue in report.issues)
    assert target.read_bytes() == b"previous"


@pytest.mark.parametrize("page_rotation", [0, 90])
@pytest.mark.parametrize("image_rotation", [3, -17, 90])
def test_native_pdf_crop_is_preserved_on_two_public_json_cycles(tmp_path, page_rotation, image_rotation):
    model = DocumentModel(
        sections=[Section(
            blocks=[Image("own", box=Box(70, 80, 108, 72, rotation=image_rotation),
                          crop=ImageCrop(left=0.1, top=0.2, right=0.1, bottom=0.1))],
            properties={"pdf": {"source_rotation": page_rotation}},
        )],
        resources={"own": Resource("own", ResourceKind.RASTER_IMAGE, "image/png", data=own_pixels())},
    )
    source = tmp_path / "cropped-source.pdf"
    assert write_document(model, source).success
    original = source.read_bytes()
    with fitz.open(source) as pdf:
        expected = pdf[0].get_pixmap(alpha=False).samples
    current = source
    for cycle in range(2):
        imported = read_document(current)
        assert imported.success
        model = document_from_json(document_to_json(imported.document))
        image = model.sections[0].blocks[0].content[0]
        assert image.crop is not None
        assert (image.crop.left, image.crop.top, image.crop.right, image.crop.bottom) == pytest.approx(
            (0.1, 0.2, 0.1, 0.1), abs=1e-7,
        )
        current = tmp_path / f"cropped-{cycle}.pdf"
        report = write_document(model, current)
        assert report.success, report.issues
        with fitz.open(current) as pdf:
            assert pdf[0].rotation == page_rotation
            assert pdf[0].get_pixmap(alpha=False).samples == expected
    assert source.read_bytes() == original
