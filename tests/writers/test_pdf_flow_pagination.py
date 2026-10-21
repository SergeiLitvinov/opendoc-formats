"""Own DOCX body images must move intact instead of hiding trailing content."""

import math
from io import BytesIO

import pymupdf
import pytest
from docx import Document
from docx.shared import Inches, Pt
from opendoc_model import Image, ImageCrop, Paragraph, document_from_json, document_to_json

from opendoc_formats.readers.docx import read_docx_model
from opendoc_formats.writers.pdf_writer import write_pdf_model
from tests.writers.test_pdf_docx_rasters import own_pixels


def own_source(path, lines=4, width=240):
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Pt(width), Pt(300)
    page.left_margin = page.right_margin = page.top_margin = page.bottom_margin = Pt(20)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    table = document.add_table(rows=2, cols=2)
    for index, cell in enumerate(table._cells):
        cell.text = f'Own cell {index}'
    for index in range(lines):
        document.add_paragraph(f'Own preceding line {index}')
    document.add_picture(BytesIO(own_pixels()), width=Inches(1), height=Inches(.65))
    document.add_paragraph('Own first caption')
    document.add_picture(BytesIO(own_pixels()), width=Inches(1), height=Inches(1))
    document.add_paragraph('Own second caption')
    document.add_paragraph('Own final paragraph must remain visible.')
    document.save(path)
    return read_docx_model(path)


@pytest.mark.parametrize('alignment', ['left', 'center', 'right'])
@pytest.mark.parametrize('angle', [0, 3])
@pytest.mark.parametrize('width', [200, 240])
def test_overflow_body_images_keep_following_text_and_geometry(tmp_path, alignment, angle, width):
    model = own_source(tmp_path / 'own.docx', width=width)
    for block in model.sections[0].blocks:
        if isinstance(block, Paragraph) and any(isinstance(run, Image) for run in block.content):
            block.alignment = alignment
            for run in block.content:
                if isinstance(run, Image):
                    run.box.rotation = angle
                    run.crop = ImageCrop(left=.1, top=.2, right=.1, bottom=.1)
    expected = document_to_json(model)
    reference = None
    for index in range(2):
        model = document_from_json(document_to_json(model))
        target = tmp_path / f'cycle-{index}.pdf'
        report = write_pdf_model(model, target)
        assert report.success, report.issues
        assert report.metrics.get('pdf_flow_page_breaks')
        assert document_to_json(model) == expected
        with pymupdf.open(target) as pdf:
            extracted = ''.join(page.get_text() for page in pdf)
            text = ' '.join(extracted.split())
            if width == 200:
                assert 'Own final paragraph must remain visible.' not in extracted
            for caption in ('Own first caption', 'Own second caption', 'Own final paragraph must remain visible.'):
                assert text.count(caption) == 1
            images = []
            for number, page in enumerate(pdf):
                for image in page.get_image_info():
                    matrix = pymupdf.Matrix(image['transform'])
                    corners = [pymupdf.Point(x, y) * matrix for x, y in ((.1, .2), (.9, .2), (.9, .9), (.1, .9))]
                    images.append((number, (min(p.x for p in corners), min(p.y for p in corners),
                                            max(p.x for p in corners), max(p.y for p in corners))))
            assert len(images) == 2
            for (_, rectangle), height in zip(images, (46.8, 72)):
                radians = math.radians(angle)
                raster_width = 72 * math.cos(radians) + height * math.sin(radians)
                height = height * math.cos(radians) + 72 * math.sin(radians)
                assert rectangle[2] - rectangle[0] == pytest.approx(raster_width, abs=.01)
                assert rectangle[3] - rectangle[1] == pytest.approx(height, abs=.01)
                assert 20 <= rectangle[1] < rectangle[3] <= 280
                if alignment == 'center':
                    assert (rectangle[0]+rectangle[2])/2 == pytest.approx(width/2, abs=.01)
                elif alignment == 'left':
                    assert rectangle[0] == pytest.approx(20, abs=.01)
                else:
                    assert rectangle[2] == pytest.approx(width-20, abs=.01)
            for number, rectangle in images:
                for word in pdf[number].get_text('words'):
                    assert not (pymupdf.Rect(rectangle) & pymupdf.Rect(word[:4])).get_area()
            pixels = [page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False).samples for page in pdf]
            assert reference is None or pixels == reference
            reference = pixels


def test_reproducer_without_repagination_refuses_publication(tmp_path, monkeypatch):
    from opendoc_formats.writers import pdf_flow

    model = own_source(tmp_path / 'own.docx')
    target = tmp_path / 'previous.pdf'
    target.write_bytes(b'previous artifact')
    monkeypatch.setattr(pdf_flow, 'next_flow_break', lambda *args: None)
    report = write_pdf_model(model, target)
    assert not report.success
    assert any(issue.feature == 'pdf-text-content' for issue in report.issues)
    assert target.read_bytes() == b'previous artifact'


def test_retry_cancellation_keeps_destination_and_input(tmp_path, monkeypatch):
    from threading import Event

    from opendoc_formats import ExportOptions, write_document
    from opendoc_formats.writers import pdf_flow

    model = own_source(tmp_path / 'own.docx')
    original = document_to_json(model)
    target = tmp_path / 'previous.pdf'
    target.write_bytes(b'previous artifact')
    stop = Event()
    native = pdf_flow.next_flow_break

    def first(*args):
        retry = native(*args)
        assert retry is not None
        stop.set()
        return retry

    monkeypatch.setattr(pdf_flow, 'next_flow_break', first)
    report = write_document(model, target, options=ExportOptions(cancelled=stop.is_set))
    assert not report.success
    assert any(issue.feature == 'export.cancelled' for issue in report.issues)
    assert target.read_bytes() == b'previous artifact'
    assert document_to_json(model) == original
    assert not list(tmp_path.glob('.*.partial'))


def test_repagination_has_finite_work_and_closes_discarded_documents(tmp_path, monkeypatch):
    import fitz
    from opendoc_model import DocumentModel, Section, TextRun

    from opendoc_formats.writers import pdf_flow

    rendered = []

    class Backend:
        def __init__(self, **kwargs):
            pass

        def write_with_links(self, *args, **kwargs):
            pdf = pymupdf.open()
            pdf.new_page()
            rendered.append(pdf)
            return pdf

    monkeypatch.setattr(fitz, 'Story', Backend)
    monkeypatch.setattr(pdf_flow, 'next_flow_break', lambda html, *args: (html, f'own-{len(rendered)}'))
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun('Own text')])])])
    target = tmp_path / 'previous.pdf'
    target.write_bytes(b'previous artifact')
    report = write_pdf_model(model, target)
    assert not report.success
    assert any('exceeds 32 retries' in issue.message for issue in report.issues), report.issues
    assert len(rendered) == 33 and all(pdf.is_closed for pdf in rendered)
    assert target.read_bytes() == b'previous artifact'


def test_same_paragraph_overflow_cannot_repeat_indefinitely():
    from opendoc_formats.writers.pdf_flow import next_flow_break

    html = '<!--opendoc-pdf-raster-1-break--><p>Own</p>'
    positions = {('opendoc-pdf-raster-1', 0): (20, 230, 92, 302)}
    retry = next_flow_break(html, positions, 280, set())
    assert retry is not None
    with pytest.raises(ValueError, match='still overflows'):
        next_flow_break(retry[0], positions, 280, {retry[1]})
