"""Own fixed-layout route: editable text, native geometry and atomic refusal."""

import builtins
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from opendoc_model import Box, DocumentModel, Paragraph, Section, TextRun, document_from_json, document_to_json

from opendoc_formats import read_document, write_document
from opendoc_formats.writers.pdf_to_docx import PdfModelConverter, create_converter


def columns(path):
    with pymupdf.open() as native:
        page = native.new_page(width=595, height=842)
        for x, label in [(54, 'Left'), (315, 'Right')]:
            for index in range(5):
                page.insert_text((x, 90 + index * 28), f'{label} column line {index + 1}',
                                 fontsize=12, fontname='Times-Roman')
        page.insert_text((54, 42), 'Own two-column source', fontsize=17, fontname='Times-Roman')
        native.save(path)


def ruled_table(path):
    with pymupdf.open() as native:
        page = native.new_page(width=595, height=842)
        for y in (100, 130, 160, 190):
            page.draw_line((60, y), (360, y), color=(0, 0, 0), width=.75)
        for x in (60, 210, 360):
            page.draw_line((x, 100), (x, 190), color=(0, 0, 0), width=.75)
        for row in range(3):
            for column in range(2):
                page.insert_text((67 + column * 150, 119 + row * 30), f'Cell {row + 1}-{column + 1}',
                                 fontsize=11, fontname='Times-Roman')
        native.save(path)


def test_columns_do_not_turn_into_cells_cutting_heading_words(tmp_path):
    source = tmp_path / 'columns.pdf'
    columns(source)
    model = read_document(source).document
    assert len(model.sections[0].blocks) == 11
    assert all(isinstance(block, Paragraph) for block in model.sections[0].blocks)
    assert model.sections[0].blocks[0].plain_text == 'Own two-column source'


def test_own_converter_uses_no_pdf2docx_and_keeps_text_editable(tmp_path, monkeypatch):
    source, output = tmp_path / 'columns.pdf', tmp_path / 'columns.docx'
    columns(source)
    before = source.read_bytes()
    importer = builtins.__import__

    def no_external_converter(name, *args, **kwargs):
        if name.split('.')[0] in {'pdf2docx', 'cv2', 'fire'}:
            raise AssertionError('External converter must not be imported')
        return importer(name, *args, **kwargs)

    monkeypatch.setattr(builtins, '__import__', no_external_converter)
    report = create_converter('opendoc').convert(source, output)
    assert report.success, report.issues
    assert report.metrics['engine'] == 'opendoc'
    assert not report.metrics['assessment_complete']
    native = Document(output)
    assert len(native.paragraphs) == 11
    assert len(native.element.xpath('.//wp:anchor')) == 11
    assert len(native.element.xpath('.//w:txbxContent')) == 11
    assert 'Own two-column source' in native.element.xpath('.//w:txbxContent//w:t/text()')
    assert not native.inline_shapes
    assert source.read_bytes() == before
    assert report.to_dict()['input_path'] == str(source)


def test_page_frames_survive_two_json_and_native_cycles(tmp_path):
    model = DocumentModel(source_format='pdf', sections=[Section(blocks=[
        Paragraph([TextRun('Editable left')], box=Box(40, 80, 160, 20)),
        Paragraph([TextRun('Editable right')], box=Box(300, 80, 160, 20)),
    ])])
    before = document_to_json(model)
    for index in range(2):
        output = tmp_path / f'cycle-{index}.docx'
        report = write_document(document_from_json(document_to_json(model)), output)
        assert report.success, report.issues
        model = read_document(output).document
        assert [(block.box.x, block.box.y, block.box.width, block.box.height) for block in model.sections[0].blocks] == [
            (40, 80, 160, 20), (300, 80, 160, 20)]
        assert [block.plain_text for block in model.sections[0].blocks] == ['Editable left', 'Editable right']
    original = document_from_json(before)
    assert original.sections[0].blocks[0].box.x == 40


@pytest.mark.parametrize('rotation', [20, 90])
def test_unsupported_rotated_frame_refuses_atomically(tmp_path, rotation):
    model = DocumentModel(source_format='pdf', sections=[Section(blocks=[
        Paragraph([TextRun('Own rotated text')], box=Box(40, 80, 160, 20, rotation=rotation))])])
    output = tmp_path / 'previous.docx'
    output.write_bytes(b'previous destination')
    report = write_document(model, output)
    assert not report.success and output.read_bytes() == b'previous destination'
    assert any(issue.feature == 'docx.paragraph-frame' for issue in report.issues)


def test_failed_import_has_serializable_report_and_preserves_destination(tmp_path):
    source, output = tmp_path / 'broken.pdf', tmp_path / 'previous.docx'
    source.write_bytes(b'not a PDF')
    output.write_bytes(b'previous destination')
    report = PdfModelConverter().convert(source, output)
    assert not report.success and report.error
    assert report.to_dict()['issues']
    assert output.read_bytes() == b'previous destination'


def test_converter_refuses_overwriting_the_pdf_source(tmp_path):
    source = tmp_path / 'source.pdf'
    columns(source)
    before = source.read_bytes()
    assert not PdfModelConverter().convert(source, source).success
    assert source.read_bytes() == before


def test_source_change_in_quality_check_rejects_before_publication(tmp_path):
    from opendoc_formats import ExportOptions

    source, output = tmp_path / 'source.pdf', tmp_path / 'previous.docx'
    columns(source)
    output.write_bytes(b'previous destination')

    def gate(report, staged):
        assert staged.is_file() and report.success
        source.write_bytes(source.read_bytes() + b'\n% modified during conversion\n')
        return True

    report = PdfModelConverter(export_options=ExportOptions(quality_gate=gate)).convert(source, output)
    assert not report.success
    assert output.read_bytes() == b'previous destination'
    assert any(issue.feature == 'pdf-docx.source-changed' for issue in report.issues)


def test_cancelled_import_never_reads_pdf_or_replaces_destination(tmp_path):
    from opendoc_formats import ImportOptions

    source, output = tmp_path / 'source.pdf', tmp_path / 'previous.docx'
    source.write_bytes(b'not a PDF')
    output.write_bytes(b'previous destination')
    report = PdfModelConverter(import_options=ImportOptions(cancelled=lambda: True)).convert(source, output)
    assert not report.success
    assert output.read_bytes() == b'previous destination'
    assert any(issue.feature == 'pdf-docx.cancelled' for issue in report.issues)


@pytest.mark.parametrize('edit', ['fill', 'reflection', 'scale', 'layer', 'outer-text'])
def test_unknown_native_text_shape_edits_are_not_discarded(tmp_path, edit):
    from docx.oxml import OxmlElement

    from opendoc_formats.ooxml.positioned_text import WP, WPS, A, positioned_text_source

    source, output = tmp_path / 'source.pdf', tmp_path / 'source.docx'
    columns(source)
    assert PdfModelConverter().convert(source, output).success
    native = Document(output)
    paragraph = native.paragraphs[0]
    anchor = next(paragraph._p.iter(WP + 'anchor'))
    if edit == 'fill':
        next(anchor.iter(WPS + 'spPr')).append(OxmlElement('a:solidFill'))
    elif edit == 'reflection':
        next(anchor.iter(A + 'xfrm')).set('flipH', '1')
    elif edit == 'scale':
        next(anchor.iter(A + 'xfrm')).find(A + 'ext').set('cx', '100')
    elif edit == 'layer':
        anchor.set('behindDoc', '1')
    else:
        paragraph.add_run('New visible outer text')
    with pytest.raises(ValueError, match='positioned text|Positioned text'):
        positioned_text_source(paragraph)


@pytest.mark.parametrize('conflicting_replica', [False, True])
def test_word_compatibility_replica_is_never_imported_as_second_text(tmp_path, conflicting_replica):
    from copy import deepcopy

    from lxml import etree

    from opendoc_formats.ooxml.positioned_text import MC, WPS, W, positioned_text_source

    source, output = tmp_path / 'source.pdf', tmp_path / 'source.docx'
    columns(source)
    assert PdfModelConverter().convert(source, output).success
    native = Document(output)
    paragraph = native.paragraphs[0]
    drawing = next(paragraph._p.iter(W + 'drawing'))
    content = next(drawing.iter(W + 'txbxContent'))
    replica = deepcopy(content)
    owner = drawing.getparent()
    alternative = etree.Element(MC + 'AlternateContent')
    choice = etree.SubElement(alternative, MC + 'Choice', nsmap={'wps': WPS[1:-1]}, Requires='wps')
    fallback = etree.SubElement(alternative, MC + 'Fallback')
    picture = etree.SubElement(fallback, W + 'pict')
    textbox = etree.SubElement(picture, '{urn:schemas-microsoft-com:vml}textbox')
    textbox.append(replica)
    if conflicting_replica:
        next(replica.iter(W + 't')).text = 'Conflicting visible text'
    owner.replace(drawing, alternative)
    choice.append(drawing)
    if conflicting_replica:
        with pytest.raises(ValueError, match='replica disagrees'):
            positioned_text_source(paragraph)
    else:
        native.save(output)
        imported = read_document(output)
        assert imported.success, imported.issues
        assert not any(issue.code == 'docx.text-boxes' for issue in imported.issues)
        assert imported.document.sections[0].blocks[0].plain_text == 'Own two-column source'


@pytest.mark.parametrize('page_rotation', [0, 90, 180, 270])
@pytest.mark.parametrize('text_rotation', [0, 90])
def test_rotated_native_text_and_page_use_model_geometry(tmp_path, page_rotation, text_rotation):
    source, output = tmp_path / 'rotated.pdf', tmp_path / 'rotated.docx'
    with pymupdf.open() as native:
        page = native.new_page(width=400, height=500)
        page.insert_text((150, 220), 'Editable rotation', fontsize=12, fontname='Times-Roman', rotate=text_rotation)
        page.set_rotation(page_rotation)
        native.save(source)
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    model = read_document(output).document
    block = model.sections[0].blocks[0]
    assert block.plain_text == 'Editable rotation'
    assert block.box.rotation % 360 == (page_rotation - text_rotation) % 360
    expected_size = (500, 400) if page_rotation % 180 else (400, 500)
    assert (model.sections[0].page.width.pt, model.sections[0].page.height.pt) == expected_size
    for index in range(2):
        output = tmp_path / f'rotated-{index}.docx'
        report = write_document(document_from_json(document_to_json(model)), output)
        assert report.success, report.issues
        model = read_document(output).document
        assert model.sections[0].blocks[0].box == block.box
        assert model.sections[0].blocks[0].plain_text == block.plain_text


def test_embedded_pdf_font_is_a_model_resource_and_embedded_in_word(tmp_path):
    import zipfile

    from opendoc_model import ResourceKind

    from opendoc_formats.fonts import FontResolver

    faces = [face for face in FontResolver.system().faces if face.embeddable and ord('A') in face.glyphs]
    if not faces:
        pytest.skip('No embeddable font available for own fixture')
    face = faces[0]
    source, output = tmp_path / 'font.pdf', tmp_path / 'font.docx'
    with pymupdf.open() as native:
        page = native.new_page()
        page.insert_font(fontname='OwnEmbedded', fontfile=str(face.path))
        page.insert_text((54, 90), 'A native embedded font', fontsize=12, fontname='OwnEmbedded')
        native.save(source)
    model = read_document(source).document
    fonts = [resource for resource in model.resources.values() if resource.kind is ResourceKind.FONT]
    assert len(fonts) == 1 and fonts[0].data
    restored = document_from_json(document_to_json(model))
    assert restored.resources[fonts[0].id].data == fonts[0].data
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    assert not any(issue.feature in {'font-unavailable', 'font-glyphs', 'font-resource'} for issue in report.issues)
    resolutions = report.metrics['font_resolution']['resolutions']
    assert any(item['exact'] and 'font_resources_' in item['path'] for item in resolutions.values())
    assert all(not Path(item['path']).exists() for item in resolutions.values())
    with zipfile.ZipFile(output) as archive:
        assert any(name.startswith('word/fonts/') and name.endswith('.odttf') for name in archive.namelist())


def test_uninstalled_font_program_survives_two_docx_json_cycles(tmp_path):
    import zipfile

    from opendoc_model import ResourceKind

    from tests.writers.test_pdf_unicode import _font

    font = tmp_path / 'own.ttf'
    _font(font)
    source, output = tmp_path / 'font.pdf', tmp_path / 'font.docx'
    with pymupdf.open() as native:
        page = native.new_page()
        page.insert_font(fontname='Own', fontfile=str(font))
        page.insert_text((54, 90), 'f i t', fontsize=12, fontname='Own')
        native.save(source)
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    font.unlink()  # No file or installed-font fallback can supply the program.
    for index in range(2):
        imported = read_document(output)
        assert imported.success, imported.issues
        model = document_from_json(document_to_json(imported.document))
        assert any(resource.kind is ResourceKind.FONT and resource.data for resource in model.resources.values())
        assert model.sections[0].blocks[0].plain_text == 'f i t'
        output = tmp_path / f'font-cycle-{index}.docx'
        report = write_document(model, output)
        assert report.success, report.issues
        assert not any(issue.feature in {'font-unavailable', 'font-glyphs', 'font-resource'} for issue in report.issues)
        with zipfile.ZipFile(output) as archive:
            assert any(name.startswith('word/fonts/') and name.endswith('.odttf') for name in archive.namelist())


def test_positioned_table_and_native_borders_survive_two_native_cycles(tmp_path):
    from opendoc_model import Table, get_borders

    source, output = tmp_path / 'table.pdf', tmp_path / 'table.docx'
    ruled_table(source)
    report = PdfModelConverter().convert(source, output)
    assert report.success, report.issues
    assert not any(issue.feature == 'image' for issue in report.issues)
    for index in range(2):
        model = read_document(output).document
        table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
        assert (table.box.x, table.box.y, table.box.width, table.box.height) == (60, 100, 300, 90)
        assert table.rows[0].cells[0].blocks[0].plain_text == 'Cell 1-1'
        assert table.properties['grid_widths_twips'] == [3000, 3000]
        assert table.rows[0].properties['docx_height_twips'] == 600
        assert get_borders(table.rows[0].cells[0]).sides['top'].width_pt == .75
        output = tmp_path / f'table-cycle-{index}.docx'
        report = write_document(document_from_json(document_to_json(model)), output)
        assert report.success, report.issues
    native = Document(output)
    assert len(native.tables) == 1
    assert native.tables[0].cell(2, 1).text == 'Cell 3-2'
    assert not native.element.xpath('.//wp:anchor')
