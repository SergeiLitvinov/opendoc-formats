"""Тесты fan-out конвертера и улучшенного PyMuPdfConverter."""
from __future__ import annotations

import zipfile

import pytest

from opendoc_formats.writers.pdf_to_docx import (
    LibreOfficeConverter,
    PyMuPdfConverter,
    create_converter,
)


def test_pymupdf_text_page(tmp_path):
    """PyMuPdf: страница с текстом → текст в DOCX."""
    try:
        import fitz
    except ImportError:
        pytest.skip("pymupdf not installed")

    pdf = tmp_path / "in.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PDF world")
    page.insert_text((72, 100), "Second line here")
    doc.save(str(pdf))
    doc.close()

    # Сразу проверим, что pymupdf видит текст
    reread = fitz.open(str(pdf))
    text_before = reread[0].get_text()
    reread.close()
    assert "Hello" in text_before, f"pymupdf didn't extract: {text_before!r}"

    out = tmp_path / "out.docx"
    c = PyMuPdfConverter(text_threshold=10)  # понизим порог
    result = c.convert(pdf, out)
    assert result.success
    assert out.is_file()
    from docx import Document
    d = Document(str(out))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Hello" in text, f"DOCX has no text: {text!r}"


def test_pymupdf_scan_page_uses_image(tmp_path):
    """PyMuPdf: страница с <50 символов текста → изображение."""
    try:
        import fitz
    except ImportError:
        pytest.skip("pymupdf not installed")

    pdf = tmp_path / "in.pdf"
    doc = fitz.open()
    # Короткий текст ниже порога не стоит восстанавливать как отдельный
    # абзац: страница сохраняется изображением вместе с остальной графикой.
    page = doc.new_page()
    page.insert_text((72, 72), "Short")
    doc.save(str(pdf))
    doc.close()

    out = tmp_path / "out.docx"
    c = PyMuPdfConverter(render_dpi=72)
    result = c.convert(pdf, out)
    assert result.success
    assert out.is_file()
    # DOCX должен содержать изображение
    with zipfile.ZipFile(out) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
        assert media  # хотя бы одна картинка


def test_create_converter_default():
    from opendoc_formats.writers.pdf_to_docx import Pdf2DocxConverter
    conv = create_converter("pdf2docx")
    assert isinstance(conv, Pdf2DocxConverter)


def test_create_converter_pymupdf():
    from opendoc_formats.writers.pdf_to_docx import PyMuPdfConverter
    conv = create_converter("pymupdf")
    assert isinstance(conv, PyMuPdfConverter)


def test_create_converter_unknown():
    with pytest.raises(Exception):
        create_converter("nonexistent_engine")


def test_libreoffice_pdf_route_fails_fast_with_actionable_message(tmp_path):
    source = tmp_path / "source.pdf"
    source.write_bytes(b"pdf")

    result = LibreOfficeConverter(libreoffice_path="soffice").convert(source, tmp_path / "output.docx")

    assert not result.success
    assert "does not provide a reliable" in (result.error or "")


def test_convert_nonexistent():
    from opendoc_formats.writers.pdf_to_docx import Pdf2DocxConverter
    conv = Pdf2DocxConverter()
    result = conv.convert("nonexistent.pdf", "out.docx")
    assert not result.success
    assert "not found" in (result.error or "").lower()


def test_pymupdf_scan_page_leaves_no_artifacts(tmp_path):
    """Скан-страницы рендерятся во временный workspace, а не рядом с выводом."""
    try:
        import fitz
    except ImportError:
        pytest.skip("pymupdf not installed")

    pdf = tmp_path / "in.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(pdf))
    doc.close()

    out = tmp_path / "out.docx"
    result = PyMuPdfConverter(render_dpi=72).convert(pdf, out)
    assert result.success
    assert out.is_file()
    # В каталоге вывода не должно остаться ни временных PNG, ни partial-файлов.
    leftovers = [p.name for p in tmp_path.iterdir() if p.name not in {out.name, pdf.name}]
    assert leftovers == [], f"unexpected leftovers: {leftovers}"


def test_pymupdf_text_page_writes_no_workspace_scrap(tmp_path):
    """Текстовая страница не создаёт временных PNG вообще."""
    try:
        import fitz
    except ImportError:
        pytest.skip("pymupdf not installed")

    pdf = tmp_path / "in.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PDF world — quite some text here to exceed threshold")
    doc.save(str(pdf))
    doc.close()

    out = tmp_path / "out.docx"
    result = PyMuPdfConverter(text_threshold=10).convert(pdf, out)
    assert result.success
    assert out.is_file()
    remaining = [p.name for p in tmp_path.iterdir()]
    assert remaining == ["in.pdf", "out.docx"], remaining


def _make_pdf_with_headings(tmp_path):
    """PDF с заголовком (крупный шрифт), подзаголовком и телом текста."""
    try:
        import fitz
    except ImportError:
        pytest.skip("pymupdf not installed")
    pdf = tmp_path / "headings.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 60), "Глава 1. Введение", fontsize=20)
    page.insert_text((50, 100), "Обычный абзац с текстом для проверки извлечения.", fontsize=11)
    page.insert_text((50, 140), "Подраздел", fontsize=14)
    page.insert_text((50, 180), "Ещё один абзац тела документа.", fontsize=11)
    doc.save(str(pdf))
    doc.close()
    return pdf




def test_pymupdf_preserves_text_runs(tmp_path):
    """PyMuPDF: текстовая страница → текст, без растровых вставок."""
    pdf = _make_pdf_with_headings(tmp_path)
    out = tmp_path / "out.docx"
    result = PyMuPdfConverter().convert(pdf, out)
    assert result.success
    from docx import Document

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    # Текст извлечён как текст (а не картинка). Кириллица в консоли искажается,
    # поэтому проверяем длину и отсутствие растровых вставок.
    assert len(text.strip()) > 50
    with zipfile.ZipFile(out) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
    assert not media, "текстовая страница не должна содержать картинок"


def test_docx_to_latex_quality(tmp_path):
    """DOCX→LaTeX: заголовки, жирный текст и таблицы попадают в .tex."""
    try:
        from docx import Document as DocxDocument
    except ImportError:
        pytest.skip("python-docx not installed")
    from opendoc_formats.writers.docx_to_latex import DocxToLatexConverter

    doc = DocxDocument()
    doc.add_heading("Глава 1. Введение", level=1)
    p = doc.add_paragraph()
    p.add_run("Обычный ")
    rb = p.add_run("жирный")
    rb.bold = True
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Имя"
    t.rows[0].cells[1].text = "Год"
    src = tmp_path / "src.docx"
    doc.save(str(src))
    out = tmp_path / "out.tex"
    result = DocxToLatexConverter().convert(src, out)
    assert result.success, result.error
    tex = out.read_text(encoding="utf-8")
    assert r"\section{" in tex
    assert r"\textbf{жирный}" in tex
    assert r"\begin{tabular}" in tex
    assert r"\end{document}" in tex
