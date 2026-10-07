"""Chart kind and numeric data must survive mutation, JSON and repeated export."""

import pytest

pytest.importorskip("pptx")
from opendoc_model.document_codec import load_document, save_document
from opendoc_model.document_model import DocumentModel, Paragraph, Section, TextRun
from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE as K
from pptx.util import Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model


def roundtrip(model, tmp_path, index):
    json_path, output = tmp_path / "model.json", tmp_path / f"result{index}.pptx"
    save_document(model, json_path)
    report = write_pptx_model(load_document(json_path), output)
    assert report.success
    return read_pptx_model(output), Presentation(output), report


@pytest.mark.parametrize(
    "kind",
    [
        K.COLUMN_STACKED,
        K.COLUMN_STACKED_100,
        K.BAR_STACKED,
        K.BAR_STACKED_100,
        K.LINE_STACKED,
        K.LINE_STACKED_100,
        K.AREA_STACKED,
        K.AREA_STACKED_100,
    ],
)
def test_stacked_series_and_gaps(tmp_path, kind):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B", "C"]
    data.add_series("First", [1, None, 3])
    data.add_series("Second", [4, 5, 6])
    slide.shapes.add_chart(kind, Pt(20), Pt(20), Pt(400), Pt(300), data)
    presentation.save(source)
    model = read_pptx_model(source)
    chart_data = model.sections[0].blocks[0].properties["pptx"]["chart"]
    assert chart_data["series"][0]["values"] == ["1", None, "3"]
    chart_data["series"][1]["values"][2] = 27
    for index in range(2):
        model, target, _ = roundtrip(model, tmp_path, index)
        chart = target.slides[0].shapes[0].chart
        assert chart.chart_type == kind
        assert list(chart.series[0].values) == [1, None, 3]
        assert list(chart.series[1].values) == [4, 5, 27]


@pytest.mark.parametrize("bubble", [False, True])
def test_numeric_series_mutation(tmp_path, bubble):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = BubbleChartData() if bubble else XyChartData()
    for name, points in (("First", [(1, 2), (3, 4)]), ("Second", [(8, 9)])):
        series = data.add_series(name)
        for x, y in points:
            series.add_data_point(x, y, 7) if bubble else series.add_data_point(x, y)
    slide.shapes.add_chart(K.BUBBLE if bubble else K.XY_SCATTER, Pt(20), Pt(20), Pt(400), Pt(300), data)
    presentation.save(source)
    model = read_pptx_model(source)
    series = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"]
    assert series[0]["x_values"] == ["1", "3"]
    series[0]["x_values"][1] = 33
    series[0]["values"][0] = 22
    if bubble:
        series[0]["bubble_sizes"][1] = 17
    for index in range(2):
        model, target, _ = roundtrip(model, tmp_path, index)
        assert target.slides[0].shapes[0].has_chart
        series = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"]
        assert list(map(float, series[0]["x_values"])) == [1, 33]
        assert list(map(float, series[0]["values"])) == [22, 4]
        assert list(map(float, series[1]["x_values"])) == [8]
        if bubble:
            assert list(map(float, series[0]["bubble_sizes"])) == [7, 17]
        from io import BytesIO
        from zipfile import ZipFile

        from lxml import etree

        blob = target.slides[0].shapes[0].chart.part.chart_workbook.xlsx_part.blob
        with ZipFile(BytesIO(blob)) as workbook:
            root = etree.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
            ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            numbers = {float(value.text) for value in root.iter(ns + "v")}
            assert {22, 33} <= numbers
            if bubble:
                assert 17 in numbers


@pytest.mark.parametrize("invalid", ["nan", "inf", None])
def test_invalid_xy_is_not_silently_published_as_chart(tmp_path, invalid):
    data = {"chart_type": "scatterChart", "series": [{"name": "Series", "x_values": [invalid], "values": [2]}]}
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Summary")], properties={"pptx": {"chart": data}})])])
    _, target, report = roundtrip(model, tmp_path, 0)
    assert not target.slides[0].shapes[0].has_chart
    assert any(issue.feature == "charts" and issue.severity.value == "loss" for issue in report.issues)


def test_xy_axes_are_not_misidentified_as_secondary_y(tmp_path):
    from lxml import etree

    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = XyChartData()
    data.add_series("Measured").add_data_point(10, 20)
    chart = slide.shapes.add_chart(K.XY_SCATTER, Pt(10), Pt(10), Pt(400), Pt(300), data).chart
    chart.category_axis.minimum_scale, chart.category_axis.maximum_scale = 1, 100
    chart.category_axis.has_title = True
    chart.category_axis.axis_title.text_frame.text = "Frequency"
    chart.value_axis.maximum_scale = 80
    chart.value_axis.major_unit = 20
    chart.value_axis.reverse_order = True
    etree.SubElement(
        chart.category_axis._element.scaling, "{http://schemas.openxmlformats.org/drawingml/2006/chart}logBase", val="10"
    )
    presentation.save(source)
    model = read_pptx_model(source)
    axes = model.sections[0].blocks[0].properties["pptx"]["chart"]["axes"]
    assert set(axes) == {"x", "y"}
    axes["y"]["max"] = 120
    for index in range(2):
        model, target, _ = roundtrip(model, tmp_path, index)
        chart = target.slides[0].shapes[0].chart
        assert chart.category_axis.minimum_scale == 1
        assert chart.category_axis.axis_title.text_frame.text == "Frequency"
        assert chart.value_axis.maximum_scale == 120
        assert chart.value_axis.major_unit == 20
        assert chart.value_axis.reverse_order
        axes = model.sections[0].blocks[0].properties["pptx"]["chart"]["axes"]
        assert axes["x"]["log_base"] == 10


@pytest.mark.parametrize(
    "kind", [K.XY_SCATTER_LINES, K.XY_SCATTER_LINES_NO_MARKERS, K.XY_SCATTER_SMOOTH, K.XY_SCATTER_SMOOTH_NO_MARKERS]
)
def test_scatter_line_style(tmp_path, kind):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = XyChartData()
    series = data.add_series("Curve")
    series.add_data_point(1, 2)
    series.add_data_point(2, 4)
    slide.shapes.add_chart(kind, Pt(20), Pt(20), Pt(400), Pt(300), data)
    presentation.save(source)
    model = read_pptx_model(source)
    for index in range(2):
        model, target, _ = roundtrip(model, tmp_path, index)
        assert target.slides[0].shapes[0].chart.chart_type == kind


def test_combo_is_not_silently_flattened(tmp_path):
    data = {
        "chart_type": "barChart",
        "combo_types": ["barChart", "lineChart"],
        "categories": ["A"],
        "series": [{"name": "S", "values": [1]}],
    }
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Summary")], properties={"pptx": {"chart": data}})])])
    _, target, report = roundtrip(model, tmp_path, 0)
    assert not target.slides[0].shapes[0].has_chart
    assert any(issue.feature == "charts" for issue in report.issues)
