"""Mixed native plots keep workbook columns and primary/secondary axis ownership."""

from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pytest

pytest.importorskip("pptx")
from lxml import etree
from opendoc_model.document_codec import load_document, save_document
from pptx import Presentation
from pptx.chart.axis import ValueAxis
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE as K
from pptx.oxml import parse_xml
from pptx.util import Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def source_combo(path, *, secondary):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Count", [1, 2])
    data.add_series("Revenue", [1000, 2000])
    chart = slide.shapes.add_chart(K.COLUMN_CLUSTERED, Pt(20), Pt(20), Pt(400), Pt(300), data).chart
    area = chart._chartSpace.find(".//" + C + "plotArea")
    bars = area.find(C + "barChart")
    line = parse_xml(f'<c:lineChart xmlns:c="{C[1:-1]}"><c:grouping val="standard"/></c:lineChart>')
    line.append(bars.findall(C + "ser")[1])
    axes = [area.find(C + tag) for tag in ("catAx", "valAx")]
    if secondary:
        extra = [deepcopy(axis) for axis in axes]
        for index, axis in enumerate(extra):
            axis.find(C + "axId").set("val", str(501 + index))
            axis.find(C + "crossAx").set("val", str(502 - index))
            line.append(etree.Element(C + "axId", val=str(501 + index)))
            area.append(axis)
        extra[1].find(C + "axPos").set("val", "r")
        ValueAxis(extra[1]).maximum_scale = 15000
        ValueAxis(extra[1]).major_unit = 3000
        ValueAxis(extra[1]).has_title = True
        ValueAxis(extra[1]).axis_title.text_frame.text = "Revenue scale"
    else:
        for axis in axes:
            line.append(etree.Element(C + "axId", val=axis.find(C + "axId").get("val")))
    area.insert(list(area).index(bars) + 1, line)
    ValueAxis(axes[1]).maximum_scale = 10
    presentation.save(path)


@pytest.mark.parametrize("secondary", [False, True])
@pytest.mark.parametrize("second_type", ["lineChart", "areaChart", "barChart"])
def test_import_mutate_combo_twice(tmp_path, secondary, second_type):
    path = tmp_path / "source.pptx"
    source_combo(path, secondary=secondary)
    model = read_pptx_model(path)
    data = model.sections[0].blocks[0].properties["pptx"]["chart"]
    assert data["combo_types"] == ["barChart", "lineChart"]
    data["combo_types"][1] = second_type
    data["series"][1]["chart_type"] = second_type
    if second_type in ("barChart", "areaChart"):
        data["series"][1]["plot"]["grouping"] = "stacked"
    data["series"][1]["values"] = [6000, 9000]
    if secondary:
        assert data["series"][1]["axis"] == "secondary_value"
        data["axes"]["secondary_value"]["max"] = 24000
    for index in range(2):
        saved, output = tmp_path / "model.json", tmp_path / f"output{index}.pptx"
        save_document(model, saved)
        report = write_pptx_model(load_document(saved), output)
        assert report.success
        assert not any(issue.feature == "charts" for issue in report.issues), report.to_dict()
        chart = Presentation(output).slides[0].shapes[0].chart
        assert len(chart.plots) == 2
        area = chart._chartSpace.find(".//" + C + "plotArea")
        plots = [element for element in area if element.tag.endswith("Chart")]
        bars, line = plots
        assert bars.tag == C + "barChart" and line.tag == C + second_type
        assert bars is not None and line is not None
        bar_ids = [node.get("val") for node in bars.findall(C + "axId")]
        line_ids = [node.get("val") for node in line.findall(C + "axId")]
        assert (bar_ids != line_ids) is secondary
        for axis in area.findall(C + "valAx"):
            expected = 10 if axis.find(C + "axId").get("val") == bar_ids[1] else 24000
            assert ValueAxis(axis).maximum_scale == expected
        for axis in area.findall(C + "catAx") + area.findall(C + "valAx"):
            partner = axis.find(C + "crossAx").get("val")
            assert partner in [node.get("val") for node in area.findall("*/" + C + "axId")]
        with ZipFile(BytesIO(chart.part.chart_workbook.xlsx_part.blob)) as workbook:
            sheet = etree.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
            ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            values = {cell.get("r"): cell.find(ns + "v").text for cell in sheet.iter(ns + "c")}
            assert float(values["B2"]) == 1 and float(values["C2"]) == 6000
        model = read_pptx_model(output)
        data = model.sections[0].blocks[0].properties["pptx"]["chart"]
        assert data["combo_types"] == ["barChart", second_type]
        if second_type in ("barChart", "areaChart"):
            assert data["series"][1]["plot"]["grouping"] == "stacked"
        assert list(map(float, data["series"][1]["values"])) == [6000, 9000]
        if secondary:
            assert data["axes"]["secondary_value"]["title"] == "Revenue scale"


def test_incompatible_category_domains_are_reported(tmp_path):
    path = tmp_path / "source.pptx"
    source_combo(path, secondary=True)
    model = read_pptx_model(path)
    data = model.sections[0].blocks[0].properties["pptx"]["chart"]
    data["series"][1]["categories"] = ["X", "Y"]
    output = tmp_path / "fallback.pptx"
    report = write_pptx_model(model, output)
    assert report.success and not report.lossless
    assert not Presentation(output).slides[0].shapes[0].has_chart


def test_moving_series_to_secondary_axis(tmp_path):
    source = tmp_path / "source.pptx"
    source_combo(source, secondary=False)
    model = read_pptx_model(source)
    data = model.sections[0].blocks[0].properties["pptx"]["chart"]
    data["series"][1]["axis"] = "secondary_value"
    data["axes"]["secondary_value"] = {"min": 0, "max": 7000, "major_unit": 1000}
    output = tmp_path / "moved.pptx"
    assert write_pptx_model(model, output).success
    reread = read_pptx_model(output).sections[0].blocks[0].properties["pptx"]["chart"]
    assert reread["series"][1]["axis"] == "secondary_value"
    assert reread["axes"]["secondary_value"]["max"] == 7000


def test_same_type_on_two_axes_without_combo_hint(tmp_path):
    source = tmp_path / "source.pptx"
    source_combo(source, secondary=True)
    model = read_pptx_model(source)
    data = model.sections[0].blocks[0].properties["pptx"]["chart"]
    data.pop("combo_types")
    data["series"][1]["chart_type"] = "barChart"
    output = tmp_path / "same.pptx"
    report = write_pptx_model(model, output)
    assert report.success and not any(issue.feature == "charts" for issue in report.issues)
    reread = read_pptx_model(output).sections[0].blocks[0].properties["pptx"]["chart"]
    assert reread["combo_types"] == ["barChart", "barChart"]
    assert reread["series"][1]["axis"] == "secondary_value"


@pytest.mark.parametrize("secondary", [False, True])
@pytest.mark.parametrize("grouping", ["clustered", "stacked", "percentStacked"])
def test_horizontal_plots_keep_axes_and_workbook_twice(tmp_path, secondary, grouping):
    source = tmp_path / "source.pptx"
    source_combo(source, secondary=secondary)
    model = read_pptx_model(source)
    data = model.sections[0].blocks[0].properties["pptx"]["chart"]
    data["combo_types"] = ["barChart", "barChart"]
    data["bar_direction"] = "bar"
    for series in data["series"]:
        series["chart_type"] = "barChart"
        series["plot"]["bar_direction"] = "bar"
    data["series"][1]["plot"]["grouping"] = grouping
    data["series"][1]["values"] = [3500, 7000]
    for index in range(2):
        saved, output = tmp_path / "model.json", tmp_path / f"horizontal{index}.pptx"
        save_document(model, saved)
        report = write_pptx_model(load_document(saved), output)
        assert report.success and not any(issue.feature == "charts" for issue in report.issues), report.to_dict()
        chart = Presentation(output).slides[0].shapes[0].chart
        area = chart._chartSpace.find(".//" + C + "plotArea")
        plots = area.findall(C + "barChart")
        assert len(plots) == 2
        assert all(plot.find(C + "barDir").get("val") == "bar" for plot in plots)
        assert plots[1].find(C + "grouping").get("val") == grouping
        axes = area.findall(C + "valAx")
        assert [axis.find(C + "axPos").get("val") for axis in axes] == (["b", "t"] if secondary else ["b"])
        assert ValueAxis(axes[0]).maximum_scale == 10
        if secondary:
            assert ValueAxis(axes[1]).maximum_scale == 15000
        with ZipFile(BytesIO(chart.part.chart_workbook.xlsx_part.blob)) as workbook:
            sheet = etree.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
            ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            values = {cell.get("r"): cell.find(ns + "v").text for cell in sheet.iter(ns + "c")}
            assert float(values["C2"]) == 3500 and float(values["C3"]) == 7000
        model = read_pptx_model(output)
        data = model.sections[0].blocks[0].properties["pptx"]["chart"]
        assert data["series"][1]["plot"]["bar_direction"] == "bar"
        assert data["series"][1].get("axis", "value") == ("secondary_value" if secondary else "value")


@pytest.mark.parametrize("second_kind,second_direction", [("lineChart", "bar"), ("barChart", "col")])
def test_incompatible_horizontal_orientation_is_reported(tmp_path, second_kind, second_direction):
    source = tmp_path / "source.pptx"
    source_combo(source, secondary=True)
    model = read_pptx_model(source)
    data = model.sections[0].blocks[0].properties["pptx"]["chart"]
    data["combo_types"][1] = second_kind
    data["series"][0]["plot"]["bar_direction"] = "bar"
    data["series"][1]["chart_type"] = second_kind
    data["series"][1]["plot"]["bar_direction"] = second_direction
    output = tmp_path / "unsupported.pptx"
    report = write_pptx_model(model, output)
    assert report.success and not report.lossless
    assert any(issue.feature == "charts" for issue in report.issues)
    assert not Presentation(output).slides[0].shapes[0].has_chart
