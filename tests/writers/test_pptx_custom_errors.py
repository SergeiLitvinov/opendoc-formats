"""Пользовательские погрешности не смещают точки при пропусках и изменениях."""

from __future__ import annotations

import pytest

pytest.importorskip("pptx")
from lxml import etree
from opendoc_model.document_codec import document_from_json, document_to_json
from opendoc_model.document_model import DocumentModel, Paragraph, Section, TextRun
from pptx import Presentation

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.readers.pptx_numeric_cache import C, indexed_numbers
from opendoc_formats.writers.pptx_writer import write_pptx_model


def _model(kind, settings):
    series = {"name": "Data", "values": [2, 4, 7, 9], "error_bars": settings}
    if kind in {"scatterChart", "bubbleChart"}:
        series["x_values"] = [1, 3, 5, 8]
    if kind == "bubbleChart":
        series["bubble_sizes"] = [3, 2, 1, 4]
    chart = {"chart_type": kind, "categories": ["A", "B", "C", "D"], "series": [series]}
    return DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Chart")], properties={"pptx": {"chart": chart}})])])


@pytest.mark.parametrize("kind", ["barChart", "scatterChart", "bubbleChart"])
@pytest.mark.parametrize("mode", ["both", "plus", "minus"])
def test_custom_error_mutation(tmp_path, kind, mode):
    values = {"plus": [0, None, 0.5, None], "minus": [0.1, 0.2, None, 0.4]}
    sides = ("plus", "minus") if mode == "both" else (mode,)
    settings = {"value_type": "cust", "bar_type": mode, **{side: values[side] for side in sides}}
    model = _model(kind, settings)
    for cycle in range(3):
        output = tmp_path / f"custom{cycle}.pptx"
        report = write_pptx_model(document_from_json(document_to_json(model)), output)
        assert report.success and not any("error_bars" in issue.message for issue in report.issues)
        native = Presentation(output).slides[0].shapes[0].chart.series[0]._element
        model = read_pptx_model(output)
        series = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]
        if cycle == 2:
            assert "error_bars" not in series
            assert native.find(C + "errBars") is None
            break
        for side in sides:
            actual = [None if item is None else float(item) for item in series["error_bars"][side]]
            assert actual == values[side]
            cache = native.find(f"{C}errBars/{C}{side}/{C}numLit")
            assert cache.find(C + "ptCount").get("val") == "4"
            assert [int(point.get("idx")) for point in cache.findall(C + "pt")] == [
                i for i, value in enumerate(values[side]) if value is not None
            ]
        if cycle == 0:
            for side in sides:
                values[side] = [0.3, None, 0, 0.8]
                series["error_bars"][side] = values[side]
        else:
            del series["error_bars"]


def test_import_reference_cache_preserves_order_and_trailing_gaps(tmp_path):
    output = tmp_path / "cached.pptx"
    assert write_pptx_model(_model("barChart", {"value_type": "cust", "bar_type": "plus", "plus": [1, 2, 3, 4]}), output).success
    presentation = Presentation(output)
    plus = presentation.slides[0].shapes[0].chart.series[0]._element.find(f"{C}errBars/{C}plus")
    plus.clear()
    reference = etree.SubElement(plus, C + "numRef")
    etree.SubElement(reference, C + "f").text = "Sheet1!$C$2:$C$5"
    cache = etree.SubElement(reference, C + "numCache")
    etree.SubElement(cache, C + "ptCount", val="4")
    for index, value in ((2, "0.5"), (0, "0")):
        etree.SubElement(etree.SubElement(cache, C + "pt", idx=str(index)), C + "v").text = value
    presentation.save(output)
    model = read_pptx_model(output)
    errors = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]["error_bars"]
    assert errors["plus"] == ["0", None, "0.5", None]
    target = tmp_path / "materialized.pptx"
    assert write_pptx_model(model, target).success
    errors = read_pptx_model(target).sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]["error_bars"]
    assert errors["plus"] == ["0.0", None, "0.5", None]


@pytest.mark.parametrize(
    "settings",
    [
        {"value_type": "cust", "plus": [1, 2, 3, 4]},
        {"value_type": "cust", "bar_type": "plus", "plus": [1, 2]},
        {"value_type": "cust", "bar_type": "plus", "plus": [1, -2, 3, 4]},
        {"value_type": "cust", "bar_type": "plus", "plus": [1, True, 3, 4]},
        {"value_type": "cust", "bar_type": "plus", "plus": [1, float("inf"), 3, 4]},
        {"value_type": "cust", "bar_type": "plus", "plus": [1, 2, 3, 4], "value": 2},
        {"value_type": "fixedVal", "value": 2, "plus": [1, 2, 3, 4]},
    ],
)
def test_invalid_custom_arrays_are_rejected_atomically(tmp_path, settings):
    import math

    output = tmp_path / "invalid.pptx"
    report = write_pptx_model(_model("barChart", settings), output)
    if any(type(value) is float and not math.isfinite(value) for value in settings.get("plus", [])):
        assert not report.success and not output.exists()
        assert any(issue.feature == "model" and "number must be finite" in issue.message for issue in report.issues)
        return
    assert report.success and any("error_bars" in issue.message for issue in report.issues)
    assert Presentation(output).slides[0].shapes[0].chart.series[0]._element.find(C + "errBars") is None


@pytest.mark.parametrize("count,indices", [(100001, []), (4, [-1]), (4, [0, 0]), (4, [4])])
def test_invalid_cache_is_bounded(count, indices):
    node = etree.Element(C + "plus")
    cache = etree.SubElement(node, C + "numLit")
    etree.SubElement(cache, C + "ptCount", val=str(count))
    for index in indices:
        etree.SubElement(etree.SubElement(cache, C + "pt", idx=str(index)), C + "v").text = "1"
    assert indexed_numbers(node) == []


def test_sparse_errors_remain_visible_in_html():
    from opendoc_formats.writers.html_writer import _error_bar_offsets, _error_bar_parts

    errors = {"value_type": "cust", "plus": [0.2, None, 0.6], "minus": [0.1, None, 0.3], "bar_type": "plus"}
    item = {"error_bars": errors, "color": "#123456"}
    assert _error_bar_offsets(item, errors, 0, 5) == (0.2, 0.1)
    assert _error_bar_offsets(item, errors, 1, 5) == (None, None)
    assert _error_bar_offsets(item, errors, 2, 5) == (0.6, 0.3)
    parts = _error_bar_parts(item, 2, 5, 50, 50, minimum=0, maximum=10, top=0, height=100)
    assert 'y1="44"' in parts[0] and 'y2="50"' in parts[0]
