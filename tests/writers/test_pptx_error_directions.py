"""Независимые направления погрешностей XY и пузырьковых диаграмм."""

from __future__ import annotations

import pytest

pytest.importorskip("pptx")
from opendoc_model.document_codec import document_from_json, document_to_json
from opendoc_model.document_model import DocumentModel, Paragraph, Section, TextRun
from pptx import Presentation

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.html_writer import _error_bar_parts
from opendoc_formats.writers.pptx_writer import write_pptx_model

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _model(kind, errors):
    series = {"name": "Data", "values": [2, 4, 7], "x_values": [1, 3, 6], "bubble_sizes": [1, 2, 3], "error_bars": errors}
    data = {"chart_type": kind, "categories": ["A", "B", "C"], "series": [series]}
    return DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Chart")], properties={"pptx": {"chart": data}})])])


@pytest.mark.parametrize("kind", ["scatterChart", "bubbleChart"])
@pytest.mark.parametrize("reverse", [False, True])
def test_two_directions_mutation_and_removal(tmp_path, kind, reverse):
    errors = [
        {"direction": "x", "value_type": "fixedVal", "value": 0.5, "bar_type": "plus", "color": "#118833"},
        {"direction": "y", "value_type": "cust", "plus": [0.2, None, 0.6], "minus": [0, 0.1, None], "color": "#882211"},
    ]
    if reverse:
        errors.reverse()
    model = _model(kind, errors)
    for cycle in range(3):
        output = tmp_path / f"directions{cycle}.pptx"
        report = write_pptx_model(document_from_json(document_to_json(model)), output)
        assert report.success and not any("error_bars" in issue.message for issue in report.issues)
        native = Presentation(output).slides[0].shapes[0].chart.series[0]._element
        assert len(native.findall(C + "errBars")) == (1 if cycle == 2 else 2)
        model = read_pptx_model(output)
        series = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]
        restored = series["error_bars"]
        if cycle == 2:
            assert isinstance(restored, dict) and restored["direction"] == "y"
            assert restored["plus"] == ["0.2", None, "0.6"]
            break
        assert isinstance(restored, list)
        axes = {item["direction"]: item for item in restored}
        assert list(axes) == (["y", "x"] if reverse else ["x", "y"])
        assert axes["x"]["value"] == (0.5 if cycle == 0 else 1.5)
        assert axes["x"]["color"] == "#118833"
        assert axes["x"]["bar_type"] == "plus"
        assert axes["y"]["color"] == "#882211"
        assert axes["y"]["minus"] == ["0.0", "0.1", None]
        if cycle == 0:
            axes["x"]["value"] = 1.5
        else:
            series["error_bars"] = [axes["y"]]


@pytest.mark.parametrize(
    "kind,errors",
    [
        ("scatterChart", [{"direction": "x"}, {"direction": "x"}]),
        ("scatterChart", [{"direction": "x"}, {"direction": "y", "value": -1}]),
        ("scatterChart", [{"direction": "x"}, {"direction": "y"}, {"direction": "z"}]),
        ("scatterChart", []),
        ("scatterChart", [{"direction": "x"}, None]),
        ("barChart", [{"direction": "x"}, {"direction": "y"}]),
    ],
)
def test_invalid_direction_collection_is_atomic(tmp_path, kind, errors):
    output = tmp_path / "invalid.pptx"
    report = write_pptx_model(_model(kind, errors), output)
    assert report.success and any("error_bars" in issue.message for issue in report.issues)
    assert not Presentation(output).slides[0].shapes[0].chart.series[0]._element.findall(C + "errBars")


@pytest.mark.parametrize("direction,color", [("x", "#118833"), ("y", "#882211")])
def test_html_selects_matching_direction(direction, color):
    errors = [
        {"direction": "x", "value_type": "fixedVal", "value": 1, "color": "#118833"},
        {"direction": "y", "value_type": "fixedVal", "value": 2, "color": "#882211"},
    ]
    parts = _error_bar_parts(
        {"error_bars": errors, "color": "#000000"}, 0, 5, 50, 50, minimum=0, maximum=10, top=0, height=100, direction=direction
    )
    assert len(parts) == 3 and all(color in part for part in parts)
