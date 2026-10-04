"""Native topology and geometry survive edits and two serialized round-trips."""

import math
from copy import deepcopy

import pytest

pytest.importorskip("pptx")
from lxml import etree
from opendoc.document_codec import load_document, save_document
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE
from pptx.enum.shapes import MSO_CONNECTOR_TYPE as CONNECTOR
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def points(block):
    box = block.box
    transform = block.properties.get("pptx", {}).get("transform")
    if transform:
        a, b, c, d, e, f = transform["matrix"]
        width, height = transform["width_pt"], transform["height_pt"]
    else:
        theta = math.radians(box.rotation)
        a, b, c, d = math.cos(theta), math.sin(theta), -math.sin(theta), math.cos(theta)
        width, height = box.width, box.height
        e = box.x + width / 2 - (a * width + c * height) / 2
        f = box.y + height / 2 - (b * width + d * height) / 2
    return [(a * x + c * y + e, b * x + d * y + f) for x, y in ((0, 0), (width, 0), (width, height), (0, height))]


def cycle(model, tmp_path, index):
    saved = tmp_path / "model.json"
    save_document(model, saved)
    output = tmp_path / f"result{index}.pptx"
    report = write_pptx_model(load_document(saved), output)
    assert report.success
    return read_pptx_model(output), Presentation(output), report


@pytest.mark.parametrize("kind", [CONNECTOR.STRAIGHT, CONNECTOR.ELBOW, CONNECTOR.CURVE])
@pytest.mark.parametrize("end", [(260, 40), (20, 240), (260, 240), (5, 10)])
def test_connector_native_geometry_and_arrows(tmp_path, kind, end):
    path = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    connector = slide.shapes.add_connector(kind, Pt(20), Pt(40), Pt(end[0]), Pt(end[1]))
    line = connector._element.spPr.get_or_add_ln()
    etree.SubElement(line, A + "prstDash", val="dash")
    etree.SubElement(line, A + "headEnd", type="triangle", w="lg", len="sm")
    etree.SubElement(line, A + "tailEnd", type="oval")
    presentation.save(path)
    model = read_pptx_model(path)
    assert len(model.sections[0].blocks) == 1
    expected = points(model.sections[0].blocks[0])
    for index in range(2):
        model, target, report = cycle(model, tmp_path, index)
        shape = target.slides[0].shapes[0]
        assert shape.shape_type == MSO_SHAPE_TYPE.LINE
        assert shape._element.spPr.find(A + "prstGeom").get("prst") == kind.xml_value
        line = shape._element.spPr.find(A + "ln")
        assert line.find(A + "headEnd").get("type") == "triangle"
        assert line.find(A + "prstDash").get("val") == "dash"
        for actual, wanted in zip(points(model.sections[0].blocks[0]), expected):
            assert actual == pytest.approx(wanted, abs=0.001)
        assert not any(issue.feature == "vector_graphics" for issue in report.issues)


def test_connections_remap_ids_and_missing_target_is_reported(tmp_path):
    path = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    left = slide.shapes.add_shape(SHAPE.RECTANGLE, Pt(20), Pt(20), Pt(60), Pt(40))
    right = slide.shapes.add_shape(SHAPE.RECTANGLE, Pt(200), Pt(20), Pt(60), Pt(40))
    connector = slide.shapes.add_connector(CONNECTOR.STRAIGHT, Pt(80), Pt(40), Pt(200), Pt(40))
    connector.begin_connect(left, 3)
    connector.end_connect(right, 1)
    presentation.save(path)
    model = read_pptx_model(path)
    model.sections[0].blocks.reverse()  # IDs change; references must not follow output position.
    for index in range(2):
        model, target, _ = cycle(model, tmp_path, index)
        connector, right, left = list(target.slides[0].shapes)
        links = connector._element.find(P + "nvCxnSpPr/" + P + "cNvCxnSpPr")
        assert links.find(A + "stCxn").get("id") == str(left.shape_id)
        assert links.find(A + "endCxn").get("id") == str(right.shape_id)
    model.sections[0].blocks.pop()
    _, target, report = cycle(model, tmp_path, 3)
    assert any(issue.feature == "relationships" and issue.severity.value == "loss" for issue in report.issues)
    links = target.slides[0].shapes[0]._element.find(P + "nvCxnSpPr/" + P + "cNvCxnSpPr")
    assert links.find(A + "stCxn") is None


def test_nested_groups_flips_geometry_and_model_translation(tmp_path):
    path = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    outer.name = "Внешняя"
    inner = outer.shapes.add_group_shape()
    inner.name = "Внутренняя"
    shape = inner.shapes.add_shape(SHAPE.CHEVRON, Pt(10), Pt(30), Pt(100), Pt(50))
    shape.text = "Before"
    shape.adjustments[0] = 0.35
    shape.rotation = 20
    outer.rotation = 30
    outer._element.grpSpPr.xfrm.set("flipH", "1")
    presentation.save(path)
    model = read_pptx_model(path)
    block = model.sections[0].blocks[0]
    expected = [(x + 17, y - 8) for x, y in points(block)]
    block.box.x += 17
    block.box.y -= 8
    block.content[0].text = "After"
    for index in range(2):
        model, target, report = cycle(model, tmp_path, index)
        outer = target.slides[0].shapes[0]
        assert outer.shape_type == MSO_SHAPE_TYPE.GROUP and outer.name == "Внешняя"
        inner = outer.shapes[0]
        assert inner.shape_type == MSO_SHAPE_TYPE.GROUP and inner.name == "Внутренняя"
        assert inner.shapes[0].adjustments[0] == pytest.approx(0.35)
        assert inner.shapes[0].text == "After"
        for actual, wanted in zip(points(model.sections[0].blocks[0]), expected):
            assert actual == pytest.approx(wanted, abs=0.001)
        assert not any(issue.feature in ("page_geometry", "groups") for issue in report.issues)


def test_custom_path_mutation_and_unsafe_xml_rejection(tmp_path):
    path = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    builder = slide.shapes.build_freeform(0, 0, scale=Pt(1))
    builder.add_line_segments([(100, 0), (50, 80), (0, 0)], close=True)
    builder.convert_to_shape(Pt(30), Pt(40))
    presentation.save(path)
    model = read_pptx_model(path)
    assert len(model.sections[0].blocks) == 1
    meta = model.sections[0].blocks[0].properties["pptx"]["shape"]
    root = etree.fromstring(meta["geometry_xml"].encode())
    root.find(".//" + A + "lnTo/" + A + "pt").set("x", "120")
    meta["geometry_xml"] = etree.tostring(root, encoding="unicode")
    bad = deepcopy(model)
    for index in range(2):
        model, target, _ = cycle(model, tmp_path, index)
        shape = target.slides[0].shapes[0]
        assert shape.shape_type == MSO_SHAPE_TYPE.FREEFORM
        assert shape._element.spPr.find(".//" + A + "lnTo/" + A + "pt").get("x") == "120"
    bad.sections[0].blocks[0].properties["pptx"]["shape"]["geometry_xml"] = '<!DOCTYPE x [<!ENTITY x "unsafe">]><x/>'
    _, _, report = cycle(bad, tmp_path, 3)
    assert any(issue.feature == "vector_graphics" for issue in report.issues)


def test_shear_is_reported_and_ambiguous_connection_not_rebound(tmp_path):
    from opendoc.document_model import Box, DocumentModel, Paragraph, Section, TextRun

    block = Paragraph(
        [TextRun("Shear")],
        box=Box(10, 10, 100, 50),
        properties={
            "pptx": {
                "object_id": "2",
                "transform": {"matrix": [1, 0, 0.5, 1, 10, 10], "width_pt": 100, "height_pt": 50},
            }
        },
    )
    connector = Paragraph(
        box=Box(10, 60, 100, 0),
        properties={
            "pptx": {
                "shape": {"kind": "cxnSp", "connections": {"stCxn": {"id": "2", "idx": "0"}}},
            }
        },
    )
    model = DocumentModel(sections=[Section(blocks=[block, deepcopy(block), connector])])
    _, target, report = cycle(model, tmp_path, 0)
    assert {"page_geometry", "relationships"} <= {issue.feature for issue in report.issues}
    assert target.slides[0].shapes[2]._element.find(".//" + A + "stCxn") is None


def test_groups_keep_stacking_order_and_slide_local_connections(tmp_path):
    path = tmp_path / "source.pptx"
    presentation = Presentation()
    for number in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Pt(10), Pt(10), Pt(80), Pt(20)).text = "Behind"
        group = slide.shapes.add_group_shape()
        left = group.shapes.add_shape(SHAPE.RECTANGLE, Pt(20), Pt(100), Pt(80), Pt(40))
        left.text = str(number)
        connector = group.shapes.add_connector(CONNECTOR.ELBOW, Pt(100), Pt(120), Pt(200), Pt(120))
        connector.begin_connect(left, 3)
        slide.shapes.add_textbox(Pt(10), Pt(200), Pt(80), Pt(20)).text = "Above"
    presentation.save(path)
    model = read_pptx_model(path)
    for index in range(2):
        model, target, _ = cycle(model, tmp_path, index)
        for number, slide in enumerate(target.slides):
            behind, group, above = list(slide.shapes)
            assert behind.text == "Behind" and above.text == "Above"
            assert group.shapes[0].text == str(number)
            start = group.shapes[1]._element.find(".//" + A + "stCxn")
            assert start.get("id") == str(group.shapes[0].shape_id)


def test_connector_can_reference_a_group(tmp_path):
    path = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_shape(SHAPE.RECTANGLE, Pt(30), Pt(40), Pt(80), Pt(40))
    connector = slide.shapes.add_connector(CONNECTOR.STRAIGHT, Pt(110), Pt(60), Pt(200), Pt(60))
    parent = connector._element.find(P + "nvCxnSpPr/" + P + "cNvCxnSpPr")
    etree.SubElement(parent, A + "stCxn", id=str(group.shape_id), idx="3")
    presentation.save(path)
    model = read_pptx_model(path)
    for index in range(2):
        model, target, report = cycle(model, tmp_path, index)
        group, connector = list(target.slides[0].shapes)
        assert connector._element.find(".//" + A + "stCxn").get("id") == str(group.shape_id)
        assert not any(issue.feature == "relationships" for issue in report.issues)
