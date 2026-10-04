"""Независимые подписи диаграмм и рядов после повторного редактирования."""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("pptx")
from lxml import etree
from opendoc.document_codec import load_document, save_document
from opendoc.document_model import DocumentModel, Paragraph, Section, TextRun
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _chart(model: DocumentModel) -> dict:
    return model.sections[0].blocks[0].properties["pptx"]["chart"]


def _cycle(model: DocumentModel, tmp_path: Path, index: int) -> tuple:
    model_path = tmp_path / "model.json"
    save_document(model, model_path)
    output = tmp_path / f"result{index}.pptx"
    report = write_pptx_model(load_document(model_path), output)
    assert report.success
    return read_pptx_model(output), Presentation(output).slides[0].shapes[0].chart, report


@pytest.mark.parametrize("kind", [XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT])
def test_labels_without_show_flags_and_series_overrides(tmp_path: Path, kind: XL_CHART_TYPE) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("First", [2, 5])
    chart = slide.shapes.add_chart(kind, Pt(10), Pt(10), Pt(400), Pt(300), data).chart
    labels = chart.plots[0]._element.get_or_add_dLbls()
    for child in list(labels):
        labels.remove(child)
    etree.SubElement(labels, C + "numFmt", formatCode="0.00", sourceLinked="0")
    etree.SubElement(labels, C + "dLblPos", val="ctr")
    etree.SubElement(labels, C + "separator").text = ""
    source = tmp_path / "source.pptx"
    presentation.save(source)
    model = read_pptx_model(source)
    assert _chart(model)["data_labels"] == {
        "num_format": "0.00",
        "num_format_linked": False,
        "position": "ctr",
        "separator": "",
    }
    for index in range(2):
        settings = _chart(model)["data_labels"]
        settings.update(show_value=index == 0, show_category=True, separator="\n" if index == 0 else "; ")
        expected = dict(settings)
        series_labels = {"show_value": False, "show_series": True, "position": "ctr"}
        _chart(model)["series"][0]["data_labels"] = series_labels
        model, target, _ = _cycle(model, tmp_path, index)
        assert _chart(model)["data_labels"] == expected
        assert _chart(model)["series"][0]["data_labels"] == series_labels
        assert target.plots[0].data_labels.show_category_name
        assert target.plots[0].data_labels.show_value == (index == 0)
        assert target.plots[0].data_labels.number_format == "0.00"
        assert not target.plots[0].data_labels.number_format_is_linked


def test_combo_plot_labels_remain_independent(tmp_path: Path) -> None:
    data = {
        "chart_type": "barChart",
        "categories": ["A", "B"],
        "combo_types": ["barChart", "lineChart"],
        "data_labels": {"show_value": True, "position": "inEnd"},
        "series": [
            {"name": "Bars", "values": [1, 2], "chart_type": "barChart"},
            {
                "name": "Line",
                "values": [3, 4],
                "chart_type": "lineChart",
                "axis": "secondary_value",
                "plot": {"data_labels": {"show_category": True, "position": "t"}},
            },
        ],
    }
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Chart")], properties={"pptx": {"chart": data}})])])
    for index in range(2):
        _chart(model)["series"][1]["plot"]["data_labels"]["separator"] = str(index)
        model, target, _ = _cycle(model, tmp_path, index)
        assert _chart(model)["data_labels"] == {"show_value": True, "position": "inEnd"}
        assert _chart(model)["series"][1]["plot"]["data_labels"] == {
            "show_category": True,
            "position": "t",
            "separator": str(index),
        }
        assert target.plots[0].data_labels.show_value
        assert target.plots[1].data_labels.show_category_name
        assert not target.plots[1].data_labels.show_value


@pytest.mark.parametrize("settings", [None, {"show_value": "false"}, {"position": "unknown"}, {"separator": 3}])
def test_invalid_label_settings_reported(tmp_path: Path, settings: object) -> None:
    data = {"chart_type": "pieChart", "categories": ["A"], "series": [{"values": [1], "data_labels": settings}]}
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Chart")], properties={"pptx": {"chart": data}})])])
    _, _, report = _cycle(model, tmp_path, 0)
    assert any("подпис" in issue.message and issue.severity.value == "loss" for issue in report.issues)


@pytest.mark.parametrize("kind", ["scatterChart", "bubbleChart", "pieChart"])
def test_extended_flags_and_removing_series_override(tmp_path: Path, kind: str) -> None:
    labels = {"show_value": True, "show_legend_key": False}
    if kind == "pieChart":
        labels.update(show_percent=True, show_leader_lines=True)
    elif kind == "bubbleChart":
        labels["show_bubble_size"] = True
    data = {
        "chart_type": kind,
        "categories": ["A", "B"],
        "data_labels": labels,
        "series": [{"values": [2, 3], "x_values": [1, 2], "bubble_sizes": [4, 5], "data_labels": {"hidden": True}}],
    }
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Chart")], properties={"pptx": {"chart": data}})])])
    model, _, _ = _cycle(model, tmp_path, 0)
    assert _chart(model)["data_labels"] == labels
    assert _chart(model)["series"][0]["data_labels"] == {"hidden": True}
    del _chart(model)["series"][0]["data_labels"]
    _chart(model)["data_labels"]["show_value"] = False
    model, target, _ = _cycle(model, tmp_path, 1)
    assert "data_labels" not in _chart(model)["series"][0]
    assert target.plots[0]._element.find(C + "dLbls/" + C + "showVal").get("val") == "0"
    assert target.series[0]._element.find(C + "dLbls") is None
