"""Merged PPTX cells project once onto logical Model, HTML and native grids."""

from zipfile import ZipFile

import pytest
from lxml import etree, html
from opendoc_model import DocumentLimits, Paragraph, TableCell, document_from_json, document_to_json
from pptx import Presentation
from pptx.util import Pt

from opendoc_formats import ExportOptions, PptxProfile, read_document, write_document
from opendoc_formats.ooxml.pptx_source_text import KEY, source_digest
from opendoc_formats.ooxml.table_grid import native_grid
from tests.writers.test_pptx_source_cells import CELL_KEY, A, own_table

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def merged_source(root):
    source = root/'logical.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(5, 4, Pt(20), Pt(30), Pt(500), Pt(300)).table
    for index, width in enumerate((80, 120, 140, 160)):
        table.columns[index].width = Pt(width)
    for first, last, text in (((0, 0), (1, 1), 'Rectangle'), ((0, 2), (0, 3), 'Horizontal'),
                              ((2, 0), (3, 3), 'Whole row pair')):
        table.cell(*first).merge(table.cell(*last))
        table.cell(*first).text = text
    for r, c, text in ((1, 2, 'After vertical'), (1, 3, 'Right'), (4, 0, 'Bottom left'), (4, 3, 'Bottom right')):
        table.cell(r, c).text = text
    edge = etree.SubElement(table.cell(1, 1)._tc.get_or_add_tcPr(), A+'lnB', w=str(Pt(3)))
    etree.SubElement(etree.SubElement(edge, A+'solidFill'), A+'srgbClr', val='123456')
    deck.save(source)
    return source


def test_logical_cells_survive_json_and_native_rebuild_twice(tmp_path):
    source = merged_source(tmp_path)
    original = source.read_bytes()
    model = read_document(source).document
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        table = model.sections[0].blocks[0]
        assert [len(row.cells) for row in table.rows] == [2, 2, 1, 0, 4]
        width, placements = native_grid(table)
        assert width == 4
        assert [(r, c) for r, c, _ in placements] == [(0, 0), (0, 2), (1, 2), (1, 3), (2, 0),
                                                       (4, 0), (4, 1), (4, 2), (4, 3)]
        assert table.properties.grid_widths_twips == [1600, 2400, 2800, 3200]
        before = document_to_json(model)
        output = tmp_path/f'rebuilt-{cycle}.pptx'
        report = write_document(model, output)
        assert report.success, report.issues
        native = Presentation(output).slides[0].shapes[0].table
        assert len(native.columns) == 4 and len(native.rows) == 5
        assert [col.width for col in native.columns] == [Pt(v) for v in (80, 120, 140, 160)]
        assert native.cell(0, 0).span_width == 2 and native.cell(0, 0).span_height == 2
        assert native.cell(0, 2).text == 'Horizontal' and native.cell(1, 2).text == 'After vertical'
        assert native.cell(2, 0).span_width == 4 and native.cell(2, 0).span_height == 2
        assert all(native.cell(3, c).is_spanned for c in range(4))
        edge = native.cell(1, 1)._tc.find(A+'tcPr/'+A+'lnB')
        assert edge.get('w') == str(Pt(3)) and edge.find('.//'+A+'srgbClr').get('val') == '123456'
        assert document_to_json(model) == before
        model = read_document(output).document
    assert source.read_bytes() == original


def test_logical_cells_export_to_html_and_docx(tmp_path):
    model = read_document(merged_source(tmp_path)).document
    output = tmp_path/'logical.html'
    report = write_document(model, output)
    assert report.success, report.issues
    table = html.fromstring(output.read_bytes()).xpath('//table')[0]
    assert [len(row.xpath('./td')) for row in table.xpath('.//tr')] == [2, 2, 1, 0, 4]
    origin = table.xpath('.//tr')[0].xpath('./td')[0]
    assert (origin.get('rowspan'), origin.get('colspan')) == ('2', '2')
    output = tmp_path/'logical.docx'
    report = write_document(model, output)
    assert report.success, report.issues
    with ZipFile(output) as archive:
        table = etree.fromstring(archive.read('word/document.xml')).find('.//'+W+'tbl')
    assert len(table.find(W+'tblGrid')) == 4
    assert [col.get(W+'w') for col in table.find(W+'tblGrid')] == ['1600', '2400', '2800', '3200']
    rows = table.findall(W+'tr')
    assert [len(row.findall(W+'tc')) for row in rows] == [2, 3, 1, 1, 4]
    assert rows[0].find(W+'tc/'+W+'tcPr/'+W+'gridSpan').get(W+'val') == '2'
    assert rows[1].find(W+'tc/'+W+'tcPr/'+W+'vMerge') is not None
    assert rows[3].find(W+'tc/'+W+'tcPr/'+W+'gridSpan').get(W+'val') == '4'


def test_source_cells_after_merges_keep_physical_binding(tmp_path):
    source = merged_source(tmp_path)
    model = document_from_json(document_to_json(read_document(source).document))
    table = model.sections[0].blocks[0]
    table.rows[0].cells[1].blocks[0].content[0].text = 'Edited horizontal'
    table.rows[1].cells[0].blocks[0].content[0].text = 'Edited after vertical'
    output = tmp_path/'edited.pptx'
    report = write_document(model, output, options=ExportOptions(pptx_profile=PptxProfile(layout='source')))
    assert report.success and report.metrics['cell_text_edits'] == 2, report.issues
    native = Presentation(output).slides[0].shapes[0].table
    assert native.cell(0, 2).text == 'Edited horizontal'
    assert native.cell(1, 2).text == 'Edited after vertical'
    with ZipFile(source) as old, ZipFile(output) as new:
        assert all(old.read(name) == new.read(name) for name in old.namelist() if name != 'ppt/slides/slide1.xml')


def test_legacy_physical_v3_snapshot_still_edits_and_rebuilds(tmp_path):
    model = read_document(own_table(tmp_path)).document
    table = model.sections[0].blocks[0]
    placed = {(r, c): cell for r, c, cell in native_grid(table)[1]}
    for ri, row in enumerate(table.rows):
        row.cells = [placed.get((ri, ci)) or TableCell([Paragraph([])], properties={CELL_KEY: {'row': ri, 'column': ci}})
                     for ci in range(3)]
    table.properties.pop('opendoc_formats.pptx_table_grid')
    table.properties.pop('opendoc_formats.pptx_continuations')
    model.metadata[KEY] = {'version': 3, 'digest': source_digest(model)}
    model = document_from_json(document_to_json(model))
    model.sections[0].blocks[0].rows[0].cells[2].blocks[0].content[0].text = 'Legacy edited'
    for layout in ('source', 'rebuild'):
        output = tmp_path/f'legacy-{layout}.pptx'
        report = write_document(model, output, options=ExportOptions(pptx_profile=PptxProfile(layout=layout)))
        assert report.success, report.issues
        native = Presentation(output).slides[0].shapes[0].table
        assert len(native.columns) == 3 and native.cell(0, 0).span_width == 2
        assert native.cell(0, 2).text.startswith('Legacy edited')


@pytest.mark.parametrize('damage', ['hidden-text', 'missing-continuation'])
def test_unproved_merge_keeps_physical_cells_and_reports_location(tmp_path, damage):
    source = merged_source(tmp_path)
    deck = Presentation(source)
    cell = deck.slides[0].shapes[0].table.cell(1, 1)
    if damage == 'hidden-text':
        cell.text = 'Own hidden continuation'
    else:
        cell._tc.attrib.pop('vMerge')
    deck.save(source)
    result = read_document(source)
    assert result.success, result.issues
    table = result.document.sections[0].blocks[0]
    assert [len(row.cells) for row in table.rows] == [4]*5
    assert table.properties['opendoc_formats.pptx_table_grid']['layout'] == 'physical-unverified'
    assert any(issue.feature == 'pptx.table-grid' and issue.location for issue in result.issues)
    if damage == 'hidden-text':
        assert table.rows[1].cells[1].blocks[0].content[0].text == 'Own hidden continuation'
    unchanged = tmp_path/'unchanged.pptx'
    report = write_document(result.document, unchanged, options=ExportOptions(pptx_profile=PptxProfile(layout='source')))
    assert report.success, report.issues
    with ZipFile(source) as old, ZipFile(unchanged) as new:
        assert all(old.read(name) == new.read(name) for name in old.namelist())
    table.rows[0].cells[0].blocks[0].content[0].text = 'Unproved edit'
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(result.document, output, options=ExportOptions(pptx_profile=PptxProfile(layout='source')))
    assert not report.success and output.read_bytes() == b'previous'


def test_stale_continuation_style_is_diagnosed_after_merge_edit(tmp_path):
    model = read_document(merged_source(tmp_path)).document
    table = model.sections[0].blocks[0]
    table.rows[0].cells[0].row_span = 1
    report = write_document(model, tmp_path/'changed.pptx')
    assert report.success, report.issues
    assert any(issue.feature == 'pptx.table-grid' for issue in report.issues)


def test_continuation_border_does_not_migrate_to_a_different_origin(tmp_path):
    model = read_document(merged_source(tmp_path)).document
    table = model.sections[0].blocks[0]
    # The former horizontal owner occupies the rectangle's old coordinate.
    table.rows[0].cells.reverse()
    output = tmp_path/'reordered.pptx'
    report = write_document(model, output)
    assert report.success, report.issues
    assert any(issue.feature == 'pptx.table-grid' for issue in report.issues)
    native = Presentation(output).slides[0].shapes[0].table
    edge = native.cell(1, 1)._tc.find(A+'tcPr/'+A+'lnB')
    assert edge is None or edge.find('.//'+A+'srgbClr') is None or edge.find('.//'+A+'srgbClr').get('val') != '123456'


def test_logical_native_grid_respects_selected_budget_before_publication(tmp_path):
    from opendoc_model import DocumentModel, Section, Table, TableRow

    model = DocumentModel(sections=[Section(blocks=[Table([TableRow([TableCell(
        [Paragraph([])], column_span=10000,
    )])])])])
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(document_limits=DocumentLimits(max_nodes=1000)))
    assert not report.success, report.issues
    assert output.read_bytes() == b'previous'
    assert any('grid' in issue.message.lower() and 'budget' in issue.message.lower() for issue in report.issues)
