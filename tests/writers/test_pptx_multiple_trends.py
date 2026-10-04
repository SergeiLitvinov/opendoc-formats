"""Несколько именованных трендов одного ряда остаются независимыми."""

from __future__ import annotations

import pytest

pytest.importorskip("pptx")
from opendoc.document_codec import document_from_json, document_to_json
from opendoc.document_model import DocumentModel, Paragraph, Section, TextRun
from pptx import Presentation

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.html_writer import write_html_model
from opendoc_formats.writers.pptx_writer import write_pptx_model

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _model(kind, trends):
    data = {
        "chart_type": kind,
        "categories": ["A", "B", "C", "D"],
        "series": [
            {
                "name": "Data",
                "values": [2, 4, 6, 9],
                "x_values": [1, 3, 6, 8],
                "bubble_sizes": [1, 2, 3, 4],
                "trendline": trends,
                "error_bars": {"value_type": "fixedVal", "value": 0.1},
            }
        ],
    }
    return DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Chart")], properties={"pptx": {"chart": data}})])])


@pytest.mark.parametrize("kind", ["lineChart", "barChart", "scatterChart", "bubbleChart"])
def test_multiple_trends_reorder_edit_and_delete(tmp_path, kind):
    trends = [
        {"type": "linear", "name": "Линия A & B", "color": "#118833", "show_equation": True},
        {"type": "poly", "order": 2, "name": "Аппроксимация", "color": "#882211", "show_r_squared": True},
    ]
    model = _model(kind, trends)
    for cycle in range(4):
        output = tmp_path / f"multi{cycle}.pptx"
        report = write_pptx_model(document_from_json(document_to_json(model)), output)
        assert report.success and not any("trendline не перенесены" in issue.message for issue in report.issues)
        native = Presentation(output).slides[0].shapes[0].chart.series[0]._element
        count = 2 if cycle < 2 else 3 - cycle
        assert len(native.findall(C + "trendline")) == count
        assert native.find(C + "errBars") is not None
        model = read_pptx_model(output)
        series = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]
        if cycle == 3:
            assert "trendline" not in series
            break
        result = series["trendline"]
        if cycle == 0:
            assert [item["name"] for item in result] == ["Линия A & B", "Аппроксимация"]
            assert result[0]["show_equation"] and result[1]["show_r_squared"]
            result[1]["order"], result[1]["name"] = 3, "Изменено"
            result.reverse()
        elif cycle == 1:
            assert [item["name"] for item in result] == ["Изменено", "Линия A & B"]
            assert [item["color"] for item in result] == ["#882211", "#118833"]
            assert result[0]["order"] == 3
            series["trendline"] = [result[0]]
        else:
            assert isinstance(result, dict) and result["name"] == "Изменено" and result["order"] == 3
            del series["trendline"]


@pytest.mark.parametrize(
    "trends",
    [[], [None], [{"type": "linear"}, {"type": "poly", "order": 9}], [{"type": "linear", "name": 3}], [{"type": "linear"}] * 65],
)
def test_invalid_trend_list_keeps_series_and_errors(tmp_path, trends):
    output = tmp_path / "invalid.pptx"
    report = write_pptx_model(_model("lineChart", trends), output)
    assert report.success and any("trendline не перенесены" in issue.message for issue in report.issues)
    native = Presentation(output).slides[0].shapes[0].chart.series[0]._element
    assert not native.findall(C + "trendline")
    assert native.find(C + "errBars") is not None


def test_html_renders_both_trends_of_same_type(tmp_path):
    output = tmp_path / "trends.html"
    model = _model("lineChart", [{"type": "linear", "color": "#118833"}, {"type": "linear", "color": "#882211"}])
    assert write_html_model(model, output).success
    html = output.read_text(encoding="utf-8")
    assert html.count('stroke-dasharray="6 4"') == 2
    assert 'stroke="#118833" stroke-width="2"' in html
    assert 'stroke="#882211" stroke-width="2"' in html
