"""Regressions isolated from the independent PowerPoint corpus acceptance."""

import pytest

pytest.importorskip("pptx")
from opendoc_model.document_codec import document_from_json, document_to_json
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.util import Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def test_shape_font_reference_and_explicit_run_color_survive_mutation(tmp_path):
    source = tmp_path / "source.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Pt(20), Pt(20), Pt(300), Pt(70))
    shape.text = "Inherited white"
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "Explicit red"
    run.font.color.rgb = RGBColor(255, 0, 0)
    deck.save(source)
    model = read_pptx_model(source)
    for cycle in range(2):
        runs = model.sections[0].blocks[0].content
        assert runs[0].style.color.to_hex().upper() == "#FFFFFF"
        assert runs[1].style.color.to_hex().upper() == ("#FF0000" if cycle == 0 else "#00FF00")
        runs[0].text = "Edited white text"
        if cycle == 0:
            from opendoc_model.color import ColorValue

            runs[1].style.color = ColorValue.from_hex("#00FF00")
        target = tmp_path / f"roundtrip{cycle}.pptx"
        assert write_pptx_model(document_from_json(document_to_json(model)), target).success
        exported = Presentation(target).slides[0].shapes[0]
        assert exported.auto_shape_type == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        assert str(exported.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
        model = read_pptx_model(target)


@pytest.mark.parametrize("textbox,zero_effect", [(True, False), (False, False), (False, True)])
def test_shape_without_effect_keeps_type_and_top_anchor_across_roundtrips(tmp_path, textbox, zero_effect):
    source = tmp_path / "source.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    geometry = (Pt(20), Pt(20), Pt(300), Pt(70))
    shape = slide.shapes.add_textbox(*geometry) if textbox else slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, *geometry)
    style = shape._element.find(P + "style")
    if style is not None:
        if zero_effect:
            style.find(A + "effectRef").set("idx", "0")
        else:
            shape._element.remove(style)
    shape.text_frame._txBody.bodyPr.attrib.pop("anchor", None)
    shape.text = "No added theme shadow or centered text"
    deck.save(source)
    model = read_pptx_model(source)
    for cycle in range(2):
        model.sections[0].blocks[0].content[0].text = f"Edited {cycle}"
        target = tmp_path / f"roundtrip{cycle}.pptx"
        assert write_pptx_model(document_from_json(document_to_json(model)), target).success
        exported = Presentation(target).slides[0].shapes[0]
        assert (exported.shape_type == MSO_SHAPE_TYPE.TEXT_BOX) is textbox
        style = exported._element.find(P + "style")
        if zero_effect:
            assert style.find(A + "effectRef").get("idx") == "0"
        else:
            assert style is None
        assert exported.text_frame._txBody.bodyPr.get("anchor") == "t"
        assert exported.text == f"Edited {cycle}"
        model = read_pptx_model(target)
