"""Наследование заполнителей и материализация стилей после изменения текста."""

from __future__ import annotations

import pytest

pytest.importorskip("pptx")
from lxml import etree
from opendoc.document_codec import document_from_json, document_to_json
from pptx import Presentation

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.readers.pptx_placeholder import A, P, placeholder_chain
from opendoc_formats.writers.pptx_writer import write_pptx_model


@pytest.mark.parametrize("level", [0, 7])
@pytest.mark.parametrize("master_geometry", [False, True])
def test_placeholder_three_level_style_mutation(tmp_path, level, master_geometry):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    layout = slide.slide_layout.placeholders[1]
    master = next(
        shape
        for shape in slide.slide_layout.slide_master.placeholders
        if shape._element.find(f"{P}nvSpPr/{P}nvPr/{P}ph").get("type") == "body"
    )
    target = slide.placeholders[1]
    for shape in (layout, target):
        shape._element.find(f"{P}nvSpPr/{P}nvPr/{P}ph").set("idx", "42")
    target._element.find(f"{P}nvSpPr/{P}nvPr/{P}ph").set("type", "body")
    if master_geometry:
        properties = layout._element.find(P + "spPr")
        transform = properties.find(A + "xfrm")
        if transform is not None:
            properties.remove(transform)
    for shape, size, attrs in ((master, "3300", {"lIns": "254000", "anchor": "b"}), (layout, None, {"rIns": "381000"})):
        body = shape._element.find(P + "txBody")
        props = body.find(A + "bodyPr")
        props.attrib.clear()
        props.attrib.update(attrs)
        for child in list(props):
            props.remove(child)
        etree.SubElement(props, A + ("normAutofit" if shape is master else "spAutoFit"))
        styles = body.find(A + "lstStyle")
        styles.clear()
        paragraph = etree.SubElement(styles, A + f"lvl{level + 1}pPr", algn="r" if shape is master else "ctr")
        rpr = etree.SubElement(paragraph, A + "defRPr")
        if size:
            rpr.set("sz", size)
            etree.SubElement(rpr, A + "latin", typeface="Arial")
            etree.SubElement(paragraph, A + "buChar", char="◆")
        else:
            rpr.set("b", "1")
            etree.SubElement(paragraph, A + "buNone")
        for old in body.findall(A + "p"):
            body.remove(old)
        etree.SubElement(body, A + "p")
        if shape is layout:
            local = etree.SubElement(body.find(A + "p"), A + "pPr", lvl=str(level))
            spacing = etree.SubElement(local, A + "spcAft")
            etree.SubElement(spacing, A + "spcPts", val="1200")
    target.text = "Inherited"
    body = target._element.find(P + "txBody")
    body.find(A + "bodyPr").attrib.clear()
    body.find(A + "bodyPr").set("lIns", "127000")
    paragraph = target.text_frame.paragraphs[0]
    paragraph.level = level
    paragraph.runs[0].font.italic = True
    source = tmp_path / "source.pptx"
    presentation.save(source)
    model = read_pptx_model(source)
    for cycle in range(3):
        block = next(block for block in model.sections[0].blocks if block.plain_text.startswith("Inherited"))
        run = block.content[0]
        assert run.style.font_size.pt == (33 if cycle < 2 else 29)
        assert run.style.font_family == "Arial"
        assert run.style.bold is True and run.style.italic is True
        meta = block.properties["pptx"]["paragraphs"][0]
        assert meta["alignment"] == "center" and meta["bullet_none"]
        assert meta["space_after_pt"] == 12
        assert "bullet_char" not in meta
        frame = block.properties["pptx"]["text_frame"]
        assert frame["lIns"] == 10 and frame["rIns"] == 30 and frame["anchor"] == "b"
        assert frame["autofit"]["mode"] == "spAutoFit"
        assert block.box.width > 0 and block.box.height > 0
        if cycle == 2:
            break
        run.text = "Inherited edited"
        if cycle == 1:
            from opendoc.document_model import Length

            run.style.font_size = Length(29)
        model = document_from_json(document_to_json(model))
        output = tmp_path / f"roundtrip{cycle}.pptx"
        assert write_pptx_model(model, output).success
        model = read_pptx_model(output)


def test_placeholder_ambiguous_idx_does_not_inherit():
    from copy import deepcopy

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    layout = slide.slide_layout._element.find(f"{P}cSld/{P}spTree")
    master = slide.slide_layout.slide_master._element.find(f"{P}cSld/{P}spTree")
    target = slide.placeholders[1]._element
    assert len(placeholder_chain(target, layout, master)) == 2
    layout.append(deepcopy(slide.slide_layout.placeholders[1]._element))
    assert placeholder_chain(target, layout, master) == []
