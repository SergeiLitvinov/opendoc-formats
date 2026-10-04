"""Геометрия наборов диаграммы сохраняется после повторных правок модели."""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("pptx")
from lxml import etree
from opendoc.document_codec import load_document, save_document
from opendoc.document_model import DocumentModel, Paragraph, Section, TextRun
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _data(model: DocumentModel) -> dict:
    return model.sections[0].blocks[0].properties["pptx"]["chart"]


def _model(data: dict) -> DocumentModel:
    return DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Chart")], properties={"pptx": {"chart": data}})])])


def _cycle(model: DocumentModel, tmp_path: Path, index: int) -> tuple:
    json_path = tmp_path / "model.json"
    save_document(model, json_path)
    output = tmp_path / f"result{index}.pptx"
    report = write_pptx_model(load_document(json_path), output)
    assert report.success
    return read_pptx_model(output), Presentation(output).slides[0].shapes[0].chart, report


@pytest.mark.parametrize(
    "kind", [XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_STACKED, XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT]
)
def test_import_and_mutate_plot_geometry(tmp_path: Path, kind: XL_CHART_TYPE) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("First", [3, 5])
    chart = slide.shapes.add_chart(kind, Pt(10), Pt(10), Pt(400), Pt(300), data).chart
    plot = chart.plots[0]
    bars = kind in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_STACKED)
    initial = {"gap_width": 37, "overlap": -25} if bars else {"first_slice_angle": 73}
    if kind == XL_CHART_TYPE.DOUGHNUT:
        initial["hole_size"] = 65
    tags = {"gap_width": "gapWidth", "overlap": "overlap", "first_slice_angle": "firstSliceAng", "hole_size": "holeSize"}
    if bars:
        plot.gap_width, plot.overlap = 37, -25
    else:
        for key, value in initial.items():
            old = plot._element.find(C + tags[key])
            if old is not None:
                plot._element.remove(old)
            plot._element.insert_element_before(etree.Element(C + tags[key], val=str(value)), "c:holeSize", "c:extLst")
    plot.vary_by_categories = True
    source = tmp_path / "source.pptx"
    presentation.save(source)
    model = read_pptx_model(source)
    for key, value in initial.items():
        assert _data(model)[key] == value
    for index in range(2):
        expected = (
            {"gap_width": 0 if index == 0 else 500, "overlap": -100 if index == 0 else 100}
            if bars
            else {
                "first_slice_angle": 0 if index == 0 else 360,
            }
        )
        if kind == XL_CHART_TYPE.DOUGHNUT:
            expected["hole_size"] = 10 if index == 0 else 90
        expected["vary_colors"] = bool(index)
        _data(model).update(expected)
        model, target, _ = _cycle(model, tmp_path, index)
        for key, value in expected.items():
            assert _data(model)[key] == value
        assert target.plots[0].vary_by_categories == bool(index)
        if bars:
            assert target.plots[0].gap_width == expected["gap_width"]
            assert target.plots[0].overlap == expected["overlap"]
        assert list(target.series[0].values) == [3, 5]


def test_combo_spacing_is_independent_and_conflicts_reported(tmp_path: Path) -> None:
    model = _model(
        {
            "chart_type": "barChart",
            "combo_types": ["barChart", "barChart"],
            "categories": ["A"],
            "gap_width": 30,
            "overlap": -10,
            "series": [
                {"name": "Primary", "values": [1], "chart_type": "barChart", "plot_index": 0},
                {
                    "name": "Secondary",
                    "values": [2],
                    "chart_type": "barChart",
                    "plot_index": 1,
                    "axis": "secondary_value",
                    "plot": {"gap_width": 80, "overlap": 20},
                },
                {
                    "name": "Third",
                    "values": [3],
                    "chart_type": "barChart",
                    "plot_index": 1,
                    "axis": "secondary_value",
                    "plot": {"gap_width": 90, "overlap": 20},
                },
            ],
        }
    )
    model, target, report = _cycle(model, tmp_path, 0)
    assert any("Противоречивое" in issue.message for issue in report.issues)
    assert [plot.gap_width for plot in target.plots] == [30, 80]
    for series in _data(model)["series"][1:]:
        assert series["plot"]["gap_width"] == 80
        series["plot"].update(gap_width=120, overlap=-40)
    model, target, report = _cycle(model, tmp_path, 1)
    assert not any("Противоречивое" in issue.message for issue in report.issues)
    assert [plot.gap_width for plot in target.plots] == [30, 120]
    assert [plot.overlap for plot in target.plots] == [-10, -40]


@pytest.mark.parametrize(
    "key,value",
    [
        ("gap_width", 501),
        ("overlap", -101),
        ("overlap", 0.5),
        ("gap_width", True),
        ("gap_width", "nan"),
        ("vary_colors", "false"),
        ("hole_size", 50),
    ],
)
def test_invalid_or_inapplicable_plot_settings_are_reported(tmp_path: Path, key: str, value: object) -> None:
    model = _model({"chart_type": "barChart", "categories": ["A"], "series": [{"values": [1]}], key: value})
    _, target, report = _cycle(model, tmp_path, 0)
    assert any(key in issue.message and issue.severity.value == "loss" for issue in report.issues)
    assert list(target.series[0].values) == [1]
