"""Оформление отдельных точек остаётся изменяемым после JSON/PPTX round-trip."""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("pptx")
from opendoc_model.document_codec import load_document, save_document
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model


@pytest.mark.parametrize("kind", [XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.COLUMN_CLUSTERED])
def test_mutate_point_fill_and_explosion(tmp_path: Path, kind: XL_CHART_TYPE) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B", "C"]
    data.add_series("Series", [1, None, 3])
    chart = slide.shapes.add_chart(kind, Pt(20), Pt(20), Pt(400), Pt(300), data).chart
    chart.series[0].points[2].format.fill.solid()
    chart.series[0].points[2].format.fill.fore_color.rgb = RGBColor.from_string("123456")
    source = tmp_path / "source.pptx"
    presentation.save(source)
    model = read_pptx_model(source)
    for index, color in enumerate(("#EE2211", "#118833")):
        series = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]
        points = series["data_points"]
        assert points["2"]["color"] == ("#123456" if index == 0 else "#EE2211")
        points["2"] = {"color": color}
        if kind != XL_CHART_TYPE.COLUMN_CLUSTERED:
            points["2"]["explosion"] = 25 + index
        model_path = tmp_path / "model.json"
        save_document(model, model_path)
        output = tmp_path / f"result{index}.pptx"
        report = write_pptx_model(load_document(model_path), output)
        assert report.success
        target = Presentation(output).slides[0].shapes[0].chart
        assert str(target.series[0].points[2].format.fill.fore_color.rgb) == color[1:]
        assert list(target.series[0].values) == [1, None, 3]
        model = read_pptx_model(output)
        points = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]["data_points"]
        assert set(points) == {"2"}
        assert points["2"]["color"] == color
        if kind != XL_CHART_TYPE.COLUMN_CLUSTERED:
            assert points["2"]["explosion"] == 25 + index


@pytest.mark.parametrize("points", [{"-1": {"color": "#FF0000"}}, {"3": {}}, {"1.5": {}}, {"0": None}, {"0": {"explosion": -5}}])
def test_invalid_point_settings_are_reported(tmp_path: Path, points: dict) -> None:
    from opendoc_model.document_model import DocumentModel, Paragraph, Section, TextRun

    data = {
        "chart_type": "pieChart",
        "categories": ["A"],
        "series": [{"name": "Series", "values": [1], "data_points": points}],
    }
    model = DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Chart")], properties={"pptx": {"chart": data}})])])
    output = tmp_path / "result.pptx"
    report = write_pptx_model(model, output)
    assert report.success
    assert any(
        issue.severity.value == "loss" and ("точки" in issue.message or "сектора" in issue.message) for issue in report.issues
    )
    assert Presentation(output).slides[0].shapes[0].has_chart
