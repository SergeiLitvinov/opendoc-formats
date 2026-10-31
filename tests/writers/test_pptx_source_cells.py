"""Native cell edits retain physical spans, caret formatting and unknown neighbors."""

from io import BytesIO
from threading import Event
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from lxml import etree
from opendoc_model import TextRun, document_from_json, document_to_json
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from opendoc_formats import ExportOptions, PptxProfile, read_document, write_document
from opendoc_formats.ooxml.pptx_source_cells import CELL_KEY
from opendoc_formats.ooxml.pptx_source_text import KEY, source_digest
from opendoc_formats.ooxml.table_grid import native_grid

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
OPTIONS = ExportOptions(pptx_profile=PptxProfile(layout='source'))


def own_table(root):
    source = root/'table.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(3, 3, Pt(20), Pt(40), Pt(500), Pt(360)).table
    table.cell(0, 0).merge(table.cell(1, 1))
    table.cell(0, 0).text = 'Merged origin'
    styled = table.cell(0, 2)
    styled.margin_left, styled.margin_top = Pt(11), Pt(9)
    styled.fill.solid()
    styled.fill.fore_color.rgb = RGBColor(220, 230, 240)
    paragraph = styled.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    first = paragraph.add_run()
    first.text, first.font.bold, first.font.size = 'Bold', True, Pt(19)
    link = paragraph.add_run()
    link.text = ' Link'
    link.hyperlink.address = 'https://example.invalid/inert'
    paragraph.add_line_break()
    paragraph.add_run().text = 'After break'
    styled.text_frame.add_paragraph()
    styled.text_frame.add_paragraph().text = 'Last paragraph'
    table.cell(2, 0).text = 'Neighbor cell'
    empty = table.cell(2, 2).text_frame.paragraphs[0]._p
    end = etree.SubElement(empty, A+'endParaRPr', sz='2300', b='1', lang='en-US')
    etree.SubElement(etree.SubElement(end, A+'solidFill'), A+'srgbClr', val='225588')
    slide.shapes.add_textbox(Pt(20), Pt(10), Pt(200), Pt(20)).text = 'Neighbor shape'
    deck.save(source)
    buffer = BytesIO()
    with ZipFile(source) as old, ZipFile(buffer, 'w', ZIP_DEFLATED) as new:
        for item in old.infolist():
            new.writestr(item, old.read(item))
        new.writestr('own/opaque.bin', b'own inert table neighbor\x00\xff')
    source.write_bytes(buffer.getvalue())
    return source


def model_at(source):
    result = read_document(source)
    assert result.success, result.issues
    return result.document


def cell(model, row, column):
    table = model.sections[0].blocks[0]
    if table.properties.get('opendoc_formats.pptx_table_grid', {}).get('layout') == 'logical-origins':
        return next(cell for r, c, cell in native_grid(table)[1] if (r, c) == (row, column))
    return table.rows[row].cells[column]


def edit_cells(model, cycle):
    cell(model, 0, 0).blocks[0].content[0].text = f'Merged edit {cycle}'
    paragraph = cell(model, 0, 2).blocks[0]
    for index, text in ((0, 'Bold edit'), (1, ' Link edit'), (3, 'After break edit')):
        paragraph.content[index].text = text
    cell(model, 0, 2).blocks[2].content[0].text = 'Last paragraph edit'
    paragraph = cell(model, 2, 2).blocks[0]
    if paragraph.content:
        paragraph.content[0].text = f'New cell {cycle}'
    else:
        paragraph.content.append(TextRun(f'New cell {cycle}'))


def test_cell_edits_keep_native_table_structure_and_two_json_cycles(tmp_path):
    source = own_table(tmp_path)
    original = source.read_bytes()
    model = model_at(source)
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        edit_cells(model, cycle)
        model.sections[0].blocks[0].box.x += 10
        before = document_to_json(model)
        output = tmp_path/f'cell-{cycle}.pptx'
        report = write_document(model, output, options=OPTIONS)
        assert report.success, report.issues
        assert report.metrics['cell_text_edits'] >= 2
        assert report.metrics['geometry_edits'] == 1
        with ZipFile(source) as old, ZipFile(output) as new:
            assert set(old.namelist()) == set(new.namelist())
            for name in old.namelist():
                if name != 'ppt/slides/slide1.xml':
                    assert new.read(name) == old.read(name)
            old_table = etree.fromstring(old.read('ppt/slides/slide1.xml')).find('.//'+A+'tbl')
            new_table = etree.fromstring(new.read('ppt/slides/slide1.xml')).find('.//'+A+'tbl')
            assert etree.tostring(old_table.find(A+'tblGrid')) == etree.tostring(new_table.find(A+'tblGrid'))
            for old_row, new_row in zip(old_table.findall(A+'tr'), new_table.findall(A+'tr'), strict=True):
                assert old_row.attrib == new_row.attrib
                for old_cell, new_cell in zip(old_row.findall(A+'tc'), new_row.findall(A+'tc'), strict=True):
                    assert old_cell.attrib == new_cell.attrib
                    assert etree.tostring(old_cell.find(A+'tcPr')) == etree.tostring(new_cell.find(A+'tcPr'))
        native = Presentation(output).slides[0].shapes[0].table
        assert native.cell(0, 0).is_merge_origin and native.cell(1, 1).is_spanned
        styled = native.cell(0, 2).text_frame.paragraphs[0]
        assert styled.runs[0].font.bold and styled.runs[0].font.size.pt == 19
        assert styled.runs[1].hyperlink.address == 'https://example.invalid/inert'
        inserted = native.cell(2, 2).text_frame.paragraphs[0].runs[0]
        assert inserted.font.size.pt == 23 and inserted.font.bold
        assert str(inserted.font.color.rgb) == '225588'
        assert native.cell(2, 0).text == 'Neighbor cell'
        assert document_to_json(model) == before
        model = model_at(output)
        assert cell(model, 2, 2).blocks[0].content[0].text == f'New cell {cycle}'
    assert source.read_bytes() == original


@pytest.mark.parametrize('change', ['continuation', 'row-order', 'cell-order', 'span', 'fill', 'new-style', 'break', 'nul'])
def test_unsupported_cell_edit_refuses_atomically(tmp_path, change):
    model = model_at(own_table(tmp_path))
    if change == 'continuation':
        model.sections[0].blocks[0].properties['opendoc_formats.pptx_continuations'][0]['properties']['fill'] = '#123456'
    elif change == 'row-order':
        model.sections[0].blocks[0].rows.reverse()
    elif change == 'cell-order':
        model.sections[0].blocks[0].rows[2].cells.reverse()
    elif change == 'span':
        cell(model, 0, 0).column_span = 1
    elif change == 'fill':
        cell(model, 0, 2).properties['fill'] = '#112233'
    elif change == 'new-style':
        cell(model, 2, 2).blocks[0].content.append(TextRun('Styled insertion'))
        cell(model, 2, 2).blocks[0].content[0].style.bold = True
    else:
        cell(model, 0, 0).blocks[0].content[0].text = 'New'+ ('\n' if change == 'break' else '\x00')+'text'
    before = document_to_json(model)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=OPTIONS)
    assert not report.success, report.issues
    assert output.read_bytes() == b'previous' and document_to_json(model) == before


@pytest.mark.parametrize('version', [1, 2])
def test_old_binding_versions_do_not_enable_cell_editing(tmp_path, version):
    model = model_at(own_table(tmp_path))
    for row in model.sections[0].blocks[0].rows:
        for item in row.cells:
            item.properties.pop(CELL_KEY, None)
    model.metadata[KEY] = {'version': version, 'digest': source_digest(model, geometry=version >= 2, tables=False)}
    model = document_from_json(document_to_json(model))
    unchanged = write_document(model, tmp_path/f'old-{version}.pptx', options=OPTIONS)
    assert unchanged.success, unchanged.issues
    cell(model, 0, 0).blocks[0].content[0].text = 'Unsupported old cell edit'
    assert not write_document(model, tmp_path/'refused.pptx', options=OPTIONS).success


def test_cell_insertion_cancelled_before_publication(tmp_path, monkeypatch):
    from opendoc_formats.writers import pptx_source_writer
    model = model_at(own_table(tmp_path))
    cell(model, 2, 2).blocks[0].content.append(TextRun('New cell'))
    cell(model, 0, 0).blocks[0].content[0].text = 'First cell changed'
    before = document_to_json(model)
    stop = Event()
    original = pptx_source_writer.apply_cell_text
    def apply(binding):
        assert not stop.is_set(), 'Cell work continued after cancellation'
        edited = original(binding)
        if edited:
            stop.set()
        return edited
    monkeypatch.setattr(pptx_source_writer, 'apply_cell_text', apply)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(
        pptx_profile=PptxProfile(layout='source'), cancelled=stop.is_set,
    ))
    assert not report.success and any(issue.feature == 'export.cancelled' for issue in report.issues)
    assert output.read_bytes() == b'previous' and document_to_json(model) == before


def test_untouched_native_control_text_is_not_rejected(tmp_path):
    source = own_table(tmp_path)
    native = Presentation(source)
    native.slides[0].shapes[0].table.cell(2, 0).text_frame.paragraphs[0].runs[0]._r.find(A+'t').text = 'Own\ttab'
    native.save(source)
    model = model_at(source)
    cell(model, 0, 0).blocks[0].content[0].text = 'Other cell edit'
    report = write_document(model, tmp_path/'tab.pptx', options=OPTIONS)
    assert report.success, report.issues


@pytest.mark.parametrize('damage', ['missing-continuation', 'unknown-flag', 'overflow-span'])
def test_unproved_merge_topology_does_not_get_edit_bindings(tmp_path, damage):
    source = own_table(tmp_path)
    native = Presentation(source)
    table = native.slides[0].shapes[0].table
    if damage == 'missing-continuation':
        table.cell(1, 1)._tc.attrib.pop('vMerge')
    elif damage == 'unknown-flag':
        table.cell(1, 1)._tc.set('hMerge', 'unknown')
    else:
        table.cell(0, 0)._tc.set('gridSpan', '4')
    native.save(source)
    model = model_at(source)
    cell(model, 0, 0).blocks[0].content[0].text = 'Cannot bind'
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=OPTIONS)
    assert not report.success and output.read_bytes() == b'previous'
