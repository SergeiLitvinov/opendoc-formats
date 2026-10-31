"""Native chart resizing changes only its slide frame, retaining chart/workbook bytes."""

from copy import deepcopy
from threading import Event
from zipfile import ZipFile

import pytest
from lxml import etree
from opendoc_model import DocumentLimits, document_from_json, document_to_json
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt

from opendoc_formats import ExportOptions, PptxProfile, read_document, write_document
from opendoc_formats.ooxml.pptx_source_chart_frame import C, chart_frame
from opendoc_formats.ooxml.pptx_source_geometry import geometry_bindings
from opendoc_formats.ooxml.pptx_source_text import KEY, source_digest
from opendoc_formats.ooxml.pptx_source_xml import SourceRoots

OPTIONS = ExportOptions(pptx_profile=PptxProfile(layout='source'))


def own_chart(root, kind='column'):
    source = root/'source.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = XyChartData() if kind == 'scatter' else CategoryChartData()
    if kind == 'scatter':
        series = data.add_series('Own numeric series')
        for x, y in ((1, 2), (4, 7), (8, 3)):
            series.add_data_point(x, y)
    else:
        data.categories = ['Long category Alpha', 'Beta', 'Gamma']
        data.add_series('Own first', [1, 4, 3])
        if kind != 'pie':
            data.add_series('Own second', [2, None, 5])
    types = {'scatter': XL_CHART_TYPE.XY_SCATTER, 'pie': XL_CHART_TYPE.PIE}
    chart = slide.shapes.add_chart(types.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED),
                                   Pt(30), Pt(35), Pt(450), Pt(280), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = 'Own native chart'
    chart.chart_style = 13
    if kind == '3d':
        space = chart._chartSpace.find(C+'chart')
        view = etree.Element(C+'view3D')
        for name, value in (('rotX', '15'), ('rotY', '20'), ('depthPercent', '100')):
            etree.SubElement(view, C+name, val=value)
        space.insert(list(space).index(space.find(C+'plotArea')), view)
        chart._chartSpace.find('.//'+C+'barChart').tag = C+'bar3DChart'
    if kind == 'manual':
        area = chart._chartSpace.find('.//'+C+'plotArea')
        layout = area.find(C+'layout')
        if layout is None:
            layout = etree.Element(C+'layout')
            area.insert(0, layout)
        manual = etree.SubElement(layout, C+'manualLayout')
        etree.SubElement(manual, C+'layoutTarget', val='inner')
        for name in ('xMode', 'yMode', 'wMode', 'hMode'):
            etree.SubElement(manual, C+name, val='factor')
        for name, value in (('x', '0.15'), ('y', '0.2'), ('w', '0.7'), ('h', '0.65')):
            etree.SubElement(manual, C+name, val=value)
    slide.shapes.add_textbox(Pt(20), Pt(400), Pt(300), Pt(30)).text = 'Own untouched neighbor'
    deck.save(source)
    return source


@pytest.mark.parametrize('kind', ['column', 'scatter', 'pie', '3d', 'manual'])
def test_chart_frame_edit_preserves_chart_workbook_and_two_json_cycles(tmp_path, kind):
    source = own_chart(tmp_path, kind)
    original = source.read_bytes()
    native = Presentation(source)
    shape = native.slides[0].shapes[0]
    shape.left += Pt(10)
    shape.top += Pt(5)
    shape.width += Pt(120)
    shape.height += Pt(60)
    reference = tmp_path/'reference.pptx'
    native.save(reference)
    model = read_document(source).document
    wanted = read_document(reference).document.sections[0].blocks[0].box
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        model.sections[0].blocks[0].box = deepcopy(wanted)
        before = document_to_json(model)
        output = tmp_path/f'chart-{cycle}.pptx'
        report = write_document(model, output, options=OPTIONS)
        assert report.success, report.issues
        if cycle == 0:
            assert report.metrics['chart_frame_edits'] == 1
        with ZipFile(source) as old, ZipFile(output) as new, ZipFile(reference) as expected:
            assert set(old.namelist()) == set(new.namelist())
            for name in old.namelist():
                if name != 'ppt/slides/slide1.xml':
                    assert new.read(name) == old.read(name)
            assert any(name.startswith('ppt/embeddings/') for name in old.namelist())
            actual = etree.fromstring(new.read('ppt/slides/slide1.xml'))
            wanted_xml = etree.fromstring(expected.read('ppt/slides/slide1.xml'))
            assert etree.tostring(actual, method='c14n') == etree.tostring(wanted_xml, method='c14n')
        assert document_to_json(model) == before
        model = read_document(output).document
        assert model.sections[0].blocks[0].box == wanted
    assert source.read_bytes() == original


@pytest.mark.parametrize('version', [2, 3, 4])
def test_old_versions_allow_chart_translation_but_keep_resize_protected(tmp_path, version):
    model = read_document(own_chart(tmp_path)).document
    model.metadata[KEY] = {'version': version, 'digest': source_digest(
        model, tables=version >= 3, table_geometry=version >= 4,
    )}
    model.sections[0].blocks[0].box.x += 10
    report = write_document(model, tmp_path/'translated.pptx', options=OPTIONS)
    assert report.success, report.issues
    model.sections[0].blocks[0].box.width += 10
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    assert not write_document(model, output, options=OPTIONS).success
    assert output.read_bytes() == b'previous'


@pytest.mark.parametrize('damage', ['rotation', 'data', 'title'])
def test_unbound_chart_edit_refuses_atomically(tmp_path, damage):
    model = read_document(own_chart(tmp_path)).document
    block = model.sections[0].blocks[0]
    block.box.width += 10
    if damage == 'rotation':
        block.box.rotation += 20
    elif damage == 'data':
        block.properties['pptx']['chart']['series'][0]['values'][0] = '99'
    else:
        block.properties['pptx']['chart']['title'] = 'New title'
    before = document_to_json(model)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=OPTIONS)
    assert not report.success and output.read_bytes() == b'previous'
    assert document_to_json(model) == before


@pytest.mark.parametrize('damage', ['external', 'wrong-type', 'wrong-media', 'wrong-root', 'unknown-uri'])
def test_unproved_native_chart_frame_does_not_enable_resize(tmp_path, damage):
    model = read_document(own_chart(tmp_path)).document
    native = geometry_bindings(model, SourceRoots())[0]
    edge = next(e for e in model.package.relationships if e.source == native.part and e.relationship_type.endswith('/chart'))
    if damage == 'external':
        edge.external = True
    elif damage == 'wrong-type':
        edge.relationship_type = 'urn:own:opaque'
    elif damage == 'wrong-media':
        model.package.parts[edge.target].media_type = 'application/octet-stream'
    elif damage == 'wrong-root':
        model.package.parts[edge.target].data = b'<own/>'
    else:
        native.shape.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}graphicData').set('uri', 'urn:own:opaque')
    assert not chart_frame(native, model, SourceRoots())


def test_chart_xml_projection_obeys_selected_budget(tmp_path):
    model = read_document(own_chart(tmp_path)).document
    chart = next(part for name, part in model.package.parts.items()
                 if name.startswith('/ppt/charts/chart') and name.endswith('.xml'))
    root = etree.fromstring(chart.data)
    for _ in range(12000):
        root.append(etree.Comment('own'))
    chart.data = etree.tostring(root)
    model.metadata[KEY]['digest'] = source_digest(model, table_geometry=True)
    model.sections[0].blocks[0].box.width += 20
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(
        pptx_profile=PptxProfile(layout='source'), document_limits=DocumentLimits(max_nodes=10000),
    ))
    assert not report.success and any(issue.feature == 'pptx.source-xml-limit' for issue in report.issues)
    assert output.read_bytes() == b'previous'


def test_cancel_after_chart_proof_stops_before_geometry_mutation(tmp_path, monkeypatch):
    from opendoc_formats.writers import pptx_source_writer

    model = read_document(own_chart(tmp_path)).document
    model.sections[0].blocks[0].box.width += 20
    before = document_to_json(model)
    stop = Event()
    original = pptx_source_writer.chart_frame
    def prove(*args):
        result = original(*args)
        stop.set()
        return result
    def refused(*args, **kwargs):
        pytest.fail('Geometry work continued after cancellation')
    monkeypatch.setattr(pptx_source_writer, 'chart_frame', prove)
    monkeypatch.setattr(pptx_source_writer, 'project_geometry', refused)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(
        pptx_profile=PptxProfile(layout='source'), cancelled=stop.is_set,
    ))
    assert not report.success and any(issue.feature == 'export.cancelled' for issue in report.issues)
    assert output.read_bytes() == b'previous' and document_to_json(model) == before
