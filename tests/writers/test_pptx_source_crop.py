"""Native image crop changes retain original raster/vector payloads and effects."""

from io import BytesIO
from threading import Event
from zipfile import ZipFile

import pytest
from lxml import etree
from opendoc_model import (
    Box,
    DocumentModel,
    Image,
    ImageCrop,
    Resource,
    ResourceKind,
    Section,
    document_from_json,
    document_to_json,
)
from pptx import Presentation
from pptx.util import Pt

from opendoc_formats import ExportOptions, PptxProfile, read_document, write_document
from opendoc_formats.ooxml.pptx_crop import SIDES, A, P
from opendoc_formats.ooxml.pptx_source_text import KEY, source_digest
from tests.writers.test_pdf_docx_rasters import own_pixels

OPTIONS = ExportOptions(pptx_profile=PptxProfile(layout='source'))
SVG = (b'<svg xmlns="http://www.w3.org/2000/svg" width="120" height="180">'
       b'<rect x="3" y="3" width="114" height="174" fill="none" stroke="red" stroke-width="6"/>'
       b'<rect x="12" y="36" width="54" height="65" fill="blue"/></svg>')


def picture(deck):
    shape = deck.slides[0].shapes[0]
    return shape.shapes[0] if shape.shape_type == 6 else shape


def own_picture(root, *, svg=False, cropped=False, grouped=False):
    source = root/'source.pptx'
    if svg:
        document = DocumentModel(sections=[Section(blocks=[Image(
            'svg', box=Box(90, 90, 180, 240), properties={'fallback_resource_id': 'png'},
        )])], resources={
            'svg': Resource('svg', ResourceKind.VECTOR_IMAGE, 'image/svg+xml', data=SVG),
            'png': Resource('png', ResourceKind.RASTER_IMAGE, 'image/png', data=own_pixels()),
        })
        report = write_document(document, source)
        assert report.success, report.issues
        deck = Presentation(source)
    else:
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        owner = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
        owner.add_picture(BytesIO(own_pixels()), Pt(90), Pt(90), Pt(180), Pt(240))
    shape = picture(deck)
    if cropped:
        for side, value in zip(SIDES, (0.1, 0.2, 0.1, 0.1), strict=True):
            setattr(shape, 'crop_'+side, value)
    blip = shape._element.find(P+'blipFill/'+A+'blip')
    etree.SubElement(blip, A+'alphaModFix', amt='65000')
    deck.slides[0].shapes.add_textbox(Pt(350), Pt(40), Pt(250), Pt(40)).text = 'Own crop neighbor'
    deck.save(source)
    return source


@pytest.mark.parametrize('svg,cropped,grouped', [(False, False, False), (False, True, False),
                                                (False, True, True), (True, True, False)])
@pytest.mark.parametrize('operation', ['set', 'clear'])
def test_crop_and_geometry_keep_media_effects_and_two_json_cycles(tmp_path, svg, cropped, grouped, operation):
    source = own_picture(tmp_path, svg=svg, cropped=cropped, grouped=grouped)
    original = source.read_bytes()
    crop = ImageCrop(0.12345, 0.23456, 0.11111, 0.12345) if operation == 'set' else None
    native = Presentation(source)
    target = picture(native)
    if crop is None:
        fill = target._element.find(P+'blipFill')
        rectangle = fill.find(A+'srcRect')
        if rectangle is not None:
            fill.remove(rectangle)
    else:
        for side in SIDES:
            setattr(target, 'crop_'+side, getattr(crop, side))
    target.left += Pt(10)
    reference = tmp_path/'reference.pptx'
    native.save(reference)
    wanted = read_document(reference).document.sections[0].blocks[0]
    model = read_document(source).document
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        model.sections[0].blocks[0].crop = crop
        model.sections[0].blocks[0].box = wanted.box
        before = document_to_json(model)
        output = tmp_path/f'cropped-{cycle}.pptx'
        report = write_document(model, output, options=OPTIONS)
        assert report.success, report.issues
        with ZipFile(source) as old, ZipFile(output) as new, ZipFile(reference) as expected:
            assert set(old.namelist()) == set(new.namelist())
            for name in old.namelist():
                if name != 'ppt/slides/slide1.xml':
                    assert new.read(name) == old.read(name)
            assert any(name.startswith('ppt/media/') for name in old.namelist())
            actual = etree.fromstring(new.read('ppt/slides/slide1.xml'))
            goal = etree.fromstring(expected.read('ppt/slides/slide1.xml'))
            assert etree.tostring(actual, method='c14n') == etree.tostring(goal, method='c14n')
        assert document_to_json(model) == before
        model = read_document(output).document
        assert model.sections[0].blocks[0].crop == crop
    assert source.read_bytes() == original


@pytest.mark.parametrize('version', [2, 3, 4, 5])
def test_old_bindings_keep_geometry_but_do_not_enable_crop(tmp_path, version):
    model = read_document(own_picture(tmp_path, cropped=True)).document
    model.metadata[KEY] = {'version': version, 'digest': source_digest(
        model, tables=version >= 3, table_geometry=version >= 4,
    )}
    model.sections[0].blocks[0].box.x += 10
    report = write_document(model, tmp_path/'moved.pptx', options=OPTIONS)
    assert report.success, report.issues
    model.sections[0].blocks[0].crop.left += 0.1
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    assert not write_document(model, output, options=OPTIONS).success
    assert output.read_bytes() == b'previous'


@pytest.mark.parametrize('damage', ['tile', 'destination-rectangle', 'unknown-attribute', 'external-link'])
def test_unproved_picture_fill_keeps_unchanged_payload_but_refuses_crop(tmp_path, damage):
    source = own_picture(tmp_path, cropped=True)
    deck = Presentation(source)
    fill = picture(deck)._element.find(P+'blipFill')
    if damage == 'tile':
        fill.remove(fill.find(A+'stretch'))
        etree.SubElement(fill, A+'tile', tx='100', ty='200')
    elif damage == 'destination-rectangle':
        fill.find(A+'stretch/'+A+'fillRect').set('l', '1000')
    elif damage == 'unknown-attribute':
        fill.find(A+'srcRect').set('{urn:own}opaque', 'own')
    else:
        fill.find(A+'blip').set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link', 'own-unproved')
    deck.save(source)
    model = read_document(source).document
    output = tmp_path/'unchanged.pptx'
    report = write_document(model, output, options=OPTIONS)
    assert report.success, report.issues
    with ZipFile(source) as old, ZipFile(output) as new:
        assert all(new.read(name) == old.read(name) for name in old.namelist())
    model.sections[0].blocks[0].crop.left += 0.1
    before = document_to_json(model)
    output.write_bytes(b'previous')
    assert not write_document(model, output, options=OPTIONS).success
    assert output.read_bytes() == b'previous' and document_to_json(model) == before


def test_percent_lexemes_and_explicit_zero_fill_rect_import_without_rewriting(tmp_path):
    source = own_picture(tmp_path, cropped=True)
    deck = Presentation(source)
    fill = picture(deck)._element.find(P+'blipFill')
    rectangle = fill.find(A+'srcRect')
    rectangle.attrib.update({'l': '12.345%', 't': '20%', 'r': '.1%', 'b': '10.0%'})
    fill.find(A+'stretch/'+A+'fillRect').attrib.update({'l': '0%', 't': '0', 'r': '0', 'b': '0'})
    deck.save(source)
    result = read_document(source)
    assert result.success, result.issues
    image = result.document.sections[0].blocks[0]
    assert image.crop == ImageCrop(0.12345, 0.2, 0.001, 0.1)
    image.box.x += 10
    output = tmp_path/'moved.pptx'
    report = write_document(result.document, output, options=OPTIONS)
    assert report.success, report.issues
    with ZipFile(output) as new:
        crop = etree.fromstring(new.read('ppt/slides/slide1.xml')).find('.//'+A+'srcRect')
        assert crop.attrib == rectangle.attrib


def test_rounding_cannot_remove_entire_image(tmp_path):
    model = read_document(own_picture(tmp_path)).document
    model.sections[0].blocks[0].crop = ImageCrop(0.499999, 0, 0.499999, 0)
    assert not model.validate()
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    before = document_to_json(model)
    report = write_document(model, output, options=OPTIONS)
    assert not report.success and any('Rounded native crop' in issue.message for issue in report.issues)
    assert output.read_bytes() == b'previous' and document_to_json(model) == before


def test_native_signed_outset_is_opaque_without_crashing_public_import(tmp_path):
    source = own_picture(tmp_path, cropped=True)
    deck = Presentation(source)
    rectangle = picture(deck)._element.find(P+'blipFill/'+A+'srcRect')
    rectangle.attrib.update({'l': '-13333', 't': '26667', 'r': '13333', 'b': '13333'})
    deck.save(source)
    result = read_document(source)
    assert result.success and not result.lossless, result.issues
    model = document_from_json(document_to_json(result.document))
    image = model.sections[0].blocks[0]
    assert image.crop is None
    assert image.properties['opendoc_formats.pptx_source_rectangle']['attributes'] == dict(rectangle.attrib)
    assert any(issue.feature == 'pptx.image-crop' and issue.location for issue in result.issues)
    output = tmp_path/'outset.pptx'
    report = write_document(model, output, options=OPTIONS)
    assert report.success, report.issues
    with ZipFile(source) as old, ZipFile(output) as new:
        assert all(old.read(name) == new.read(name) for name in old.namelist())
    image.crop = ImageCrop(0.1, 0.2, 0.1, 0.1)
    output.write_bytes(b'previous')
    assert not write_document(model, output, options=OPTIONS).success
    assert output.read_bytes() == b'previous'


def test_crop_cancellation_keeps_model_and_destination(tmp_path, monkeypatch):
    from opendoc_formats.writers import pptx_source_writer

    model = read_document(own_picture(tmp_path)).document
    model.sections[0].blocks[0].crop = ImageCrop(0.1, 0.2, 0.1, 0.1)
    before = document_to_json(model)
    stop = Event()
    original = pptx_source_writer.project_image_crop
    def project(binding):
        changed = original(binding)
        stop.set()
        return changed
    monkeypatch.setattr(pptx_source_writer, 'project_image_crop', project)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(
        pptx_profile=PptxProfile(layout='source'), cancelled=stop.is_set,
    ))
    assert not report.success and any(issue.feature == 'export.cancelled' for issue in report.issues)
    assert output.read_bytes() == b'previous' and document_to_json(model) == before
