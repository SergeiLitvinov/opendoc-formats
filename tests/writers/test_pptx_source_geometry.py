"""Native source Box projection against independently edited DrawingML frames."""

from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pytest
from lxml import etree
from opendoc_model import Box, DocumentLimits, document_from_json, document_to_json
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR_TYPE
from pptx.util import Inches

from opendoc_formats import ExportOptions, PptxProfile, read_document, write_document
from opendoc_formats.ooxml.pptx_source_text import KEY, source_digest
from tests.writers.test_pdf_docx_rasters import own_pixels

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
OPTIONS = ExportOptions(pptx_profile=PptxProfile(layout='source'))


def fixture(root, kind, angle=0, flip=False, nested=False):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape() if kind != 'table' else None
    owner = group.shapes.add_group_shape() if nested else group
    if kind == 'picture':
        shape = owner.shapes.add_picture(BytesIO(own_pixels()), Inches(1), Inches(1), Inches(2), Inches(1))
    elif kind == 'table':
        shape = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(2), Inches(1))
        shape.table.cell(0, 0).text = 'Native cell'
    elif kind == 'connector':
        shape = owner.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT, Inches(1), Inches(1), Inches(3), Inches(1))
    else:
        shape = owner.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        shape.text = 'Native grouped text'
        shape.rotation = 13
    slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = 'Native neighbor'
    groups = ([group, owner] if nested else [group]) if group is not None else []
    for index, current in enumerate(groups):
        transform = current._element.find(P+'grpSpPr/'+A+'xfrm')
        for tag, attrs in (
            ('off', {'x': 1270000+index*381000, 'y': 762000}),
            ('ext', {'cx': 3810000, 'cy': 2540000}),
            ('chOff', {'x': 254000, 'y': 381000}),
            ('chExt', {'cx': 2540000, 'cy': 1270000}),
        ):
            transform.find(A+tag).attrib.update({key: str(value) for key, value in attrs.items()})
        transform.set('rot', str(round((angle if index == 0 else 7)*60000)))
        transform.set('flipH', '1' if flip else '0')
    source = root/'source.pptx'
    deck.save(source)
    return source, shape.shape_id


def block_with(model, identity):
    return next(block for block in model.sections[0].blocks if block.properties['pptx']['object_id'] == str(identity))


@pytest.mark.parametrize('kind,angle,flip,nested', [
    (kind, angle, flip, nested) for kind in ('textbox', 'picture', 'connector')
    for angle, flip, nested in ((0, False, False), (25, False, False), (90, True, False), (33, True, True))
]+[('table', 0, False, False)])
def test_native_box_edit_inverts_group_frame_and_preserves_neighbors(tmp_path, kind, angle, flip, nested):
    source, identity = fixture(tmp_path, kind, angle, flip, nested)
    original = source.read_bytes()
    expected = Presentation(source)
    shape = expected.slides[0]._element.xpath(f'.//p:cNvPr[@id="{identity}"]')[0].getparent().getparent()
    frame = shape.find(P+'xfrm') if kind == 'table' else shape.find(P+'spPr/'+A+'xfrm')
    offset = frame.find(A+'off')
    offset.set('x', str(int(offset.get('x'))+31750))
    offset.set('y', str(int(offset.get('y'))+63500))
    if kind != 'table':
        extent = frame.find(A+'ext')
        extent.set('cx', str(int(extent.get('cx'))+127000))
        extent.set('cy', str(int(extent.get('cy'))+254000))
        frame.set('rot', str(int(frame.get('rot', '0'))+17*60000))
    reference = tmp_path/'native-reference.pptx'
    expected.save(reference)
    wanted = read_document(reference).document
    model = read_document(source).document
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        block_with(model, identity).box = deepcopy(block_with(wanted, identity).box)
        if kind == 'textbox':
            block_with(model, identity).content[0].text = 'Text and geometry changed together'
        before = document_to_json(model)
        output = tmp_path/f'cycle-{cycle}.pptx'
        result = write_document(model, output, options=OPTIONS)
        assert result.success, result.issues
        if cycle == 0:
            assert result.metrics['geometry_precision']['max_slide_error_emu'] == 1
        with ZipFile(source) as old, ZipFile(output) as new, ZipFile(reference) as native:
            assert set(old.namelist()) == set(new.namelist())
            for name in old.namelist():
                if name != 'ppt/slides/slide1.xml':
                    assert new.read(name) == old.read(name)
            target_xml = etree.fromstring(new.read('ppt/slides/slide1.xml'))
            reference_xml = etree.fromstring(native.read('ppt/slides/slide1.xml'))
            if kind == 'textbox':
                reference_xml.find('.//'+A+'t').text = 'Text and geometry changed together'
            assert etree.tostring(target_xml, method='c14n') == etree.tostring(reference_xml, method='c14n')
        assert document_to_json(model) == before
        model = read_document(output).document
        actual, goal = block_with(model, identity).box, block_with(wanted, identity).box
        for field in ('x', 'y', 'width', 'height', 'rotation'):
            assert getattr(actual, field) == pytest.approx(getattr(goal, field), abs=1/12700)
    assert source.read_bytes() == original


def test_ambiguous_45_degree_resize_refuses_without_replacing_destination(tmp_path):
    source, identity = fixture(tmp_path, 'textbox', 45)
    model = read_document(source).document
    block_with(model, identity).box.width += 10
    before = document_to_json(model)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    result = write_document(model, output, options=OPTIONS)
    assert not result.success and any(issue.feature == 'pptx.source-geometry' for issue in result.issues)
    assert output.read_bytes() == b'previous' and document_to_json(model) == before


def test_graphic_frame_resize_requires_native_grid_projection(tmp_path):
    source, identity = fixture(tmp_path, 'table')
    model = read_document(source).document
    block_with(model, identity).box.width += 10
    result = write_document(model, tmp_path/'refused.pptx', options=OPTIONS)
    assert not result.success and not (tmp_path/'refused.pptx').exists()


def test_legacy_v1_text_binding_keeps_native_text_profile(tmp_path):
    source, identity = fixture(tmp_path, 'textbox', 25)
    model = read_document(source).document
    model.metadata[KEY] = {'version': 1, 'digest': source_digest(model, geometry=False)}
    model = document_from_json(document_to_json(model))
    block_with(model, identity).content[0].text = 'Legacy text binding still works'
    result = write_document(model, tmp_path/'v1.pptx', options=OPTIONS)
    assert result.success, result.issues


@pytest.mark.parametrize('node_kind', ['element', 'comment', 'processing-instruction'])
def test_mutated_source_xml_obeys_parser_budget_before_lxml_projection(tmp_path, node_kind):
    source, _identity = fixture(tmp_path, 'textbox')
    model = read_document(source).document
    part = model.package.parts['/ppt/slides/slide1.xml']
    root = etree.fromstring(part.data)
    for _ in range(12000):
        if node_kind == 'element':
            etree.SubElement(root, 'own')
        elif node_kind == 'comment':
            root.append(etree.Comment('own'))
        else:
            root.append(etree.ProcessingInstruction('own', 'inert'))
    part.data = etree.tostring(root)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(
        pptx_profile=PptxProfile(layout='source'), document_limits=DocumentLimits(max_nodes=10000),
    ))
    assert not report.success and any(issue.feature == 'pptx.source-xml-limit' for issue in report.issues)
    assert output.read_bytes() == b'previous'


def test_connected_connector_needs_endpoint_projection(tmp_path):
    source, identity = fixture(tmp_path, 'connector')
    native = Presentation(source)
    shape = native.slides[0].shapes[0].shapes[0]
    properties = shape._element.find(P+'nvCxnSpPr/'+P+'cNvCxnSpPr')
    etree.SubElement(properties, A+'stCxn', id=str(native.slides[0].shapes[1].shape_id), idx='0')
    native.save(source)
    model = read_document(source).document
    block_with(model, identity).box.x += 1
    report = write_document(model, tmp_path/'refused.pptx', options=OPTIONS)
    assert not report.success and any('endpoint' in issue.message for issue in report.issues)
    assert not (tmp_path/'refused.pptx').exists()


def test_native_mirrored_reframe_outside_source_precision_refuses(tmp_path):
    source, identity = fixture(tmp_path, 'textbox', 33, True, True)
    model = read_document(source).document
    # Recorded from the independently edited PowerPoint original. PowerPoint
    # normalizes the nested group; retaining original parent frames cannot
    # represent these bounds within one slide EMU. Never silently accept drift.
    block_with(model, identity).box = Box(
        -266.0063864743431, 96.17631599249486, 450.4752138935682, 406.7544546291835, rotation=70,
    )
    before = document_to_json(model)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=OPTIONS)
    assert not report.success and any('one slide-coordinate EMU' in issue.message for issue in report.issues)
    assert output.read_bytes() == b'previous' and document_to_json(model) == before
