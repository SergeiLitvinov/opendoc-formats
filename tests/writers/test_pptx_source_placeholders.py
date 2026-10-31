"""Local frame overrides retain native placeholder/layout/master semantics."""

from copy import deepcopy
from threading import Event
from zipfile import ZipFile

import pytest
from lxml import etree
from opendoc_model import DocumentLimits, document_from_json, document_to_json
from pptx import Presentation
from pptx.util import Pt

from opendoc_formats import ExportOptions, PptxProfile, read_document, write_document
from opendoc_formats.ooxml.pptx_source_text import KEY, source_digest

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
OPTIONS = ExportOptions(pptx_profile=PptxProfile(layout='source'))


def own_placeholder(root, master=False, angle=0):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = 'Own unchanged title'
    body = slide.placeholders[1]
    body.text = 'Own inherited body'
    layout = slide.slide_layout.placeholders[1]
    owner = layout
    if master:
        frame = layout._element.find(P+'spPr/'+A+'xfrm')
        if frame is not None:
            layout._element.find(P+'spPr').remove(frame)
        owner = next(s for s in slide.slide_layout.slide_master.placeholders
                     if s._element.find('.//'+P+'ph').get('type') == 'body')
    owner.left, owner.top, owner.width, owner.height = Pt(44), Pt(120), Pt(620), Pt(330)
    if angle:
        owner.rotation = angle
    source = root/'source.pptx'
    deck.save(source)
    return source


def body(model):
    return next(block for block in model.sections[0].blocks if block.plain_text.startswith('Own inherited'))


@pytest.mark.parametrize('master', [False, True])
@pytest.mark.parametrize('angle', [0, 17])
def test_inherited_frame_becomes_local_override_without_changing_ancestors(tmp_path, master, angle):
    source = own_placeholder(tmp_path, master, angle)
    original = source.read_bytes()
    model = read_document(source).document
    assert body(model).box.x == 44 and body(model).box.rotation == angle
    native = Presentation(source)
    target = native.slides[0].placeholders[1]
    wanted = deepcopy(body(model).box)
    wanted.x += 10
    wanted.y += 5
    wanted.width -= 40
    wanted.height -= 20
    wanted.rotation += 9
    target.left, target.top = Pt(wanted.x), Pt(wanted.y)
    target.width, target.height = Pt(wanted.width), Pt(wanted.height)
    target.rotation = wanted.rotation
    reference = tmp_path/'reference.pptx'
    native.save(reference)
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        body(model).box = deepcopy(wanted)
        body(model).content[0].text = 'Own inherited edited'
        before = document_to_json(model)
        output = tmp_path/f'override-{cycle}.pptx'
        report = write_document(model, output, options=OPTIONS)
        assert report.success, report.issues
        with ZipFile(source) as old, ZipFile(output) as new, ZipFile(reference) as expected:
            assert set(old.namelist()) == set(new.namelist())
            for name in old.namelist():
                if name != 'ppt/slides/slide1.xml':
                    assert old.read(name) == new.read(name)
            actual = etree.fromstring(new.read('ppt/slides/slide1.xml'))
            goal = etree.fromstring(expected.read('ppt/slides/slide1.xml'))
            goal.findall('.//'+A+'t')[-1].text = 'Own inherited edited'
            assert etree.tostring(actual, method='c14n') == etree.tostring(goal, method='c14n')
        assert document_to_json(model) == before
        model = read_document(output).document
        assert body(model).box == wanted
        source = output
    assert (tmp_path/'source.pptx').read_bytes() == original


def test_text_only_edit_does_not_materialize_inherited_frame(tmp_path):
    source = own_placeholder(tmp_path)
    model = read_document(source).document
    body(model).content[0].text = 'Own inherited revised'
    output = tmp_path/'text-only.pptx'
    report = write_document(model, output, options=OPTIONS)
    assert report.success and report.metrics['geometry_edits'] == 0, report.issues
    assert Presentation(output).slides[0].placeholders[1]._element.find(P+'spPr/'+A+'xfrm') is None


@pytest.mark.parametrize('version', [2, 7])
def test_old_bindings_do_not_enable_inherited_frame_changes(tmp_path, version):
    model = read_document(own_placeholder(tmp_path)).document
    model.metadata[KEY] = {'version': version, 'digest': source_digest(
        model, tables=version >= 3, table_geometry=version >= 4,
        image_crops=version >= 6, structured_text=version >= 7,
    )}
    assert write_document(model, tmp_path/'unchanged.pptx', options=OPTIONS).success
    body(model).box.x += 10
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    assert not write_document(model, output, options=OPTIONS).success
    assert output.read_bytes() == b'previous'


def test_partial_local_frame_is_opaque_instead_of_inheriting_false_extents(tmp_path):
    source = own_placeholder(tmp_path)
    native = Presentation(source)
    native.slides[0].placeholders[1].left += Pt(10)
    native.save(source)
    result = read_document(source)
    assert result.success, result.issues
    assert not any('Own inherited body' in str(block) for block in result.document.sections[0].blocks)
    assert any(i.reason == 'incomplete-local-frame' and i.location for i in result.issues)
    output = tmp_path/'partial.pptx'
    report = write_document(result.document, output, options=OPTIONS)
    assert report.success, report.issues
    with ZipFile(source) as old, ZipFile(output) as new:
        assert all(old.read(name) == new.read(name) for name in old.namelist())


def test_partial_layout_frame_is_not_replaced_with_master_geometry(tmp_path):
    source = own_placeholder(tmp_path)
    native = Presentation(source)
    frame = native.slides[0].slide_layout.placeholders[1]._element.find(P+'spPr/'+A+'xfrm')
    frame.remove(frame.find(A+'ext'))
    native.save(source)
    result = read_document(source)
    assert result.success, result.issues
    assert not any('Own inherited body' in str(block) for block in result.document.sections[0].blocks)


def test_ambiguous_layout_identity_refuses_guessed_override(tmp_path):
    source = own_placeholder(tmp_path)
    native = Presentation(source)
    layout = native.slides[0].slide_layout._element.find(P+'cSld/'+P+'spTree')
    layout.append(deepcopy(native.slides[0].slide_layout.placeholders[1]._element))
    native.save(source)
    result = read_document(source)
    assert result.success, result.issues
    assert not any('Own inherited body' in str(block) for block in result.document.sections[0].blocks)


def test_inherited_xml_obeys_selected_budget_before_local_override(tmp_path):
    model = read_document(own_placeholder(tmp_path)).document
    part = next(p for name, p in model.package.parts.items() if name == '/ppt/slideLayouts/slideLayout2.xml')
    root = etree.fromstring(part.data)
    for _ in range(12000):
        root.append(etree.Comment('own'))
    part.data = etree.tostring(root)
    model.metadata[KEY]['digest'] = source_digest(
        model, table_geometry=True, image_crops=True, structured_text=True, placeholder_geometry=True,
    )
    body(model).box.x += 10
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(
        pptx_profile=PptxProfile(layout='source'), document_limits=DocumentLimits(max_nodes=10000),
    ))
    assert not report.success and any(i.feature == 'pptx.source-xml-limit' for i in report.issues)
    assert output.read_bytes() == b'previous'


def test_cancel_after_local_override_keeps_model_and_destination(tmp_path, monkeypatch):
    from opendoc_formats.writers import pptx_source_writer

    model = read_document(own_placeholder(tmp_path)).document
    body(model).box.x += 10
    before = document_to_json(model)
    stop = Event()
    original = pptx_source_writer.project_geometry
    def project(*args, **kwargs):
        changed = original(*args, **kwargs)
        if changed:
            stop.set()
        return changed
    monkeypatch.setattr(pptx_source_writer, 'project_geometry', project)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(
        pptx_profile=PptxProfile(layout='source'), cancelled=stop.is_set,
    ))
    assert not report.success and any(i.feature == 'export.cancelled' for i in report.issues)
    assert output.read_bytes() == b'previous' and document_to_json(model) == before
