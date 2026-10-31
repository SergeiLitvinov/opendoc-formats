"""Source table dimensions preserve native topology, payloads and old profiles."""

from copy import deepcopy
from threading import Event
from zipfile import ZipFile

import pytest
from lxml import etree
from opendoc_model import document_from_json, document_to_json
from pptx import Presentation
from pptx.util import Pt

from opendoc_formats import ExportOptions, PptxProfile, read_document, write_document
from opendoc_formats.ooxml.pptx_source_text import KEY, source_digest
from tests.writers.test_pptx_logical_grid import merged_source
from tests.writers.test_pptx_source_geometry import A, P

OPTIONS = ExportOptions(pptx_profile=PptxProfile(layout='source'))


@pytest.mark.parametrize('mode,widths,heights', [
    ('box', [96, 144, 168, 192], [72]*5),
    ('points', [90, 110, 140, 160], [70, 50, 60, 60, 60]),
    ('twips', [90, 110, 140, 160], [70, 50, 60, 60, 60]),
    ('both', [90, 110, 140, 160], [70, 50, 60, 60, 60]),
    ('combined', [90, 130, 180, 200], [80, 60, 70, 70, 80]),
])
def test_source_table_dimensions_match_independent_native_grid_twice(tmp_path, mode, widths, heights):
    source = merged_source(tmp_path)
    original = source.read_bytes()
    native = Presentation(source)
    table = native.slides[0].shapes[0].table
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Pt(width)
    for row, height in zip(table.rows, heights, strict=True):
        row.height = Pt(height)
    reference = tmp_path/'reference.pptx'
    native.save(reference)
    model = read_document(source).document
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        table = model.sections[0].blocks[0]
        assert model.metadata[KEY]['version'] >= 4
        table.box.width, table.box.height = sum(widths), sum(heights)
        if mode in ('points', 'both', 'combined'):
            table.properties['column_widths_pt'] = widths
        if mode in ('twips', 'both'):
            table.properties['grid_widths_twips'] = [v*20 for v in widths]
        if mode != 'box':
            for row, height in zip(table.rows, heights, strict=True):
                row.properties['height_pt'] = height
        before = document_to_json(model)
        output = tmp_path/f'result-{cycle}.pptx'
        report = write_document(model, output, options=OPTIONS)
        assert report.success, report.issues
        if cycle == 0:
            assert report.metrics['table_grid_edits'] == 1
        with ZipFile(source) as old, ZipFile(output) as new, ZipFile(reference) as expected:
            assert set(old.namelist()) == set(new.namelist())
            for name in old.namelist():
                if name != 'ppt/slides/slide1.xml':
                    assert new.read(name) == old.read(name)
            actual = etree.fromstring(new.read('ppt/slides/slide1.xml'))
            wanted = etree.fromstring(expected.read('ppt/slides/slide1.xml'))
            assert etree.tostring(actual, method='c14n') == etree.tostring(wanted, method='c14n')
        assert document_to_json(model) == before
        model = read_document(output).document
        table = model.sections[0].blocks[0]
        assert table.properties['column_widths_pt'] == widths
        assert [row.properties['height_pt'] for row in table.rows] == heights
    assert source.read_bytes() == original


@pytest.mark.parametrize('damage', ['sum', 'zero', 'nan', 'huge-integer', 'small-column',
                                    'point-twip', 'rotation', 'sub-emu', 'shape-count'])
def test_unproved_dimensions_refuse_without_changing_input_or_destination(tmp_path, damage):
    source = merged_source(tmp_path)
    model = read_document(source).document
    table = model.sections[0].blocks[0]
    if damage == 'sum':
        table.properties['column_widths_pt'][0] += 1
    elif damage == 'zero':
        table.rows[0].properties['height_pt'] = 0
    elif damage == 'nan':
        table.rows[0].properties['height_pt'] = float('nan')
    elif damage == 'huge-integer':
        table.rows[0].properties['height_pt'] = 10**500
    elif damage == 'small-column':
        table.properties['column_widths_pt'][0] = 0.5/12700
    elif damage == 'point-twip':
        table.properties['column_widths_pt'] = [90, 110, 140, 160]
        table.properties['grid_widths_twips'] = [2000, 2000, 2800, 3200]
    elif damage == 'rotation':
        table.box.rotation += 10
    elif damage == 'sub-emu':
        table.box.width = 1/12700
    else:
        table.properties['column_widths_pt'].pop()
    # Nonfinite extension values are rejected by the Model codec itself.
    before = document_to_json(model) if damage != 'nan' else deepcopy(model)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=OPTIONS)
    assert not report.success and output.read_bytes() == b'previous', report.issues
    if damage != 'nan':
        assert document_to_json(model) == before


@pytest.mark.parametrize('version', [2, 3])
def test_legacy_bindings_keep_translation_only_for_tables(tmp_path, version):
    model = read_document(merged_source(tmp_path)).document
    model.metadata[KEY] = {'version': version, 'digest': source_digest(model, tables=version >= 3)}
    table = model.sections[0].blocks[0]
    table.box.x += 10
    report = write_document(model, tmp_path/'moved.pptx', options=OPTIONS)
    assert report.success, report.issues
    table.box.width += 100
    assert not write_document(model, tmp_path/'resized.pptx', options=OPTIONS).success


def test_v4_snapshot_keeps_table_resize_after_new_chart_binding_version(tmp_path):
    model = read_document(merged_source(tmp_path)).document
    model.metadata[KEY] = {'version': 4, 'digest': source_digest(model, table_geometry=True)}
    model = document_from_json(document_to_json(model))
    table = model.sections[0].blocks[0]
    table.box.width += 100
    table.box.height += 60
    output = tmp_path/'v4.pptx'
    report = write_document(model, output, options=OPTIONS)
    assert report.success and report.metrics['table_grid_edits'] == 1, report.issues
    native = Presentation(output).slides[0].shapes[0].table
    assert [column.width for column in native.columns] == [Pt(v) for v in (96, 144, 168, 192)]
    assert [row.height for row in native.rows] == [Pt(72)]*5


def test_table_grid_changes_and_cell_text_share_one_source_dom(tmp_path):
    source = merged_source(tmp_path)
    model = read_document(source).document
    table = model.sections[0].blocks[0]
    table.box.width += 100
    table.rows[1].cells[0].blocks[0].content[0].text = 'Editable after resize'
    output = tmp_path/'edited.pptx'
    report = write_document(model, output, options=OPTIONS)
    assert report.success, report.issues
    assert report.metrics['table_grid_edits'] == 1 and report.metrics['cell_text_edits'] == 1
    native = Presentation(output).slides[0].shapes[0].table
    assert native.cell(1, 2).text == 'Editable after resize'
    assert [column.width for column in native.columns] == [Pt(v) for v in (96, 144, 168, 192)]


def test_table_grid_cancellation_keeps_input_and_destination(tmp_path, monkeypatch):
    from opendoc_formats.writers import pptx_source_writer

    model = read_document(merged_source(tmp_path)).document
    model.sections[0].blocks[0].box.width += 100
    before = document_to_json(model)
    stop = Event()
    original = pptx_source_writer.project_table_grid
    def project(binding):
        result = original(binding)
        stop.set()
        return result
    monkeypatch.setattr(pptx_source_writer, 'project_table_grid', project)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(
        pptx_profile=PptxProfile(layout='source'), cancelled=stop.is_set,
    ))
    assert not report.success and any(issue.feature == 'export.cancelled' for issue in report.issues)
    assert output.read_bytes() == b'previous' and document_to_json(model) == before


def test_unproved_original_grid_extent_does_not_enable_resize(tmp_path):
    source = merged_source(tmp_path)
    native = Presentation(source)
    native.slides[0].shapes[0]._element.find(P+'xfrm/'+A+'ext').set('cx', str(Pt(501)))
    native.save(source)
    model = read_document(source).document
    model.sections[0].blocks[0].box.width += 100
    report = write_document(model, tmp_path/'refused.pptx', options=OPTIONS)
    assert not report.success


def test_fractional_resize_distributes_rounding_without_grid_drift(tmp_path):
    model = read_document(merged_source(tmp_path)).document
    table = model.sections[0].blocks[0]
    table.box.width += 1/12700
    table.box.height += 1/12700
    output = tmp_path/'fractional.pptx'
    report = write_document(model, output, options=OPTIONS)
    assert report.success, report.issues
    native = Presentation(output).slides[0].shapes[0]
    widths = [column.width for column in native.table.columns]
    heights = [row.height for row in native.table.rows]
    assert sum(widths) == native.width == Pt(500)+1
    assert sum(heights) == native.height == Pt(300)+1
    assert all(abs(value-original * native.width / Pt(500)) <= 1
               for value, original in zip(widths, [Pt(v) for v in (80, 120, 140, 160)], strict=True))
    assert max(heights)-min(heights) == 1
