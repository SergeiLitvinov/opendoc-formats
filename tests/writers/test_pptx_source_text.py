"""Native source text edits preserve unknown neighbors and refuse stale projections."""

from copy import deepcopy
from threading import Event
from zipfile import ZipFile

import pytest
from lxml import etree
from opendoc_model import document_from_json, document_to_json
from pptx import Presentation
from pptx.util import Inches, Pt

from opendoc_formats import ExportOptions, PptxProfile, read_document, write_document
from opendoc_formats.ooxml.pptx_source_text import KEY, source_digest
from tests.readers.test_pptx_source_package import own_pptx, rewrite

OPTIONS = ExportOptions(pptx_profile=PptxProfile(layout='source'))
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def text_model(path):
    result = read_document(path)
    assert result.success, result.issues
    return result.document


@pytest.mark.parametrize('edited', [False, True])
def test_original_package_and_source_text_survive_two_native_cycles(tmp_path, edited):
    source = own_pptx(tmp_path)
    original = source.read_bytes()
    model = text_model(source)
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        if edited:
            model.sections[0].blocks[0].content[0].text = f'Own native text {cycle}'
        before = document_to_json(model)
        output = tmp_path/f'cycle-{cycle}.pptx'
        result = write_document(model, output, options=OPTIONS)
        assert result.success and not result.lossless, result.issues
        assert any(issue.feature == 'pptx.source-container' for issue in result.issues)
        with ZipFile(source) as old, ZipFile(output) as new:
            assert set(old.namelist()) == set(new.namelist())
            changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
            assert changed == (['ppt/slides/slide1.xml'] if edited else [])
            if edited:
                expected = etree.fromstring(old.read(changed[0]))
                expected.find('.//'+A+'t').text = f'Own native text {cycle}'
                actual = etree.fromstring(new.read(changed[0]))
                assert etree.tostring(expected, method='c14n') == etree.tostring(actual, method='c14n')
        assert Presentation(output).slides[0].shapes[0].text == (
            f'Own native text {cycle}' if edited else 'Own source text'
        )
        assert document_to_json(model) == before
        model = text_model(output)
        assert '/own/unknown.bin' in model.package.parts
    assert source.read_bytes() == original


@pytest.mark.parametrize('change', ['box', 'style', 'link', 'notes', 'order', 'raw', 'binding', 'newline', 'nul'])
def test_unsupported_or_stale_change_refuses_atomically(tmp_path, change):
    source = own_pptx(tmp_path)
    model = text_model(source)
    block = model.sections[0].blocks[0]
    if change == 'box':
        model.metadata[KEY] = {'version': 1, 'digest': source_digest(model, geometry=False)}
        block.box.x += 10
    elif change == 'style':
        block.content[0].style.bold = True
    elif change == 'link':
        block.content[0].link = 'https://example.invalid/new'
    elif change == 'notes':
        model.sections[0].properties['notes'] = 'New notes'
    elif change == 'order':
        model.sections.append(deepcopy(model.sections[0]))
    elif change == 'raw':
        model.package.parts['/own/unknown.bin'].data += b'changed'
    elif change == 'binding':
        model.metadata[KEY]['version'] = True
    else:
        block.content[0].text = 'New'+ ('\n' if change == 'newline' else '\x00')+'text'
    before = document_to_json(model)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'own previous file')
    result = write_document(model, output, options=OPTIONS)
    assert not result.success, result.issues
    assert output.read_bytes() == b'own previous file'
    assert document_to_json(model) == before


def test_cancel_during_zip_write_preserves_source_model_and_destination(tmp_path, monkeypatch):
    from opendoc_formats.writers import pptx_source_writer
    source = own_pptx(tmp_path)
    model = text_model(source)
    model.sections[0].blocks[0].content[0].text = 'Edited'
    before = document_to_json(model)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'kept')
    stop = Event()
    original = pptx_source_writer.part_name
    def first(name):
        stop.set()
        return original(name)
    monkeypatch.setattr(pptx_source_writer, 'part_name', first)
    result = write_document(model, output, options=ExportOptions(
        pptx_profile=PptxProfile(layout='source'), cancelled=stop.is_set,
    ))
    assert not result.success and any(issue.feature == 'export.cancelled' for issue in result.issues)
    assert output.read_bytes() == b'kept' and document_to_json(model) == before


@pytest.mark.parametrize('profile', ['path', 'content-type', 'relationship'])
def test_source_profile_refuses_signed_text_edit(tmp_path, profile):
    source = own_pptx(tmp_path)
    name = '_xmlsignatures/sig1.xml' if profile == 'path' else 'own/signature.xml'
    updates = {name: b'<own-signature/>'}
    if profile == 'content-type':
        with ZipFile(source) as archive:
            types = etree.fromstring(archive.read('[Content_Types].xml'))
        etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                         PartName='/'+name,
                         ContentType='application/vnd.openxmlformats-package.digital-signature-xmlsignature+xml')
        updates['[Content_Types].xml'] = etree.tostring(types)
    elif profile == 'relationship':
        updates['own/_rels/unknown.bin.rels'] = (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rIdSignature" '
            'Type="http://schemas.openxmlformats.org/package/2006/relationships/digital-signature/signature" '
            'Target="signature.xml"/></Relationships>'
        ).encode()
    rewrite(source, updates)
    model = text_model(source)
    model.sections[0].blocks[0].content[0].text = 'Edited'
    output = tmp_path/'signed.pptx'
    result = write_document(model, output, options=OPTIONS)
    assert not result.success and any(issue.feature == 'pptx.source-signature' for issue in result.issues)
    assert not output.exists()


def test_grouped_multirun_hyperlinks_and_paragraphs_keep_exact_native_topology(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    frame = shape.text_frame
    first = frame.paragraphs[0].add_run()
    first.text, first.font.bold, first.font.size = 'Styled first', True, Pt(18)
    second = frame.paragraphs[0].add_run()
    second.text, second.font.italic = ' linked second', True
    second.hyperlink.address = 'https://example.invalid/inert'
    frame.add_paragraph().text = 'Next paragraph'
    slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = 'Neighbor'
    source = tmp_path/'grouped.pptx'
    deck.save(source)
    model = text_model(source)
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        content = model.sections[0].blocks[0].content
        assert [run.text for run in content] == (
            ['Styled first', ' linked second', '\n', 'Next paragraph'] if cycle == 0
            else ['Revised bold', ' Revised link', '\n', 'Revised next']
        )
        for run, text in zip((content[0], content[1], content[3]),
                             ('Revised bold', ' Revised link', 'Revised next'), strict=True):
            run.text = text
        output = tmp_path/f'grouped-{cycle}.pptx'
        report = write_document(model, output, options=OPTIONS)
        assert report.success, report.issues
        with ZipFile(source) as old, ZipFile(output) as new:
            for name in old.namelist():
                if name == 'ppt/slides/slide1.xml':
                    continue
                assert new.read(name) == old.read(name)
        target = Presentation(output)
        grouped = target.slides[0].shapes[0].shapes[0].text_frame
        assert grouped.paragraphs[0].runs[0].font.bold
        assert grouped.paragraphs[0].runs[1].font.italic
        assert grouped.paragraphs[0].runs[1].hyperlink.address == 'https://example.invalid/inert'
        assert target.slides[0].shapes[1].text == 'Neighbor'
        model = text_model(output)


@pytest.mark.parametrize('value', [True, 1, None, 'unknown'])
def test_profile_rejects_unknown_layout(value):
    with pytest.raises(ValueError):
        PptxProfile(layout=value)
