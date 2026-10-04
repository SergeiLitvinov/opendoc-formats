"""Нативные тренды и погрешности сохраняются после изменения и удаления."""

from __future__ import annotations

import pytest

pytest.importorskip("pptx")
from opendoc.document_codec import document_from_json, document_to_json
from opendoc.document_model import DocumentModel, Paragraph, Section, TextRun
from pptx import Presentation

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _model(kind, trend, errors):
    series = {"name": "Data", "values": [2, 4, 7, 9], "trendline": trend, "error_bars": errors}
    if kind == "scatterChart":
        series["x_values"] = [1, 3, 4, 8]
    chart = {"chart_type": kind, "categories": ["A", "B", "C", "D"], "series": [series]}
    return DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Chart")], properties={"pptx": {"chart": chart}})])])


@pytest.mark.parametrize("kind", ["lineChart", "barChart", "scatterChart"])
@pytest.mark.parametrize("trend_type", ["linear", "exp", "log", "poly", "power", "movingAvg"])
def test_statistics_mutation_roundtrip(tmp_path, kind, trend_type):
    trend = {"type": trend_type, "color": "#225588"}
    if trend_type == "poly":
        trend["order"] = 3
    elif trend_type == "movingAvg":
        trend["period"] = 2
    else:
        trend.update(show_equation=True, show_r_squared=True)
    errors = {"value_type": "fixedVal", "value": 0.5, "bar_type": "both", "color": "#882255"}
    model = _model(kind, trend, errors)
    for cycle in range(3):
        model = document_from_json(document_to_json(model))
        output = tmp_path / f"stats{cycle}.pptx"
        report = write_pptx_model(model, output)
        assert report.success
        assert not any("не перенесены" in issue.message for issue in report.issues)
        series_xml = Presentation(output).slides[0].shapes[0].chart.series[0]._element
        tags = [child.tag for child in series_xml]
        model = read_pptx_model(output)
        series = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]
        assert list(map(float, series["values"])) == [2, 4, 7, 9]
        if cycle == 2:
            assert "trendline" not in series and "error_bars" not in series
            break
        assert (
            tags.index(C + "trendline")
            < tags.index(C + "errBars")
            < tags.index(C + ("xVal" if kind == "scatterChart" else "cat"))
        )
        assert series["trendline"]["type"] == (trend_type if cycle == 0 else "poly")
        assert series["trendline"]["color"] == ("#225588" if cycle == 0 else "#118833")
        assert series["error_bars"]["value"] == (0.5 if cycle == 0 else 10)
        assert series["error_bars"]["value_type"] == ("fixedVal" if cycle == 0 else "percentage")
        assert series["error_bars"]["bar_type"] == ("both" if cycle == 0 else "plus")
        if cycle == 0:
            assert series["error_bars"]["color"] == "#882255"
            if trend_type == "poly":
                assert series["trendline"]["order"] == 3
            elif trend_type == "movingAvg":
                assert series["trendline"]["period"] == 2
            else:
                assert series["trendline"]["show_equation"] and series["trendline"]["show_r_squared"]
        if cycle == 0:
            series["trendline"] = {"type": "poly", "order": 4, "color": "#118833"}
            series["error_bars"] = {"value_type": "percentage", "value": 10, "bar_type": "plus"}
        else:
            del series["trendline"]
            del series["error_bars"]


@pytest.mark.parametrize("value_type", ["stdDev", "stdErr", "percentage", "fixedVal"])
@pytest.mark.parametrize("direction", ["x", "y"])
def test_xy_error_modes(tmp_path, value_type, direction):
    settings = {"value_type": value_type, "direction": direction, "bar_type": "minus"}
    if value_type != "stdErr":
        settings["value"] = 2
    model = _model("scatterChart", None, settings)
    output = tmp_path / "errors.pptx"
    assert write_pptx_model(model, output).success
    result = read_pptx_model(output).sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]
    assert result["error_bars"] == settings


@pytest.mark.parametrize(
    "key,settings",
    [
        ("trendline", {"type": "poly", "order": 7}),
        ("trendline", {"type": "movingAvg", "period": 1}),
        ("trendline", {"type": "linear", "order": 3}),
        ("trendline", {"type": "movingAvg", "show_equation": True}),
        ("error_bars", {"value_type": "cust", "plus": [1, 2]}),
        ("error_bars", {"value_type": "fixedVal", "value": float("nan")}),
        ("error_bars", {"value_type": "stdErr", "value": 1}),
        ("error_bars", {"direction": "x"}),
    ],
)
def test_invalid_statistics_reported_without_partial_xml(tmp_path, key, settings):
    import math

    model = _model("lineChart", None, None)
    model.sections[0].blocks[0].properties["pptx"]["chart"]["series"][0][key] = settings
    output = tmp_path / "invalid.pptx"
    report = write_pptx_model(model, output)
    if type(settings.get("value")) is float and not math.isfinite(settings["value"]):
        assert not report.success and not output.exists()
        assert any(issue.feature == "model" and "number must be finite" in issue.message for issue in report.issues)
        return
    assert report.success and any(key in issue.message and issue.location for issue in report.issues)
    series = Presentation(output).slides[0].shapes[0].chart.series[0]._element
    assert series.find(C + ("trendline" if key == "trendline" else "errBars")) is None


def test_combo_statistics_are_independent(tmp_path):
    model = _model("barChart", {"type": "linear"}, {"value_type": "fixedVal", "value": 1})
    chart = model.sections[0].blocks[0].properties["pptx"]["chart"]
    chart["combo_types"] = ["barChart", "lineChart"]
    chart["series"].append(
        {
            "name": "Second",
            "values": [3, 5, 8, 10],
            "chart_type": "lineChart",
            "axis": "secondary_value",
            "trendline": {"type": "poly", "order": 2},
            "error_bars": {"value_type": "stdErr"},
        }
    )
    for cycle in range(2):
        output = tmp_path / f"combo{cycle}.pptx"
        assert write_pptx_model(document_from_json(document_to_json(model)), output).success
        model = read_pptx_model(output)
        series = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"]
        assert series[0]["trendline"]["type"] == "linear"
        assert series[0]["error_bars"]["value_type"] == "fixedVal"
        assert series[1]["trendline"]["order"] == 2 + cycle
        assert series[1]["error_bars"]["value_type"] == "stdErr"
        series[1]["trendline"]["order"] = 3


def test_horizontal_bar_error_direction(tmp_path):
    model = _model("barChart", None, {"value_type": "stdDev", "value": 2})
    model.sections[0].blocks[0].properties["pptx"]["chart"]["bar_direction"] = "bar"
    output = tmp_path / "horizontal.pptx"
    assert write_pptx_model(model, output).success
    series = read_pptx_model(output).sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]
    assert series["error_bars"]["direction"] == "x"
