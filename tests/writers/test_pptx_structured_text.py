"""Ordinary native runs remain editable around protected breaks and field caches."""

from threading import Event
from zipfile import ZipFile

import pytest
from lxml import etree
from opendoc_model import Paragraph, document_from_json, document_to_json
from pptx import Presentation
from pptx.util import Pt

from opendoc_formats import ExportOptions, PptxProfile, read_document, write_document
from opendoc_formats.ooxml.pptx_source_text import KEY, source_digest

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
OPTIONS = ExportOptions(pptx_profile=PptxProfile(layout='source'))


def own_structured_text(root, kind='shape', cache='1'):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if kind == 'table':
        owner = slide.shapes.add_table(1, 1, Pt(40), Pt(40), Pt(450), Pt(300)).table.cell(0, 0)
    else:
        shapes = slide.shapes.add_group_shape().shapes if kind == 'group' else slide.shapes
        owner = shapes.add_textbox(Pt(40), Pt(40), Pt(450), Pt(300))
    paragraph = owner.text_frame.paragraphs[0]
    before = paragraph.add_run()
    before.text = 'Before '
    before.font.bold, before.font.size = True, Pt(22)
    paragraph.add_line_break()
    br = paragraph._p.find(A+'br')
    etree.SubElement(br, A+'rPr', lang='en-US', sz='1900')
    field = etree.SubElement(paragraph._p, A+'fld', id='{53EFAB24-E5A1-4D49-8B67-0113B60A3311}', type='slidenum')
    etree.SubElement(field, A+'rPr', lang='en-US', sz='2200', b='1')
    etree.SubElement(field, A+'t').text = cache
    after = paragraph.add_run()
    after.text = ' After'
    after.font.italic, after.font.size = True, Pt(22)
    after.hyperlink.address = 'https://example.invalid/inert'
    owner.text_frame.add_paragraph()
    owner.text_frame.add_paragraph().text = 'Tail'
    slide.shapes.add_textbox(Pt(40), Pt(400), Pt(500), Pt(40)).text = 'Neighbor stays'
    source = root/'source.pptx'
    deck.save(source)
    return source


def text_frame(deck):
    shape = deck.slides[0].shapes[0]
    if shape.shape_type == 6:
        shape = shape.shapes[0]
    return shape.table.cell(0, 0).text_frame if shape.has_table else shape.text_frame


def runs(model):
    block = model.sections[0].blocks[0]
    if isinstance(block, Paragraph):
        return block.content
    return [run for paragraph in block.rows[0].cells[0].blocks for run in paragraph.content]


@pytest.mark.parametrize('kind', ['shape', 'table', 'group'])
@pytest.mark.parametrize('cache', ['1', ''])
def test_edit_runs_preserves_fields_breaks_and_two_json_cycles(tmp_path, kind, cache):
    source = own_structured_text(tmp_path, kind, cache)
    original = source.read_bytes()
    model = read_document(source).document
    assert any(issue.feature == 'pptx.field' for issue in read_document(source).issues)
    for cycle in range(2):
        model = document_from_json(document_to_json(model))
        wanted = {}
        for run in runs(model):
            if run.text.startswith('Before') or run.text.startswith(' After') or run.text.startswith('Tail'):
                text = run.text.split(' edited')[0]+' edited '+str(cycle)
                wanted[run.text] = text
                run.text = text
        before = document_to_json(model)
        reference = Presentation(source)
        frame = text_frame(reference)
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                if run.text in wanted:
                    run.text = wanted[run.text]
        expected = tmp_path/f'expected-{cycle}.pptx'
        reference.save(expected)
        output = tmp_path/f'output-{cycle}.pptx'
        report = write_document(model, output, options=OPTIONS)
        assert report.success, report.issues
        with ZipFile(source) as old, ZipFile(output) as new, ZipFile(expected) as native:
            assert set(old.namelist()) == set(new.namelist())
            for name in old.namelist():
                if name != 'ppt/slides/slide1.xml':
                    assert old.read(name) == new.read(name)
            old_xml = etree.fromstring(old.read('ppt/slides/slide1.xml'))
            actual = etree.fromstring(new.read('ppt/slides/slide1.xml'))
            goal = etree.fromstring(native.read('ppt/slides/slide1.xml'))
            assert etree.tostring(actual, method='c14n') == etree.tostring(goal, method='c14n')
            for name in ('fld', 'br'):
                assert ([etree.tostring(node) for node in old_xml.iter(A+name)]
                        == [etree.tostring(node) for node in actual.iter(A+name)])
        assert document_to_json(model) == before
        assert text_frame(Presentation(output)).paragraphs[0].runs[0].font.bold
        assert text_frame(Presentation(output)).paragraphs[0].runs[1].hyperlink.address == 'https://example.invalid/inert'
        model = read_document(output).document
        source = output
    assert (tmp_path/'source.pptx').read_bytes() == original


@pytest.mark.parametrize('kind', ['shape', 'table'])
@pytest.mark.parametrize('change', ['field-cache', 'break', 'break-kind', 'field-style', 'run-count'])
def test_protected_native_structure_cannot_be_changed(tmp_path, kind, change):
    model = read_document(own_structured_text(tmp_path, kind)).document
    content = runs(model)
    if change == 'field-cache':
        next(run for run in content if run.text == '1').text = '99'
    elif change == 'field-style':
        next(run for run in content if run.text == '1').style.bold = True
    elif change in ('break', 'break-kind'):
        br = next(run for run in content if run.properties.get('pptx_break') == 'line')
        if change == 'break':
            br.text = 'Changed'
        else:
            br.properties['pptx_break'] = 'paragraph'
    else:
        block = model.sections[0].blocks[0]
        paragraph = block if isinstance(block, Paragraph) else block.rows[0].cells[0].blocks[0]
        paragraph.content.reverse()
    before = document_to_json(model)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=OPTIONS)
    assert not report.success and output.read_bytes() == b'previous'
    assert document_to_json(model) == before


@pytest.mark.parametrize('kind', ['shape', 'table'])
@pytest.mark.parametrize('version', [1, 3, 6])
def test_old_versions_keep_old_text_scope(tmp_path, kind, version):
    model = read_document(own_structured_text(tmp_path, kind)).document
    model.metadata[KEY] = {'version': version, 'digest': source_digest(
        model, geometry=version >= 2, tables=version >= 3,
        table_geometry=version >= 4, image_crops=version >= 6,
    )}
    report = write_document(model, tmp_path/'unchanged.pptx', options=OPTIONS)
    assert report.success, report.issues
    runs(model)[0].text = 'Unsupported in old version'
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    assert not write_document(model, output, options=OPTIONS).success
    assert output.read_bytes() == b'previous'


def test_unknown_native_child_does_not_get_guessed_text_slots(tmp_path):
    source = own_structured_text(tmp_path)
    native = Presentation(source)
    etree.SubElement(text_frame(native).paragraphs[0]._p, '{urn:own}opaque')
    native.save(source)
    model = read_document(source).document
    report = write_document(model, tmp_path/'unchanged.pptx', options=OPTIONS)
    assert report.success, report.issues
    runs(model)[0].text = 'No guessed binding'
    assert not write_document(model, tmp_path/'refused.pptx', options=OPTIONS).success


def test_cancel_during_structured_text_mutation_keeps_destination(tmp_path, monkeypatch):
    from opendoc_formats.writers import pptx_source_writer

    model = read_document(own_structured_text(tmp_path, 'table')).document
    runs(model)[0].text = 'First edit'
    runs(model)[-1].text = 'Second edit'
    before = document_to_json(model)
    stop = Event()
    original = pptx_source_writer.apply_cell_text
    def apply(binding):
        assert not stop.is_set()
        changed = original(binding)
        if changed:
            stop.set()
        return changed
    monkeypatch.setattr(pptx_source_writer, 'apply_cell_text', apply)
    output = tmp_path/'kept.pptx'
    output.write_bytes(b'previous')
    report = write_document(model, output, options=ExportOptions(
        pptx_profile=PptxProfile(layout='source'), cancelled=stop.is_set,
    ))
    assert not report.success and any(issue.feature == 'export.cancelled' for issue in report.issues)
    assert output.read_bytes() == b'previous' and document_to_json(model) == before
