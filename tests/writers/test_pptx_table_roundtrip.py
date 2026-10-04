"""Cell text and direct appearance remain independently editable after two round-trips."""

import pytest

pytest.importorskip("pptx")
from lxml import etree
from opendoc.document_codec import load_document, save_document
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def test_cell_text_and_appearance_mutations(tmp_path):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Pt(20), Pt(20), Pt(300), Pt(160)).table
    table.rows[0].height, table.rows[1].height = Pt(95), Pt(65)
    cell = table.cell(0, 0)
    cell.margin_left, cell.margin_top = Pt(11), Pt(9)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(10, 20, 30)
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    first = paragraph.add_run()
    first.text = "Bold"
    first.font.bold, first.font.size = True, Pt(19)
    first.font.color.rgb = RGBColor(200, 30, 20)
    second = paragraph.add_run()
    second.text = "Link"
    second.hyperlink.address = "https://example.com/"
    paragraph.add_line_break()
    paragraph.add_run().text = "Continuation"
    cell.text_frame.add_paragraph()  # Preserve an empty middle paragraph.
    last = cell.text_frame.add_paragraph()
    last.text = "Last"
    last.alignment = PP_ALIGN.CENTER
    last.space_after = Pt(6)
    line = etree.Element(A + "lnL", w=str(Pt(2.5)))
    solid = etree.SubElement(line, A + "solidFill")
    etree.SubElement(solid, A + "srgbClr", val="112233")
    etree.SubElement(line, A + "prstDash", val="dash")
    cell._tc.get_or_add_tcPr().insert(0, line)
    table.cell(1, 0).merge(table.cell(1, 1))
    table.cell(1, 0).text = "Merged"
    edge = etree.Element(A + "lnR", w=str(Pt(3)))
    etree.SubElement(etree.SubElement(edge, A + "solidFill"), A + "srgbClr", val="778899")
    table.cell(1, 1)._tc.get_or_add_tcPr().insert(0, edge)
    table.cell(0, 1).fill.background()
    presentation.save(source)
    model = read_pptx_model(source)
    block = model.sections[0].blocks[0]
    cell_model = block.rows[0].cells[0]
    assert len(cell_model.blocks) == 3
    cell_model.blocks[0].content[0].text = "Edited"
    cell_model.properties["fill"] = "#445566"
    block.rows[0].properties["height_pt"] = 110
    for index in range(2):
        saved, output = tmp_path / "model.json", tmp_path / f"result{index}.pptx"
        save_document(model, saved)
        assert write_pptx_model(load_document(saved), output).success
        table = Presentation(output).slides[0].shapes[0].table
        cell = table.cell(0, 0)
        assert len(cell.text_frame.paragraphs) == 3
        first, blank, last = cell.text_frame.paragraphs
        assert first.text == "EditedLink\vContinuation" and blank.text == ""
        assert first.runs[0].font.bold and first.runs[0].font.size == Pt(19)
        assert first.runs[1].hyperlink.address == "https://example.com/"
        assert first.alignment == PP_ALIGN.RIGHT and last.alignment == PP_ALIGN.CENTER
        assert cell.fill.fore_color.rgb == RGBColor(0x44, 0x55, 0x66)
        assert cell.margin_left == Pt(11) and cell.margin_top == Pt(9)
        assert cell.vertical_anchor == MSO_ANCHOR.MIDDLE
        border = cell._tc.find(".//" + A + "lnL")
        assert border.get("w") == str(Pt(2.5))
        assert border.find(".//" + A + "srgbClr").get("val") == "112233"
        assert border.find(A + "prstDash").get("val") == "dash"
        assert table.rows[0].height == Pt(110)
        assert table.cell(1, 0).span_width == 2
        edge = table.cell(1, 1)._tc.find(".//" + A + "lnR")
        assert edge is not None and edge.get("w") == str(Pt(3))
        assert edge.find(".//" + A + "srgbClr").get("val") == "778899"
        assert table.cell(0, 1)._tc.find(".//" + A + "noFill") is not None
        model = read_pptx_model(output)
