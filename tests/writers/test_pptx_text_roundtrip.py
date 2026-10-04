"""Paragraph boundaries, list numbering, and spacing survive editable round-trips."""

import pytest

pytest.importorskip("pptx")
from lxml import etree
from opendoc.document_codec import load_document, save_document
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@pytest.mark.parametrize("mode", ["noAutofit", "normAutofit", "spAutoFit"])
@pytest.mark.parametrize("in_table", [False, True])
def test_tabs_vertical_text_and_autofit_are_editable(tmp_path, mode, in_table):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = (
        slide.shapes.add_table(1, 1, Pt(10), Pt(10), Pt(300), Pt(200))
        if in_table
        else slide.shapes.add_textbox(Pt(10), Pt(10), Pt(300), Pt(200))
    )
    frame = shape.table.cell(0, 0).text_frame if in_table else shape.text_frame
    frame.text = "Label\tValue"
    body = frame._txBody.bodyPr
    body.set("vert", "vert270")
    for node in list(body):
        if node.tag in (A + "noAutofit", A + "normAutofit", A + "spAutoFit"):
            body.remove(node)
    attrs = {"fontScale": "85000", "lnSpcReduction": "12000"} if mode == "normAutofit" else {}
    etree.SubElement(body, A + mode, **attrs)
    props = frame.paragraphs[0]._p.get_or_add_pPr()
    props.set("defTabSz", str(Pt(36)))
    tabs = etree.SubElement(props, A + "tabLst")
    etree.SubElement(tabs, A + "tab", pos=str(Pt(90)), algn="r")
    presentation.save(source)
    model = read_pptx_model(source)
    block = model.sections[0].blocks[0]
    paragraph = block.rows[0].cells[0].blocks[0] if in_table else block
    settings = paragraph.properties["pptx"]
    assert settings["text_frame"]["autofit"]["mode"] == mode
    assert settings["text_frame"]["vert"] == "vert270"
    settings["text_frame"]["vert"] = "eaVert"
    settings["paragraphs"][0]["tabs"][0]["position_pt"] = 120
    if mode == "normAutofit":
        settings["text_frame"]["autofit"]["fontScale"] = 75000
    for index in range(2):
        saved, output = tmp_path / "model.json", tmp_path / f"output{index}.pptx"
        save_document(model, saved)
        assert write_pptx_model(load_document(saved), output).success
        shape = Presentation(output).slides[0].shapes[0]
        frame = shape.table.cell(0, 0).text_frame if in_table else shape.text_frame
        assert frame.text == "Label\tValue"
        assert frame._txBody.bodyPr.get("vert") == "eaVert"
        fit = [node for node in frame._txBody.bodyPr if node.tag in (A + "noAutofit", A + "normAutofit", A + "spAutoFit")]
        assert len(fit) == 1 and fit[0].tag == A + mode
        if mode == "normAutofit":
            assert fit[0].get("fontScale") == "75000" and fit[0].get("lnSpcReduction") == "12000"
        props = frame.paragraphs[0]._p.get_or_add_pPr()
        assert props.get("defTabSz") == str(Pt(36))
        tab = props.find(A + "tabLst/" + A + "tab")
        assert tab.get("pos") == str(Pt(120)) and tab.get("algn") == "r"
        model = read_pptx_model(output)


def test_paragraphs_lists_and_soft_breaks(tmp_path):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(Pt(20), Pt(20), Pt(400), Pt(300)).text_frame
    frame.margin_left, frame.margin_top = Pt(19), Pt(13)
    frame.word_wrap = False
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = "First"
    paragraph.add_line_break()
    paragraph.add_run().text = "Continuation"
    paragraph.alignment = PP_ALIGN.RIGHT
    paragraph.space_before = Pt(7)
    paragraph.line_spacing = 1.5
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", str(Pt(30)))
    ppr.set("indent", str(Pt(-12)))
    etree.SubElement(ppr, A + "buFont", typeface="Arial")
    etree.SubElement(ppr, A + "buChar", char="•")
    numbered = frame.add_paragraph()
    numbered.text = "Second"
    numbered.level = 2
    numbered.alignment = PP_ALIGN.CENTER
    numbered.line_spacing = Pt(22)
    numbered.space_after = Pt(9)
    etree.SubElement(numbered._p.get_or_add_pPr(), A + "buAutoNum", type="arabicPeriod", startAt="4")
    plain = frame.add_paragraph()
    plain.text = "Third"
    etree.SubElement(plain._p.get_or_add_pPr(), A + "buNone")
    presentation.save(source)
    model = read_pptx_model(source)
    block = model.sections[0].blocks[0]
    block.content[0].text = "Edited"
    block.properties["pptx"]["paragraphs"][1]["number_start"] = 8
    for index in range(2):
        saved, output = tmp_path / "model.json", tmp_path / f"output{index}.pptx"
        save_document(model, saved)
        assert write_pptx_model(load_document(saved), output).success
        target = Presentation(output).slides[0].shapes[0].text_frame
        assert len(target.paragraphs) == 3
        first, second, third = target.paragraphs
        assert first.text == "Edited\vContinuation"
        assert first.alignment == PP_ALIGN.RIGHT and second.alignment == PP_ALIGN.CENTER
        assert first.line_spacing == 1.5 and second.line_spacing == Pt(22)
        assert first.space_before == Pt(7) and second.space_after == Pt(9)
        assert first._p.get_or_add_pPr().get("indent") == str(Pt(-12))
        assert first._p.find(".//" + A + "buChar").get("char") == "•"
        assert second._p.find(".//" + A + "buAutoNum").get("startAt") == "8"
        assert second.level == 2
        assert third._p.find(".//" + A + "buNone") is not None
        assert third.alignment is None
        assert target.margin_left == Pt(19) and target.margin_top == Pt(13)
        assert target.vertical_anchor == MSO_ANCHOR.MIDDLE and not target.word_wrap
        model = read_pptx_model(output)


def test_local_list_style_and_canonical_spacing_override(tmp_path):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(Pt(10), Pt(10), Pt(300), Pt(200)).text_frame
    frame.text = "Inherited list"
    style = frame._txBody.find(A + "lstStyle")
    level = etree.SubElement(style, A + "lvl1pPr", marL=str(Pt(35)))
    etree.SubElement(level, A + "buChar", char="→")
    frame.paragraphs[0].line_spacing = Pt(28)
    plain = frame.add_paragraph()
    plain.text = "No list"
    etree.SubElement(plain._p.get_or_add_pPr(), A + "buNone")
    presentation.save(source)
    model = read_pptx_model(source)
    block = model.sections[0].blocks[0]
    block.properties["line_spacing"] = 1.2
    output = tmp_path / "output.pptx"
    assert write_pptx_model(model, output).success
    paragraphs = Presentation(output).slides[0].shapes[0].text_frame.paragraphs
    assert paragraphs[0].line_spacing == 1.2
    assert paragraphs[0]._p.find(".//" + A + "buChar").get("char") == "→"
    assert paragraphs[1]._p.find(".//" + A + "buChar") is None
    assert paragraphs[1]._p.find(".//" + A + "buNone") is not None


def test_run_override_keeps_inherited_font_and_runs_are_independent(tmp_path):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(Pt(10), Pt(10), Pt(300), Pt(200)).text_frame
    paragraph = frame.paragraphs[0]
    paragraph.font.name, paragraph.font.size, paragraph.font.bold = "Arial", Pt(23), True
    paragraph.add_run().text = "Inherited"
    paragraph.add_run().text = "Independent"
    run = paragraph.add_run()
    run.text = "Explicit normal"
    run.font.bold = False
    presentation.save(source)
    model = read_pptx_model(source)
    content = model.sections[0].blocks[0].content
    assert all(item.style.font_family == "Arial" and item.style.font_size.pt == 23 for item in content)
    assert content[2].style.bold is False
    content[0].style.font_size.pt = 31
    assert content[1].style.font_size.pt == 23
    output = tmp_path / "output.pptx"
    assert write_pptx_model(model, output).success
    restored = read_pptx_model(output).sections[0].blocks[0].content
    assert [item.style.font_size.pt for item in restored] == [31, 23, 23]


def test_placeholder_list_from_master_is_materialized(tmp_path):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    body = slide.placeholders[1]
    body.text = "Master list"
    # Select bodyStyle explicitly; title and body have independent defaults.
    p = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    level = slide.slide_layout.slide_master._element.find(p + "txStyles/" + p + "bodyStyle/" + A + "lvl1pPr")
    for node in list(level):
        if node.tag in (A + "buChar", A + "buAutoNum", A + "buNone"):
            level.remove(node)
    etree.SubElement(level, A + "buChar", char="◆")
    presentation.save(source)
    model = read_pptx_model(source)
    block = next(block for block in model.sections[0].blocks if block.plain_text == "Master list")
    assert block.properties["pptx"]["paragraphs"][0]["bullet_char"] == "◆"
    output = tmp_path / "output.pptx"
    assert write_pptx_model(model, output).success
    paragraph = next(
        shape for shape in Presentation(output).slides[0].shapes if shape.text == "Master list"
    ).text_frame.paragraphs[0]
    assert paragraph._p.find(".//" + A + "buChar").get("char") == "◆"
