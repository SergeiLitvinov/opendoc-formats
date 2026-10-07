"""Шрифты темы образца, переопределения и повторное редактирование PPTX."""

from __future__ import annotations

import pytest

pytest.importorskip("pptx")
from lxml import etree
from opendoc_model.document_codec import document_from_json, document_to_json
from opendoc_model.document_model import DocumentModel, Paragraph, Section, Table, TextRun, TextStyle
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Inches

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.readers.pptx_theme_fonts import A, slide_theme_fonts
from opendoc_formats.writers.pptx_writer import write_pptx_model


def _install_theme(owner, name, major, minor, *, override=False):
    relationship = RT.THEME_OVERRIDE if override else RT.THEME
    for rel in list(owner.rels.values()):
        if rel.reltype == relationship:
            owner.drop_rel(rel.rId)
    root = etree.Element(A + ("themeOverride" if override else "theme"), nsmap={"a": A[1:-1]})
    container = root if override else etree.SubElement(root, A + "themeElements")
    scheme = etree.SubElement(container, A + "fontScheme", name=name)
    for family, font in (("majorFont", major), ("minorFont", minor)):
        node = etree.SubElement(scheme, A + family)
        etree.SubElement(node, A + "latin", typeface=font)
        etree.SubElement(node, A + "ea", typeface="")
        etree.SubElement(node, A + "cs", typeface="")
    part = Part(
        PackURI(f"/ppt/theme/{name}.xml"),
        "application/vnd.openxmlformats-officedocument.themeOverride+xml"
        if override
        else "application/vnd.openxmlformats-officedocument.theme+xml",
        owner.package,
        etree.tostring(root),
    )
    owner.relate_to(part, relationship)


@pytest.mark.parametrize("override", ["master", "layout", "slide"])
def test_theme_fonts_materialized_and_edited(tmp_path, override):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    _install_theme(slide.slide_layout.slide_master.part, "customMaster", "Georgia", "Arial")
    if override in {"layout", "slide"}:
        _install_theme(slide.slide_layout.part, "customLayout", "Verdana", "Tahoma", override=True)
    if override == "slide":
        _install_theme(slide.part, "customSlide", "Cambria", "Calibri", override=True)
    expected = {"master": ("Georgia", "Arial"), "layout": ("Verdana", "Tahoma"), "slide": ("Cambria", "Calibri")}[override]
    slide.shapes.title.text = "Title"
    p = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    defaults = slide.slide_layout.slide_master._element.find(f"{p}txStyles/{p}titleStyle/{A}lvl1pPr/{A}defRPr")
    defaults.find(A + "latin").set("typeface", "+mj-lt")
    body = slide.placeholders[1]
    body.text = "Body"
    body.text_frame.paragraphs[0].runs[0].font.name = "+mn-lt"
    explicit = body.text_frame.paragraphs[0].add_run()
    explicit.text, explicit.font.name = "Explicit", "Courier New"
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(2), Inches(1)).table
    table.cell(0, 0).text = "Cell"
    table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.name = "+mn-lt"
    source = tmp_path / "source.pptx"
    presentation.save(source)
    model = read_pptx_model(source)
    for cycle in range(3):
        paragraphs = [block for block in model.sections[0].blocks if isinstance(block, Paragraph)]
        title = next(block for block in paragraphs if block.plain_text == "Title")
        body_block = next(block for block in paragraphs if block.plain_text.startswith("Body"))
        assert title.content[0].style.font_family == (expected[0] if cycle < 2 else "Consolas")
        assert [item.style.font_family for item in body_block.content] == [expected[1], "Courier New"]
        cell = next(block for block in model.sections[0].blocks if isinstance(block, Table)).rows[0].cells[0]
        assert cell.blocks[0].content[0].style.font_family == expected[1]
        if cycle == 2:
            break
        if cycle == 1:
            title.content[0].style.font_family = "Consolas"
        output = tmp_path / f"result{cycle}.pptx"
        assert write_pptx_model(document_from_json(document_to_json(model)), output).success
        model = read_pptx_model(output)


def test_slide_override_does_not_leak_to_sibling():
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    _install_theme(first.slide_layout.slide_master.part, "masterFonts", "Georgia", "Arial")
    _install_theme(first.part, "oneSlide", "Verdana", "Tahoma", override=True)
    assert slide_theme_fonts(first)["+mj-lt"] == "Verdana"
    assert slide_theme_fonts(second)["+mj-lt"] == "Georgia"
    assert "+mj-ea" not in slide_theme_fonts(first)


def test_unresolved_theme_font_reports_loss(tmp_path):
    model = DocumentModel(
        sections=[Section(blocks=[Paragraph(content=[TextRun("Text", style=TextStyle(font_family="+mn-ea"))])])]
    )
    report = write_pptx_model(model, tmp_path / "unresolved.pptx")
    assert report.success and not report.lossless
    assert any(issue.feature == "fonts" and issue.location for issue in report.issues)


def test_color_only_override_keeps_fonts_and_malformed_theme_is_tolerated():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    _install_theme(slide.slide_layout.slide_master.part, "masterFonts", "Georgia", "Arial")
    _install_theme(slide.part, "overrideFonts", "Verdana", "Tahoma", override=True)
    override = slide.part.part_related_by(RT.THEME_OVERRIDE)
    override._blob = f'<a:themeOverride xmlns:a="{A[1:-1]}"><a:clrScheme name="colors"/></a:themeOverride>'.encode()
    assert slide_theme_fonts(slide)["+mj-lt"] == "Georgia"
    override._blob = b"<invalid"
    assert slide_theme_fonts(slide)["+mj-lt"] == "Georgia"
