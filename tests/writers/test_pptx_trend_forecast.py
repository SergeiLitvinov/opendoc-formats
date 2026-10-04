"""Прогноз и пересечение нативных трендов сохраняются при редактировании."""

from __future__ import annotations

import pytest

pytest.importorskip("pptx")
from opendoc.document_codec import document_from_json, document_to_json
from opendoc.document_model import DocumentModel, Paragraph, Section, TextRun
from pptx import Presentation

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.html_writer import write_html_model
from opendoc_formats.writers.pptx_writer import write_pptx_model

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _model(kind, trends):
    data = {
        "chart_type": kind,
        "categories": ["A", "B", "C", "D"],
        "series": [{"name": "Data", "values": [2, 4, 6, 9], "x_values": [1, 3, 5, 8], "trendline": trends}],
    }
    return DocumentModel(sections=[Section(blocks=[Paragraph([TextRun("Chart")], properties={"pptx": {"chart": data}})])])


@pytest.mark.parametrize("chart_kind", ["lineChart", "scatterChart"])
@pytest.mark.parametrize("trend_type", ["linear", "poly", "exp", "log", "power"])
def test_forecast_independent_mutation_and_removal(tmp_path, chart_kind, trend_type):
    trend = {"type": trend_type, "forward": 2.5, "backward": 0.5, "show_equation": True}
    if trend_type in {"linear", "poly", "exp"}:
        trend["intercept"] = 1.5 if trend_type == "exp" else -1.5
    model = _model(chart_kind, [trend, {"type": "linear", "forward": 0, "intercept": 0}])
    for cycle in range(3):
        output = tmp_path / f"forecast{cycle}.pptx"
        report = write_pptx_model(document_from_json(document_to_json(model)), output)
        assert report.success and not any("trendline не перенесены" in issue.message for issue in report.issues)
        native = Presentation(output).slides[0].shapes[0].chart.series[0]._element.findall(C + "trendline")[0]
        model = read_pptx_model(output)
        trends = model.sections[0].blocks[0].properties["pptx"]["chart"]["series"][0]["trendline"]
        assert trends[1]["forward"] == 0 and trends[1]["intercept"] == 0
        if cycle == 2:
            assert not any(key in trends[0] for key in ("forward", "backward", "intercept"))
            assert native.find(C + "forward") is None
            break
        assert trends[0]["forward"] == (2.5 if cycle == 0 else 4)
        assert trends[0]["backward"] == (0.5 if cycle == 0 else 0)
        if "intercept" in trend:
            assert trends[0]["intercept"] == (trend["intercept"] if cycle == 0 else 2)
        tags = [child.tag for child in native]
        assert tags.index(C + "forward") < tags.index(C + "backward") < tags.index(C + "dispEq")
        if cycle == 0:
            trends[0].update(forward=4, backward=0)
            if "intercept" in trend:
                trends[0]["intercept"] = 2
        else:
            for key in ("forward", "backward", "intercept"):
                trends[0].pop(key, None)


@pytest.mark.parametrize(
    "settings",
    [
        {"type": "linear", "forward": -1},
        {"type": "linear", "backward": float("inf")},
        {"type": "linear", "intercept": float("nan")},
        {"type": "linear", "forward": True},
        {"type": "movingAvg", "forward": 2},
        {"type": "movingAvg", "intercept": 1},
        {"type": "log", "intercept": 1},
        {"type": "power", "intercept": 1},
        {"type": "exp", "intercept": 0},
    ],
)
def test_invalid_forecast_rejects_collection(tmp_path, settings):
    import math

    output = tmp_path / "invalid.pptx"
    report = write_pptx_model(_model("lineChart", [{"type": "linear"}, settings]), output)
    if any(type(value) is float and not math.isfinite(value) for value in settings.values()):
        assert not report.success and not output.exists()
        assert any(issue.feature == "model" and "number must be finite" in issue.message for issue in report.issues)
        return
    assert report.success and any("trendline не перенесены" in issue.message for issue in report.issues)
    assert not Presentation(output).slides[0].shapes[0].chart.series[0]._element.findall(C + "trendline")


def test_html_reports_unsupported_forecast_without_drawing_false_trend(tmp_path):
    output = tmp_path / "forecast.html"
    report = write_html_model(_model("lineChart", [{"type": "linear", "forward": 2}, {"type": "linear"}]), output)
    assert report.success and any(issue.feature == "chart-trendline" for issue in report.issues)
    assert output.read_text(encoding="utf-8").count('stroke-dasharray="6 4"') == 1
