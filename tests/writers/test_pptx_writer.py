"""Mutation round-trips must operate on native slide objects, not screenshots."""

from zipfile import ZipFile

import pytest

pytest.importorskip("pptx")
from opendoc_model.document_model import Formula, FormulaFormat, Paragraph, Table
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Pt

from opendoc_formats.readers.pptx import read_pptx_model
from opendoc_formats.writers.pptx_writer import write_pptx_model


def test_native_mutation_roundtrip(tmp_path):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Pt(20), Pt(30), Pt(160), Pt(80))
    shape.text = "Original\nSecond line"
    shape.text_frame.paragraphs[0].runs[0].font.bold = True
    table = slide.shapes.add_table(2, 2, Pt(20), Pt(140), Pt(200), Pt(80)).table
    table.cell(0, 0).text = "Cell"
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Series", [1, 2])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(250), Pt(40), Pt(300), Pt(200), data)
    slide.notes_slide.notes_text_frame.text = "Notes"
    presentation.save(source)
    model = read_pptx_model(source)
    blocks = model.sections[0].blocks
    text = next(b for b in blocks if isinstance(b, Paragraph) and "Original" in b.plain_text)
    text.content[0].text = "Changed"
    text.box.x = 44
    table_block = next(b for b in blocks if isinstance(b, Table))
    table_block.rows[0].cells[0].blocks[0].content[0].text = "Edited cell"
    chart = next(b for b in blocks if isinstance(b, Paragraph) and b.properties.get("pptx", {}).get("chart"))
    chart.properties["pptx"]["chart"]["series"][0]["values"] = [7, 9]
    formula = Formula('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                      '<m:r><m:t>x+1</m:t></m:r></m:oMath>', FormulaFormat.OMML)
    text.content.append(formula)
    for turn in range(2):
        output = tmp_path / f"result{turn}.pptx"
        report = write_pptx_model(model, output)
        assert report.success
        reopened = Presentation(output)
        assert len(reopened.slides) == 1
        shapes = reopened.slides[0].shapes
        assert shapes[0].auto_shape_type == MSO_AUTO_SHAPE_TYPE.OVAL
        assert shapes[0].left == Pt(44)
        assert shapes[0].text_frame.paragraphs[0].runs[0].font.bold
        assert next(s for s in shapes if s.has_table).table.cell(0, 0).text == "Edited cell"
        assert list(next(s for s in shapes if s.has_chart).chart.series[0].values) == [7, 9]
        assert reopened.slides[0].notes_slide.notes_text_frame.text == "Notes"
        model = read_pptx_model(output)
        paragraph = model.sections[0].blocks[0]
        assert "Changed" in paragraph.plain_text
        assert any(isinstance(item, Formula) and "x+1" in item.value for item in paragraph.content)
        with ZipFile(output) as archive:
            assert any(name.startswith("ppt/embeddings/") for name in archive.namelist())




@pytest.mark.parametrize("preview", [False, True])
def test_svg_and_crop_survive_two_roundtrips(tmp_path, preview):
    from io import BytesIO

    from opendoc_model.document_model import Box, DocumentModel, Image, ImageCrop, Resource, ResourceKind, Section
    from PIL import Image as PillowImage

    svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><rect width="10" height="10" fill="red"/></svg>'
    resources = {"vector": Resource("vector", ResourceKind.VECTOR_IMAGE, "image/svg+xml", data=svg)}
    properties = {}
    if preview:
        buffer = BytesIO()
        PillowImage.new("RGB", (10, 10), "red").save(buffer, format="PNG")
        resources["preview"] = Resource("preview", ResourceKind.RASTER_IMAGE, "image/png", data=buffer.getvalue())
        properties["fallback_resource_id"] = "preview"
    block = Image("vector", "Красный квадрат", Box(10, 20, 100, 100), properties, ImageCrop(left=0.1))
    model = DocumentModel(sections=[Section(blocks=[block])], resources=resources)
    for index in range(2):
        path = tmp_path / f"svg{index}.pptx"
        assert write_pptx_model(model, path).success
        model = read_pptx_model(path)
        image = model.sections[0].blocks[0]
        assert model.resources[image.resource_id].data == svg
        assert image.alt_text == "Красный квадрат"
        assert image.crop.left == pytest.approx(0.1)
        assert bool(image.properties.fallback_resource_id) is preview






def test_merged_table_and_hyperlink(tmp_path):
    source, output = tmp_path / "source.pptx", tmp_path / "output.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 3, Pt(20), Pt(20), Pt(300), Pt(100)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Merged"
    table.columns[0].width = Pt(60)
    run = slide.shapes.add_textbox(Pt(20), Pt(150), Pt(200), Pt(30)).text_frame.paragraphs[0].add_run()
    run.text = "Link"
    run.hyperlink.address = "https://example.com/"
    presentation.save(source)
    assert write_pptx_model(read_pptx_model(source), output).success
    shapes = Presentation(output).slides[0].shapes
    assert shapes[0].table.cell(0, 0).span_width == 2
    assert shapes[0].table.columns[0].width == Pt(60)
    assert shapes[1].text_frame.paragraphs[0].runs[0].hyperlink.address == "https://example.com/"
