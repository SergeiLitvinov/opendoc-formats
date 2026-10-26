"""Installed merged Word flow from owned canonical DTOs, without PDF engines."""

import argparse
import importlib.util
from pathlib import Path
from tempfile import TemporaryDirectory


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--crop', choices=('top', 'bottom', 'plain'))
    args = parser.parse_args()
    from docx import Document
    from opendoc_model import (
        Box,
        ColorValue,
        DocumentModel,
        Image,
        Paragraph,
        Resource,
        ResourceKind,
        Section,
        Table,
        TableCell,
        TableRow,
        TextRun,
        document_from_json,
        document_to_json,
    )

    import opendoc_formats
    from opendoc_formats import DocxProfile, ExportOptions, get_object_mapping, read_document, write_document
    from opendoc_formats.writers.docx_writer import write_docx_model

    assert 'site-packages' in Path(opendoc_formats.__file__).parts
    for name in ('fontTools', 'pymupdf', 'fitz', 'pdf2docx', 'cv2', 'numpy', 'PIL'):
        assert importlib.util.find_spec(name) is None, name
    model = DocumentModel(source_format='pdf', sections=[Section()])
    lines = [(x, top, x, bottom) for x, top, bottom in
             ((70, 150, 270), (170, 190, 270), (270, 150, 270), (370, 150, 270))]
    lines += [(left, y, right, y) for y, left, right in
              ((150, 70, 370), (190, 70, 370), (230, 170, 370), (270, 70, 370))]
    for index, (a, b, c, d) in enumerate(lines):
        key = f'owned-{index}'
        model.resources[key] = Resource(key, ResourceKind.VECTOR_IMAGE, 'application/pdf+vector', b'owned', properties={
            'items': [['l', [a, b], [c, d]]], 'width': 1, 'stroke': ColorValue.from_hex('000000').to_dict(),
        })
        model.sections[0].blocks.append(Paragraph([Image(key, box=Box(a, b, c-a, d-b))]))
    rows = []
    for specs in (
        [('Top merge', 70, 150, 270, 190, 1, 2), ('R1C3', 270, 150, 370, 190, 1, 1)],
        [('Tall merge', 70, 190, 170, 270, 2, 1), ('R2C2', 170, 190, 270, 230, 1, 1),
         ('R2C3', 270, 190, 370, 230, 1, 1)],
        [('R3C2', 170, 230, 270, 270, 1, 1), ('R3C3', 270, 230, 370, 270, 1, 1)],
    ):
        rows.append(TableRow([TableCell([Paragraph([TextRun(text)])], row_span=rs, column_span=cs,
                                        properties={'pdf_cell_bbox': [a, b, c, d]})
                              for text, a, b, c, d, rs, cs in specs]))
    table = Table(rows, properties={
        'pdf_column_edges': [70, 170, 270, 370], 'pdf_row_edges': [150, 190, 230, 270],
        'pdf_source_cell_bounds': [list(cell.properties['pdf_cell_bbox']) for row in rows for cell in row.cells],
    })
    if args.crop == 'top':
        table.rows = table.rows[:1]
    elif args.crop in ('bottom', 'plain'):
        table.rows = table.rows[1:]
        if args.crop == 'plain':
            table.rows[0].cells = table.rows[0].cells[1:]
    model.sections[0].blocks.append(table)
    original = document_to_json(model)
    with TemporaryDirectory(prefix='opendoc_merged_flow_') as folder:
        for public in (False, True):
            path = Path(folder)/f'merged-{public}.docx'
            report = (write_document(model, path, options=ExportOptions(docx_profile=DocxProfile('flow'))) if public else
                      write_docx_model(model, path, profile=DocxProfile('flow')))
            assert report.success, report.issues
            mapping = get_object_mapping(report)
            assert mapping.publication == 'published' and len(mapping.sources) == 8
            excluded = {'top': 3, 'bottom': 1, 'plain': 2}.get(args.crop, 0)
            assert sum(len(entry.source_ids) for entry in mapping.entries if entry.extra['range_exclusion']) == excluded
            if args.crop != 'plain':
                assert len(mapping.targets[0].extra['cell_topology']) == sum(len(row.cells) for row in table.rows)
            assert document_to_json(model) == original
            for cycle in range(3):
                native = Document(path).tables[0]
                if args.crop in (None, 'top'):
                    assert native.cell(0, 0)._tc is native.cell(0, 1)._tc
                    assert native.cell(0, 0).text == 'Top merge' and abs(native.cell(0, 0).width.pt-200) < .1
                if args.crop in (None, 'bottom'):
                    row = 0 if args.crop == 'bottom' else 1
                    assert native.cell(row, 0)._tc is native.cell(row+1, 0)._tc
                    assert native.cell(row, 0).text == 'Tall merge'
                if args.crop == 'plain':
                    assert [[cell.text for cell in row.cells] for row in native.rows] == [['R2C2', 'R2C3'], ['R3C2', 'R3C3']]
                if cycle < 2:
                    imported = read_document(path)
                    assert imported.success, imported.issues
                    if args.crop == 'top':
                        restored_table = next(b for b in imported.document.sections[0].blocks if isinstance(b, Table))
                        assert restored_table.properties['docx_source_grid_identity']['state'] == 'verified'
                        assert [cell.column_span for cell in restored_table.rows[0].cells] == [2, 1]
                    path = Path(folder)/f'cycle-{public}-{cycle}.docx'
                    assert write_document(document_from_json(document_to_json(imported.document)), path).success
    print('Installed merged flow: horizontal/vertical spans, widths, 8 source occurrences and two cycles without PDF engines')
    if args.crop:
        print(f'Original topology retained after {args.crop} crop; all excluded source lines accounted')


if __name__ == '__main__':
    main()
