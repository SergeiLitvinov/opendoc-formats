"""Isolated C-IMAGE-RECT-01 candidate acceptance; never a production writer.

Run with -I in an environment containing the separately fingerprinted Model draft
and Formats wheel. The prototype below projects one own picture only, refuses
all unrelated Model edits, and uses public Model APIs for every carrier operation.
"""

from __future__ import annotations

import argparse
import json
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

import opendoc_model as od
from lxml import etree
from pptx import Presentation

import opendoc_formats
from opendoc_formats import ExporterRegistry, ExporterSpec, read_document, write_document
from opendoc_formats.ooxml.pptx_crop import A, P, read_crop
from opendoc_formats.ooxml.pptx_source_geometry import geometry_bindings
from opendoc_formats.ooxml.pptx_source_xml import SourceRoots


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    for module in (od, opendoc_formats):
        assert 'site-packages' in Path(module.__file__).parts
    assert hasattr(od, 'ImageSourceRect'), 'The published Model wheel lacks the candidate API'
    result = read_document(args.source)
    assert result.success, result.issues
    document = result.document
    image = document.sections[0].blocks[0]
    assert isinstance(image, od.Image) and image.crop is None
    roots = SourceRoots()
    native = geometry_bindings(document, roots)[0]
    node = native.shape.find(P+'blipFill/'+A+'srcRect')
    insets = read_crop(node)
    original = od.ImageSourceRect(insets.left, insets.top, 1-insets.right, 1-insets.bottom,
                                 extra={'consumer_probe': {'nested': [1, 2]}})
    baseline = od.document_to_json(document)
    resource = deepcopy(document.resources)
    box = deepcopy(image.box)
    od.set_image_source_rect(image, original)
    assert image.crop is None and image.box == box and document.resources == resource
    for cycle in range(2):
        path = args.output.parent/f'candidate-model-{cycle}.json'
        report = write_document(document, path)
        assert report.success, report.issues
        document = read_document(path).document
        image = document.sections[0].blocks[0]
        assert od.get_image_source_rect(image) == original
    detached = od.get_image_source_rect(image)
    detached.extra['consumer_probe']['nested'].append(99)
    assert od.get_image_source_rect(image).extra['consumer_probe']['nested'] == [1, 2]
    previous = deepcopy(image.properties)
    try:
        od.set_image_source_rect(image, od.ImageSourceRect(0, 0, 1, 1, extra={'large': list(range(3000))}),
                                 limits=od.DocumentLimits(max_nodes=1000))
    except od.ArtifactLimitError:
        pass
    else:
        raise AssertionError('Selected source-rectangle quota was ignored')
    assert image.properties == previous
    conflicted = deepcopy(image)
    conflicted.crop = od.ImageCrop(0.1, 0.2, 0.1, 0.1)
    try:
        od.get_image_source_rect(conflicted)
    except ValueError:
        pass
    else:
        raise AssertionError('Conflicting crop/source rectangle was accepted')
    od.set_image_source_rect(conflicted, original)
    assert conflicted.crop is None and od.get_image_source_rect(conflicted) == original
    wanted = od.ImageSourceRect(-0.26667, 0.26667, 0.86667, 0.86667)
    od.set_image_source_rect(image, wanted)
    assert od.get_image_source_rect(image).extra == original.extra
    guarded = args.output.parent/'production-blocked.html'
    guarded.write_bytes(b'previous')
    report = write_document(document, guarded)
    assert not report.success and any(i.feature == 'export.image-source-rect' for i in report.issues)
    assert guarded.read_bytes() == b'previous'

    def prototype(model, path):
        cleared = deepcopy(model)
        od.set_image_source_rect(cleared.sections[0].blocks[0], None)
        assert od.document_to_json(cleared) == baseline, 'Prototype refuses unrelated Model edits'
        rectangle = od.get_image_source_rect(model.sections[0].blocks[0])
        native = geometry_bindings(model, SourceRoots())[0]
        target = native.shape.find(P+'blipFill/'+A+'srcRect')
        values = (rectangle.left, rectangle.top, 1-rectangle.right, 1-rectangle.bottom)
        assert all(-2147483648/100000 <= value <= 2147483647/100000 for value in values)
        units = [round(value*100000) for value in values]
        assert units[0]+units[2] < 100000 and units[1]+units[3] < 100000
        for side, value in zip(('l', 't', 'r', 'b'), units, strict=True):
            target.set(side, str(value))
        edited = etree.tostring(target.getroottree(), encoding='UTF-8', xml_declaration=True, standalone=True)
        with ZipFile(args.source) as source, ZipFile(path, 'w') as output:
            for item in source.infolist():
                output.writestr(item, edited if '/'+item.filename == native.part else source.read(item))
        return od.ConversionReport(path)

    def validate(path):
        assert Presentation(path).slides[0].shapes[0].crop_left == -0.26667

    registry = ExporterRegistry()
    registry.register(ExporterSpec('candidate', ('.pptx',), prototype, validator=validate, supports_image_source_rects=True))
    before = od.document_to_json(document)
    report = registry.write(document, args.output)
    assert report.success and od.document_to_json(document) == before, report.issues
    with ZipFile(args.source) as source, ZipFile(args.output) as output:
        changed = [name for name in source.namelist() if source.read(name) != output.read(name)]
        assert changed == ['ppt/slides/slide1.xml']
    results = {'candidate_only': True, 'production_api_adopted': False, 'json_cycles': 2,
               'unknown_fields_detached': True, 'selected_quota_atomic': True, 'crop_conflict_checked': True,
               'setter_migration_checked': True, 'production_guard_checked': True, 'changed_parts': changed,
               'target': od.get_image_source_rect(image).to_dict(), 'input_model_unchanged': True,
               'resources_unchanged': document.resources == resource, 'box_unchanged': image.box == box}
    (args.output.parent/'consumer-results.json').write_text(json.dumps(results, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(results))


if __name__ == '__main__':
    main()
