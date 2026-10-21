"""Exercise each installed PDF profile in a fresh, isolated interpreter."""

from __future__ import annotations

import argparse
import importlib.util
import unicodedata
from pathlib import Path
from tempfile import TemporaryDirectory


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("profile", choices=("pdf", "pdf-rich", "pdf-layout", "pdf-docx"))
    args = parser.parse_args()
    from opendoc_model import document_from_json, document_to_json

    import opendoc_formats
    from opendoc_formats import ExportOptions, ImportOptions, read_document, write_document

    assert "site-packages" in Path(opendoc_formats.__file__).resolve().parts
    corpus = Path(__file__).resolve().parents[1] / "tests/corpus/native"
    source = corpus / "text-pages.pdf"
    absent = {
        "pdf": ("pypdf", "pdfplumber", "pdf2docx", "cv2", "numpy", "docx", "lxml", "PIL"),
        "pdf-rich": ("pypdf", "pdfplumber", "pdf2docx", "cv2", "numpy", "docx", "lxml", "PIL"),
        "pdf-layout": ("pypdf", "fitz", "pdf2docx", "cv2", "numpy", "docx", "lxml"),
        "pdf-docx": ("pypdf", "pdfplumber"),
    }[args.profile]
    assert all(importlib.util.find_spec(name) is None for name in absent), absent
    from opendoc_formats.readers.pdf import get_pdf_info
    from opendoc_formats.support.io import validate_pdf

    assert validate_pdf(source)
    assert get_pdf_info(str(source)) == {'title': '', 'author': '', 'subject': ''}
    assert all(importlib.util.find_spec(name) is None for name in absent)
    with TemporaryDirectory(prefix="opendoc-formats-pdf-profile-") as directory:
        output = Path(directory)
        if args.profile in {"pdf", "pdf-rich"}:
            import fitz
            from opendoc_model import Rect2D, get_integration, get_outline, get_page_geometry
            from opendoc_model.diagnostics import ConversionReport
            from opendoc_model.inspection import compare_inspections
            from opendoc_model.text_quality_policy import TextPreservationPolicy

            from opendoc_formats.pdf import PdfDocument
            from opendoc_formats.readers.pdf import read_pdf
            from opendoc_formats.support.inspection import inspect_path

            text = read_pdf(str(source))
            assert text.engine.startswith('pymupdf') and text.pages == 2 and 'Native PDF' in text.plain
            assert not any('ModuleNotFoundError' in warning for warning in text.warnings)

            native_inspection = inspect_path(source)
            assert native_inspection.metadata["text_flow_profile"]["complete"]
            assert native_inspection.metadata["pdf_field_names_complete"]
            assert "object_inventory_scope" not in native_inspection.metadata
            quality = ConversionReport(output / "inspection.pdf")
            assert TextPreservationPolicy("flow", 0).evaluate(
                quality, compare_inspections(native_inspection, inspect_path(source))
            )
            with PdfDocument(source) as pdf:
                assert pdf.page_count == 2
                assert pdf.page_info(0).width > 0
                assert pdf.render_page(0, max_dimension=256).png.startswith(b"\x89PNG")
            result = read_document(source)
            assert result.success, result.issues
            assert "Native PDF" in str(result.document)
            from opendoc_model import (
                Box,
                DocumentModel,
                Image,
                Length,
                PageSettings,
                Paragraph,
                Resource,
                ResourceKind,
                Section,
                TextRun,
            )

            pixels = fitz.Pixmap(fitz.csRGB, (0, 0, 12, 18), False)
            pixels.clear_with(255)
            flow = DocumentModel(sections=[Section(
                page=PageSettings(width=Length(240), height=Length(300), margin_left=Length(20),
                                  margin_right=Length(20), margin_top=Length(20), margin_bottom=Length(20)),
                blocks=[*[Paragraph([TextRun(f'Own row {i}')]) for i in range(6)],
                        Paragraph([Image('own', box=Box(0, 0, 72, 46.8), properties={'placement': 'inline'})]),
                        Paragraph([TextRun('Own first caption')]),
                        Paragraph([Image('own', box=Box(0, 0, 72, 72), properties={'placement': 'inline'})]),
                        Paragraph([TextRun('Own second caption')]), Paragraph([TextRun('Own final paragraph')])],
            )], resources={'own': Resource('own', ResourceKind.RASTER_IMAGE, 'image/png', data=pixels.tobytes('png'))})
            flow_target = output / 'flow.pdf'
            flow_report = write_document(document_from_json(document_to_json(flow)), flow_target)
            assert flow_report.success, flow_report.issues
            assert flow_report.metrics.get('pdf_flow_page_breaks')
            with fitz.open(flow_target) as flow_pdf:
                flow_text = ' '.join(unicodedata.normalize('NFKC', ''.join(page.get_text() for page in flow_pdf)).split())
                assert all(flow_text.count(text) == 1 for text in (
                    'Own first caption', 'Own second caption', 'Own final paragraph',
                ))
                assert sum(len(page.get_image_info()) for page in flow_pdf) == 2
            from opendoc_formats.errors import OperationCancelledError
            from opendoc_formats.writers.pdf_docx_model import prepare_pdf_docx

            original = document_to_json(result.document)
            prepared = prepare_pdf_docx(result.document)
            assert document_to_json(prepare_pdf_docx(result.document, cancelled=lambda: False)) == document_to_json(prepared)
            polls = 0

            def stop_preparation() -> bool:
                nonlocal polls
                polls += 1
                return polls >= 12

            try:
                prepare_pdf_docx(result.document, cancelled=stop_preparation)
            except OperationCancelledError:
                pass
            else:
                raise AssertionError('PDF to Word preparation ignored cancellation')
            assert polls == 12 and document_to_json(result.document) == original
            assert all(importlib.util.find_spec(name) is None for name in absent)
            from opendoc_formats.readers.pdf_ocr_types import OcrPageResult

            class OwnOcrEngine:
                is_available = True

                def recognize_pdf_geometry(self, *args, **kwargs):
                    return [OcrPageResult(warnings=['Own installed OCR warning']), OcrPageResult()]

            warned = read_document(source, options=ImportOptions(
                pdf_mode='structure', ocr_engine_factory=lambda **kwargs: OwnOcrEngine(),
            ))
            assert warned.success and any(i.message == 'Own installed OCR warning' for i in warned.issues)
            transport = output / 'ocr-warnings.json'
            assert write_document(warned.document, transport).success
            assert any(i.message == 'Own installed OCR warning' for i in read_document(transport).issues)
            from opendoc_formats.errors import OperationCancelledError

            class CancelledOcrEngine:
                is_available = True

                def recognize_pdf_geometry(self, *args, **kwargs):
                    raise OperationCancelledError('Own installed cancellation')

            cancelled = read_document(source, options=ImportOptions(
                pdf_mode='structure', ocr_engine_factory=lambda **kwargs: CancelledOcrEngine(), cancelled=lambda: False,
            ))
            assert cancelled.document is None and not cancelled.success and cancelled.issues[0].code == 'import.cancelled'
            report = write_document(result.document, output / "document.pdf")
            assert report.success, report.to_dict()
            assert read_document(output / "document.pdf").success
            interactive_source = output / "interactive.pdf"
            with fitz.open() as fixture:
                page = fixture.new_page(width=300, height=400)
                page.insert_text((30, 60), "Interactive PDF")
                page.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(30, 40, 140, 65), "uri": "https://example.invalid"})
                page.add_text_annot((160, 90), "Own note")
                widget = fitz.Widget()
                widget.field_name = "own-field"
                widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
                widget.field_value = "Own value"
                widget.rect = fitz.Rect(30, 150, 170, 175)
                page.add_widget(widget)
                page.set_rotation(90)
                fixture.set_toc([[1, "Own outline", 1]])
                fixture.save(interactive_source)
            interactive = read_document(interactive_source)
            assert interactive.success, interactive.issues
            interaction_model = document_from_json(document_to_json(interactive.document))
            interaction = get_integration(interaction_model)
            assert len(interaction.annotations) == 2 and interaction.forms[0].value == "Own value"
            assert interaction_model.resources["pdf-original-source"].data == interactive_source.read_bytes()
            assert get_outline(interaction_model).entries[0].target.target_id == "pdf-page-1"
            assert get_page_geometry(interaction.pages[0]).media_box == Rect2D(0, -400, 300, 400)
            interaction_report = write_document(interaction_model, output / "interactive-export.pdf")
            assert interaction_report.success and not interaction_report.lossless, interaction_report.to_dict()
            with fitz.open(output / "interactive-export.pdf") as converted:
                assert converted[0].rotation == 90
                assert converted.get_toc()[0][:3] == [1, "Own outline", 1]
            vector = read_document(corpus / "vector-export.pdf")
            assert vector.success, vector.issues
            assert any(r.media_type == "application/pdf+vector" for r in vector.document.resources.values())
            restored = document_from_json(document_to_json(vector.document))
            assert restored.resources == vector.document.resources
            vector_report = write_document(vector.document, output / "vector.pdf")
            assert vector_report.success and vector_report.metrics["pdf_vectors"]["native"] == 1, vector_report.to_dict()
            from opendoc_model import Resource, ResourceKind, VisualSurrogate

            from opendoc_formats import get_object_mapping, read_pdf_origin_mapping
            from opendoc_formats.writers.html_resources import document_images

            image, _ = next(document_images(restored))
            restored.resources[image.resource_id].properties["items"] = [["unknown"]]
            pixels = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20), False)
            pixels.clear_with(128)
            restored.resources["own-visual"] = Resource(
                "own-visual", ResourceKind.RASTER_IMAGE, "image/png", data=pixels.tobytes("png")
            )
            image.visual_surrogate = VisualSurrogate("own-visual", "unsupported vector", "image/png")
            mapped_output = output / "mapped.pdf"
            mapped = write_document(restored, mapped_output)
            assert mapped.success, mapped.to_dict()
            mapping = get_object_mapping(mapped)
            assert mapping.publication == "published" and mapping.assessment_complete and mapping.targets
            previous = mapped_output.read_bytes()
            rejected = write_document(restored, mapped_output, options=ExportOptions(
                quality_gate=lambda report, staged: report.lossless,
            ))
            assert not rejected.success and mapped_output.read_bytes() == previous
            assert get_object_mapping(rejected).publication == "rejected" and not get_object_mapping(rejected).targets
            origins = read_pdf_origin_mapping(mapped_output)
            expected_origins = {item.id for item in origins.sources}
            for cycle in range(2):
                imported = read_document(mapped_output)
                assert imported.success, imported.issues
                carried = document_from_json(document_to_json(imported.document))
                mapped_output = output / f"origin-cycle-{cycle}.pdf"
                report = write_document(carried, mapped_output)
                assert report.success, report.to_dict()
                origins = read_pdf_origin_mapping(mapped_output)
                assert origins.assessment_complete and {item.id for item in origins.sources} == expected_origins
            from opendoc_model import DocumentModel, Image, Paragraph, Section, validate_object_mapping

            svg_model = DocumentModel(resources={
                'own-svg': Resource('own-svg', ResourceKind.VECTOR_IMAGE, 'image/svg+xml', data=(
                    b'<svg xmlns="http://www.w3.org/2000/svg" width="40" height="20">'
                    b'<circle cx="20" cy="10" r="8" fill="blue" opacity="0.4"/></svg>'
                )),
            }, sections=[Section(blocks=[Paragraph([Image('own-svg')]), Paragraph([Image('own-svg')])])])
            svg_output = output / 'svg.pdf'
            svg_report = write_document(svg_model, svg_output)
            assert svg_report.success, svg_report.to_dict()
            svg_mapping = get_object_mapping(svg_report)
            assert svg_mapping.assessment_complete and len(svg_mapping.targets) == 2
            validate_object_mapping(svg_mapping, source=svg_model, target=read_document(svg_output).document)
            svg_origins = {item.id for item in read_pdf_origin_mapping(svg_output).sources}
            from opendoc_formats import compose_object_mappings

            composed = svg_mapping
            for cycle in range(2):
                carried = document_from_json(document_to_json(read_document(svg_output).document))
                svg_output = output / f'svg-cycle-{cycle}.pdf'
                svg_report = write_document(carried, svg_output)
                assert svg_report.success, svg_report.to_dict()
                next_mapping = get_object_mapping(svg_report)
                source_nodes = {item.node_id: item.id for item in next_mapping.sources}
                bridges = {item.id: (source_nodes[item.node_id],) for item in composed.targets}
                composed = compose_object_mappings(
                    composed, next_mapping, bridges, source=svg_model, target=read_document(svg_output).document,
                )
                assert composed.assessment_complete and len(composed.targets) == 2
                restored = read_pdf_origin_mapping(svg_output)
                assert restored.assessment_complete and {item.id for item in restored.sources} == svg_origins
            raster_source = output / "raster-source.pdf"
            with fitz.open() as fixture:
                pixels = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 30), False)
                pixels.clear_with(128)
                pixels.set_pixel(0, 0, (255, 0, 0))
                for rotation in (0, 90):
                    page = fixture.new_page(width=120, height=180)
                    page.insert_image(page.rect, stream=pixels.tobytes("png"), keep_proportion=False)
                    page.set_cropbox(fitz.Rect(10, 20, 110, 160))
                    page.set_rotation(rotation)
                fixture.save(raster_source)

            def raster_snapshot(path):
                with fitz.open(path) as fixture:
                    return [(tuple(page.rect), page.rotation,
                             [tuple(item["bbox"]) for item in page.get_image_info()],
                             page.get_pixmap(alpha=False).samples) for page in fixture]

            expected_raster = raster_snapshot(raster_source)
            raster_path = raster_source
            for cycle in range(2):
                raster = read_document(raster_path)
                assert raster.success, raster.issues
                raster_model = document_from_json(document_to_json(raster.document))
                assert all(len(section.blocks) == 1 for section in raster_model.sections)
                raster_path = output / f"raster-{cycle}.pdf"
                raster_report = write_document(raster_model, raster_path)
                assert raster_report.success and raster_report.metrics["pdf_rasters"]["native"] == 2
                assert raster_snapshot(raster_path) == expected_raster
            assert not read_document(source, options=ImportOptions(cancelled=lambda: True)).success
            from opendoc_model import Box, DocumentModel, Image, ImageCrop, Resource, ResourceKind, Section

            cropped = DocumentModel(
                sections=[Section(blocks=[Image("own", box=Box(30, 40, 80, 100, rotation=3),
                                               crop=ImageCrop(0.1, 0.2, 0.1, 0.1))])],
                resources={"own": Resource("own", ResourceKind.RASTER_IMAGE, "image/png", data=pixels.tobytes("png"))},
            )
            crop_path = output / "crop-source.pdf"
            assert write_document(cropped, crop_path).success
            expected_crop = raster_snapshot(crop_path)
            for cycle in range(2):
                imported = read_document(crop_path)
                assert imported.success, imported.issues
                crop_model = document_from_json(document_to_json(imported.document))
                assert crop_model.sections[0].blocks[0].content[0].crop == ImageCrop(0.1, 0.2, 0.1, 0.1)
                crop_path = output / f"crop-{cycle}.pdf"
                assert write_document(crop_model, crop_path).success
                actual_crop = raster_snapshot(crop_path)
                assert len(actual_crop) == len(expected_crop)
                for actual, expected in zip(actual_crop, expected_crop, strict=True):
                    assert actual[:2] == expected[:2] and actual[3] == expected[3]
            target = output / "cancelled.pdf"
            target.write_bytes(b"existing")
            cancelled = write_document(result.document, target, options=ExportOptions(cancelled=lambda: True))
            assert not cancelled.success and target.read_bytes() == b"existing"
        elif args.profile == "pdf-layout":
            from opendoc_formats.readers.pdf import read_pdf

            text = read_pdf(str(source))
            assert text.engine == "pdfplumber" and text.pages == 2 and "Native PDF" in text.plain
            result = read_document(source)
            assert not result.success and result.issues[0].code == "import.backend-unavailable"
        else:
            from docx import Document
            from pdf2docx import Converter

            # The engine's single-worker route keeps this acceptance bounded.
            converter = Converter(str(source))
            try:
                target = output / "document.docx"
                converter.convert(str(target), multi_processing=False)
            finally:
                converter.close()
            assert "Native PDF" in "\n".join(p.text for p in Document(target).paragraphs)
            assert read_document(source).success  # pdf2docx transitively provides PyMuPDF.
    print(f"Installed profile passed: {args.profile}; excluded engines are absent")


if __name__ == "__main__":
    main()
