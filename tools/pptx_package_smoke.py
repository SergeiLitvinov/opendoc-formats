"""Installed PPTX-only profile: original opaque bytes, graph and two JSON cycles."""

import hashlib
import importlib.util
import io
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile

import opendoc_model as od
from lxml import etree
from PIL import Image as Pixels
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

import opendoc_formats
from opendoc_formats import ExportOptions, PptxProfile, read_document, write_document


def main():
    for module in (od, opendoc_formats):
        assert 'site-packages' in Path(module.__file__).parts
    for name in ('fontTools', 'docx', 'pymupdf', 'pypdf', 'bs4'):
        assert importlib.util.find_spec(name) is None, name
    with TemporaryDirectory(prefix='opendoc_pptx_package_') as directory:
        source = Path(directory)/'own.pptx'
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        textbox = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1))
        textbox.text = 'Own editable text'
        table = slide.shapes.add_table(2, 3, 0, Inches(2), Inches(4), Inches(1)).table
        table.cell(0, 0).merge(table.cell(1, 1))
        table.cell(0, 0).text = 'Own cell'
        a = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
        for paragraph in (textbox.text_frame.paragraphs[0], table.cell(0, 0).text_frame.paragraphs[0]):
            paragraph.add_line_break()
            field = etree.SubElement(paragraph._p, a+'fld',
                                     id='{419B6445-F189-4081-90FE-CF44AD924004}', type='slidenum')
            etree.SubElement(field, a+'t').text = '1'
            paragraph.add_run().text = ' Protected neighbor'
        chart_data = CategoryChartData()
        chart_data.categories = ['Alpha', 'Beta']
        chart_data.add_series('Own source series', [2, 4])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(2), Inches(3), Inches(2), chart_data)
        image = Pixels.new('RGBA', (4, 4), (10, 100, 200, 128))
        pixels = io.BytesIO()
        image.save(pixels, format='PNG')
        slide.shapes.add_picture(io.BytesIO(pixels.getvalue()), 0, 0, Inches(1), Inches(1))
        inherited = deck.slides.add_slide(deck.slide_layouts[1])
        inherited.shapes.title.text = 'Own placeholder title'
        inherited.placeholders[1].text = 'Own placeholder body'
        deck.save(source)
        buffer = io.BytesIO()
        with ZipFile(source) as original, ZipFile(buffer, 'w', ZIP_DEFLATED) as output:
            for item in original.infolist():
                output.writestr(item, original.read(item))
            output.writestr('own/unknown.bin', b'own inert payload\x00\xff')
            output.writestr('own/_rels/unknown.bin.rels', (
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rIdLink" Type="urn:own:external" TargetMode="External" '
                'Target="https://example.invalid/never-fetch"/></Relationships>'
            ))
        source.write_bytes(buffer.getvalue())
        digest = hashlib.sha256(source.read_bytes()).digest()
        result = read_document(source)
        assert result.success and not result.lossless, result.issues
        model = result.document
        assert [len(row.cells) for row in model.sections[0].blocks[1].rows] == [2, 1]
        expected = od.document_to_json(model)
        for _ in range(2):
            model = od.document_from_json(od.document_to_json(model))
            assert model.validate() == [] and od.document_to_json(model) == expected
            assert model.package.format == 'pptx'
            with ZipFile(source) as archive:
                assert set(model.package.parts) == {'/'+name for name in archive.namelist()}
                for name, part in model.package.parts.items():
                    assert part.data == archive.read(name[1:])
            assert any(edge.source == '/own/unknown.bin' and edge.external
                       and edge.target == 'https://example.invalid/never-fetch' for edge in model.package.relationships)
            records = od.get_integration(model).preservation
            assert {r.provenance.package_part for r in records if r.issue.code == 'pptx.package-part'} == set(model.package.parts)
            assert 'Own editable text' in str(model)
        assert hashlib.sha256(source.read_bytes()).digest() == digest
        model.sections[0].blocks[0].content[0].text = 'Own native edit'
        model.sections[0].blocks[0].box.x += 10
        model.sections[0].blocks[0].box.width += 12
        model.sections[0].blocks[0].box.rotation += 5
        model.sections[0].blocks[1].rows[0].cells[0].blocks[0].content[0].text = 'Own cell edit'
        model.sections[0].blocks[1].rows[0].cells[1].blocks[0].content.append(od.TextRun('Own empty cell edit'))
        model.sections[0].blocks[1].box.width += 72
        model.sections[0].blocks[1].box.height += 36
        model.sections[0].blocks[2].box.width += 24
        model.sections[0].blocks[2].box.height += 12
        model.sections[0].blocks[3].crop = od.ImageCrop(0.1, 0.2, 0.1, 0.1)
        model.sections[1].blocks[1].box.x += 10
        model.sections[1].blocks[1].box.width -= 20
        model.sections[1].blocks[1].content[0].text = 'Own placeholder edit'
        output = Path(directory)/'edited.pptx'
        report = write_document(model, output, options=ExportOptions(pptx_profile=PptxProfile(layout='source')))
        assert report.success, report.issues
        assert report.metrics['geometry_edits'] == 4
        assert report.metrics['cell_text_edits'] == 2
        assert report.metrics['table_grid_edits'] == 1
        assert report.metrics['chart_frame_edits'] == 1
        assert report.metrics['image_crop_edits'] == 1
        with ZipFile(source) as old, ZipFile(output) as new:
            assert set(old.namelist()) == set(new.namelist())
            assert [name for name in old.namelist() if old.read(name) != new.read(name)] == [
                'ppt/slides/slide1.xml', 'ppt/slides/slide2.xml',
            ]
            before_xml = etree.fromstring(old.read('ppt/slides/slide1.xml'))
            after_xml = etree.fromstring(new.read('ppt/slides/slide1.xml'))
            for tag in ('fld', 'br'):
                assert ([etree.tostring(node) for node in before_xml.iter(a+tag)]
                        == [etree.tostring(node) for node in after_xml.iter(a+tag)])
        assert 'Own native edit' in str(read_document(output).document)
        chart = Presentation(output).slides[0].shapes[2]
        assert chart.width == Inches(3)+Pt(24) and chart.height == Inches(2)+Pt(12)
        assert Presentation(output).slides[0].shapes[3].crop_left == 0.1
        placeholder = Presentation(output).slides[1].placeholders[1]
        assert placeholder.text == 'Own placeholder edit' and placeholder.left == Pt(46) and placeholder.width == Pt(628)
        rebuilt = Path(directory)/'rebuilt.pptx'
        report = write_document(read_document(output).document, rebuilt)
        assert report.success, report.issues
        table = Presentation(rebuilt).slides[0].shapes[1].table
        assert len(table.columns) == 3 and table.cell(0, 0).span_width == 2 and table.cell(0, 0).span_height == 2
        assert table.cell(1, 1).is_spanned and table.cell(0, 2).text == 'Own empty cell edit'
        assert sum(column.width for column in table.columns) == Inches(5)
        assert sum(row.height for row in table.rows) == Inches(1.5)
    print('Installed PPTX-only profile: physical parts, inert external edge and two JSON cycles passed')


if __name__ == '__main__':
    main()
